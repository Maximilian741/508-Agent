"""Write approved fixes back into an .xlsx workbook.

Surgery, not a round trip. The writer rewrites ONLY the XML parts a fix
touches and copies every other part of the package through unchanged:

* ``docProps/core.xml``        — ``dc:title`` / ``dc:language``;
* ``xl/drawings/drawingN.xml`` — ``descr`` (alt text) on one picture/chart;
* a new ``xl/tables/tableN.xml`` plus the sheet's ``<tableParts>``, its
  relationships and ``[Content_Types].xml`` — turning a clear header + data
  range into an Excel table whose first row is its header row.

Loading and re-saving with openpyxl is deliberately NOT how this works:
openpyxl's writer sets ``descr="Picture"`` on every image it saves (wiping
alt text the author wrote), and drops shapes, text boxes, slicers, sparklines
and other parts it does not model. An accessibility fix that silently
destroys existing accessibility work is not a fix.

What gets written is decided by diffing the (executor-mutated) tree against
a fresh parse of the source: an approved executor changes the tree, an
unapproved finding leaves it as parsed, so only approved fixes reach bytes.

Memory stays bounded whatever the workbook's size: a sheet part is never
parsed whole (``<tableParts>`` is spliced into its short tail while the part
is streamed), every other part is stream-copied, and verification re-reads
only the parts an edit touched.

Every entry in ``applied`` means the edit is in the saved file: each edited
part is read back with our own scanner's readers (a spliced sheet is proved
byte-identical to the source through the end of its ``<sheetData>``), the
output is reopened with openpyxl before it is accepted, and each edit is
confirmed individually. The pipeline treats XLSX
as writer-confirmed for every action (see ``_WRITER_CONFIRMED_FORMATS`` in
app/api/pipeline.py), so an approved fix that is not in ``applied`` is
neither counted nor charged. Should verification fail with the table
conversions in, the writer retries without them (they are the only
structural change) before giving up with ``failed_to_save_xlsx``.
"""

from __future__ import annotations

import logging
import os
import posixpath
import re
import shutil
import warnings
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from lxml import etree

from app.models.accessibility import (
    AccessibilityTree,
    ImageNode,
    TableCellNode,
    TableCellType,
    TableNode,
    TableRowNode,
    iter_reading_order,
)
from app.parsers.xlsx_parser import (
    CT_CORE,
    CT_TABLE,
    NS_CP,
    NS_CT,
    NS_DC,
    NS_MAIN,
    NS_PKG_REL,
    NS_R,
    REL_CORE,
    REL_TABLE,
    DataBlock,
    Package,
    SheetScan,
    WorkbookScan,
    build_tree,
    cnvpr_of,
    format_range,
    iter_drawing_objects,
    parse_xml,
    rels_part_for,
    scan_workbook,
    xlsx_work_slot,
)

logger = logging.getLogger(__name__)

# openpyxl's full (non read-only) load is the strongest reopen check, but it
# materializes every cell; when the worksheets hold more XML than this the
# read-only load plus a schema-level parse of each new table part is used
# instead (a 60,000-row export spent ~8 s in the full load alone).
_FULL_VERIFY_MAX_SHEET_XML = 12 * 1024 * 1024


@dataclass
class _AltEdit:
    node_id: str
    action: str
    drawing_part: str
    index: int
    cnvpr_id: str
    descr: Optional[str]          # None removes the attribute
    what: str


@dataclass
class _TableEdit:
    node_id: str
    sheet: SheetScan
    block: DataBlock
    columns: List[str]
    # Filled in while building the package.
    part: str = ""
    name: str = ""
    table_id: int = 0
    rid: str = ""


@dataclass
class _Edits:
    title: Optional[str] = None
    language: Optional[str] = None
    alts: List[_AltEdit] = field(default_factory=list)
    tables: List[_TableEdit] = field(default_factory=list)

    def empty(self) -> bool:
        return self.title is None and self.language is None and not self.alts and not self.tables


def write_remediated_xlsx(source_path: Path, tree: AccessibilityTree, output_path: Path) -> Dict[str, Any]:
    # Scanning the source again (to diff against) and rewriting the package
    # share the parser's per-process XLSX slots: see xlsx_work_slot.
    with xlsx_work_slot():
        return _write_remediated_xlsx(Path(source_path), tree, Path(output_path))


def _write_remediated_xlsx(source_path: Path, tree: AccessibilityTree, output_path: Path) -> Dict[str, Any]:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source_path, output_path)

    applied: List[Dict[str, Any]] = []
    skipped: List[Dict[str, Any]] = []
    root_id = tree.root.id

    try:
        scan = scan_workbook(source_path)
        reference = build_tree(source_path, scan).tree
    except Exception as exc:
        logger.warning("xlsx writer could not open the source: %s", exc)
        skipped.append({"target_id": root_id, "reason": f"failed_to_open_xlsx:{exc}"})
        return {"applied": applied, "skipped": skipped}

    edits = _collect_edits(tree, reference, scan, skipped)
    if edits.empty():
        return {"applied": applied, "skipped": skipped}

    tmp = output_path.with_name(output_path.name + ".tmp")
    try:
        problems = _build_and_verify(source_path, tmp, edits, scan)
        if problems and edits.tables:
            # The table conversion is the only structural edit; keep the
            # metadata/alt fixes rather than lose everything to it.
            logger.warning("xlsx table conversion failed verification, retrying without it: %s", problems)
            for te in edits.tables:
                skipped.append(
                    {
                        "target_id": te.node_id,
                        "action": "ADD_TABLE_HEADERS",
                        "reason": "output_failed_verification: the converted table did not reopen cleanly, so it was left as it was",
                    }
                )
            edits.tables = []
            if edits.empty():
                tmp.unlink(missing_ok=True)
                return {"applied": applied, "skipped": skipped}
            problems = _build_and_verify(source_path, tmp, edits, scan)
        if problems:
            raise RuntimeError("; ".join(problems)[:500])
        os.replace(str(tmp), str(output_path))
    except Exception as exc:
        try:
            tmp.unlink(missing_ok=True)
        except Exception:
            pass
        # Restore the untouched source so nothing half-written is left behind.
        try:
            shutil.copyfile(source_path, output_path)
        except Exception:
            pass
        logger.exception("Failed to save remediated xlsx: %s", exc)
        return {"applied": [], "skipped": skipped + [{"target_id": root_id, "reason": f"failed_to_save_xlsx:{exc}"}]}

    if edits.title is not None:
        applied.append(
            {"kind": "document_title", "action": "SET_DOCUMENT_TITLE", "target_id": root_id,
             "summary": f"Set the workbook title to {edits.title!r}"}
        )
    if edits.language is not None:
        applied.append(
            {"kind": "document_language", "action": "SET_DOCUMENT_LANGUAGE", "target_id": root_id,
             "summary": f"Set the workbook language to {edits.language!r}"}
        )
    for ae in edits.alts:
        applied.append(
            {"kind": "image_alt_text", "action": ae.action, "target_id": ae.node_id,
             "summary": (f"Removed alt text from the decorative {ae.what}" if ae.descr is None
                         else f"Alt text on the {ae.what}: {ae.descr!r}")}
        )
    for te in edits.tables:
        applied.append(
            {"kind": "excel_table", "action": "ADD_TABLE_HEADERS", "target_id": te.node_id,
             "summary": (f"Formatted {te.block.ref} on sheet {te.sheet.name!r} as an Excel table "
                         f"({te.name}) whose first row is its header row: " + ", ".join(te.columns[:8])
                         + (", …" if len(te.columns) > 8 else ""))}
        )
    return {"applied": applied, "skipped": skipped}


# ---------------------------------------------------------------------------
# What changed
# ---------------------------------------------------------------------------


def _collect_edits(tree: AccessibilityTree, reference: AccessibilityTree, scan: WorkbookScan, skipped: List[Dict[str, Any]]) -> _Edits:
    edits = _Edits()
    root, ref_root = tree.root, reference.root

    new_title = (root.metadata.properties or {}).get("title")
    old_title = (ref_root.metadata.properties or {}).get("title")
    if isinstance(new_title, str) and new_title.strip() and new_title.strip() != (old_title or "").strip():
        if _XML_UNSAFE_RE.search(new_title):
            skipped.append({"target_id": root.id, "action": "SET_DOCUMENT_TITLE",
                            "reason": "title_refused: it contains characters a workbook cannot store"})
        else:
            edits.title = new_title.strip()

    new_lang = (root.metadata.language or "").strip()
    old_lang = (ref_root.metadata.language or "").strip()
    if new_lang and new_lang != old_lang:
        if _XML_UNSAFE_RE.search(new_lang):
            skipped.append({"target_id": root.id, "action": "SET_DOCUMENT_LANGUAGE",
                            "reason": "language_refused: it contains characters a workbook cannot store"})
        else:
            edits.language = new_lang

    ref_nodes = {n.id: n for n in iter_reading_order(ref_root)}
    sheets = {s.index: s for s in scan.sheets}

    for node in iter_reading_order(root):
        ref = ref_nodes.get(node.id)
        if isinstance(node, ImageNode) and isinstance(ref, ImageNode):
            _collect_alt(node, ref, edits, skipped)
        elif isinstance(node, TableNode) and isinstance(ref, TableNode):
            _collect_table(node, ref, sheets, edits, skipped)
    return edits


# "Image xlsx-s1-chart1 — ..." / "Chart xlsx-s2-img1 - ...": a label built
# from our own node id, the shape an offline suggestion takes when it only
# glued an id in front of a caption.
_ID_LABEL_ALT_RE = re.compile(
    r"^(?:image|picture|figure|graphic|photo|chart)\s+[a-z]+(?:-[a-z0-9]+)+\s+[—–-]\s+",
    re.IGNORECASE,
)
# Characters XML 1.0 cannot carry in an attribute or element: writing one
# would make lxml refuse the whole part and fail every other fix with it.
_XML_UNSAFE_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def _alt_problem(text: str, node_id: str) -> Optional[str]:
    """Why ``text`` must not be written as this object's alt text, or None."""
    from app.analyzers.image_analyzer import is_nondescriptive_alt

    if _XML_UNSAFE_RE.search(text):
        return "the suggested description contains characters a workbook cannot store"
    if node_id and node_id.lower() in text.lower():
        return "the suggested description contains our internal id for the object instead of words about it"
    if _ID_LABEL_ALT_RE.match(text):
        return "the suggested description is a label, not a description"
    if is_nondescriptive_alt(text):
        return "the suggested description is a file name or placeholder"
    return None


def _collect_alt(node: ImageNode, ref: ImageNode, edits: _Edits, skipped: List[Dict[str, Any]]) -> None:
    new = (node.alt_text or "").strip()
    old = (ref.alt_text or "").strip()
    if new == old:
        return
    props = ref.metadata.properties or {}
    drawing = props.get("drawing_part")
    if not drawing or props.get("object_index") is None:
        skipped.append({"target_id": node.id, "reason": "image_not_addressable"})
        return
    kind = "chart" if props.get("object_kind") == "chart" else "picture"
    what = f"{kind} at {props.get('anchor_cell') or 'its position'} on sheet {props.get('sheet_name')!r}"
    if new:
        problem = _alt_problem(new, node.id)
        if problem:
            # The last gate before the customer's file: whatever produced the
            # text, this is not a description a person would accept, so the
            # object keeps its finding and nothing is written or charged.
            skipped.append({"target_id": node.id, "action": "GENERATE_ALT_TEXT", "reason": f"alt_text_refused: {problem}"})
            return
        edits.alts.append(
            _AltEdit(node.id, "GENERATE_ALT_TEXT", drawing, int(props["object_index"]), str(props.get("cnvpr_id") or ""), new, what)
        )
    elif node.is_decorative and old:
        edits.alts.append(
            _AltEdit(node.id, "REMOVE_DECORATIVE_ALT_TEXT", drawing, int(props["object_index"]), str(props.get("cnvpr_id") or ""), None, what)
        )


def _collect_table(node: TableNode, ref: TableNode, sheets: Dict[int, SheetScan], edits: _Edits, skipped: List[Dict[str, Any]]) -> None:
    ref_props = ref.metadata.properties or {}
    if ref_props.get("header_detection") != "candidate":
        return
    rows = [r for r in node.children if isinstance(r, TableRowNode)]
    ref_rows = [r for r in ref.children if isinstance(r, TableRowNode)]
    if not rows or not ref_rows:
        return
    first = rows[0]
    cells = [c for c in first.children if isinstance(c, TableCellNode)]
    promoted = bool(cells) and all(c.cell_type == TableCellType.HEADER for c in cells)
    if not promoted:
        return  # the fix was not approved (or not run) for this block
    if first.id != ref_rows[0].id or any((c.metadata.properties or {}).get("synthesized") for c in cells):
        # A made-up "Column 1 | Column 2" row is not a header the author
        # wrote; an Excel table built on it would announce placeholders.
        skipped.append(
            {"target_id": node.id, "action": "ADD_TABLE_HEADERS",
             "reason": "no_clear_header_row: the first row is not a row of headings, so no table was created"}
        )
        return
    sheet = sheets.get(int(ref_props.get("sheet_index") or 0))
    block = None
    if sheet is not None:
        block = next((b for b in sheet.blocks if b.ref == ref_props.get("cell_range") and b.state == "candidate"), None)
    if sheet is None or block is None:
        skipped.append({"target_id": node.id, "action": "ADD_TABLE_HEADERS", "reason": "range_not_found"})
        return
    columns = []
    for c in range(block.c1, block.c2 + 1):
        info = sheet.grid.get((block.r1, c))
        columns.append(info.text if info is not None else "")
    if any(not c.strip() for c in columns):
        skipped.append({"target_id": node.id, "action": "ADD_TABLE_HEADERS", "reason": "header_cell_empty"})
        return
    edits.tables.append(_TableEdit(node_id=node.id, sheet=sheet, block=block, columns=columns))


# ---------------------------------------------------------------------------
# Building the package
# ---------------------------------------------------------------------------


def _xml_bytes(root: etree._Element) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _build_and_verify(source: Path, dest: Path, edits: _Edits, scan: WorkbookScan) -> List[str]:
    replacements, additions, splices = _plan_parts(source, edits, scan)
    try:
        _rewrite_zip(source, dest, replacements, additions, splices)
    except _SpliceError as exc:
        # Only a table conversion splices a sheet: report it like a failed
        # verification so the caller retries without the tables and the
        # other fixes still land.
        return [f"could not add the table to its sheet: {exc}"]
    return _verify(dest, edits, source)


class _SpliceError(ValueError):
    """A sheet part we cannot find our way through to add <tableParts>."""


# ---------------------------------------------------------------------------
# Adding <tableParts> to a sheet part without parsing it
#
# A sheet part can be tens of megabytes of <row>/<c> elements; building an
# lxml tree of it just to append one element cost ~10x its size (a 4 MB
# ledger export peaked at 950 MB). <tableParts> lives AFTER <sheetData> (only
# <extLst> may follow it), so the part is streamed: everything up to the end
# of <sheetData> is copied byte for byte, and only the short tail after it is
# edited. Verification then proves the head is byte-identical to the source
# and re-reads the tail.
# ---------------------------------------------------------------------------

_SPLICE_CHUNK = 1 << 20
_SPLICE_KEEP = 4096                 # longer than any token searched for
_SPLICE_MAX_PROLOG = 1 << 20        # the root tag must start within this
_NAME = rb"[A-Za-z_][\w.\-]*"
_ATTRS = rb"(?:\s+[^\s=<>/]+\s*=\s*(?:\"[^\"]*\"|'[^']*'))*"
_PROLOG_ITEM_RE = re.compile(rb"\s*(?:<\?.*?\?>|<!--.*?-->|<!DOCTYPE[^>\[]*(?:\[.*?\])?\s*>)", re.S)
_ROOT_TAG_RE = re.compile(rb"\s*(?P<tag><(?P<qname>(?:" + _NAME + rb":)?worksheet)" + _ATTRS + rb"\s*>)")
# A tag whose name is (prefix:)sheetData, matched at the '<' before a
# "sheetData" found with bytes.find (a regex scan of every '<' in a 60 MB
# part was ~0.4 s; find() is ~20x faster).
_SHEETDATA_TAG_RE = re.compile(
    rb"<(?P<close>/?)(?:" + _NAME + rb":)?sheetData(?:\s[^<>]*?)?(?P<self>/?)>"
)
_OPENERS = ((b"<!--", b"-->"), (b"<![CDATA[", b"]]>"), (b"<?", b"?>"))
_TAIL_TOKEN_RE = re.compile(
    rb"<!--.*?-->|<!\[CDATA\[.*?\]\]>|<\?.*?\?>"
    rb"|<(?P<close>/?)(?P<prefix>(?:" + _NAME + rb":)?)(?P<local>" + _NAME + rb")(?P<attrs>" + _ATTRS + rb")\s*(?P<self>/?)>",
    re.S,
)
_COUNT_ATTR_RE = re.compile(rb"(\scount\s*=\s*)(\"[^\"]*\"|'[^']*')")


class _SheetSplit:
    """The pieces of a streamed sheet part: its root start tag, how long the
    head (everything through the end of <sheetData>) is and its digest, and
    the tail after it."""

    def __init__(self) -> None:
        self.root_tag = b""
        self.qname = b""
        self.head_len = 0
        self.head_digest = ""
        self.tail = b""


def _split_sheet(src, emit=None) -> _SheetSplit:
    """Stream a sheet part from ``src``; pass the head to ``emit`` (when
    given) as it goes, and return the split. Raises ValueError when the part
    is not a worksheet this can find its way through."""
    import hashlib

    out = _SheetSplit()
    digest = hashlib.sha256()
    buf = bytearray()
    eof = False

    def more() -> bool:
        nonlocal eof
        chunk = src.read(_SPLICE_CHUNK)
        if chunk:
            buf.extend(chunk)
            return True
        eof = True
        return False

    def flush(n: int) -> None:
        if n <= 0:
            return
        piece = bytes(buf[:n])
        digest.update(piece)
        out.head_len += n
        if emit is not None:
            emit(piece)
        del buf[:n]

    # The root start tag (after an optional BOM, the XML declaration and any
    # comments or processing instructions).
    while True:
        pos = 3 if buf[:3] == b"\xef\xbb\xbf" else 0
        while True:
            m = _PROLOG_ITEM_RE.match(buf, pos)
            if m is None or m.end() == pos:
                break
            pos = m.end()
        root = _ROOT_TAG_RE.match(buf, pos)
        if root is not None:
            break
        if eof or len(buf) > _SPLICE_MAX_PROLOG:
            raise ValueError("the sheet part does not start with a worksheet element")
        more()
    out.root_tag = bytes(root.group("tag"))
    out.qname = bytes(root.group("qname"))

    # The end of <sheetData>, skipping comments, CDATA and PIs (the only
    # places a '<' can appear that is not markup: text and attribute values
    # cannot hold a literal '<').
    pos = root.end()
    closer: Optional[bytes] = None
    end = -1
    while True:
        keep_from = -1          # where to cut when more data is needed
        if closer is not None:
            i = buf.find(closer, pos)
            if i >= 0:
                pos = i + len(closer)
                closer = None
                continue
            keep_from = max(pos, len(buf) - len(closer) + 1)
        else:
            hit = buf.find(b"sheetData", pos)
            bang = buf.find(b"<!", pos, hit if hit >= 0 else len(buf))
            quest = buf.find(b"<?", pos, hit if hit >= 0 else len(buf))
            special = min((x for x in (bang, quest) if x >= 0), default=-1)
            if special >= 0:
                opener = next((o for o in _OPENERS if buf.startswith(o[0], special)), None)
                if opener is not None:
                    closer = opener[1]
                    pos = special + len(opener[0])
                    continue
                if len(buf) - special >= len(b"<![CDATA[") or eof:
                    raise ValueError("unexpected markup in the sheet data")
                keep_from = special  # a comment/CDATA opener cut at the chunk edge
            elif hit >= 0:
                lt = buf.rfind(b"<", max(0, hit - _SPLICE_KEEP), hit)
                m = _SHEETDATA_TAG_RE.match(buf, lt) if lt >= pos else None
                if m is not None and (m.group("close") or m.group("self")):
                    end = m.end()
                    break
                if m is not None:
                    pos = m.end()              # the start tag: keep going
                    continue
                if lt >= pos and buf.find(b">", hit) < 0 and not eof:
                    keep_from = lt             # the tag is cut at the chunk edge
                else:
                    pos = hit + len(b"sheetData")  # the word in text, not a tag
                    continue
            else:
                keep_from = max(pos, len(buf) - _SPLICE_KEEP)
        if eof:
            raise ValueError("the sheet part has no end of <sheetData>")
        keep_from = min(keep_from, len(buf))
        flush(keep_from)
        pos = max(0, pos - keep_from)
        more()
    flush(end)
    rest = src.read()
    out.tail = bytes(buf) + rest
    out.head_digest = digest.hexdigest()
    return out


def _root_nsmap(split: _SheetSplit) -> Dict[Optional[str], str]:
    root = parse_xml(split.root_tag + b"</" + split.qname + b">")
    if root.tag != f"{{{NS_MAIN}}}worksheet":
        raise ValueError("the sheet part's root is not a SpreadsheetML worksheet")
    return dict(root.nsmap)


@dataclass
class _TailChild:
    prefix: bytes
    local: bytes
    start: int                    # offset of its start tag
    start_end: int                # offset just past its start tag
    end: int                      # offset of its end tag (== start_end when self-closing)
    selfclosing: bool
    children: int = 0


def _tail_children(tail: bytes) -> Tuple[List[_TailChild], int]:
    """The worksheet's direct children in ``tail`` (which starts inside the
    worksheet, just after </sheetData>) and the offset of its end tag."""
    children: List[_TailChild] = []
    depth = 1
    current: Optional[_TailChild] = None
    for m in _TAIL_TOKEN_RE.finditer(tail):
        local = m.group("local")
        if local is None:
            continue  # comment / CDATA / PI
        if m.group("close"):
            depth -= 1
            if depth == 0:
                if local != b"worksheet":
                    raise ValueError("unbalanced sheet tail")
                return children, m.start()
            if depth == 1 and current is not None:
                current.end = m.start()
                children.append(current)
                current = None
            continue
        selfclosing = bool(m.group("self"))
        if depth == 1:
            child = _TailChild(m.group("prefix"), local, m.start(), m.end(), m.end(), selfclosing)
            if selfclosing:
                children.append(child)
            else:
                current = child
        elif depth == 2 and current is not None:
            current.children += 1
        if not selfclosing:
            depth += 1
    raise ValueError("the sheet tail does not close the worksheet")


def _free_prefix(nsmap: Dict[Optional[str], str], want: str) -> str:
    if want not in nsmap:
        return want
    n = 1
    while f"{want}{n}" in nsmap:
        n += 1
    return f"{want}{n}"


def _add_table_parts(split: _SheetSplit, rids: List[str]) -> bytes:
    """The sheet tail with a <tablePart> for each of ``rids`` added (to the
    existing <tableParts>, or a new one before the worksheet's <extLst> or its
    end tag), checked by re-reading it before it is returned."""
    nsmap = _root_nsmap(split)
    main = (split.qname.split(b":")[0] + b":") if b":" in split.qname else b""
    r_prefix = next((p for p, uri in nsmap.items() if uri == NS_R and p), None)
    r_decl = b""
    if r_prefix is None:
        r_prefix = _free_prefix(nsmap, "r")
        r_decl = b' xmlns:%s="%s"' % (r_prefix.encode(), NS_R.encode())
    rp = r_prefix.encode()
    tail = split.tail
    children, root_end = _tail_children(tail)
    existing = next((c for c in children if c.local == b"tableParts"), None)
    if existing is None:
        parts = b"".join(b'<%stablePart %s:id="%s"/>' % (main, rp, rid.encode()) for rid in rids)
        insert = b'<%stableParts count="%d"%s>%s</%stableParts>' % (main, len(rids), r_decl, parts, main)
        ext = next((c for c in children if c.local == b"extLst"), None)
        at = ext.start if ext is not None else root_end
        new_tail = tail[:at] + insert + tail[at:]
    else:
        count = existing.children + len(rids)
        prefix = existing.prefix
        parts = b"".join(b'<%stablePart %s:id="%s"%s/>' % (prefix, rp, rid.encode(), r_decl) for rid in rids)
        start_tag = tail[existing.start:existing.start_end]
        if existing.selfclosing:
            start_tag = start_tag[: start_tag.rfind(b"/")] + b">"
        if _COUNT_ATTR_RE.search(start_tag):
            start_tag = _COUNT_ATTR_RE.sub(lambda m: m.group(1) + b'"%d"' % count, start_tag, count=1)
        else:
            start_tag = start_tag[:-1].rstrip() + b' count="%d">' % count
        if existing.selfclosing:
            new_tail = (tail[:existing.start] + start_tag + parts + b"</%stableParts>" % prefix
                        + tail[existing.start_end:])
        else:
            new_tail = (tail[:existing.start] + start_tag + tail[existing.start_end:existing.end]
                        + parts + tail[existing.end:])
    problem = _tail_problem(split.root_tag, new_tail, rids)
    if problem:
        raise ValueError(problem)
    return new_tail


def _tail_problem(root_tag: bytes, tail: bytes, rids: List[str]) -> Optional[str]:
    """Why ``tail`` (after the root start tag) is not a well-formed end of a
    worksheet whose <tableParts> lists every one of ``rids`` with a matching
    count and sits last before an optional <extLst>, or None."""
    import io

    tp_tag = f"{{{NS_MAIN}}}tableParts"
    part_tag = f"{{{NS_MAIN}}}tablePart"
    ext_tag = f"{{{NS_MAIN}}}extLst"
    rid_attr = f"{{{NS_R}}}id"
    found: Optional[Tuple[int, List[str]]] = None
    after: List[str] = []
    try:
        depth = 0
        for event, el in etree.iterparse(io.BytesIO(root_tag + tail), events=("start", "end"),
                                         resolve_entities=False, no_network=True):
            if event == "start":
                depth += 1
                continue
            depth -= 1
            if depth != 1:
                continue
            if found is not None:
                after.append(el.tag)
            if el.tag == tp_tag:
                if found is not None:
                    return "two <tableParts> elements"
                raw_count = (el.get("count") or "").strip()
                found = (
                    int(raw_count) if raw_count.isdigit() else -1,
                    [p.get(rid_attr) or "" for p in el if p.tag == part_tag],
                )
            el.clear()
    except etree.XMLSyntaxError as exc:
        return f"the sheet no longer parses after adding its table: {exc}"
    if found is None:
        return "the sheet has no <tableParts> after adding its table"
    count, listed = found
    if any(rid not in listed for rid in rids):
        return "the sheet's <tableParts> does not list the new table"
    if count != len(listed):
        return "the sheet's <tableParts> count does not match its entries"
    if any(tag != ext_tag for tag in after):
        return "<tableParts> is not the last element of the sheet"
    return None


def _plan_parts(
    source: Path, edits: _Edits, scan: WorkbookScan
) -> Tuple[Dict[str, bytes], List[Tuple[str, bytes]], Dict[str, List[str]]]:
    replacements: Dict[str, bytes] = {}
    additions: List[Tuple[str, bytes]] = []
    # sheet part (lower-cased zip name) -> relationship ids of the tables to
    # list in its <tableParts>; applied while the zip is rewritten.
    splices: Dict[str, List[str]] = {}
    with zipfile.ZipFile(str(source)) as zf:
        pkg = Package(zf)

        def current(part: str) -> Optional[bytes]:
            key = part.lower()
            if key in replacements:
                return replacements[key]
            for name, data in additions:
                if name.lower() == key:
                    return data
            return pkg.read(part) if pkg.has(part) else None

        def put(part: str, data: bytes) -> None:
            real = pkg.name(part)
            if real is not None:
                replacements[real.lower()] = data
                return
            for i, (name, _old) in enumerate(additions):
                if name.lower() == part.lower():
                    additions[i] = (name, data)
                    return
            additions.append((part, data))

        content_types = parse_xml(pkg.read("[Content_Types].xml"))

        def add_override(part: str, ctype: str) -> None:
            pn = "/" + part.lstrip("/")
            for el in content_types:
                if (el.get("PartName") or "").lower() == pn.lower():
                    el.set("ContentType", ctype)
                    return
            ov = etree.SubElement(content_types, f"{{{NS_CT}}}Override")
            ov.set("PartName", pn)
            ov.set("ContentType", ctype)

        # -- core properties -------------------------------------------------
        if edits.title is not None or edits.language is not None:
            core_part = scan.core_part
            created = core_part is None
            if created:
                core_part = "docProps/core.xml"
                core = etree.Element(
                    f"{{{NS_CP}}}coreProperties",
                    nsmap={"cp": NS_CP, "dc": NS_DC, "dcterms": "http://purl.org/dc/terms/",
                           "dcmitype": "http://purl.org/dc/dcmitype/", "xsi": "http://www.w3.org/2001/XMLSchema-instance"},
                )
            else:
                core = parse_xml(current(core_part) or b"")
            for tag, value in (("title", edits.title), ("language", edits.language)):
                if value is None:
                    continue
                el = core.find(f"{{{NS_DC}}}{tag}")
                if el is None:
                    el = etree.SubElement(core, f"{{{NS_DC}}}{tag}")
                el.text = value
            put(core_part, _xml_bytes(core))
            if created:
                rels = parse_xml(current("_rels/.rels") or b'<Relationships xmlns="%s"/>' % NS_PKG_REL.encode())
                rel = etree.SubElement(rels, f"{{{NS_PKG_REL}}}Relationship")
                rel.set("Id", _unique_rid(rels))
                rel.set("Type", REL_CORE)
                rel.set("Target", core_part)
                put("_rels/.rels", _xml_bytes(rels))
                add_override(core_part, CT_CORE)

        # -- alt text --------------------------------------------------------
        for ae in edits.alts:
            data = current(ae.drawing_part)
            if data is None:
                raise ValueError(f"drawing part {ae.drawing_part} is missing")
            root = parse_xml(data)
            target = None
            for i, (obj, _anchor) in enumerate(iter_drawing_objects(root)):
                if i == ae.index:
                    target = cnvpr_of(obj)
                    break
            if target is None or (ae.cnvpr_id and (target.get("id") or "") != ae.cnvpr_id):
                raise ValueError(f"object {ae.index} in {ae.drawing_part} no longer matches")
            if ae.descr is None:
                target.attrib.pop("descr", None)
            else:
                target.set("descr", ae.descr)
            put(ae.drawing_part, _xml_bytes(root))

        # -- tables ----------------------------------------------------------
        used_parts = set(scan.table_parts)
        used_names = set(scan.table_names) | set(scan.defined_names)
        next_id = scan.max_table_id
        by_sheet: Dict[str, List[_TableEdit]] = {}
        for te in edits.tables:
            by_sheet.setdefault(te.sheet.part, []).append(te)
        for sheet_part, tes in by_sheet.items():
            real_sheet = pkg.name(sheet_part)
            if real_sheet is None:
                raise ValueError(f"sheet part {sheet_part} is missing")
            rels_part = rels_part_for(sheet_part)
            rels_data = current(rels_part)
            rels = parse_xml(rels_data) if rels_data else etree.Element(f"{{{NS_PKG_REL}}}Relationships", nsmap={None: NS_PKG_REL})
            # The sheet part itself is NOT parsed here: its <tableParts> is
            # spliced in while the zip is rewritten (see _split_sheet).
            rids = splices.setdefault(real_sheet.lower(), [])
            for te in tes:
                n = 1
                while f"xl/tables/table{n}.xml" in used_parts:
                    n += 1
                te.part = f"xl/tables/table{n}.xml"
                used_parts.add(te.part)
                m = 1
                while f"table{m}" in used_names:
                    m += 1
                te.name = f"Table{m}"
                used_names.add(te.name.casefold())
                next_id += 1
                te.table_id = next_id
                additions.append((te.part, _table_xml(te)))
                add_override(te.part, CT_TABLE)
                rid = _unique_rid(rels)
                rel = etree.SubElement(rels, f"{{{NS_PKG_REL}}}Relationship")
                rel.set("Id", rid)
                rel.set("Type", REL_TABLE)
                rel.set("Target", posixpath.relpath(te.part, posixpath.dirname(sheet_part)))
                te.rid = rid
                rids.append(rid)
            put(rels_part, _xml_bytes(rels))

        put("[Content_Types].xml", _xml_bytes(content_types))
    return replacements, additions, splices


def _unique_rid(rels: etree._Element) -> str:
    taken = {el.get("Id") for el in rels}
    n = len(taken) + 1
    while f"rId{n}" in taken:
        n += 1
    return f"rId{n}"


def _table_xml(te: _TableEdit) -> bytes:
    t = etree.Element(f"{{{NS_MAIN}}}table", nsmap={None: NS_MAIN})
    t.set("id", str(te.table_id))
    t.set("name", te.name)
    t.set("displayName", te.name)
    t.set("ref", te.block.ref)
    t.set("totalsRowShown", "0")
    cols = etree.SubElement(t, f"{{{NS_MAIN}}}tableColumns")
    cols.set("count", str(len(te.columns)))
    for i, name in enumerate(te.columns, start=1):
        col = etree.SubElement(cols, f"{{{NS_MAIN}}}tableColumn")
        col.set("id", str(i))
        col.set("name", name)
    # No style name and no filter buttons: the sheet looks exactly as it did.
    # What changes is that the first row is now DECLARED the header row,
    # which is what a screen reader announces as it moves along the data.
    style = etree.SubElement(t, f"{{{NS_MAIN}}}tableStyleInfo")
    for attr in ("showFirstColumn", "showLastColumn", "showRowStripes", "showColumnStripes"):
        style.set(attr, "0")
    return _xml_bytes(t)


def _rewrite_zip(
    source: Path,
    dest: Path,
    replacements: Dict[str, bytes],
    additions: List[Tuple[str, bytes]],
    splices: Optional[Dict[str, List[str]]] = None,
) -> None:
    """Copy ``source`` to ``dest`` member by member, streamed (a copied part
    is never held whole in memory), with ``replacements`` swapped in,
    ``splices`` applied to their sheet parts, and ``additions`` appended."""
    splices = splices or {}
    now = datetime.now().timetuple()[:6]
    with zipfile.ZipFile(str(source)) as zin, zipfile.ZipFile(str(dest), "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            key = info.filename.lower()
            # Some producers write zero (pre-1980) timestamps, which zipfile
            # refuses to write back; the date carries no meaning in OOXML.
            stamp = info.date_time if info.date_time[0] >= 1980 else (1980, 1, 1, 0, 0, 0)
            zi = zipfile.ZipInfo(info.filename, date_time=stamp)
            zi.compress_type = zipfile.ZIP_DEFLATED if not info.is_dir() else zipfile.ZIP_STORED
            zi.external_attr = info.external_attr
            data = replacements.get(key)
            if data is not None:
                zout.writestr(zi, data)
            elif key in splices:
                with zin.open(info) as src, zout.open(zi, "w") as dst:
                    try:
                        split = _split_sheet(src, dst.write)
                        new_tail = _add_table_parts(split, splices[key])
                    except (ValueError, etree.LxmlError, UnicodeError) as exc:
                        raise _SpliceError(f"{info.filename}: {exc}") from exc
                    if split.head_len + len(split.tail) != info.file_size:
                        raise _SpliceError(f"streaming {info.filename} lost bytes")
                    dst.write(new_tail)
            elif info.is_dir():
                zout.writestr(zi, b"")
            else:
                with zin.open(info) as src, zout.open(zi, "w") as dst:
                    shutil.copyfileobj(src, dst, _SPLICE_CHUNK)
        for name, data in additions:
            zi = zipfile.ZipInfo(name, date_time=now)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(zi, data)


# ---------------------------------------------------------------------------
# Verification
# ---------------------------------------------------------------------------


def _sheet_xml_bytes(path: Path) -> int:
    with zipfile.ZipFile(str(path)) as zf:
        return sum(i.file_size for i in zf.infolist() if i.filename.lower().startswith("xl/worksheets/"))


def _openpyxl_load(path: Path):
    """Load ``path`` the way verification does (full, or read-only when the
    worksheets are large). Raises whatever openpyxl raises."""
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        # A file handle, not the path: openpyxl refuses a name that does not
        # end in .xlsx, and the candidate output is a .tmp file.
        with open(path, "rb") as fh:
            if _sheet_xml_bytes(path) <= _FULL_VERIFY_MAX_SHEET_XML:
                return openpyxl.load_workbook(fh)
            wb = openpyxl.load_workbook(fh, read_only=True)
            _ = wb.sheetnames
            return wb


def _table_parts_problems(path: Path, edits: _Edits) -> List[str]:
    """Each new table part, read by openpyxl's own table model on its own."""
    from openpyxl.worksheet.table import Table as _XlTable

    problems: List[str] = []
    with zipfile.ZipFile(str(path)) as zf:
        for te in edits.tables:
            try:
                tbl = _XlTable.from_tree(etree.fromstring(zf.read(te.part)))
            except Exception as exc:
                problems.append(f"openpyxl cannot read the new table {te.name}: {exc}")
                continue
            if tbl.ref != te.block.ref or [c.name for c in tbl.tableColumns] != te.columns:
                problems.append(f"openpyxl reads {te.name} differently from what we wrote")
    return problems


def _verify(path: Path, edits: _Edits, source: Optional[Path] = None) -> List[str]:
    problems: List[str] = []
    try:
        with zipfile.ZipFile(str(path)) as zf:
            bad = zf.testzip()
            if bad:
                problems.append(f"corrupt member {bad}")
                return problems
    except Exception as exc:
        return [f"not a readable zip: {exc}"]

    # 1. Our own scanner's readers read every edit back from the parts it
    # touched. (Every other part was streamed through unchanged, and the zip
    # test above checked each one against its CRC. Re-scanning the whole
    # workbook here used to hold a second copy of every sheet's window.)
    try:
        problems.extend(_read_back_problems(path, edits, source))
    except Exception as exc:
        return [f"reading the edits back failed: {exc}"]
    if problems:
        return problems

    # 2. openpyxl reopens it (the bar for "a workbook Excel-family tools can
    # open"), and sees the new tables where we put them.
    try:
        wb = _openpyxl_load(path)
    except Exception as exc:
        # openpyxl is stricter than Excel: some producers write files Excel
        # opens and openpyxl refuses (a font "family" above 14 is enough).
        # When it refuses the customer's OWN file too, our edits are not what
        # it objects to, and failing here would turn every fix on that
        # workbook into a 422. The edits were already read back above by our
        # scanner; the new table parts are still checked by openpyxl's own
        # table model on their own.
        if source is not None and not _openpyxl_opens(source):
            logger.info("openpyxl cannot read the source workbook either (%s); verified with our own reader", exc)
            return _table_parts_problems(path, edits)
        return [f"openpyxl could not reopen the workbook: {exc}"]
    try:
        if getattr(wb, "read_only", False):
            problems.extend(_table_parts_problems(path, edits))
        else:
            for te in edits.tables:
                ws = wb[te.sheet.name]
                tbl = ws.tables.get(te.name)
                if tbl is None or tbl.ref != te.block.ref:
                    problems.append(f"openpyxl does not see {te.name} at {te.block.ref}")
                elif [c.name for c in tbl.tableColumns] != te.columns:
                    problems.append(f"openpyxl reads different column names for {te.name}")
            if edits.title is not None and (wb.properties.title or "") != edits.title:
                problems.append("openpyxl reads a different title")
    except Exception as exc:
        problems.append(f"openpyxl could not read the workbook back: {exc}")
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return problems


def _read_back_problems(path: Path, edits: _Edits, source: Optional[Path]) -> List[str]:
    from app.parsers.xlsx_parser import (
        REL_OFFICE_DOCUMENT,
        _core_properties,
        _read_drawing,
        _read_table,
        parse_range,
    )

    problems: List[str] = []
    with zipfile.ZipFile(str(path)) as zf:
        pkg = Package(zf)
        root_rels = pkg.rels("")
        if not any(r.type == REL_OFFICE_DOCUMENT for r in root_rels.values()):
            problems.append("the package no longer points at its workbook")
        if edits.title is not None or edits.language is not None:
            core_part = next((r.target for r in root_rels.values() if r.type == REL_CORE and not r.external), None)
            title, language = _core_properties(pkg, core_part)
            if edits.title is not None and title != edits.title:
                problems.append("title did not persist")
            if edits.language is not None and language != edits.language:
                problems.append("language did not persist")
        drawings: Dict[str, Dict[int, Any]] = {}
        for ae in edits.alts:
            key = ae.drawing_part.lower()
            if key not in drawings:
                drawings[key] = {d.index: d for d in _read_drawing(pkg, ae.drawing_part)}
            obj = drawings[key].get(ae.index)
            if obj is None or (obj.descr or None) != (ae.descr or None):
                problems.append(f"alt text did not persist on {ae.node_id}")
        if edits.tables:
            ct_root = parse_xml(pkg.read("[Content_Types].xml"))
            overrides = {
                (el.get("PartName") or "").lower(): el.get("ContentType")
                for el in ct_root if el.tag == f"{{{NS_CT}}}Override"
            }
            by_sheet: Dict[str, List[_TableEdit]] = {}
            for te in edits.tables:
                by_sheet.setdefault(te.sheet.part, []).append(te)
            for sheet_part, tes in by_sheet.items():
                split = _checked_split(pkg, source, sheet_part)
                problem = split if isinstance(split, str) else _tail_problem(
                    split.root_tag, split.tail, [te.rid for te in tes]
                )
                if problem:
                    problems.append(f"sheet {tes[0].sheet.name!r}: {problem}")
            for te in edits.tables:
                rel = pkg.rels(te.sheet.part).get(te.rid)
                tdef = None
                if rel is not None and not rel.external and rel.type == REL_TABLE and rel.target.lower() == te.part.lower():
                    tdef = _read_table(pkg, te.part, te.rid)
                if (
                    tdef is None
                    or tdef.header_rows < 1
                    or tdef.ref != parse_range(te.block.ref)
                    or tdef.display_name != te.name
                    or tdef.columns != te.columns
                ):
                    problems.append(f"table {te.name} did not persist as a declared header row")
                elif overrides.get("/" + te.part.lower()) != CT_TABLE:
                    problems.append(f"table {te.name} has no content type")
    return problems


def _checked_split(pkg: Package, source: Optional[Path], sheet_part: str):
    """The output sheet part's split, after proving everything up to the end
    of its <sheetData> is byte-identical to the source's; or a problem."""
    with pkg.open(sheet_part) as fh:
        out = _split_sheet(fh)
    if source is not None:
        with zipfile.ZipFile(str(source)) as zs:
            src_pkg = Package(zs)
            with src_pkg.open(sheet_part) as fh:
                orig = _split_sheet(fh)
        if (orig.head_len, orig.head_digest, orig.root_tag) != (out.head_len, out.head_digest, out.root_tag):
            return f"the cells of sheet part {sheet_part} changed while adding its table"
    return out


def _openpyxl_opens(path: Path) -> bool:
    try:
        _openpyxl_load(path).close()
        return True
    except Exception:
        return False


__all__ = ["write_remediated_xlsx"]
