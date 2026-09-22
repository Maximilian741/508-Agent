"""DOCX extraction helpers used by scan/fix routes.

Two complementary entry points are exposed:

* :meth:`DOCXParser.parse` returns the dict-shaped detection payload used by
  the legacy read-only ``/documents`` routes.
* :meth:`DOCXParser.parse_to_tree` builds an :class:`AccessibilityTree` so the
  document can flow through the same analyzer + executor pipeline as PDFs.
"""

from __future__ import annotations

import base64
import re
import zipfile
from pathlib import Path
from typing import Any, Callable, Dict, Iterator, List, Optional, Tuple
from urllib.parse import urlparse

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

from app.parsers.document_id import derive_document_id
from app.models.accessibility import (
    AccessibilityTree,
    ContentKind,
    DocumentNode,
    HeadingNode,
    ImageNode,
    LinkNode,
    ListItemNode,
    ListNode,
    NodeContent,
    NodeMetadata,
    ParagraphNode,
    ParserResult,
    SectionNode,
    TableCellNode,
    TableCellType,
    TableHeaderScope,
    TableNode,
    TableRowNode,
)


# Decided once per process: is there a vision provider that could ever read
# inlined image bytes? Under the heuristic provider nothing consumes them, and
# base64-inflating every image into the tree cost tens of MB per request.
try:
    from app.ai.semantic_inference import vision_provider_configured as _vpc

    _WANT_IMAGE_BYTES = bool(_vpc())
except Exception:  # pragma: no cover - never let the AI module break parsing
    _WANT_IMAGE_BYTES = True


class DOCXParser:
    def parse(self, file_path: str) -> Dict[str, object]:
        doc = Document(file_path)
        core = doc.core_properties
        w_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        generic_link_labels = {
            "click here",
            "here",
            "read more",
            "learn more",
            "more",
            "link",
            "this",
        }

        headings: List[Dict[str, object]] = []
        empty_heading_sections: List[int] = []
        prev_level: Optional[int] = None
        skipped_jumps: List[Dict[str, object]] = []
        style_cache: Dict[Any, str] = {}
        for idx, para in enumerate(iter_body_paragraphs(doc), start=1):
            style_name = paragraph_style_name(para, style_cache)
            if not style_name.lower().startswith("heading"):
                continue
            level = 1
            try:
                tail = style_name.split(" ", 1)[1]
                level = max(1, min(6, int(tail)))
            except Exception:
                level = 1
            text = (para.text or "").strip()
            if not text:
                empty_heading_sections.append(idx)
            headings.append({"level": level, "text": text, "section": idx})
            if prev_level is not None and level > prev_level + 1:
                skipped_jumps.append({"from": prev_level, "to": level, "section": idx, "text": text})
            prev_level = level

        image_count = 0
        alt_missing = 0
        for rel in doc.part.rels.values():
            reltype = str(rel.reltype)
            if "image" in reltype:
                image_count += 1
                alt_missing += 1

        hyperlink_count = 0
        generic_links: List[Dict[str, object]] = []
        invalid_links: List[Dict[str, object]] = []
        for idx, para in enumerate(iter_body_paragraphs(doc), start=1):
            hyperlink_nodes = list(para._p.iterfind(f".//{w_ns}hyperlink"))
            if not hyperlink_nodes:
                continue
            for hyperlink in hyperlink_nodes:
                text_bits = [node.text or "" for node in hyperlink.iterfind(f".//{w_ns}t")]
                link_text = "".join(text_bits).strip()
                hyperlink_count += 1
                rid = hyperlink.get(f"{rel_ns}id")
                anchor = hyperlink.get(f"{w_ns}anchor")
                target = ""
                if rid and rid in doc.part.rels:
                    rel = doc.part.rels[rid]
                    target = str(getattr(rel, "target_ref", "") or "")
                elif anchor:
                    target = f"#{anchor}"
                if not link_text:
                    continue
                normalized = link_text.lower().strip()
                if normalized in generic_link_labels:
                    generic_links.append({"section": idx, "text": link_text})
                invalid_reason = _invalid_link_reason(target)
                if invalid_reason:
                    invalid_links.append({"section": idx, "text": link_text, "target": target, "reason": invalid_reason})

        tables = len(list(iter_body_tables(doc)))
        tables_missing_headers: List[int] = []
        table_header_scope_flags: List[Dict[str, object]] = []
        generic_headers = {"column", "column 1", "column 2", "header", "n/a", "na", "value"}
        for table_index, table in enumerate(iter_body_tables(doc), start=1):
            _rows = list(iter_table_rows(table))
            if not _rows:
                tables_missing_headers.append(table_index)
                continue
            header_cells = _rows[0].cells
            header_text = [(cell.text or "").strip() for cell in header_cells]
            if not any(header_text):
                tables_missing_headers.append(table_index)
                continue
            normalized_headers = [text.lower() for text in header_text if text]
            unique_headers = set(normalized_headers)
            if len(header_text) > 1 and len(unique_headers) <= 1:
                table_header_scope_flags.append(
                    {"table": table_index, "reason": "duplicate_or_single_header_label", "headers": header_text[:10]}
                )
            elif any(text.lower() in generic_headers for text in header_text if text):
                table_header_scope_flags.append(
                    {"table": table_index, "reason": "generic_header_labels", "headers": header_text[:10]}
                )
        return {
            "documentType": "docx",
            "title": (core.title or "").strip(),
            "language": (getattr(core, "language", None) or "").strip(),
            "headings": headings,
            "emptyHeadingSections": empty_heading_sections,
            "headingJumps": skipped_jumps,
            "imageCount": image_count,
            "missingAltCount": alt_missing,
            "hyperlinkCount": hyperlink_count,
            "genericLinks": generic_links,
            "invalidLinks": invalid_links,
            "tables": tables,
            "tablesMissingHeaders": tables_missing_headers,
            "tableHeaderScopeFlags": table_header_scope_flags,
            "outlineCount": len(headings),
        }


    def parse_to_tree(self, file_path: str) -> ParserResult:
        """Build an :class:`AccessibilityTree` from a DOCX file.

        The tree captures: document title/language, paragraphs, headings,
        images (with embedded base64 + mime), tables (with header detection),
        lists, and hyperlinks.  Reading order in DOCX flows linearly so the
        children of the synthetic root section preserve the source ordering.
        """
        return self.parse_document(Document(file_path), file_path)

    def parse_document(
        self,
        doc,
        file_path: str,
        register: Optional[Callable[[str, Any], None]] = None,
    ) -> ParserResult:
        """Build the tree from an already-open python-docx ``Document``.

        ``register(node_id, element)`` — when given — is called for every node
        id this walk mints, with the live python-docx/lxml object that id
        names (a Paragraph, a ``<wp:docPr>``, a link element, a ``(row, cell)``
        pair, a Table …). The docx writer opens the copy it is about to edit,
        runs THIS walk over it with a register callback, and edits exactly the
        objects it was handed. Parser and writer therefore cannot disagree
        about which element ``docx-img-7`` is: there is one walk, not a walk
        and a hand-maintained mirror of it. (Six mirrors had drifted: alt text
        for an image after a picture in a heading landed on the NEXT picture.)
        """

        path = Path(file_path)
        reg = register or _no_register
        theme_colors = _docx_theme_colors(file_path)
        styles = DocxStyleResolver(doc)
        core = doc.core_properties

        title = (core.title or "").strip()
        language = (getattr(core, "language", None) or "").strip()
        if not language:
            # dc:language is rarely set, but Word writes the document language
            # into styles.xml docDefaults <w:lang w:val="en-US"/> on save — and
            # per the writer's own docstring THAT is where screen readers and
            # Word's Accessibility Checker read it. Reading only dc:language
            # flagged DOCUMENT_LANGUAGE_MISSING on essentially every Word
            # document, and the "fix" then overwrote en-US with a less
            # specific en. Same source of truth for detector and fixer now.
            language = _docx_default_lang(doc) or ""
        properties: Dict[str, Any] = {"filename": path.name}
        if title:
            properties["title"] = title
        ff_total, ff_unlabeled, ff_derivable = _docx_form_field_counts(doc)
        if ff_total:
            properties["form_fields_total"] = ff_total
            properties["form_fields_unlabeled"] = ff_unlabeled
            # How many unlabeled controls we can confidently auto-label from
            # nearby text (drives FILL_FORM_FIELD_LABELS; the rest stay manual).
            properties["form_fields_derivable"] = ff_derivable
        root = DocumentNode(
            id="doc-1",
            content=NodeContent(kind=ContentKind.NONE),
            metadata=NodeMetadata(
                language=language or None,
                source_format="docx",
                properties=properties,
            ),
            children=[],
            accessibility_flags=[],
        )

        body_section = SectionNode(
            id="docx-body",
            content=NodeContent(kind=ContentKind.TEXT, text="Document"),
            metadata=NodeMetadata(source_format="docx", properties={}),
            children=[],
            accessibility_flags=[],
        )
        root.children.append(body_section)
        ids = _IdCounter()
        ids.bookmarks = _bookmark_names(doc)

        # Pre-load embedded images so we can attach bytes to the matching
        # <w:drawing> nodes encountered while iterating paragraphs.
        image_blobs = _collect_image_blobs(doc.part)
        body_partname = _partname(doc.part)

        # Paragraph + heading walker (preserves order; lists handled as groups).
        list_collector: List[ListItemNode] = []
        list_marker: Optional[str] = None

        para_style_cache: Dict[Any, str] = {}
        for paragraph in iter_body_paragraphs(doc):
            p_el = paragraph._p  # noqa: SLF001
            style_name = paragraph_style_name(paragraph, para_style_cache)
            text = paragraph_text(p_el).strip()
            # A Heading-styled paragraph is a heading even when it is numbered
            # (w:numPr on the paragraph: "1.2 Scope" with Word's own
            # numbering). It used to be swallowed as a list item, which broke
            # the outline and invented heading-level jumps.
            # A paragraph is also a heading when its OUTLINE LEVEL says so —
            # a template's own "Agency Heading 1" (based on Heading 1) or a
            # direct "Outline level: Level 2". Word's navigation pane, screen
            # readers and PDF export all treat it as one; reading only the
            # style NAME reported such documents as having no headings.
            heading_level = _heading_level_from_style(style_name) or _outline_heading_level(p_el, styles)

            # Links and pictures belong to EVERY kind of paragraph — a
            # "click here" in a bullet, a logo set in the title line. They
            # used to be read only from plain paragraphs, and the writer's
            # image index (which counted every paragraph) then disagreed with
            # the parser from the first such picture onward.
            link_nodes = _link_nodes_from_p(
                p_el, doc.part, ids, reg=reg, context_text=text,
                extra_props={"docx_story": "body", "docx_part": body_partname},
            )
            image_nodes = _image_nodes_from_p(
                p_el, image_blobs, ids, "docx-img", reg,
                story="body", partname=body_partname,
                context_fn=lambda _p=paragraph, _t=text: _image_context_for_paragraph(_p, _t),
            )

            if not heading_level and _is_list_paragraph(paragraph):
                marker = _detect_list_marker(paragraph)
                if list_marker is None:
                    list_marker = marker
                if list_marker != marker:
                    body_section.children.append(_finalize_list(ids, list_collector, list_marker))
                    list_collector = []
                    list_marker = marker
                li_id = ids("docx-li")
                reg(li_id, paragraph)
                list_collector.append(
                    ListItemNode(
                        id=li_id,
                        content=NodeContent(kind=ContentKind.TEXT, text=text or "•"),
                        metadata=NodeMetadata(source_format="docx", properties=_text_color_props(paragraph, theme_colors)),
                        children=[*link_nodes, *image_nodes],
                        accessibility_flags=[],
                    )
                )
                continue

            if list_collector:
                body_section.children.append(_finalize_list(ids, list_collector, list_marker))
                list_collector = []
                list_marker = None

            if heading_level:
                h_id = ids("docx-h")
                reg(h_id, paragraph)
                h_props = dict(_text_color_props(paragraph, theme_colors) or {})
                h_props["heading_visual"] = _heading_visual(p_el, styles, text)
                body_section.children.append(
                    HeadingNode(
                        id=h_id,
                        level=heading_level,
                        content=NodeContent(kind=ContentKind.TEXT, text=text or "Heading"),
                        metadata=NodeMetadata(source_format="docx", properties=h_props),
                        children=[*link_nodes, *image_nodes],
                        accessibility_flags=[],
                    )
                )
                continue

            # Hyperlink runs first.
            for link in link_nodes:
                body_section.children.append(link)

            # Inline images. Each carries nearby human text: ``caption`` only
            # when it is a REAL caption (Word's Caption style next to the
            # picture, or a numbered "Figure 3:" label), ``nearby_text`` for
            # anything else, so an alt-text provider can tell the two apart.
            for image in image_nodes:
                body_section.children.append(image)

            if text and not link_nodes:
                para_props = _text_color_props(paragraph, theme_colors)
                is_fake_heading = _looks_like_fake_heading(paragraph, style_name, text)
                if is_fake_heading:
                    # Visually a heading (Title/Subtitle style, or short
                    # all-bold large text) but NOT a real Heading style —
                    # invisible to screen-reader navigation. Flagged by
                    # TextStyledAsHeadingAnalyzer.
                    para_props = dict(para_props or {})
                    para_props["looks_like_heading"] = True
                    # What it LOOKS like (effective size / weight / numbering),
                    # so PROMOTE_HEADING can place it in the document's own
                    # outline instead of defaulting to a sibling H1.
                    para_props["heading_visual"] = _heading_visual(p_el, styles, text)
                # A line that looks like a heading is NOT also a fake-list item.
                # A big/bold numbered section header ("1. Introduction") is a
                # heading, not a bullet — tagging it both ways lets two fixes
                # (PROMOTE_HEADING + FIX_LIST_STRUCTURE) target the same node,
                # where only one can win in the writer (the other is silently
                # dropped yet still scored/charged). Promotion takes precedence.
                fake_sig = None if is_fake_heading else _fake_list_signature(text)
                if fake_sig:
                    kind, char, ordinal = fake_sig
                    para_props = dict(para_props or {})
                    para_props["fake_list_kind"] = kind
                    para_props["fake_list_char"] = char
                    if ordinal is not None:
                        para_props["fake_list_ordinal"] = ordinal
                p_id = ids("docx-p")
                reg(p_id, paragraph)
                body_section.children.append(
                    ParagraphNode(
                        id=p_id,
                        content=NodeContent(kind=ContentKind.TEXT, text=text),
                        metadata=NodeMetadata(source_format="docx", properties=para_props),
                        children=[],
                        accessibility_flags=[],
                    )
                )

        if list_collector:
            body_section.children.append(_finalize_list(ids, list_collector, list_marker))

        # Group consecutive fake-list paragraphs into runs (one flag each).
        _group_fake_list_runs(body_section.children)

        # Text boxes: w:txbxContent is invisible to doc.paragraphs, so
        # sidebar/callout content (very common in government documents) would
        # otherwise never be analyzed. Paragraphs mint a distinct docx-tbp /
        # docx-tblink id space so the body's docx-p / docx-link counters are
        # untouched. Links here ARE remediable (registered like any other).
        for tb_p in _iter_text_box_paragraphs(doc.element.body):
            tb_text = paragraph_text(tb_p).strip()
            tb_links = _link_nodes_from_p(
                tb_p, doc.part, ids, prefix="docx-tblink", reg=reg, context_text=tb_text,
                extra_props={"in_text_box": True, "docx_story": "text_box", "docx_part": body_partname},
            )
            body_section.children.extend(tb_links)
            if tb_text and not tb_links:
                body_section.children.append(
                    ParagraphNode(
                        id=ids("docx-tbp"),
                        content=NodeContent(kind=ContentKind.TEXT, text=tb_text),
                        metadata=NodeMetadata(
                            source_format="docx",
                            properties={"in_text_box": True},
                        ),
                        children=[],
                        accessibility_flags=[],
                    )
                )

        # Footnotes / endnotes: separate package parts, also invisible to
        # doc.paragraphs. Same distinct-id-space pattern as text boxes
        # (docx-fnp / docx-fnlink); link targets resolve against the NOTE
        # part's own relationships.
        for note_part, note_root in _iter_note_parts(doc):
            reg(_NOTE_PARTS_KEY, (note_part, note_root))
            for n_p in _note_paragraphs(note_root):
                fn_text = paragraph_text(n_p).strip()
                fn_links = _link_nodes_from_p(
                    n_p, note_part, ids, prefix="docx-fnlink", reg=reg, context_text=fn_text,
                    extra_props={"in_footnote": True, "docx_story": "footnote",
                                 "docx_part": _partname(note_part)},
                )
                body_section.children.extend(fn_links)
                if fn_text and not fn_links:
                    body_section.children.append(
                        ParagraphNode(
                            id=ids("docx-fnp"),
                            content=NodeContent(kind=ContentKind.TEXT, text=fn_text),
                            metadata=NodeMetadata(
                                source_format="docx",
                                properties={"in_footnote": True},
                            ),
                            children=[],
                            accessibility_flags=[],
                        )
                    )

        # Tables (linearly after paragraphs is acceptable for the v1 flow).
        top_tables = list(iter_body_tables(doc))
        for table in top_tables:
            body_section.children.append(
                _table_to_node(table, ids, image_blobs, reg, body_partname)
            )

        # Page headers and footers: separate parts (word/header1.xml …) that
        # doc.paragraphs never opens. The agency logo in the letterhead — on
        # every printed page — got no finding and no alt, a grey footer was
        # never measured, and a re-scan called the file clean. Their own id
        # space (docx-hfimg / docx-hflink / docx-hfp) keeps body ids stable.
        furniture = _header_footer_section(doc, ids, reg, theme_colors)
        if furniture is not None:
            root.children.append(furniture)

        # Ensure unique ids.
        raw_metadata = {
            "filename": path.name,
            "title": title,
            "language": language,
            "table_count": len(top_tables),
        }
        return ParserResult(
            document_id=derive_document_id(path),
            format="docx",
            tree=AccessibilityTree(root=root, metadata=raw_metadata),
            raw_metadata=raw_metadata,
        )


# ----- DOCX → AccessibilityTree helpers --------------------------------------


_DOCX_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_DOCX_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_DRAWING_NS = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
_DRAWINGML_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_PIC_NS = "{http://schemas.openxmlformats.org/drawingml/2006/picture}"


# w:themeColor attribute value -> clrScheme element name.
_THEME_COLOR_MAP = {
    "dark1": "dk1", "text1": "dk1",
    "light1": "lt1", "background1": "lt1",
    "dark2": "dk2", "text2": "dk2",
    "light2": "lt2", "background2": "lt2",
    "accent1": "accent1", "accent2": "accent2", "accent3": "accent3",
    "accent4": "accent4", "accent5": "accent5", "accent6": "accent6",
    "hyperlink": "hlink", "followedHyperlink": "folHlink",
}


def _docx_theme_colors(file_path: str) -> Dict[str, str]:
    """Map theme colour-scheme names (dk1, lt1, accent1…) to RGB hex."""
    out: Dict[str, str] = {}
    try:
        with zipfile.ZipFile(file_path) as z:
            names = [n for n in z.namelist() if n.startswith("word/theme/theme") and n.endswith(".xml")]
            if not names:
                return {}
            xml = z.read(sorted(names)[0])
        root = etree.fromstring(xml)
        a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        scheme = root.find(f".//{a}clrScheme")
        if scheme is None:
            return {}
        for child in scheme:
            name = etree.QName(child).localname
            srgb = child.find(f"{a}srgbClr")
            sysclr = child.find(f"{a}sysClr")
            if srgb is not None and srgb.get("val"):
                out[name] = srgb.get("val").upper()
            elif sysclr is not None and sysclr.get("lastClr"):
                out[name] = sysclr.get("lastClr").upper()
    except Exception:
        return {}
    return out


def _apply_tint_shade(hex_color: str, tint_hex: Optional[str], shade_hex: Optional[str]) -> str:
    """Apply WordprocessingML w:themeTint / w:themeShade to a base RGB hex.

    themeTint blends toward white (keep ``tint/255`` of the colour); themeShade
    blends toward black (multiply by ``shade/255``).
    """
    try:
        r = int(hex_color[0:2], 16)
        g = int(hex_color[2:4], 16)
        b = int(hex_color[4:6], 16)
    except (ValueError, IndexError):
        return hex_color
    if tint_hex:
        try:
            t = int(tint_hex, 16) / 255.0
            r = round(r * t + 255 * (1 - t)); g = round(g * t + 255 * (1 - t)); b = round(b * t + 255 * (1 - t))
        except ValueError:
            pass
    if shade_hex:
        try:
            s = int(shade_hex, 16) / 255.0
            r = round(r * s); g = round(g * s); b = round(b * s)
        except ValueError:
            pass
    return f"{max(0, min(255, r)):02X}{max(0, min(255, g)):02X}{max(0, min(255, b)):02X}"


def _run_color_hex(run, theme_colors: Dict[str, str]) -> Optional[str]:
    """The run's effective text colour as RGB hex: explicit sRGB, or a resolved
    theme colour (with tint/shade). ``None`` for auto/inherited/unknown."""
    try:
        rgb = run.font.color.rgb  # RGBColor only when explicitly sRGB
    except Exception:
        rgb = None
    if rgb is not None:
        return str(rgb)
    # Theme colour via the underlying w:color element.
    try:
        color_el = run._element.find(f".//{_DOCX_NS}rPr/{_DOCX_NS}color")
        if color_el is None:
            return None
        tc = color_el.get(f"{_DOCX_NS}themeColor")
        if not tc or not theme_colors:
            return None
        base = theme_colors.get(_THEME_COLOR_MAP.get(tc, tc))
        if not base:
            return None
        return _apply_tint_shade(
            base,
            color_el.get(f"{_DOCX_NS}themeTint"),
            color_el.get(f"{_DOCX_NS}themeShade"),
        )
    except Exception:
        return None


def _explicit_run_colors(paragraph, theme_colors: Optional[Dict[str, str]] = None) -> List[Dict[str, Any]]:
    """Per-run text colours (for contrast analysis): explicit sRGB *and* resolved
    theme colours. Auto/inherited/unknown colours are skipped, never guessed."""
    theme_colors = theme_colors or {}
    out: List[Dict[str, Any]] = []
    for run in getattr(paragraph, "runs", []) or []:
        if not (run.text or "").strip():
            continue
        hex6 = _run_color_hex(run, theme_colors)
        if not hex6:
            continue
        size_pt = None
        try:
            if run.font.size is not None:
                size_pt = float(run.font.size.pt)
        except Exception:
            size_pt = None
        out.append({"c": hex6, "sz": size_pt, "b": bool(run.font.bold) if run.font.bold is not None else False})
    return out


def _paragraph_bg(paragraph) -> Optional[str]:
    """Explicit paragraph shading fill (``w:shd@w:fill``) if a real colour."""
    try:
        shd = paragraph._p.find(f"{_DOCX_NS}pPr/{_DOCX_NS}shd")
        if shd is not None:
            fill = shd.get(f"{_DOCX_NS}fill")
            if fill and fill.lower() not in ("auto",):
                return fill
    except Exception:
        pass
    return None


def _text_color_props(paragraph, theme_colors: Optional[Dict[str, str]] = None) -> Dict[str, Any]:
    """Build the ``metadata.properties`` carrying contrast inputs (or empty)."""
    colors = _explicit_run_colors(paragraph, theme_colors)
    if not colors:
        return {}
    props: Dict[str, Any] = {"explicit_text_colors": colors}
    bg = _paragraph_bg(paragraph)
    if bg:
        props["bg_color"] = bg
    return props


def _text_excluding_controls(el) -> str:
    """Visible ``w:t`` text under ``el``, EXCLUDING any nested ``w:sdt`` content.

    A field label must come from real label text, never from another content
    control's value/placeholder (which is what hijacked the 2nd field's label
    in a "Name: [ctrl] Date: [ctrl]" paragraph or a table cell that itself
    contains a control).
    """
    parts: List[str] = []

    def _walk(node) -> None:
        for child in node:
            local = etree.QName(child).localname
            if local == "sdt":
                continue  # skip a nested control's own text
            if local == "t":
                if child.text:
                    parts.append(child.text)
            else:
                _walk(child)

    try:
        _walk(el)
    except Exception:
        return ""
    return "".join(parts)


_GENERIC_SDT_PROMPTS = {
    "click or tap here to enter text",
    "click here to enter text",
    "choose an item",
    "enter text",
    "select an item",
    "choose a date",
}
# Short tokens that are field artifacts, not names.
_LABEL_STOPWORDS = {"n/a", "na", "tbd", "required", "optional", "yes", "no"}
_HEADING_PREFIX_RE = re.compile(r"^(section|part|chapter|article|step|appendix)\b", re.IGNORECASE)


def _clean_form_label(text: str, max_len: int = 60) -> Optional[str]:
    """Normalize candidate label text, or None if it isn't label-like.

    A real field label is a short noun phrase, optionally ending in a colon —
    not a sentence, question, instruction, heading or generic prompt.
    Deliberately conservative: a wrong label is worse than none.
    """
    t = (text or "").strip()
    # Strip leading bullet/asterisk/dash/colon artifacts ("* Required", "- Name").
    t = re.sub(r"^[\s•\*\-–—·:]+", "", t)
    # Strip a single trailing colon and surrounding space.
    t = t.rstrip().rstrip(":").strip()
    if not t or len(t) > max_len:
        return None
    if not any(c.isalpha() for c in t):
        return None
    low = t.lower()
    if low in _GENERIC_SDT_PROMPTS or low in _LABEL_STOPWORDS:
        return None
    # Sentences / questions / instructions are not labels.
    if t.endswith((".", "?", "!")):
        return None
    if "." in t and " " in t:  # an internal period with spaces reads as prose
        return None
    if _HEADING_PREFIX_RE.match(t):  # "Section 4 Employment" is a heading
        return None
    if len(t.split()) > 6:  # a field label is short
        return None
    return t


def _derive_sdt_label(sdt) -> Optional[str]:
    """Best-confidence accessible label for an unlabeled content control.

    Two high-precision sources, in order:
      1. Inline — the label text in the SAME paragraph immediately before this
         control, i.e. only text SINCE the previous control boundary (so the
         2nd control in "Name: [ ] Date: [ ]" derives "Date", not "Name ...").
      2. Table — the text of the cell immediately left of the control's cell,
         excluding any control content living in that cell.
    Returns None (leave for manual review) when neither yields a clean label,
    so we never invent a misleading name.
    """
    try:
        parent = sdt.getparent()
        if parent is None:
            return None
        # 1. Inline preceding text in the same paragraph, since the last control.
        if etree.QName(parent).localname == "p":
            preceding: List[str] = []
            for child in parent:
                if child is sdt:
                    break
                if etree.QName(child).localname == "sdt":
                    preceding = []  # an earlier control ends the previous label
                    continue
                txt = _text_excluding_controls(child)
                if txt:
                    preceding.append(txt)
            label = _clean_form_label("".join(preceding))
            if label:
                return label
        # 2. Table: the cell to the left in the same row.
        cell = sdt
        while cell is not None and etree.QName(cell).localname != "tc":
            cell = cell.getparent()
        if cell is not None:
            row = cell.getparent()
            if row is not None and etree.QName(row).localname == "tr":
                prev_cell = None
                for tc in row:
                    if etree.QName(tc).localname != "tc":
                        continue
                    if tc is cell:
                        break
                    prev_cell = tc
                if prev_cell is not None:
                    label = _clean_form_label(_text_excluding_controls(prev_cell))
                    if label:
                        return label
    except Exception:
        return None
    return None


def _docx_form_field_counts(doc) -> "tuple[int, int, int]":
    """Return ``(total, unlabeled, derivable)`` content controls (``w:sdt``).

    A content control's accessible label is its ``w:alias`` (the title). One
    with no alias has no accessible name. ``derivable`` is how many of the
    unlabeled ones we can confidently auto-label from nearby text (the rest
    stay manual). Reuses the same root-metadata keys as the PDF AcroForm check
    so a single analyzer flags both.
    """
    total = 0
    unlabeled = 0
    derivable = 0
    try:
        body = doc.element.body
        for sdt in body.iter(f"{_DOCX_NS}sdt"):
            total += 1
            alias = sdt.find(f"{_DOCX_NS}sdtPr/{_DOCX_NS}alias")
            val = alias.get(f"{_DOCX_NS}val") if alias is not None else None
            if not val or not str(val).strip():
                unlabeled += 1
                if _derive_sdt_label(sdt):
                    derivable += 1
    except Exception:
        return (total, unlabeled, derivable)
    return (total, unlabeled, derivable)
_REL_IMAGE_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


class _IdCounter:
    def __init__(self) -> None:
        self._counts: Dict[str, int] = {}
        # Walk-wide context the id counter already travels with: the
        # document's bookmark names, for resolving internal links. None ->
        # unknown (then any internal link is taken as valid, never guessed
        # broken).
        self.bookmarks: Optional[set] = None

    def __call__(self, prefix: str) -> str:
        i = self._counts.get(prefix, 0) + 1
        self._counts[prefix] = i
        return f"{prefix}-{i}"


def _no_register(_node_id: str, _obj: Any) -> None:
    return None


# Registry key under which the walk reports each (note_part, parsed_root) pair
# — the writer must re-serialize a note part it edited, and only the walk
# knows which parsed root it handed out.
_NOTE_PARTS_KEY = "__docx_note_parts__"

_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_W_T = qn("w:t")
_W_TAB = qn("w:tab")
_W_BR = qn("w:br")
_W_CR = qn("w:cr")
_W_R = qn("w:r")
_W_P = qn("w:p")
_W_TBL = qn("w:tbl")
_W_TXBX = qn("w:txbxContent")
# Subtrees whose w:t is NOT visible paragraph text: tracked deletions and
# moved-away text, text boxes (walked on their own), property blocks, and the
# legacy VML copy of a drawing Word keeps in mc:Fallback (a duplicate of the
# mc:Choice content it actually renders).
_TEXT_SKIP = {
    qn("w:del"), qn("w:moveFrom"), _W_TXBX, qn("w:pPr"), qn("w:rPr"),
    qn("w:sdtPr"), qn("w:instrText"), qn("w:delText"), _MC_FALLBACK,
}


def paragraph_text(p_el) -> str:
    """The text a reader sees in ``<w:p>`` — what Word renders, not what
    python-docx's ``paragraph.text`` happens to read.

    ``paragraph.text`` joins only direct ``w:r`` / ``w:hyperlink`` children,
    so a tracked insertion (``w:ins``), an inline content control
    (``w:sdt``), a smart tag or a HYPERLINK field vanished: "work up to
    <ins>three</ins> days" read as "work up to  days", and a paragraph that is
    one tracked insertion read as empty and was never analyzed. Tracked
    DELETIONS stay out (they are not in the document a reader gets).

    Shared by the parser and the writer's id pairing — both decide which
    paragraphs mint a ``docx-p`` id from this text.
    """
    parts: List[str] = []

    def walk(node) -> None:
        for child in node:
            tag = child.tag
            if tag == _W_T:
                if child.text:
                    parts.append(child.text)
            elif tag == _W_TAB:
                parts.append("\t")
            elif tag in (_W_BR, _W_CR):
                parts.append("\n")
            elif tag in _TEXT_SKIP or not isinstance(tag, str):
                continue
            else:
                walk(child)

    try:
        walk(p_el)
    except Exception:  # pragma: no cover - never let text extraction kill a parse
        return ""
    return "".join(parts)


def _partname(part) -> str:
    try:
        return str(part.partname)
    except Exception:  # pragma: no cover - defensive
        return ""


def _snippet_around(text: str, needle: str, width: int = 200) -> Optional[str]:
    """``text`` cut to <= ``width`` chars and still containing ``needle`` —
    the surrounding sentence a UI shows with the offending words highlighted."""
    t = " ".join((text or "").split())
    if not t:
        return None
    if len(t) <= width:
        return t
    n = " ".join((needle or "").split())
    i = t.find(n) if n else -1
    if i < 0:
        return t[: width - 1].rstrip() + "…"
    start = max(0, i - (width - len(n)) // 2)
    end = min(len(t), start + width)
    start = max(0, end - width)
    out = t[start:end].strip()
    if start > 0:
        out = "…" + out[1:]
    if end < len(t):
        out = out[:-1] + "…"
    return out


# ----- pictures --------------------------------------------------------------

_WP_INLINE = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
_WP_ANCHOR = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor"
_WP_DOCPR = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
_A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_R_LINK = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link"


def _inside(el, stop, tags) -> bool:
    """True when an ancestor of ``el`` strictly below ``stop`` has a tag in ``tags``."""
    anc = el.getparent()
    while anc is not None and anc is not stop:
        if anc.tag in tags:
            return True
        anc = anc.getparent()
    return False


def iter_paragraph_drawings(p_el) -> Iterator[Tuple[Any, str, Any]]:
    """``(wrapper, rId, docPr)`` for every picture drawn inside ``<w:p>``.

    One entry per ``wp:inline`` / ``wp:anchor`` that shows an image, in
    document order, with THAT drawing's own ``<wp:docPr>`` (where Word keeps
    its alt text). A picture inside a text box anchored in this paragraph is
    its own entry with its own docPr — the text box's anchor no longer also
    claims the picture's blip and gets the alt written onto the text box.
    Drawings in ``mc:Fallback`` (the legacy duplicate) are skipped.
    """
    wrappers = (_WP_INLINE, _WP_ANCHOR)
    for wrapper in p_el.iter(*wrappers):
        if _inside(wrapper, p_el, (_MC_FALLBACK,)):
            continue
        blip = None
        for cand in wrapper.iter(_A_BLIP):
            if not _inside(cand, wrapper, wrappers):
                blip = cand
                break
        if blip is None:
            continue
        rid = blip.get(_R_EMBED) or blip.get(_R_LINK)
        if not rid:
            continue
        doc_pr = wrapper.find(_WP_DOCPR)
        if doc_pr is None:
            continue
        yield wrapper, rid, doc_pr


def _image_nodes_from_p(
    p_el,
    blobs: Dict[str, Tuple[Optional[str], Optional[str]]],
    ids: "_IdCounter",
    prefix: str,
    reg: Callable[[str, Any], None],
    *,
    story: str,
    partname: str,
    context_fn: Optional[Callable[[], Optional[Tuple[str, str]]]] = None,
    extra_props: Optional[Dict[str, Any]] = None,
) -> List[ImageNode]:
    """ImageNodes for the pictures in ``p_el`` (see ``iter_paragraph_drawings``).

    ``context_fn`` returns ``(kind, text)`` — kind ``"caption"`` for a real
    caption, ``"nearby"`` for other text near the picture — evaluated only
    when the paragraph actually has a picture.
    """
    out: List[ImageNode] = []
    context: Optional[Tuple[str, str]] = None
    context_done = False
    for _wrapper, rid, doc_pr in iter_paragraph_drawings(p_el):
        alt_text = (doc_pr.get("descr") or doc_pr.get("title") or "").strip()
        is_decorative = (doc_pr.get("hidden") or "").lower() in {"1", "true"}
        b64, mime = blobs.get(rid, (None, None))
        properties: Dict[str, Any] = {"image_rid": rid, "docx_story": story, "docx_part": partname}
        if extra_props:
            properties.update(extra_props)
        if b64:
            properties["image_b64"] = b64
            properties["image_mime"] = mime or "image/png"
        if not context_done:
            context = context_fn() if context_fn else None
            context_done = True
        if context:
            kind, text = context
            properties["caption" if kind == "caption" else "nearby_text"] = text
            properties["snippet"] = _snippet_around(text, "")
        node_id = ids(prefix)
        reg(node_id, doc_pr)
        if is_decorative and alt_text:
            out.append(
                ImageNode.model_construct(
                    id=node_id,
                    node_type=ImageNode.type_value(),
                    content=NodeContent(kind=ContentKind.NONE),
                    metadata=NodeMetadata(source_format="docx", properties=properties),
                    children=[],
                    accessibility_flags=[],
                    is_decorative=True,
                    alt_text=alt_text,
                )
            )
        else:
            out.append(
                ImageNode(
                    id=node_id,
                    content=NodeContent(kind=ContentKind.NONE),
                    metadata=NodeMetadata(source_format="docx", properties=properties),
                    children=[],
                    accessibility_flags=[],
                    is_decorative=is_decorative,
                    alt_text=alt_text or None,
                )
            )
    return out


# ----- effective formatting (style chain) ---------------------------------------


def _toggle_on(el) -> bool:
    """Value of an OOXML on/off element (w:b, w:i …): absent -> False,
    present with no val / true / 1 / on -> True."""
    if el is None:
        return False
    val = (el.get(qn("w:val")) or "").strip().lower()
    return val not in ("0", "false", "off", "none")


class DocxStyleResolver:
    """Effective paragraph/run formatting, resolved the way Word does it:
    direct formatting -> character style chain -> paragraph style chain ->
    ``docDefaults``. Read straight from ``styles.xml`` (no python-docx
    ``paragraph.style``, which costs a full styles scan per call).

    Used by the parser to measure how prominent a heading-looking line is,
    and by the writer to keep a promoted line looking exactly as it did.
    """

    MAX_CHAIN = 25

    def __init__(self, doc) -> None:
        self._doc = doc
        self.reload()

    def reload(self) -> None:
        styles_el = self._doc.styles.element
        self._styles: Dict[str, Any] = {}
        self.default_paragraph_style_id: Optional[str] = None
        for st in styles_el.iterfind(qn("w:style")):
            sid = st.get(qn("w:styleId"))
            if not sid:
                continue
            self._styles.setdefault(sid, st)
            if (
                self.default_paragraph_style_id is None
                and st.get(qn("w:type")) == "paragraph"
                and (st.get(qn("w:default")) or "").lower() in ("1", "true", "on")
            ):
                self.default_paragraph_style_id = sid
        dd = styles_el.find(qn("w:docDefaults"))
        self._dd_rpr = dd.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}") if dd is not None else None
        self._dd_ppr = dd.find(f"{qn('w:pPrDefault')}/{qn('w:pPr')}") if dd is not None else None
        self._chain_cache: Dict[Optional[str], List[Any]] = {}

    def style_element(self, style_id: Optional[str]):
        return self._styles.get(style_id) if style_id else None

    def chain(self, style_id: Optional[str]) -> List[Any]:
        if style_id in self._chain_cache:
            return self._chain_cache[style_id]
        out: List[Any] = []
        seen = set()
        sid = style_id
        while sid and sid in self._styles and sid not in seen and len(out) < self.MAX_CHAIN:
            seen.add(sid)
            st = self._styles[sid]
            out.append(st)
            based = st.find(qn("w:basedOn"))
            sid = based.get(qn("w:val")) if based is not None else None
        self._chain_cache[style_id] = out
        return out

    def paragraph_style_id(self, p_el) -> Optional[str]:
        """The paragraph's style id; an unknown or missing id falls back to
        the default paragraph style, as Word renders it."""
        pPr = p_el.find(qn("w:pPr"))
        ps = pPr.find(qn("w:pStyle")) if pPr is not None else None
        sid = ps.get(qn("w:val")) if ps is not None else None
        if not sid or sid not in self._styles:
            sid = self.default_paragraph_style_id
        return sid

    def run_element(self, r_el, p_style_id: Optional[str], tag: str):
        """The ``rPr/{tag}`` element that decides this run's formatting."""
        rpr = r_el.find(qn("w:rPr"))
        if rpr is not None:
            el = rpr.find(tag)
            if el is not None:
                return el
            rs = rpr.find(qn("w:rStyle"))
            if rs is not None:
                for st in self.chain(rs.get(qn("w:val"))):
                    el = st.find(f"{qn('w:rPr')}/{tag}")
                    if el is not None:
                        return el
        return self.style_run_element(p_style_id, tag)

    def style_run_element(self, p_style_id: Optional[str], tag: str):
        for st in self.chain(p_style_id):
            el = st.find(f"{qn('w:rPr')}/{tag}")
            if el is not None:
                return el
        if self._dd_rpr is not None:
            return self._dd_rpr.find(tag)
        return None

    def paragraph_element(self, p_el, p_style_id: Optional[str], tag: str, direct: bool = True):
        """The ``pPr/{tag}`` element that decides this paragraph's formatting."""
        if direct:
            pPr = p_el.find(qn("w:pPr"))
            if pPr is not None:
                el = pPr.find(tag)
                if el is not None:
                    return el
        for st in self.chain(p_style_id):
            el = st.find(f"{qn('w:pPr')}/{tag}")
            if el is not None:
                return el
        if self._dd_ppr is not None:
            return self._dd_ppr.find(tag)
        return None

    def run_size_pt(self, r_el, p_style_id: Optional[str]) -> float:
        el = self.run_element(r_el, p_style_id, qn("w:sz"))
        try:
            return int(el.get(qn("w:val"))) / 2.0 if el is not None else 10.0
        except (TypeError, ValueError):
            return 10.0  # Word's size when nothing in the hierarchy sets one

    def run_bold(self, r_el, p_style_id: Optional[str]) -> bool:
        return _toggle_on(self.run_element(r_el, p_style_id, qn("w:b")))


def visible_runs(p_el) -> List[Any]:
    """The ``w:r`` elements whose text a reader sees in this paragraph (same
    exclusions as :func:`paragraph_text`)."""
    out: List[Any] = []

    def walk(node) -> None:
        for child in node:
            tag = child.tag
            if tag == _W_R:
                if paragraph_text(child).strip():
                    out.append(child)
            elif tag in _TEXT_SKIP or not isinstance(tag, str):
                continue
            else:
                walk(child)

    walk(p_el)
    return out


_OUTLINE_KEYWORD_RE = re.compile(r"^(part|book|volume|chapter|unit|module)\b", re.IGNORECASE)
_OUTLINE_KEYWORD_RANK = {"part": 3, "book": 3, "volume": 3, "chapter": 2, "unit": 2, "module": 2}
_OUTLINE_NUMBER_RE = re.compile(r"^(\d{1,3}(?:\.\d{1,3})*)(?:[.)]|\s)")


def _heading_visual(p_el, styles: "DocxStyleResolver", text: str) -> Dict[str, Any]:
    """How prominent a heading (real or fake) LOOKS: effective size (largest
    visible run), whether every visible run is bold, and any outline cue in
    the text ("Part 3", "2.1"). PROMOTE_HEADING places a fake heading in the
    document's own outline from these — it never guesses a level."""
    sid = styles.paragraph_style_id(p_el)
    runs = visible_runs(p_el)
    size: Optional[float] = None
    bold = bool(runs)
    for r in runs:
        s = styles.run_size_pt(r, sid)
        size = s if size is None else max(size, s)
        if not styles.run_bold(r, sid):
            bold = False
    t = (text or "").strip()
    kw = _OUTLINE_KEYWORD_RE.match(t)
    num = _OUTLINE_NUMBER_RE.match(t)
    return {
        "size_pt": size,
        "bold": bold,
        "keyword_rank": _OUTLINE_KEYWORD_RANK.get(kw.group(1).lower(), 0) if kw else 0,
        "number_depth": len(num.group(1).split(".")) if num else 0,
    }


# ----- page headers / footers --------------------------------------------------


def iter_header_footer_parts(doc) -> List[Tuple[str, Any, str]]:
    """``[("header"|"footer", part, variant)]`` — variant is the reference's
    ``w:type`` ("default", "first" or "even") — for every header/footer part
    Word actually PRINTS, in section order, each part ONCE (sections that
    "link to previous" share a part and must not be audited twice).

    Printed means: the default variant always; the first-page variant only
    in a section with "Different first page" (``w:titlePg``); the even-page
    variant only when the document has "Different odd & even pages"
    (``w:evenAndOddHeaders`` in settings). Word keeps the part and its
    reference when either option is switched OFF, so a template that once
    had a first-page letterhead still carries it — auditing that part
    reported (and charged to fix) a logo no reader ever meets. A section
    with no reference of a variant inherits the previous section's, as Word
    does. Unreferenced parts are ignored, like Word ignores them.

    The parser and the writer both walk headers through this one function,
    so the ``docx-hf*`` ids they mint cannot disagree.
    """
    out: List[Tuple[str, Any, str]] = []
    seen: set = set()
    hdr, ftr = qn("w:headerReference"), qn("w:footerReference")
    rid_attr = qn("r:id")
    try:
        related = doc.part.related_parts
        sect_prs = [
            sp for sp in doc.element.body.iter(qn("w:sectPr"))
            # A tracked change to section properties keeps the OLD sectPr
            # inside w:sectPrChange; it is history, not a section.
            if sp.getparent() is None or sp.getparent().tag != qn("w:sectPrChange")
        ]
    except Exception:  # pragma: no cover - defensive
        return out
    even_on = False
    try:
        # Read the settings part only if it exists — ``doc.settings`` would
        # CREATE one, and the writer runs this walk on the copy it saves.
        from docx.opc.constants import RELATIONSHIP_TYPE as _RT

        settings_el = doc.part.part_related_by(_RT.SETTINGS).element
        even_on = _toggle_on(settings_el.find(qn("w:evenAndOddHeaders")))
    except Exception:  # no settings part: Word's default (off)
        even_on = False
    effective: Dict[Tuple[str, str], Any] = {}
    for sp in sect_prs:
        own: List[Tuple[str, str]] = []
        for ref in sp:
            if ref.tag == hdr:
                kind = "header"
            elif ref.tag == ftr:
                kind = "footer"
            else:
                continue
            variant = ref.get(qn("w:type")) or "default"
            part = related.get(ref.get(rid_attr)) if hasattr(related, "get") else None
            if part is None or getattr(part, "element", None) is None:
                continue
            effective[(kind, variant)] = part
            own.append((kind, variant))
        printed = {"default"}
        if _toggle_on(sp.find(qn("w:titlePg"))):
            printed.add("first")
        if even_on:
            printed.add("even")
        # This section's own references first (document order), then what it
        # inherits — a fixed order both walks reproduce.
        inherited = [
            (k, v) for k in ("header", "footer") for v in ("default", "first", "even")
            if (k, v) not in own
        ]
        for kind, variant in own + inherited:
            if variant not in printed:
                continue
            part = effective.get((kind, variant))
            if part is None:
                continue
            key = _partname(part) or id(part)
            if key in seen:
                continue
            seen.add(key)
            out.append((kind, part, variant))
    return out


def iter_story_paragraphs(root_el) -> List[Any]:
    """Every ``<w:p>`` in a header/footer story, in document order — body
    paragraphs AND table-cell paragraphs (a letterhead is usually a layout
    table: logo | agency name) — except paragraphs inside text boxes (their
    anchor paragraph owns their pictures) and mc:Fallback duplicates."""
    return [
        p for p in root_el.iter(_W_P)
        if not _inside(p, root_el, (_W_TXBX, _MC_FALLBACK))
    ]


class _StoryParent:
    """Minimal parent for a python-docx Paragraph living in a header/footer
    part: ``Paragraph.part`` must resolve to THAT part."""

    def __init__(self, part) -> None:
        self.part = part


def _header_footer_section(doc, ids: "_IdCounter", reg, theme_colors) -> Optional[SectionNode]:
    section = SectionNode(
        id="docx-page-furniture",
        content=NodeContent(kind=ContentKind.TEXT, text="Page headers and footers"),
        metadata=NodeMetadata(source_format="docx", properties={"page_furniture": True}),
        children=[],
        accessibility_flags=[],
    )
    for kind, part, variant in iter_header_footer_parts(doc):
        try:
            root_el = part.element
            blobs = _collect_image_blobs(part)
            partname = _partname(part)
            paragraphs = iter_story_paragraphs(root_el)
            story_text = " ".join(
                t for t in (" ".join(paragraph_text(p).split()) for p in paragraphs) if t
            )
        except Exception:  # pragma: no cover - a broken header never kills the parse
            continue
        parent = _StoryParent(part)
        for p_el in paragraphs:
            text = paragraph_text(p_el).strip()
            base = {"docx_story": kind, "docx_part": partname, "page_furniture": kind,
                    "docx_story_variant": variant}
            links = _link_nodes_from_p(
                p_el, part, ids, prefix="docx-hflink", reg=reg, context_text=text, extra_props=base,
            )
            section.children.extend(links)
            ctx = " ".join(text.split()) or story_text
            section.children.extend(
                _image_nodes_from_p(
                    p_el, blobs, ids, "docx-hfimg", reg, story=kind, partname=partname,
                    context_fn=(lambda _c=ctx: ("nearby", _c[:200]) if _c else None),
                    extra_props={"page_furniture": kind, "docx_story_variant": variant},
                )
            )
            if text and not links:
                section.children.append(_story_paragraph_node(p_el, parent, text, base, ids, reg, theme_colors))
            # Text boxes anchored here (a footer's "Privacy: click here"
            # callout, a letterhead address block). Their pictures already
            # belong to this anchor paragraph; their words and links do not,
            # and were never read. Fallback copies are skipped as in the body.
            for tb_p in _iter_text_box_paragraphs(p_el):
                tb_text = paragraph_text(tb_p).strip()
                tb_base = dict(base, in_text_box=True)
                tb_links = _link_nodes_from_p(
                    tb_p, part, ids, prefix="docx-hflink", reg=reg, context_text=tb_text, extra_props=tb_base,
                )
                section.children.extend(tb_links)
                if tb_text and not tb_links:
                    section.children.append(
                        _story_paragraph_node(tb_p, parent, tb_text, tb_base, ids, reg, theme_colors)
                    )
    return section if section.children else None


def _story_paragraph_node(p_el, parent, text: str, base: Dict[str, Any], ids, reg, theme_colors) -> ParagraphNode:
    """A header/footer line as a ParagraphNode (with its contrast inputs),
    registered so the writer can recolour exactly this paragraph."""
    para = Paragraph(p_el, parent)
    props = dict(_text_color_props(para, theme_colors) or {})
    props.update(base)
    hf_id = ids("docx-hfp")
    reg(hf_id, para)
    return ParagraphNode(
        id=hf_id,
        content=NodeContent(kind=ContentKind.TEXT, text=text),
        metadata=NodeMetadata(source_format="docx", properties=props),
        children=[],
        accessibility_flags=[],
    )


def _iter_sdt_aware(parent_el, want_tag: str):
    """Direct children of ``parent_el`` with tag ``want_tag``, in document
    order, DESCENDING into block-level content controls (w:sdt/w:sdtContent,
    possibly nested) — and into nothing else.

    Word templates — government forms especially — wrap whole sections,
    paragraphs and table rows in content controls. python-docx's
    ``doc.paragraphs`` / ``doc.tables`` / ``table.rows`` read only DIRECT
    children (``./w:p`` etc.), so everything inside an SDT was invisible: a
    form built from content controls analyzed with fewer findings than the
    same document without them, and its fixable issues were never even
    detected. A manual walk (not ``.iter()``) so we never descend into a
    nested table's subtree and double-count its paragraphs.
    """
    sdt = qn("w:sdt")
    sdt_content = qn("w:sdtContent")
    for child in parent_el:
        tag = getattr(child, "tag", None)
        if tag == want_tag:
            yield child
        elif tag == sdt:
            content = child.find(sdt_content)
            if content is not None:
                yield from _iter_sdt_aware(content, want_tag)


def iter_body_paragraphs(doc):
    """Every body-level Paragraph in document order, SDT-descended.

    The writer's id-pairing indexes MUST iterate with this same helper —
    parser and writer agree on ``docx-p-N`` ids only because they walk the
    body identically.
    """
    from docx.text.paragraph import Paragraph

    for el in _iter_sdt_aware(doc.element.body, qn("w:p")):
        yield Paragraph(el, doc._body)  # noqa: SLF001


def iter_body_tables(doc):
    """Every body-level Table in document order, SDT-descended (see above)."""
    from docx.table import Table

    for el in _iter_sdt_aware(doc.element.body, qn("w:tbl")):
        yield Table(el, doc._body)  # noqa: SLF001


def iter_table_rows(table):
    """Every row of ``table`` in order, including SDT-wrapped rows.

    ``table.rows`` reads direct ``w:tr`` only; a repeating-section content
    control wraps its rows in w:sdt and they vanished from analysis. (An
    SDT-wrapped individual CELL is still out of scope — rare, and cell
    addressing runs through python-docx's grid logic we don't reimplement.)
    """
    from docx.table import _Row

    for tr in _iter_sdt_aware(table._tbl, qn("w:tr")):  # noqa: SLF001
        yield _Row(tr, table)


def paragraph_style_name(paragraph, cache: Dict[Any, str]) -> str:
    """``paragraph.style.name`` without python-docx's per-paragraph cost.

    ``paragraph.style`` resolves the *default* paragraph style — the case for
    most body text, which carries no explicit ``w:pStyle`` — by walking every
    style element in ``styles.xml`` and reading attributes off each one.
    Profiling a 200-page document showed that single property at 89% of the
    writer's wall time (5.7s of 6.4s, 1.7M attribute reads), and the parser
    pays it again. At 500 pages that is ~40s across the two, most of the way
    to a proxy timeout, for a lookup whose answer is the same for every
    body paragraph in the file.

    So: read the ``w:pStyle`` id straight off the XML and resolve it through
    a per-document ``cache`` (style-id -> name); resolve the implicit default
    ONCE and cache it under ``None``. Falls back to the slow property for any
    id the cache cannot resolve, so the answer is always identical to what
    python-docx would have said — the parser and writer must agree on ids,
    and both call this.
    """
    pPr = paragraph._p.find(qn("w:pPr"))  # noqa: SLF001
    style_id = None
    if pPr is not None:
        pStyle = pPr.find(qn("w:pStyle"))
        if pStyle is not None:
            style_id = pStyle.get(qn("w:val"))
    if style_id in cache:
        return cache[style_id]
    try:
        name = (paragraph.style.name or "") if paragraph.style else ""
    except Exception:
        name = ""
    cache[style_id] = name
    return name


def _outline_heading_level(p_el, styles: "DocxStyleResolver") -> int:
    """Heading level from the paragraph's effective ``w:outlineLvl`` (direct,
    else its style chain): 0-8 are heading levels 1-9 (capped at 6 like the
    style-name rule), 9 is body text. 0 when it is not a heading."""
    try:
        el = styles.paragraph_element(p_el, styles.paragraph_style_id(p_el), qn("w:outlineLvl"))
        if el is None:
            return 0
        val = int(el.get(qn("w:val")))
    except (TypeError, ValueError):
        return 0
    if 0 <= val <= 8:
        return min(6, val + 1)
    return 0


def _heading_level_from_style(style: str) -> int:
    if not style:
        return 0
    if not style.lower().startswith("heading"):
        return 0
    parts = style.split()
    if len(parts) < 2:
        return 1
    try:
        return max(1, min(6, int(parts[1])))
    except Exception:
        return 1


def _is_list_paragraph(paragraph) -> bool:
    return is_list_p(paragraph._p)  # noqa: SLF001


def is_list_p(p_el) -> bool:
    """True when ``<w:p>`` carries its own Word list numbering. ``numId 0`` is
    Word's explicit "no numbering" (it switches off a style's numbering), so
    it is NOT a list item."""
    pPr = p_el.find(f"{_DOCX_NS}pPr")
    if pPr is None:
        return False
    numPr = pPr.find(f"{_DOCX_NS}numPr")
    if numPr is None:
        return False
    num_id = numPr.find(f"{_DOCX_NS}numId")
    if num_id is not None and (num_id.get(f"{_DOCX_NS}val") or "").strip() == "0":
        return False
    return True


def _detect_list_marker(paragraph) -> str:
    pPr = paragraph._p.find(f"{_DOCX_NS}pPr")
    if pPr is None:
        return "bullet"
    numPr = pPr.find(f"{_DOCX_NS}numPr")
    if numPr is None:
        return "bullet"
    numId = numPr.find(f"{_DOCX_NS}numId")
    return f"num-{numId.get(f'{_DOCX_NS}val')}" if numId is not None else "bullet"


def _finalize_list(ids: _IdCounter, items: List[ListItemNode], marker: Optional[str]) -> ListNode:
    return ListNode(
        id=ids("docx-list"),
        ordered=bool(marker and marker.startswith("num")),
        marker=marker,
        content=NodeContent(kind=ContentKind.NONE),
        metadata=NodeMetadata(source_format="docx"),
        children=list(items),
        accessibility_flags=[],
    )


def _collect_image_blobs(part) -> Dict[str, Tuple[Optional[str], Optional[str]]]:
    """Return ``{rId: (base64, mime)}`` for the images ``part`` embeds (the
    document part, or a header/footer part — each has its own rIds)."""

    blobs: Dict[str, Tuple[Optional[str], Optional[str]]] = {}
    if not _WANT_IMAGE_BYTES:
        # Nothing reads the bytes without a vision provider; only the mime
        # would be recorded, and that is recoverable from the rId.
        return blobs
    for rel_id, rel in part.rels.items():
        try:
            if "image" not in str(rel.reltype):
                continue
            blob = rel.target_part.blob
        except Exception:
            continue
        if not blob or len(blob) > 2_000_000:
            continue
        target = str(getattr(rel, "target_ref", "") or "").lower()
        if target.endswith(".jpg") or target.endswith(".jpeg"):
            mime = "image/jpeg"
        elif target.endswith(".gif"):
            mime = "image/gif"
        elif target.endswith(".bmp"):
            mime = "image/bmp"
        else:
            mime = "image/png"
        blobs[rel_id] = ((base64.b64encode(blob).decode("ascii") if _WANT_IMAGE_BYTES else None), mime)
    return blobs


_FLD_HYPERLINK_RE = re.compile(r"HYPERLINK\s+(?:\"([^\"]+)\"|(\S+))")
# HYPERLINK \l "bookmark" — the \l switch names a location in this document.
_FLD_LOCAL_RE = re.compile(r"\\l\s+(?:\"([^\"]+)\"|(\S+))")


def _bookmark_names(doc) -> set:
    """Every bookmark name in the main document (what w:anchor links target)."""
    try:
        return {
            bm.get(qn("w:name"))
            for bm in doc.element.body.iter(qn("w:bookmarkStart"))
            if bm.get(qn("w:name"))
        }
    except Exception:  # pragma: no cover - defensive
        return set()


def _link_elements_in_paragraph(p) -> List[Tuple[str, Any]]:
    """Return ``[(kind, element), …]`` for every link the parser emits from a
    ``<w:p>``: ``w:hyperlink`` with visible text, plus HYPERLINK
    ``w:fldSimple`` fields with visible text — in document order.

    This is the single source of truth for "what counts as a link"; the docx
    writer indexes links (and decides which paragraphs consumed a ``docx-p``
    id) through this same helper, so parser/writer id minting cannot drift.
    """

    out: List[Tuple[str, Any]] = []
    for el in p.iter(f"{_DOCX_NS}hyperlink", f"{_DOCX_NS}fldSimple"):
        # A link nested inside a TEXT BOX below this paragraph belongs to the
        # text-box walk (docx-tblink id space) — emitting it here too would
        # double-count it and desync both id spaces. (When ``p`` itself lives
        # inside a text box, the ancestor walk stops at ``p`` before reaching
        # the txbxContent above it, so text-box paragraphs still match.)
        # (mc:Fallback is the legacy duplicate of content Word renders from
        # mc:Choice — counting it would report every such link twice.)
        if _inside(el, p, (_W_TXBX, _MC_FALLBACK)):
            continue
        text = "".join(
            (t.text or "") for t in el.iterfind(f".//{_DOCX_NS}t")
        ).strip()
        if not text:
            continue
        if el.tag == f"{_DOCX_NS}hyperlink":
            out.append(("hyperlink", el))
        else:
            instr = el.get(f"{_DOCX_NS}instr") or ""
            if "HYPERLINK" in instr:
                out.append(("fldsimple", el))
    return out


def _link_nodes_from_p(
    p_el,
    part,
    ids: _IdCounter,
    prefix: str = "docx-link",
    extra_props: Optional[Dict[str, Any]] = None,
    reg: Optional[Callable[[str, Any], None]] = None,
    context_text: Optional[str] = None,
) -> List[LinkNode]:
    """Build LinkNodes from a raw ``<w:p>`` element (body, cell, text box,
    note or header/footer). ``context_text`` (the paragraph's own text)
    becomes each link's ``snippet`` — the sentence around "click here"."""
    nodes: List[LinkNode] = []
    reg = reg or _no_register
    links = _link_elements_in_paragraph(p_el)
    if links and context_text is None:
        context_text = paragraph_text(p_el)
    for kind, el in links:
        text = "".join(
            (t.text or "") for t in el.iterfind(f".//{_DOCX_NS}t")
        ).strip()
        if kind == "hyperlink":
            rid = el.get(f"{_DOCX_REL_NS}id")
            target = ""
            if rid and rid in part.rels:
                target = str(part.rels[rid].target_ref or "")
            anchor = (el.get(f"{_DOCX_NS}anchor") or "").strip()
        else:
            instr = el.get(f"{_DOCX_NS}instr") or ""
            m_local = _FLD_LOCAL_RE.search(instr)
            instr_wo_local = _FLD_LOCAL_RE.sub(" ", instr)
            m = _FLD_HYPERLINK_RE.search(instr_wo_local)
            target = (m.group(1) or m.group(2)) if m else ""
            if target.startswith("\\"):
                target = ""  # a field switch, not a destination
            anchor = ((m_local.group(1) or m_local.group(2)) if m_local else "").strip()
        # An internal link — every Word TOC entry and cross-reference is a
        # w:hyperlink with w:anchor="_Toc…" and NO r:id — used to come out
        # with target None, so LINK_TARGET_BROKEN fired on every line of every
        # table of contents. It points at a bookmark: valid when that bookmark
        # exists, genuinely broken when it does not.
        if anchor:
            if target:
                if "#" not in target:
                    target = f"{target}#{anchor}"
            elif ids.bookmarks is None or anchor in ids.bookmarks or anchor.lower() == "_top":
                target = f"#{anchor}"
        props: Dict[str, Any] = {"link_kind": kind}
        if extra_props:
            props.update(extra_props)
        snippet = _snippet_around(context_text or "", text)
        if snippet:
            props["snippet"] = snippet
            # The exact words to highlight inside the snippet (whitespace
            # normalised the same way the snippet is).
            highlight = " ".join(text.split())
            if highlight and highlight in snippet:
                props["highlight"] = highlight
        node_id = ids(prefix)
        reg(node_id, el)
        nodes.append(
            LinkNode(
                id=node_id,
                target=target or None,
                content=NodeContent(kind=ContentKind.TEXT, text=text),
                metadata=NodeMetadata(source_format="docx", properties=props),
                children=[],
                accessibility_flags=[],
            )
        )
    return nodes


def _hyperlink_nodes_in_paragraph(paragraph, ids: _IdCounter) -> List[LinkNode]:
    return _link_nodes_from_p(paragraph._p, paragraph.part, ids)


def _iter_text_box_paragraphs(body_el) -> List[Any]:
    """All ``<w:p>`` elements living inside text boxes (``w:txbxContent``),
    in document order. Text boxes are invisible to ``doc.paragraphs`` —
    sidebars and callouts would otherwise never be analyzed.

    A modern text box is stored twice: the DrawingML copy Word renders
    (mc:Choice) and a VML copy for old readers (mc:Fallback). Walking both
    reported every sidebar link and line twice, so the Fallback copy is
    skipped."""
    out: List[Any] = []
    for tx in body_el.iter(_W_TXBX):
        if _inside(tx, body_el, (_MC_FALLBACK,)):
            continue
        out.extend(tx.iterfind(f"{_DOCX_NS}p"))
    return out


_NOTE_TYPE_SKIP = {"separator", "continuationSeparator", "continuationNotice"}


def _iter_note_parts(doc) -> List[Tuple[Any, Any]]:
    """``[(part, parsed_root)]`` for the footnotes and endnotes parts (in
    that fixed order) when present. Notes live in separate package parts that
    ``doc.paragraphs`` never opens — their content was previously invisible.
    Shared with the docx writer so note link ids (``docx-fnlink-N``) mint
    identically; the writer re-serializes the parsed root back into the
    part's blob after rewriting."""
    from docx.opc.constants import RELATIONSHIP_TYPE as _RT

    out: List[Tuple[Any, Any]] = []
    for rt in (_RT.FOOTNOTES, _RT.ENDNOTES):
        try:
            part = doc.part.part_related_by(rt)
        except KeyError:
            continue
        # A part python-docx loaded as XML already has a live element (and
        # serializes it on save); a plain blob part is parsed here and the
        # writer writes the root back into its blob.
        live = getattr(part, "element", None)
        if live is not None and isinstance(getattr(live, "tag", None), str):
            out.append((part, live))
            continue
        try:
            root = etree.fromstring(part.blob)
        except Exception:
            continue
        out.append((part, root))
    return out


def _note_paragraphs(root) -> List[Any]:
    """Real note ``<w:p>`` elements; separator/continuation stubs skipped."""
    out: List[Any] = []
    for note in root.iter(f"{_DOCX_NS}footnote", f"{_DOCX_NS}endnote"):
        if (note.get(f"{_DOCX_NS}type") or "") in _NOTE_TYPE_SKIP:
            continue
        out.extend(note.iterfind(f".//{_DOCX_NS}p"))
    return out


# ----- Fake-list detection ----------------------------------------------------
#
# Plain paragraphs typed as "- item" / "1. item" with no Word numbering are
# read by screen readers as disconnected prose — no list semantics, no item
# count, no nesting.  The parser marks candidates here; a post-pass groups
# consecutive candidates into runs; ListStructureAnalyzer flags each run; and
# the FIX_LIST_STRUCTURE executor + docx writer convert them into REAL Word
# lists (w:numPr + numbering.xml).  Precision over recall: en/em dashes are
# excluded (dialogue), alpha ordinals are excluded ("A. Smith"), and numbered
# runs must count 1, 2, 3… from 1.

_FAKE_BULLET_RE = re.compile(r"^([-*•·])\s+\S")
_FAKE_DECIMAL_RE = re.compile(r"^(\d{1,3})[.)]\s+\S")


def _fake_list_signature(text: str) -> Optional[Tuple[str, str, Optional[int]]]:
    """Return ``(kind, prefix_char, ordinal)`` when ``text`` is typed like a
    list item, else None. ``kind`` is "bullet" or "decimal"."""

    m = _FAKE_BULLET_RE.match(text)
    if m:
        return ("bullet", m.group(1), None)
    m = _FAKE_DECIMAL_RE.match(text)
    if m:
        return ("decimal", m.group(1), int(m.group(1)))
    return None


def strip_fake_list_prefix(text: str) -> str:
    """Remove the literal typed marker ("- ", "1. ", "2) "…) from ``text``.

    Shared with the executor and the docx writer so the three stay in
    lockstep about what counts as a marker.
    """

    out = re.sub(r"^[-*•·]\s+", "", text, count=1)
    if out != text:
        return out
    return re.sub(r"^\d{1,3}[.)]\s+", "", text, count=1)


def _group_fake_list_runs(body_children: List[Any]) -> None:
    """Mark runs of >=2 consecutive same-kind fake-list ParagraphNodes.

    The FIRST node of each run gets ``fake_list_run_ids`` (all member ids,
    itself included) — the analyzer flags that node, so one issue surfaces
    per typed list. Numbered runs must be sequential starting at 1, which
    keeps prose like "1986. It was…" or stray numbered sentences out.
    """

    run: List[Any] = []
    run_kind: Optional[str] = None
    run_char: Optional[str] = None
    next_ordinal: Optional[int] = None

    def flush() -> None:
        nonlocal run, run_kind, run_char, next_ordinal
        if len(run) >= 2:
            first = run[0]
            props = dict(first.metadata.properties or {})
            props["fake_list_run_ids"] = [n.id for n in run]
            first.metadata.properties = props
        run = []
        run_kind = None
        run_char = None
        next_ordinal = None

    for child in body_children:
        sig = None
        if isinstance(child, ParagraphNode):
            props = child.metadata.properties or {}
            kind = props.get("fake_list_kind")
            if kind:
                sig = (kind, props.get("fake_list_char"), props.get("fake_list_ordinal"))
        if sig is None:
            flush()
            continue
        kind, char, ordinal = sig
        if run and (kind != run_kind or (kind == "bullet" and char != run_char)):
            flush()
        if kind == "decimal":
            expected = 1 if not run else next_ordinal
            if ordinal != expected:
                flush()
                if ordinal != 1:
                    continue  # mid-sequence stray ("1986. …") — not a list start
            next_ordinal = (ordinal or 0) + 1
        if not run:
            run_kind, run_char = kind, char
        run.append(child)
    flush()


def _looks_like_fake_heading(paragraph, style_name: str, text: str) -> bool:
    """True when a plain paragraph is visually presented as a heading.

    The classic title-page failure: 24pt bold text typed as a normal paragraph
    (or Word's Title/Subtitle styles, which are NOT Heading 1-9 and don't enter
    the navigation outline). Conservative on purpose:

    * Title / Subtitle styles always count — they're unambiguous.
    * Otherwise the text must be SHORT (<= 60 chars, <= 8 words), not end like
      a sentence, have every visible run bold, AND carry an explicit font size
      >= 14pt on some run. Ordinary bold emphasis inside body text fails the
      size requirement; bold labels fail nothing else often enough that the
      size requirement is what keeps precision high.
    """
    sn = (style_name or "").strip().lower()
    if sn in ("title", "subtitle"):
        return True

    t = (text or "").strip()
    if not t or len(t) > 60 or len(t.split()) > 8:
        return False
    if t.endswith((".", "!", "?", ";", ":", ",")):
        return False

    saw_text_run = False
    max_size_pt = 0.0
    for run in paragraph.runs:
        if not (run.text or "").strip():
            continue
        saw_text_run = True
        if not run.bold:
            return False
        try:
            if run.font.size is not None:
                max_size_pt = max(max_size_pt, float(run.font.size.pt))
        except Exception:
            pass
    return saw_text_run and max_size_pt >= 14.0


def _cell_text_is_bold(cell) -> bool:
    """True when the cell has visible text and every run carrying it is bold —
    the common 'visual header' convention (a bold first row)."""
    saw_text = False
    for para in cell.paragraphs:
        for run in para.runs:
            if (run.text or "").strip():
                saw_text = True
                if not run.bold:
                    return False
    return saw_text


def _docx_row_is_header(row) -> bool:
    """A row is a genuine header row when Word marks it as one (``w:tblHeader``,
    the *Repeat as header row* property) or every populated cell is bold."""
    trPr = row._tr.find(qn("w:trPr"))
    if trPr is not None:
        th = trPr.find(qn("w:tblHeader"))
        if th is not None and th.get(qn("w:val")) not in ("0", "false", "off"):
            return True
    populated = [c for c in row.cells if (c.text or "").strip()]
    return bool(populated) and all(_cell_text_is_bold(c) for c in populated)


_FIGURE_LABEL_RE = re.compile(
    r"^(figure|fig\.?|chart|graph|photo|photograph|image|map|exhibit|diagram|illustration|plate)"
    r"\s*[0-9IVXivx]+[a-z]?\s*[.:)\-–—]?\s+\S",
    re.IGNORECASE,
)


def _is_figure_label(text: str) -> bool:
    return bool(_FIGURE_LABEL_RE.match(" ".join((text or "").split())))


def _image_context_for_paragraph(paragraph, own_text: str) -> Optional[Tuple[str, str]]:
    """``(kind, text)`` describing a picture in ``paragraph``, or None.

    kind ``"caption"`` — a REAL caption, i.e. text a human wrote *as the
    picture's caption*: Word's Caption-styled paragraph right after (or
    before) the picture, or a numbered figure label ("Figure 2: Revenue by
    region") in or next to it.

    kind ``"nearby"`` — other human text near the picture (its own sentence,
    or the nearest preceding paragraph with at least three words). This is
    context, NOT a description: it is often a form placeholder, a list item
    or the next section's first line, so it is recorded under a different
    key and an alt-text provider can refuse to turn it into alt text.

    Never a filename.
    """
    p_el = paragraph._p  # noqa: SLF001
    nxt = p_el.getnext()
    prv = p_el.getprevious()
    # A caption ABOVE the picture only counts when it is not the caption
    # BELOW the previous picture ([pic1][Figure 1][pic2] — "Figure 1" is
    # pic1's, and pic2 has none).
    prv_ok = prv is not None and prv.tag == _W_P and not (
        prv.getprevious() is not None and prv.getprevious().tag == _W_P
        and next(iter_paragraph_drawings(prv.getprevious()), None) is not None
    )
    candidates = [s for s in (nxt, prv if prv_ok else None) if s is not None and s.tag == _W_P]
    for sib in candidates:
        cap = _paragraph_caption_text(sib)
        if cap:
            return ("caption", " ".join(cap.split())[:200])
    own = " ".join((own_text or "").split())
    if own and _is_figure_label(own):
        return ("caption", own[:200])
    for sib in candidates:
        txt = " ".join(paragraph_text(sib).split())
        if _is_figure_label(txt):
            return ("caption", txt[:200])
    if len(own.split()) >= 3:
        return ("nearby", own[:200])
    prev = prv
    hops = 0
    while prev is not None and hops < 4:
        if prev.tag == _W_P:
            txt = " ".join(paragraph_text(prev).split())
            if len(txt.split()) >= 3:
                return ("nearby", txt[:200])
            hops += 1
        prev = prev.getprevious()
    return None


def _docx_default_lang(doc) -> Optional[str]:
    """The document's run-default language from styles.xml, or None.

    Order: docDefaults/rPrDefault/rPr/w:lang@w:val, then the Normal style's
    rPr/w:lang. These are what Word writes and what assistive tech reads.
    """
    try:
        styles_el = doc.styles.element
    except Exception:
        return None
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is not None:
        lang = dd.find(f"./{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:lang')}")
        if lang is not None:
            val = (lang.get(qn("w:val")) or "").strip()
            if val:
                return val
    for st in styles_el.iterfind(qn("w:style")):
        if (st.get(qn("w:styleId")) or "").lower() == "normal":
            lang = st.find(f"./{qn('w:rPr')}/{qn('w:lang')}")
            if lang is not None:
                val = (lang.get(qn("w:val")) or "").strip()
                if val:
                    return val
    return None


_CAPTION_STYLE_VALS = {"caption"}  # Word built-in "Caption" paragraph style id


def _paragraph_caption_text(p_el) -> Optional[str]:
    """Return the text of ``p_el`` iff it is a ``Caption``-styled ``<w:p>``.

    A Word table caption is a paragraph styled "Caption" sitting adjacent to the
    table — that is the programmatic caption Word's Accessibility Checker and
    screen readers read (a plain sentence above a table is NOT associated with
    it). Returns ``None`` for anything that isn't a non-empty Caption paragraph.
    """
    if p_el is None or p_el.tag != qn("w:p"):
        return None
    pPr = p_el.find(qn("w:pPr"))
    if pPr is None:
        return None
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        return None
    if (pStyle.get(qn("w:val")) or "").strip().lower() not in _CAPTION_STYLE_VALS:
        return None
    text = "".join(t.text or "" for t in p_el.iterfind(f".//{qn('w:t')}")).strip()
    return text or None


def _docx_table_caption(table) -> Optional[str]:
    """The caption text for a Word table, if a Caption paragraph is adjacent.

    Word convention puts a table's caption immediately ABOVE it; some authors
    place it below. We check both immediate XML siblings of the ``<w:tbl>``.

    Guard for two adjacent tables ``[tableA][caption][tableB]``: a Caption
    paragraph that is itself immediately followed by another ``<w:tbl>`` is, by
    the caption-above convention, tableB's caption — so tableA must not absorb
    it via ``getnext`` (which would mask tableA's own missing caption).

    Scope: only top-level body tables (``doc.tables``) — like every other table
    fix, tables nested inside cells and tables in headers/footers are out of
    scope.
    """
    tbl = table._tbl
    prev_cap = _paragraph_caption_text(tbl.getprevious())
    if prev_cap:
        return prev_cap
    nxt = tbl.getnext()
    nxt_cap = _paragraph_caption_text(nxt)
    if nxt_cap and nxt is not None:
        after = nxt.getnext()
        if after is not None and after.tag == qn("w:tbl"):
            return None  # that caption belongs to the following table
    return nxt_cap


def _table_to_node(
    table,
    ids: _IdCounter,
    blobs: Optional[Dict[str, Tuple[Optional[str], Optional[str]]]] = None,
    reg: Optional[Callable[[str, Any], None]] = None,
    partname: str = "",
    nested: bool = False,
) -> TableNode:
    """TableNode for a Word table — its rows, cells, the links and pictures
    in each cell, and any table NESTED in a cell (as a child of that cell,
    minted in its own ``docx-n*`` id space so top-level ids never shift).

    Nested tables used to be dropped entirely, so TABLE_NESTED could never
    fire on a Word file and an inner data table's missing headers were never
    reported. Pictures in cells (the logo cell of a form) were never read.
    """
    reg = reg or _no_register
    blobs = blobs or {}
    pfx = "docx-n" if nested else "docx-"
    # Decide once whether row 0 is a header. Word marks real headers with
    # w:tblHeader or styles them bold; if neither AND the table is a clear data
    # grid (>=3 rows, >=2 cols), type row 0 as DATA so TABLE_MISSING_HEADERS can
    # fire. Small/ambiguous tables keep the legacy header assumption to avoid
    # false positives on layout tables.
    all_rows = list(iter_table_rows(table))
    n_cols = max((len(r.cells) for r in all_rows), default=0)
    looks_like_data_table = len(all_rows) >= 3 and n_cols >= 2
    first_is_header = _docx_row_is_header(all_rows[0]) if all_rows else False
    treat_row0_as_header = first_is_header or not looks_like_data_table

    # Merged cells repeat the same underlying <w:tc> across the grid; what is
    # inside it must only be emitted once. Keyed on the element itself (held
    # in the set, so its proxy stays alive) — keying on id() of a proxy that
    # was garbage-collected could match a DIFFERENT cell and silently drop
    # its links.
    seen_tcs: set = set()

    rows: List[TableRowNode] = []
    first_row_text: List[str] = []
    for row_index, row in enumerate(all_rows):
        cells: List[TableCellNode] = []
        for cell in row.cells:
            text = (cell.text or "").strip()
            if row_index == 0:
                first_row_text.append(" ".join(text.split()))
            is_header_cell = row_index == 0 and bool(text) and treat_row0_as_header
            cell_type = TableCellType.HEADER if is_header_cell else TableCellType.DATA
            cell_children: List[Any] = []
            tc = cell._tc  # noqa: SLF001
            if tc not in seen_tcs:
                seen_tcs.add(tc)
                # Hyperlinks (incl. fldSimple fields) and pictures inside the
                # cell get their own nodes so analysis/remediation reaches them.
                for cell_paragraph in cell.paragraphs:
                    cp_el = cell_paragraph._p  # noqa: SLF001
                    cp_text = paragraph_text(cp_el)
                    cell_children.extend(
                        _link_nodes_from_p(
                            cp_el, cell_paragraph.part, ids, reg=reg, context_text=cp_text,
                            extra_props={"docx_story": "table", "docx_part": partname},
                        )
                    )
                    cell_children.extend(
                        _image_nodes_from_p(
                            cp_el, blobs, ids, "docx-cimg", reg, story="table", partname=partname,
                            context_fn=lambda _p=cell_paragraph, _t=cp_text: _image_context_for_paragraph(_p, _t),
                        )
                    )
                for inner_el in _iter_sdt_aware(tc, _W_TBL):
                    from docx.table import Table as _Table

                    cell_children.append(
                        _table_to_node(_Table(inner_el, cell), ids, blobs, reg, partname, nested=True)
                    )
            cell_id = ids(f"{pfx}cell")
            reg(cell_id, (row, cell))
            cells.append(
                TableCellNode(
                    id=cell_id,
                    cell_type=cell_type,
                    header_scope=TableHeaderScope.COLUMN if cell_type == TableCellType.HEADER else TableHeaderScope.NONE,
                    content=NodeContent(kind=ContentKind.TEXT, text=text or " "),
                    metadata=NodeMetadata(source_format="docx"),
                    children=cell_children,
                    accessibility_flags=[],
                )
            )
        row_id = ids(f"{pfx}row")
        reg(row_id, row)
        rows.append(
            TableRowNode(
                id=row_id,
                content=NodeContent(kind=ContentKind.NONE),
                metadata=NodeMetadata(source_format="docx"),
                children=cells,
                accessibility_flags=[],
            )
        )
    caption = _docx_table_caption(table)
    table_props: Dict[str, Any] = {"caption": caption} if caption else {}
    if nested:
        table_props["nested_table"] = True
    # Location: the caption if the table has one, else its first row — the
    # words a person scanning the document recognises the table by.
    snippet = caption or " | ".join(t for t in first_row_text if t)
    if snippet:
        table_props["snippet"] = _snippet_around(snippet, "")
    table_id = ids(f"{pfx}table")
    reg(table_id, table)
    return TableNode(
        id=table_id,
        content=NodeContent(kind=ContentKind.NONE),
        metadata=NodeMetadata(source_format="docx", properties=table_props),
        children=rows,
        accessibility_flags=[],
    )


def _invalid_link_reason(target: str) -> str:
    value = (target or "").strip()
    if not value:
        return "missing_target"
    if value.startswith("#"):
        return ""
    lowered = value.lower()
    if lowered in {"http://", "https://", "www.", "mailto:"}:
        return "placeholder_target"
    parsed = urlparse(value)
    if parsed.scheme in {"http", "https", "mailto"}:
        if parsed.scheme in {"http", "https"} and not parsed.netloc:
            return "missing_host"
        return ""
    if parsed.scheme == "" and parsed.path:
        return ""
    return "unsupported_scheme"
