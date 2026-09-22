"""PPTX extraction helpers used by scan/fix routes.

Same dual-API pattern as :mod:`app.parsers.docx_parser`: a legacy
dict-returning :meth:`PPTXParser.parse` plus a new
:meth:`PPTXParser.parse_to_tree` that emits an :class:`AccessibilityTree`.
"""

from __future__ import annotations

import base64
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import urlparse

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

try:
    from pptx.enum.dml import MSO_FILL
except Exception:  # pragma: no cover - defensive
    MSO_FILL = None

from app.parsers.document_id import derive_document_id
from app.models.accessibility import (
    AccessibilityTree,
    ContentKind,
    DocumentNode,
    HeadingNode,
    ImageNode,
    LinkNode,
    NodeContent,
    NodeMetadata,
    ParagraphNode,
    ParserResult,
    SectionNode,
    TableCellNode,
    TableCellType,
    TableHeaderScope,
    TableNode,
    TableRowNode,
)


# Decided once per process: is there a vision provider that could ever read
# inlined image bytes? Under the heuristic provider nothing consumes them, and
# base64-inflating every image into the tree cost tens of MB per request.
try:
    from app.ai.semantic_inference import vision_provider_configured as _vpc

    _WANT_IMAGE_BYTES = bool(_vpc())
except Exception:  # pragma: no cover - never let the AI module break parsing
    _WANT_IMAGE_BYTES = True


class PPTXParser:
    def parse(self, file_path: str) -> Dict[str, object]:
        prs = Presentation(file_path)
        slide_details: List[Dict[str, object]] = []
        slides_missing_titles: List[int] = []
        total_images = 0
        missing_alt = 0
        total_tables = 0
        tables_missing_headers: List[Dict[str, object]] = []
        table_header_scope_flags: List[Dict[str, object]] = []
        hyperlink_count = 0
        generic_links: List[Dict[str, object]] = []
        invalid_links: List[Dict[str, object]] = []
        reading_order_warnings: List[Dict[str, object]] = []
        generic_link_labels = {
            "click here",
            "here",
            "read more",
            "learn more",
            "more",
            "link",
            "this",
        }
        generic_headers = {"column", "column 1", "column 2", "header", "n/a", "na", "value"}

        for idx, slide in enumerate(prs.slides, start=1):
            title = slide.shapes.title.text.strip() if slide.shapes.title and slide.shapes.title.text else ""
            if not title:
                slides_missing_titles.append(idx)
            slide_images = 0
            slide_tables = 0
            shape_positions: List[tuple[float, float]] = []
            for shape in _iter_shapes_recursive(slide.shapes):
                left = float(getattr(shape, "left", 0) or 0)
                top = float(getattr(shape, "top", 0) or 0)
                shape_positions.append((top, left))
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    slide_images += 1
                    alt_text = _shape_descr(shape)
                    if not alt_text:
                        missing_alt += 1
                if shape.has_table:
                    slide_tables += 1
                    table = shape.table
                    if len(table.rows) == 0:
                        tables_missing_headers.append({"slide": idx, "table": slide_tables})
                    else:
                        first_row = table.rows[0]
                        header_text = [(cell.text or "").strip() for cell in first_row.cells]
                        if not any(header_text):
                            tables_missing_headers.append({"slide": idx, "table": slide_tables})
                        else:
                            normalized_headers = [text.lower() for text in header_text if text]
                            unique_headers = set(normalized_headers)
                            if len(header_text) > 1 and len(unique_headers) <= 1:
                                table_header_scope_flags.append(
                                    {
                                        "slide": idx,
                                        "table": slide_tables,
                                        "reason": "duplicate_or_single_header_label",
                                        "headers": header_text[:10],
                                    }
                                )
                            elif any(text.lower() in generic_headers for text in header_text if text):
                                table_header_scope_flags.append(
                                    {
                                        "slide": idx,
                                        "table": slide_tables,
                                        "reason": "generic_header_labels",
                                        "headers": header_text[:10],
                                    }
                                )
                link_text = (getattr(shape, "text", None) or "").strip()
                shape_link = None
                try:
                    shape_link = getattr(shape.click_action.hyperlink, "address", None)
                except Exception:
                    shape_link = None
                if shape_link:
                    hyperlink_count += 1
                    if link_text and link_text.lower().strip() in generic_link_labels:
                        generic_links.append({"slide": idx, "text": link_text})
                    invalid_reason = _invalid_link_reason(str(shape_link))
                    if invalid_reason:
                        invalid_links.append(
                            {"slide": idx, "text": link_text or "(shape link)", "target": str(shape_link), "reason": invalid_reason}
                        )
                if hasattr(shape, "text_frame") and shape.text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run_link = getattr(getattr(run, "hyperlink", None), "address", None)
                            if not run_link:
                                continue
                            hyperlink_count += 1
                            run_text = (run.text or "").strip()
                            if run_text and run_text.lower().strip() in generic_link_labels:
                                generic_links.append({"slide": idx, "text": run_text})
                            invalid_reason = _invalid_link_reason(str(run_link))
                            if invalid_reason:
                                invalid_links.append(
                                    {"slide": idx, "text": run_text or "(run link)", "target": str(run_link), "reason": invalid_reason}
                                )
            total_images += slide_images
            total_tables += slide_tables
            if len(shape_positions) > 2:
                disorder = 0
                for i in range(1, len(shape_positions)):
                    if shape_positions[i][0] < shape_positions[i - 1][0]:
                        disorder += 1
                if disorder > 0:
                    reading_order_warnings.append({"slide": idx, "disorderCount": disorder})
            slide_details.append(
                {
                    "slide": idx,
                    "title": title,
                    "images": slide_images,
                    "tables": slide_tables,
                }
            )

        title = (prs.core_properties.title or "").strip()
        language = (getattr(prs.core_properties, "language", None) or "").strip()
        headings = [{"level": 1, "text": s["title"], "slide": s["slide"]} for s in slide_details if s["title"]]
        return {
            "documentType": "pptx",
            "title": title,
            "language": language,
            "slideCount": len(prs.slides),
            "slides": slide_details,
            "slidesMissingTitles": slides_missing_titles,
            "headings": headings,
            "imageCount": total_images,
            "missingAltCount": missing_alt,
            "hyperlinkCount": hyperlink_count,
            "genericLinks": generic_links,
            "invalidLinks": invalid_links,
            "tables": total_tables,
            "tablesMissingHeaders": tables_missing_headers,
            "tableHeaderScopeFlags": table_header_scope_flags,
            "readingOrderWarnings": reading_order_warnings,
        }


    def parse_to_tree(self, file_path: str) -> ParserResult:
        path = Path(file_path)
        prs = Presentation(file_path)
        theme_colors = _pptx_theme_colors(file_path)
        title = (prs.core_properties.title or "").strip()
        language = (getattr(prs.core_properties, "language", None) or "").strip()
        properties: Dict[str, Any] = {"filename": path.name}
        if title:
            properties["title"] = title
        root = DocumentNode(
            id="doc-1",
            content=NodeContent(kind=ContentKind.NONE),
            metadata=NodeMetadata(
                language=language or None,
                source_format="pptx",
                properties=properties,
            ),
            children=[],
            accessibility_flags=[],
        )
        ids = _IdCounter()
        slides_missing_titles = 0
        # Slide size in points, for the location contract (so a finding can be
        # drawn where it is on the slide). python-pptx answers None for a deck
        # without <p:sldSz>; PowerPoint then uses the 10in x 7.5in default.
        try:
            _sw = int(prs.slide_width or 9144000)
            _sh = int(prs.slide_height or 6858000)
        except (TypeError, ValueError):
            _sw, _sh = 9144000, 6858000
        page_size = [round(_sw / _EMU_PER_PT, 2), round(_sh / _EMU_PER_PT, 2)]

        for slide_index, slide in enumerate(prs.slides, start=1):
            slide_title = ""
            title_shape = None
            try:
                title_shape = slide.shapes.title
                if title_shape is not None and title_shape.text:
                    slide_title = title_shape.text.strip()
            except Exception:
                title_shape = None
                slide_title = ""
            if not slide_title:
                slides_missing_titles += 1

            # python-pptx builds a fresh proxy per access, so identity (`is`)
            # is unreliable — match the title placeholder by shape id.
            title_shape_id = None
            if slide_title and title_shape is not None:
                try:
                    title_shape_id = title_shape.shape_id
                except Exception:
                    title_shape_id = None

            section_props: dict = {"slide_number": slide_index, "page_size": list(page_size)}
            if not slide_title:
                # Per-slide marker so the analyzer can flag THIS slide (the
                # issue count then reflects how many slides lack titles).
                section_props["missing_title"] = True
            if _slide_reading_order_inverted(slide, title_shape_id):
                # Tree order != visual order: a screen reader would read this
                # slide bottom-then-top. Flagged by ReadingOrderAnalyzer.
                section_props["reading_order_inverted"] = True
            section = SectionNode(
                id=ids(f"slide-{slide_index}-section"),
                content=NodeContent(kind=ContentKind.TEXT, text=slide_title or f"Slide {slide_index}"),
                metadata=NodeMetadata(
                    page=slide_index,
                    source_format="pptx",
                    properties=section_props,
                ),
                children=[],
                accessibility_flags=[],
            )
            root.children.append(section)

            if slide_title:
                _h_props: dict = {
                    "order_hint": (
                        float(getattr(title_shape, "top", 0) or 0),
                        float(getattr(title_shape, "left", 0) or 0),
                    )
                }
                _h_props.update(_location_props(_abs_rect_emu(title_shape), page_size))
                _h_cc = _contrast_props_for_shape(title_shape, theme_colors)
                if _h_cc:
                    _h_props.update(_h_cc)
                section.children.append(
                    HeadingNode(
                        id=ids(f"slide-{slide_index}-h"),
                        level=1,
                        content=NodeContent(kind=ContentKind.TEXT, text=slide_title),
                        metadata=NodeMetadata(
                            page=slide_index,
                            source_format="pptx",
                            properties=_h_props,
                        ),
                        children=[],
                        accessibility_flags=[],
                    )
                )

            for shape in _iter_shapes_recursive(slide.shapes):
                # The title placeholder already became the slide's HeadingNode;
                # emitting its text again as a ParagraphNode would double-count
                # it (its hyperlinks, if any, are still collected below).
                is_title_shape = (
                    title_shape_id is not None
                    and getattr(shape, "shape_id", None) == title_shape_id
                )
                rect = _abs_rect_emu(shape)
                if rect is not None:
                    # Slide coordinates, also for a shape inside a moved or
                    # scaled group (python-pptx reports those in the group's
                    # own child space, so "topmost" compared apples/oranges).
                    left, top = float(rect[0]), float(rect[1])
                else:
                    top = float(getattr(shape, "top", 0) or 0)
                    left = float(getattr(shape, "left", 0) or 0)
                loc = _location_props(rect, page_size)
                shape_meta_props = {"order_hint": (top, left), **loc}
                if _image_kind(shape):
                    section.children.append(_picture_to_image_node(shape, slide_index, ids, shape_meta_props))
                    continue
                if shape.has_table:
                    table_node = _table_to_node_pptx(shape.table, slide_index, ids)
                    if loc:
                        table_node.metadata.properties = {**(table_node.metadata.properties or {}), **loc}
                    section.children.append(table_node)
                    continue
                if hasattr(shape, "text_frame") and shape.text_frame:
                    text = (shape.text or "").strip()
                    if not text:
                        continue
                    if not is_title_shape:
                        _props = dict(shape_meta_props)
                        _cc = _contrast_props_for_shape(shape, theme_colors)
                        if _cc:
                            _props.update(_cc)
                        pid = ids(f"slide-{slide_index}-p")
                        fake_kind = _fake_list_kind_for_shape(shape, text)
                        if fake_kind:
                            # Typed "- item" lines in a plain text box (no real
                            # bullet formatting). ListStructureAnalyzer flags it;
                            # FIX_LIST_STRUCTURE + the pptx writer convert the
                            # lines to real a:buChar/a:buAutoNum bullets.
                            _props["fake_list_run_ids"] = [pid]
                            _props["fake_list_kind"] = fake_kind
                        section.children.append(
                            ParagraphNode(
                                id=pid,
                                content=NodeContent(kind=ContentKind.TEXT, text=text),
                                metadata=NodeMetadata(
                                    page=slide_index,
                                    source_format="pptx",
                                    properties=_props,
                                ),
                                children=[],
                                accessibility_flags=[],
                            )
                        )
                    # Hyperlinks within runs get their own LinkNode (adjacent
                    # same-address runs are coalesced so the text isn't doubled).
                    for paragraph in shape.text_frame.paragraphs:
                        for href, runs in _iter_hyperlink_groups(paragraph):
                            text = "".join((r.text or "") for r in runs).strip() or "link"
                            section.children.append(
                                LinkNode(
                                    id=ids(f"slide-{slide_index}-link"),
                                    target=str(href),
                                    content=NodeContent(kind=ContentKind.TEXT, text=text),
                                    # The link's box is the shape it sits in.
                                    metadata=NodeMetadata(page=slide_index, source_format="pptx", properties=dict(loc)),
                                    children=[],
                                    accessibility_flags=[],
                                )
                            )

        # Record per-deck slide-title coverage for the analyzer.
        root.metadata.properties["slides_total"] = len(prs.slides)
        root.metadata.properties["slides_missing_titles"] = slides_missing_titles

        raw_metadata = {
            "filename": path.name,
            "title": title,
            "language": language,
            "slide_count": len(prs.slides),
        }
        return ParserResult(
            document_id=derive_document_id(path),
            format="pptx",
            tree=AccessibilityTree(root=root, metadata=raw_metadata),
            raw_metadata=raw_metadata,
        )


# ----- PPTX → AccessibilityTree helpers --------------------------------------


class _IdCounter:
    def __init__(self) -> None:
        self._counts: Dict[str, int] = {}

    def __call__(self, prefix: str) -> str:
        i = self._counts.get(prefix, 0) + 1
        self._counts[prefix] = i
        return f"{prefix}-{i}"


def _iter_shapes_recursive(shapes):
    """Yield shapes depth-first, descending into group shapes.

    A group is a container, not content, so we yield its children (recursively)
    rather than the group itself. Without this, pictures/tables nested in a
    group are invisible to the analyzer — a common real-world false negative
    ("no missing alt") for grouped images.

    ``parse_to_tree`` and the writer's ``_index_shapes_by_parser_id`` MUST use
    this same walker so the ids minted in iteration order stay aligned.
    """
    for shape in shapes:
        try:
            is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
        except Exception:
            is_group = False
        if is_group:
            try:
                yield from _iter_shapes_recursive(shape.shapes)
            except Exception:
                continue
        else:
            yield shape


_P_GRP_SP_TAG = "{http://schemas.openxmlformats.org/presentationml/2006/main}grpSp"
_P_GRP_SP_PR_TAG = "{http://schemas.openxmlformats.org/presentationml/2006/main}grpSpPr"
_EMU_PER_PT = 12700.0


def _abs_rect_emu(shape) -> Optional[Tuple[int, int, int, int]]:
    """``(left, top, width, height)`` of ``shape`` in SLIDE coordinates (EMU).

    python-pptx reports a grouped shape's position in its group's CHILD
    coordinate space (``a:chOff``/``a:chExt``), which is only the slide's own
    space when the group was never moved or resized. Map the rectangle out
    through every enclosing group's transform. Placeholders get their
    inherited layout position from python-pptx. None when the shape has no
    usable geometry.
    """
    try:
        left, top = int(shape.left), int(shape.top)
        width, height = int(shape.width), int(shape.height)
    except (TypeError, ValueError, AttributeError):
        return None
    x0, y0, x1, y1 = float(left), float(top), float(left + width), float(top + height)
    try:
        parent = shape._element.getparent()  # noqa: SLF001
    except Exception:
        parent = None
    while parent is not None and parent.tag == _P_GRP_SP_TAG:
        xfrm = parent.find(f"{_P_GRP_SP_PR_TAG}/{_A_NS}xfrm")
        if xfrm is None:
            break
        off, ext = xfrm.find(f"{_A_NS}off"), xfrm.find(f"{_A_NS}ext")
        ch_off, ch_ext = xfrm.find(f"{_A_NS}chOff"), xfrm.find(f"{_A_NS}chExt")
        if off is None or ext is None or ch_off is None or ch_ext is None:
            break
        try:
            ox, oy = float(off.get("x")), float(off.get("y"))
            ex, ey = float(ext.get("cx")), float(ext.get("cy"))
            cx, cy = float(ch_off.get("x")), float(ch_off.get("y"))
            cex, cey = float(ch_ext.get("cx")), float(ch_ext.get("cy"))
        except (TypeError, ValueError):
            break
        sx = ex / cex if cex else 1.0
        sy = ey / cey if cey else 1.0
        x0, x1 = ox + (x0 - cx) * sx, ox + (x1 - cx) * sx
        y0, y1 = oy + (y0 - cy) * sy, oy + (y1 - cy) * sy
        parent = parent.getparent()
    return int(round(x0)), int(round(y0)), int(round(x1 - x0)), int(round(y1 - y0))


def _location_props(rect: Optional[Tuple[int, int, int, int]], page_size: Optional[List[float]]) -> Dict[str, Any]:
    """``bbox``/``page_size`` for the shared location contract: points, origin
    at the slide's BOTTOM-left (the same convention as a PDF page), so one
    highlight overlay serves both formats."""
    if rect is None or page_size is None:
        return {}
    left, top, width, height = rect
    slide_h = page_size[1] * _EMU_PER_PT
    return {
        "bbox": [
            round(left / _EMU_PER_PT, 2),
            round((slide_h - (top + height)) / _EMU_PER_PT, 2),
            round((left + width) / _EMU_PER_PT, 2),
            round((slide_h - top) / _EMU_PER_PT, 2),
        ],
        "page_size": list(page_size),
    }


_P_PIC_TAG = "{http://schemas.openxmlformats.org/presentationml/2006/main}pic"
_P_GRAPHIC_FRAME_TAG = "{http://schemas.openxmlformats.org/presentationml/2006/main}graphicFrame"
_GRAPHIC_DATA_TAG = "{http://schemas.openxmlformats.org/drawingml/2006/main}graphicData"
_IMAGE_LIKE_GRAPHIC_URIS = {
    "http://schemas.openxmlformats.org/drawingml/2006/chart": "chart",
    "http://schemas.openxmlformats.org/drawingml/2006/diagram": "smartart",
}
_DECORATIVE_EXT_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"


def _image_kind(shape) -> Optional[str]:
    """What kind of image a shape is, or None when it is not one.

    ``shape_type == PICTURE`` alone missed most real-deck images: a picture
    dropped into a content/picture PLACEHOLDER (the "insert picture" icon on
    a layout — the usual way pictures get onto slides) is a ``p:pic`` whose
    shape_type is PLACEHOLDER, and a chart or SmartArt graphic needs alt text
    just as much. So: every ``p:pic`` (pictures, filled picture placeholders,
    video/audio posters) and every chart/SmartArt graphic frame.

    ``parse_to_tree`` and the writer's ``_index_shapes_by_parser_id`` MUST
    both use this so the minted image ids stay aligned.
    """
    try:
        el = shape._element  # noqa: SLF001
    except Exception:
        return None
    tag = getattr(el, "tag", None)
    if tag == _P_PIC_TAG:
        return "picture"
    if tag == _P_GRAPHIC_FRAME_TAG:
        gd = el.find(f".//{_GRAPHIC_DATA_TAG}")
        if gd is not None:
            return _IMAGE_LIKE_GRAPHIC_URIS.get(gd.get("uri") or "")
    return None


def _shape_marked_decorative(shape) -> bool:
    """True when PowerPoint's "Mark as decorative" is set on the shape.

    That is the Office 2017 ``adec:decorative`` extension under ``p:cNvPr`` —
    the same mechanism the writer uses. A decorative image is skipped by
    screen readers, so it needs no alt text; flagging it as missing/vague alt
    (the ``descr`` PowerPoint leaves behind is often "image.png") was a false
    positive, including on our own output after a mark-decorative fix.
    """
    try:
        for nv in shape._element.iter():  # noqa: SLF001
            tag = nv.tag
            if isinstance(tag, str) and tag.endswith("}cNvPr"):
                for ext in nv.iter(f"{_A_NS}ext"):
                    if ext.get("uri") != _DECORATIVE_EXT_URI:
                        continue
                    for child in ext:
                        if isinstance(child.tag, str) and child.tag.endswith("}decorative"):
                            return (child.get("val") or "").strip().lower() in ("1", "true")
                return False
    except Exception:
        return False
    return False


def _shape_bg_hex(shape) -> Optional[str]:
    """Explicit solid-fill background colour of a shape, or None.

    PPTX backgrounds are often theme/slide-level and not reliably knowable, so
    we only report a background when the shape itself has an explicit solid RGB
    fill. Text whose background we can't determine is left unevaluated (we never
    assume white on a slide — that would false-flag light text on dark slides).
    """
    if MSO_FILL is None:
        return None
    try:
        fill = shape.fill
        if fill.type == MSO_FILL.SOLID:
            rgb = fill.fore_color.rgb
            if rgb is not None:
                return str(rgb)
    except Exception:
        pass
    return None


_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


_RO_MARGIN_EMU = 182880  # 0.2 inch — vertical gap before "stacked" counts


def _slide_reading_order_inverted(slide, title_shape_id) -> bool:
    """True when the slide's text shapes are stacked vertically but appear in
    REVERSE order in the shape tree — screen readers follow tree order, so the
    slide reads bottom-then-top (WCAG 1.3.2).

    Precision guards: top-level shapes only (group children carry
    group-relative geometry); substantial text only (>=12 chars, so page
    numbers and tiny labels never count); and only NON-overlapping vertical
    separation flags — side-by-side columns overlap vertically, and column
    order is a legitimate authoring choice."""

    geo: List[Tuple[float, float]] = []
    try:
        shapes = list(slide.shapes)
    except Exception:
        return False
    for sh in shapes:
        if title_shape_id is not None and getattr(sh, "shape_id", None) == title_shape_id:
            continue
        if not (hasattr(sh, "text_frame") and getattr(sh, "has_text_frame", False)):
            continue
        try:
            txt = (sh.text or "").strip()
        except Exception:
            continue
        if len(txt) < 12:
            continue
        try:
            top = float(sh.top or 0)
            height = float(sh.height or 0)
        except (TypeError, ValueError):
            continue
        if height <= 0:
            continue
        geo.append((top, height))
    for i in range(len(geo)):
        for j in range(i + 1, len(geo)):
            # Earlier-in-tree shape sits ENTIRELY below a later one.
            if geo[i][0] > geo[j][0] + geo[j][1] + _RO_MARGIN_EMU:
                return True
    return False


def _paragraph_has_real_bullet(paragraph) -> bool:
    """True when the a:pPr explicitly carries bullet formatting."""
    pPr = paragraph._p.find(f"{_A_NS}pPr")
    if pPr is None:
        return False
    return (
        pPr.find(f"{_A_NS}buChar") is not None
        or pPr.find(f"{_A_NS}buAutoNum") is not None
    )


def _fake_list_kind_for_shape(shape, text: str) -> Optional[str]:
    """Return "bullet"/"decimal" when a plain TEXT BOX contains a typed
    fake-list run (>=2 consecutive lines like "- item" / sequential "1. item").

    Scoped to non-placeholder shapes: placeholders inherit bullet formatting
    from the layout/master (invisible at paragraph level), so typed markers
    there can't be judged safely. Paragraphs that already carry explicit
    bullet formatting never count.
    """

    from app.parsers.docx_parser import _fake_list_signature  # shared predicate

    try:
        if getattr(shape, "is_placeholder", False):
            return None
    except Exception:
        return None
    try:
        paragraphs = list(shape.text_frame.paragraphs)
    except Exception:
        return None

    run_len = 0
    run_kind: Optional[str] = None
    run_char: Optional[str] = None
    next_ord: Optional[int] = None
    for p in paragraphs:
        line = "".join((r.text or "") for r in p.runs).strip() or (p.text or "").strip()
        sig = _fake_list_signature(line) if line else None
        if sig is None or _paragraph_has_real_bullet(p):
            if run_len >= 2:
                return run_kind
            run_len = 0
            run_kind = run_char = next_ord = None
            continue
        kind, char, ordinal = sig
        if run_len and (kind != run_kind or (kind == "bullet" and char != run_char)):
            if run_len >= 2:
                return run_kind
            run_len = 0
            run_kind = run_char = next_ord = None
        if kind == "decimal":
            expected = 1 if not run_len else next_ord
            if ordinal != expected:
                if run_len >= 2:
                    return run_kind
                run_len = 0
                run_kind = run_char = next_ord = None
                if ordinal != 1:
                    continue
            next_ord = (ordinal or 0) + 1
        if not run_len:
            run_kind, run_char = kind, char
        run_len += 1
    return run_kind if run_len >= 2 else None

# a:schemeClr val -> theme clrScheme element name (standard colour map).
_PPTX_SCHEME_MAP = {
    "tx1": "dk1", "dk1": "dk1", "bg1": "lt1", "lt1": "lt1",
    "tx2": "dk2", "dk2": "dk2", "bg2": "lt2", "lt2": "lt2",
    "accent1": "accent1", "accent2": "accent2", "accent3": "accent3",
    "accent4": "accent4", "accent5": "accent5", "accent6": "accent6",
    "hlink": "hlink", "folHlink": "folHlink",
}


def _pptx_theme_colors(file_path: str) -> Dict[str, str]:
    """Map theme colour-scheme names (dk1, lt1, accent1…) to RGB hex."""
    out: Dict[str, str] = {}
    try:
        with zipfile.ZipFile(file_path) as z:
            names = [n for n in z.namelist() if n.startswith("ppt/theme/theme") and n.endswith(".xml")]
            if not names:
                return {}
            xml = z.read(sorted(names)[0])
        root = etree.fromstring(xml)
        scheme = root.find(f".//{_A_NS}clrScheme")
        if scheme is None:
            return {}
        for child in scheme:
            name = etree.QName(child).localname
            srgb = child.find(f"{_A_NS}srgbClr")
            sysclr = child.find(f"{_A_NS}sysClr")
            if srgb is not None and srgb.get("val"):
                out[name] = srgb.get("val").upper()
            elif sysclr is not None and sysclr.get("lastClr"):
                out[name] = sysclr.get("lastClr").upper()
    except Exception:
        return {}
    return out


def _apply_pptx_color_mods(hex_color: str, scheme_el) -> str:
    """Apply a:lumMod / a:lumOff / a:tint / a:shade modifiers (RGB approximation)."""
    try:
        r = float(int(hex_color[0:2], 16))
        g = float(int(hex_color[2:4], 16))
        b = float(int(hex_color[4:6], 16))
    except (ValueError, IndexError):
        return hex_color
    for mod in scheme_el:
        tag = etree.QName(mod).localname
        try:
            val = int(mod.get("val", "0")) / 100000.0
        except (TypeError, ValueError):
            continue
        if tag == "lumMod":
            r, g, b = r * val, g * val, b * val
        elif tag == "lumOff":
            r, g, b = r + 255 * val, g + 255 * val, b + 255 * val
        elif tag == "tint":  # toward white
            r, g, b = r + (255 - r) * val, g + (255 - g) * val, b + (255 - b) * val
        elif tag == "shade":  # toward black
            r, g, b = r * val, g * val, b * val
    return f"{max(0, min(255, round(r))):02X}{max(0, min(255, round(g))):02X}{max(0, min(255, round(b))):02X}"


def _run_color_hex_pptx(run, theme_colors: Dict[str, str]) -> Optional[str]:
    """Run text colour as RGB hex: explicit sRGB, or a resolved theme scheme
    colour (with lum/tint/shade modifiers). ``None`` if unknown."""
    try:
        rgb = run.font.color.rgb
    except Exception:
        rgb = None
    if rgb is not None:
        return str(rgb)
    if not theme_colors:
        return None
    try:
        scheme_el = run._r.find(f"{_A_NS}rPr/{_A_NS}solidFill/{_A_NS}schemeClr")
        if scheme_el is None:
            return None
        val = scheme_el.get("val")
        base = theme_colors.get(_PPTX_SCHEME_MAP.get(val, val))
        if not base:
            return None
        return _apply_pptx_color_mods(base, scheme_el)
    except Exception:
        return None


def _contrast_props_for_shape(shape, theme_colors: Optional[Dict[str, str]] = None) -> Dict[str, Any]:
    """Contrast inputs for a text shape: only when BOTH a run colour (explicit or
    resolved theme) and an explicit shape background are known."""
    theme_colors = theme_colors or {}
    bg = _shape_bg_hex(shape)
    if not bg:
        return {}
    colors: List[Dict[str, Any]] = []
    try:
        tf = shape.text_frame
    except Exception:
        return {}
    for para in tf.paragraphs:
        for run in para.runs:
            if not (run.text or "").strip():
                continue
            hex6 = _run_color_hex_pptx(run, theme_colors)
            if not hex6:
                continue
            size_pt = None
            try:
                if run.font.size is not None:
                    size_pt = float(run.font.size.pt)
            except Exception:
                size_pt = None
            colors.append({"c": hex6, "sz": size_pt, "b": bool(run.font.bold) if run.font.bold is not None else False})
    if not colors:
        return {}
    return {"explicit_text_colors": colors, "bg_color": bg}


def _iter_hyperlink_groups(paragraph):
    """Yield ``(href, [runs])`` for each hyperlink in a paragraph.

    Consecutive runs sharing the same address are coalesced into ONE group —
    PowerPoint frequently splits a single visual hyperlink across multiple runs,
    and treating each as its own link would duplicate the rewritten text. Both
    the parser and the writer use this so link ids stay aligned.
    """
    groups: List = []
    prev_href = None
    for run in getattr(paragraph, "runs", []) or []:
        try:
            # python-pptx's ``run.hyperlink`` does get_or_add_rPr(): merely
            # LOOKING for a link stamped an empty <a:rPr/> into every run the
            # writer indexed, so slides nobody touched came back changed.
            r_pr = run._r.rPr  # noqa: SLF001
            if r_pr is None or r_pr.find(f"{_A_NS}hlinkClick") is None:
                href = None
            else:
                href = run.hyperlink.address
        except Exception:
            href = None
        if href:
            if href == prev_href and groups:
                groups[-1][1].append(run)
            else:
                groups.append((href, [run]))
        prev_href = href
    return groups


def _shape_descr(shape) -> str:
    """Read a shape's alt text from <p:cNvPr @descr>.

    python-pptx's ``Picture`` does not expose ``alternative_text`` in current
    versions, so reading that attribute returns nothing and every image looks
    like it is missing alt text. The real value lives on the ``cNvPr`` element.
    """
    try:
        for nv in shape._element.iter():
            tag = nv.tag
            if isinstance(tag, str) and tag.endswith("}cNvPr"):
                return (nv.get("descr") or "").strip()
    except Exception:
        return ""
    return ""


def _chart_title(shape) -> Optional[str]:
    """A chart's own title text — the author's caption for it — or None."""
    try:
        chart = shape.chart
        if chart.has_title:
            text = " ".join((chart.chart_title.text_frame.text or "").split())
            return text[:200] or None
    except Exception:
        return None
    return None


def _picture_to_image_node(shape, slide_index: int, ids: "_IdCounter", extra_props: Dict[str, Any]) -> ImageNode:
    alt_text = _shape_descr(shape)
    kind = _image_kind(shape) or "picture"
    is_decorative = _shape_marked_decorative(shape)
    if is_decorative:
        # Screen readers skip a decorative image; whatever descr PowerPoint
        # left behind is never announced, so it is neither missing nor vague.
        alt_text = ""
    image_b64: Optional[str] = None
    image_mime: Optional[str] = None
    try:
        blob = shape.image.blob
        if _WANT_IMAGE_BYTES and blob and len(blob) <= 2_000_000:
            image_b64 = base64.b64encode(blob).decode("ascii")
            content_type = getattr(shape.image, "content_type", None)
            image_mime = str(content_type) if content_type else "image/png"
    except Exception:
        pass

    properties: Dict[str, Any] = dict(extra_props)
    if image_b64:
        properties["image_b64"] = image_b64
        properties["image_mime"] = image_mime or "image/png"
    properties["shape_id"] = getattr(shape, "shape_id", None)
    if kind != "picture":
        properties["image_kind"] = kind
    if kind == "chart" and not is_decorative:
        title = _chart_title(shape)
        if title:
            properties["caption"] = title
    metadata = NodeMetadata(page=slide_index, source_format="pptx", properties=properties)
    if is_decorative and alt_text:
        return ImageNode.model_construct(
            id=ids(f"slide-{slide_index}-img"),
            node_type=ImageNode.type_value(),
            content=NodeContent(kind=ContentKind.NONE),
            metadata=metadata,
            children=[],
            accessibility_flags=[],
            is_decorative=True,
            alt_text=alt_text,
        )
    return ImageNode(
        id=ids(f"slide-{slide_index}-img"),
        content=NodeContent(kind=ContentKind.NONE),
        metadata=metadata,
        children=[],
        accessibility_flags=[],
        is_decorative=is_decorative,
        alt_text=alt_text or None,
    )


def _table_to_node_pptx(table, slide_index: int, ids: "_IdCounter") -> TableNode:
    rows: List[TableRowNode] = []
    row_objects = list(table.rows)
    # PowerPoint marks a header row with the "Header Row" band (the firstRow flag
    # on <a:tblPr>, exposed as Table.first_row). Type row 0 as a header only when
    # that band is on, OR the table is too small/ambiguous to be a clear data
    # grid. A data grid (>=3 rows, >=2 cols) with the header band OFF has no
    # header — type row 0 as DATA so TABLE_MISSING_HEADERS fires. (Mirrors the
    # DOCX rule; conservative — small/ambiguous tables keep the header assumption.)
    n_cols = max((len(r.cells) for r in row_objects), default=0)
    looks_like_data_table = len(row_objects) >= 3 and n_cols >= 2
    has_header_band = bool(getattr(table, "first_row", False))
    treat_row0_as_header = has_header_band or not looks_like_data_table
    for row_index, row in enumerate(row_objects):
        cells: List[TableCellNode] = []
        for cell in row.cells:
            text = (cell.text or "").strip()
            is_header_cell = row_index == 0 and bool(text) and treat_row0_as_header
            cell_type = TableCellType.HEADER if is_header_cell else TableCellType.DATA
            cells.append(
                TableCellNode(
                    id=ids(f"slide-{slide_index}-cell"),
                    cell_type=cell_type,
                    header_scope=TableHeaderScope.COLUMN if cell_type == TableCellType.HEADER else TableHeaderScope.NONE,
                    content=NodeContent(kind=ContentKind.TEXT, text=text or " "),
                    metadata=NodeMetadata(page=slide_index, source_format="pptx"),
                    children=[],
                    accessibility_flags=[],
                )
            )
        rows.append(
            TableRowNode(
                id=ids(f"slide-{slide_index}-row"),
                content=NodeContent(kind=ContentKind.NONE),
                metadata=NodeMetadata(page=slide_index, source_format="pptx"),
                children=cells,
                accessibility_flags=[],
            )
        )
    return TableNode(
        id=ids(f"slide-{slide_index}-table"),
        content=NodeContent(kind=ContentKind.NONE),
        metadata=NodeMetadata(page=slide_index, source_format="pptx"),
        children=rows,
        accessibility_flags=[],
    )


def _invalid_link_reason(target: str) -> str:
    value = (target or "").strip()
    if not value:
        return "missing_target"
    if value.startswith("#"):
        return ""
    lowered = value.lower()
    if lowered in {"http://", "https://", "www.", "mailto:"}:
        return "placeholder_target"
    parsed = urlparse(value)
    if parsed.scheme in {"http", "https", "mailto"}:
        if parsed.scheme in {"http", "https"} and not parsed.netloc:
            return "missing_host"
        return ""
    if parsed.scheme == "" and parsed.path:
        return ""
    return "unsupported_scheme"
