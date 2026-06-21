"""Smoke: FIX_CONTRAST for DOCX — recolour low-contrast runs to nearest AA shade.

Real python-docx round trip proving the auto-fix AND the honesty invariant:
the failing run's colour is rewritten to an AA-passing shade in the output
bytes, re-parsing clears LOW_CONTRAST_TEXT, a run that already passes (in the
same paragraph) is left untouched, and an approved recolour whose paragraph
can't be resolved is recorded in skipped (never credited/charged).

Usage:
    python -m app.devtools.smoke_docx_contrast
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_dcx_')}/s.db")

from docx import Document  # noqa: E402
from docx.shared import RGBColor  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.api.pipeline import _action_persists  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.docx_writer import write_remediated_docx  # noqa: E402

CONTRAST = "LOW_CONTRAST_TEXT"


def _build_docx(path: Path) -> None:
    doc = Document()
    # Paragraph 1: a single grey run (#999999 ~2.8:1 on white) — must be fixed.
    p1 = doc.add_paragraph()
    r = p1.add_run("Grey text that fails AA on white and must be auto-darkened.")
    r.font.color.rgb = RGBColor.from_string("999999")
    # Paragraph 2: TWO runs — one failing grey, one passing near-black. Only the
    # grey one may be recoloured (mixed-colour FP guard).
    p2 = doc.add_paragraph()
    bad = p2.add_run("Light grey fails. ")
    bad.font.color.rgb = RGBColor.from_string("AAAAAA")
    good = p2.add_run("Dark text already passes.")
    good.font.color.rgb = RGBColor.from_string("111111")
    doc.save(str(path))


def _flag_count(tree, code):
    n = 0

    def walk(node):
        nonlocal n
        n += sum(1 for f in node.accessibility_flags if f.code.value == code)
        for ch in node.children:
            walk(ch)

    walk(tree.root)
    return n


def _run_hexes(path: Path):
    """All explicit run colours in the doc, uppercased."""
    doc = Document(str(path))
    out = []
    for p in doc.paragraphs:
        for run in p.runs:
            try:
                rgb = run.font.color.rgb
            except Exception:
                rgb = None
            if rgb is not None:
                out.append(str(rgb).upper())
    return out


def main() -> int:
    failures = 0

    def check(name, cond, extra=""):
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="dcx_smoke_"))
    policy = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)

    check("honesty matrix: FIX_CONTRAST persists for docx", _action_persists("FIX_CONTRAST", "docx"))

    src = tmp / "contrast.docx"
    _build_docx(src)

    # ---- detection ----
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    n_flags = _flag_count(res.tree, CONTRAST)
    check("two low-contrast paragraphs flagged", n_flags == 2, f"got {n_flags}")

    # ---- remediate ----
    plans = plan_remediations(res.tree, policy)
    execs = execute_plans(res.tree, plans)
    fix_ok = [e for e in execs if e.action_code.value == "FIX_CONTRAST" and e.status.value == "success"]
    check("FIX_CONTRAST executed for both paragraphs", len(fix_ok) == 2, f"got {len(fix_ok)}")

    out = tmp / "contrast.fixed.docx"
    result = write_remediated_docx(src, res.tree, out)
    applied_contrast = [a for a in result["applied"] if a.get("action") == "FIX_CONTRAST"]
    check("writer applied FIX_CONTRAST to both paragraphs", len(applied_contrast) == 2, str(result["applied"]))

    hexes = _run_hexes(out)
    check("failing greys (#999999/#AAAAAA) gone from output", "999999" not in hexes and "AAAAAA" not in hexes, str(hexes))
    check("passing dark run (#111111) left untouched", "111111" in hexes, str(hexes))

    # ---- re-analysis: flag clears ----
    res2 = parse_to_tree(str(out))
    run_analyzers(res2.tree)
    check("re-parse: ALL low-contrast flags cleared", _flag_count(res2.tree, CONTRAST) == 0,
          f"got {_flag_count(res2.tree, CONTRAST)}")

    # ---- honesty: unresolved paragraph not credited ----
    res3 = parse_to_tree(str(src))
    run_analyzers(res3.tree)
    # Find a flagged node, keep its fix map, but give it a bogus id so the
    # writer's paragraph-id lookup misses.
    flagged = None

    def find(node):
        nonlocal flagged
        if flagged is None and any(f.code.value == CONTRAST for f in node.accessibility_flags):
            flagged = node
        for ch in node.children:
            find(ch)

    find(res3.tree.root)
    plans3 = plan_remediations(res3.tree, policy)
    execute_plans(res3.tree, plans3)
    flagged.id = "docx-p-DOESNOTEXIST"  # break the locator after execution
    out3 = tmp / "contrast.broken.docx"
    res_broken = write_remediated_docx(src, res3.tree, out3)
    applied_ids = {a.get("target_id") for a in res_broken["applied"] if a.get("action") == "FIX_CONTRAST"}
    skipped_reasons = {s.get("reason") for s in res_broken["skipped"]}
    check("unresolved paragraph NOT credited as FIX_CONTRAST", "docx-p-DOESNOTEXIST" not in applied_ids)
    check("unresolved paragraph recorded in skipped",
          any("contrast" in str(r) for r in skipped_reasons), str(skipped_reasons))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
