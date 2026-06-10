"""Smoke: malformed uploads fail CLEAN (4xx + JSON), never 500 — and the
analyze endpoint survives parallel load.

A public upload endpoint gets garbage all day: empty files, truncated
office files, renamed extensions, corrupted zip members. Every one of those
must produce a structured 4xx (no stack traces, no 500s, no traceback text
in the body), and a burst of concurrent valid analyzes must all succeed.

Usage:
    python -m app.devtools.smoke_malformed_inputs
"""

from __future__ import annotations

import concurrent.futures
import io
import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_mf_')}/s.db"

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.main import app

    client = TestClient(app, raise_server_exceptions=False)
    r = client.post(
        "/auth/sign-in",
        json={"email": "fuzz@example.com", "displayName": "Fuzz", "password": "fuzzpass123"},
    )
    auth = {"Authorization": f"Bearer {r.json()['token']}"}

    tmp = Path(tempfile.mkdtemp(prefix="fuzz_"))

    # A valid docx, used whole (concurrency) and truncated (corruption case).
    d = Document()
    d.core_properties.title = "Valid"
    d.add_heading("Valid Document", level=1)
    d.add_paragraph("Perfectly ordinary paragraph for the load test.")
    valid_docx = tmp / "valid.docx"
    d.save(str(valid_docx))
    valid_bytes = valid_docx.read_bytes()

    # A docx whose document.xml is corrupted XML (zip itself is fine).
    broken_xml = tmp / "broken_xml.docx"
    with zipfile.ZipFile(valid_docx) as zin, zipfile.ZipFile(broken_xml, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = data[: len(data) // 2] + b"<unclosed><<<garbage"
            zout.writestr(item, data)

    def post(name: str, payload: bytes, mime: str = DOCX_MIME):
        return client.post(
            "/pipeline/analyze",
            files={"file": (name, payload, mime)},
            headers=auth,
        )

    def clean_4xx(label: str, resp) -> None:
        check(f"{label}: 4xx not 5xx", 400 <= resp.status_code < 500, f"got {resp.status_code}")
        body = resp.text or ""
        check(f"{label}: structured JSON error", body.startswith("{") and "detail" in body, body[:120])
        check(f"{label}: no traceback leaked", "Traceback" not in body and "File \"" not in body)

    clean_4xx("empty .docx", post("empty.docx", b""))
    clean_4xx("garbage .docx", post("garbage.docx", b"\x00\x01\x02 not a zip at all" * 40))
    clean_4xx("garbage .pdf", post("garbage.pdf", b"not a pdf preamble" * 60, PDF_MIME))
    clean_4xx("garbage .pptx", post("garbage.pptx", b"\xff\xfe junk" * 50))
    clean_4xx("truncated .docx", post("truncated.docx", valid_bytes[:200]))
    clean_4xx("docx bytes named .pdf", post("renamed.pdf", valid_bytes, PDF_MIME))
    # A real %PDF preamble inside a .docx name: zip sniff must reject it.
    clean_4xx("pdf bytes named .docx", post("renamed.docx", b"%PDF-1.7\n1 0 obj\n<<>>\nendobj\n%%EOF"))
    clean_4xx("unsupported .txt", post("notes.txt", b"hello world", "text/plain"))
    clean_4xx("corrupted document.xml", post("broken_xml.docx", broken_xml.read_bytes()))

    # Valid control still works after all that abuse.
    ok = post("valid.docx", valid_bytes)
    check("valid docx still 200 after fuzzing", ok.status_code == 200, ok.text[:200])

    # --- Concurrency: 8 parallel analyzes all succeed ------------------------
    def one(i: int) -> int:
        return post(f"load_{i}.docx", valid_bytes).status_code

    with concurrent.futures.ThreadPoolExecutor(max_workers=8) as pool:
        statuses = list(pool.map(one, range(8)))
    check("8 concurrent analyzes all 200", statuses == [200] * 8, str(statuses))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
