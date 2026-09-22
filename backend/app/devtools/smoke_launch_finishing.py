"""Smoke: launch-finishing pass behaviors.

Pins the fixes from the pre-sale finishing audit:
1. PASSWORD RESET — request → token email → reset → old password dead, new
   one works; tokens are single-use; a reset token can NOT verify an email
   (cross-purpose guard); bogus/expired tokens 410; the request endpoint never
   discloses whether an account exists.
2. STARTER-GRANT GATE — when SMTP is configured, /auth/grant-starter requires
   a verified email (403 verify_email_first); without SMTP (dev) it grants.
3. PRICING ALignment — the admin MRR table matches the public billing page
   (team=$99, business=$499).

Usage:
    python -m app.devtools.smoke_launch_finishing
"""

from __future__ import annotations

import os
import sys
import tempfile

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_finish_')}/s.db"
os.environ.pop("SMTP_HOST", None)  # start un-configured (dev behavior)

from fastapi.testclient import TestClient  # noqa: E402


def main() -> int:
    failures = 0

    def check(name: str, cond: bool) -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name)
        if not cond:
            failures += 1

    from app.main import app

    client = TestClient(app)

    # --- account setup -----------------------------------------------------
    EMAIL = "finish-pass@example.com"
    r = client.post(
        "/auth/sign-in",
        json={"email": EMAIL, "displayName": "Finish", "password": "originalpw1"},
    )
    check("sign-up 200", r.status_code == 200)
    token = r.json().get("token")
    auth = {"Authorization": f"Bearer {token}"}

    # --- 2. grant gate (dev: no SMTP -> grant works) ------------------------
    r = client.post("/auth/grant-starter", headers=auth)
    check("grant-starter works without SMTP configured", r.status_code == 200 and r.json().get("granted") is True)

    # With SMTP configured and email unverified -> 403 verify_email_first.
    os.environ["SMTP_HOST"] = "smtp.example.com"
    # Sign-up now mails the verification link when SMTP is configured; record
    # it instead of letting smtplib dial smtp.example.com from a smoke.
    import app.api.auth as _auth_mod

    _real_send = _auth_mod.send_email
    _auth_mod.send_email = lambda **kw: True
    try:
        r2 = client.post(
            "/auth/sign-in",
            json={"email": "second-user@example.com", "displayName": "U2", "password": "anotherpw1"},
        )
        auth2 = {"Authorization": f"Bearer {r2.json()['token']}"}
        r = client.post("/auth/grant-starter", headers=auth2)
        check(
            "grant-starter blocked for unverified email when SMTP configured",
            r.status_code == 403 and "verify_email_first" in r.text,
        )
    finally:
        os.environ.pop("SMTP_HOST", None)
        _auth_mod.send_email = _real_send

    # --- 1. password reset --------------------------------------------------
    # Request: response identical whether or not the account exists.
    r_real = client.post("/auth/request-password-reset", json={"email": EMAIL})
    r_fake = client.post("/auth/request-password-reset", json={"email": "nobody@example.com"})
    check("reset request 200 for real account", r_real.status_code == 200 and r_real.json().get("queued") is True)
    check(
        "reset request indistinguishable for unknown account (no enumeration)",
        r_fake.status_code == r_real.status_code and r_fake.json() == r_real.json(),
    )

    # Grab the token from the DB (mailer logs to console in dev).
    from app.db.session_sqlalchemy import session_scope
    from app.db.models import EmailVerifyTokenRow, UserRow
    from sqlalchemy import select

    with session_scope() as session:
        user = session.execute(select(UserRow).where(UserRow.email == EMAIL)).scalar_one()
        row = session.execute(
            select(EmailVerifyTokenRow).where(EmailVerifyTokenRow.user_id == user.id)
        ).scalar_one_or_none()
        reset_token = row.token if row else None
    check("reset token minted with pr_ prefix", bool(reset_token) and reset_token.startswith("pr_"))

    # Cross-purpose guard: the reset token must NOT verify the email.
    r = client.get(f"/auth/verify-email?token={reset_token}")
    check("reset token cannot verify an email (410)", r.status_code == 410)

    # Weak password rejected by schema.
    r = client.post("/auth/reset-password", json={"token": reset_token, "password": "short"})
    check("reset with short password -> 422", r.status_code == 422)

    # Real reset succeeds.
    r = client.post("/auth/reset-password", json={"token": reset_token, "password": "brandnewpw9"})
    check("reset 200", r.status_code == 200 and r.json().get("reset") is True)

    # Old password dead; new password works.
    r = client.post("/auth/sign-in", json={"email": EMAIL, "password": "originalpw1"})
    check("old password no longer signs in (401)", r.status_code == 401)
    r = client.post("/auth/sign-in", json={"email": EMAIL, "password": "brandnewpw9"})
    check("new password signs in (200)", r.status_code == 200)

    # Token is single-use.
    r = client.post("/auth/reset-password", json={"token": reset_token, "password": "thirdpw123"})
    check("used token rejected (410)", r.status_code == 410)

    # Bogus token rejected.
    r = client.post("/auth/reset-password", json={"token": "pr_deadbeefdeadbeef", "password": "whatever123"})
    check("bogus token rejected (410)", r.status_code == 410)

    # --- 3. price alignment --------------------------------------------------
    from app.api.stripe_billing import plan_price_usd

    check("admin MRR price: team == $99", plan_price_usd("team") == 99)
    check("admin MRR price: business == $499", plan_price_usd("business") == 499)

    # --- 4. engine-accuracy fixes --------------------------------------------
    import tempfile as _tf
    from pathlib import Path as _Path

    from docx import Document as _Docx
    from docx.oxml import OxmlElement as _Ox
    from docx.oxml.ns import qn as _qn

    from app.analyzers.registry import run_analyzers as _run
    from app.api.pipeline import _action_persists as _persists
    from app.parsers import parse_to_tree as _parse
    from app.services.remediation_engine import RemediationEngine as _Engine
    from app.services.remediation_planner import (
        RemediationPolicy as _Policy,
        plan_remediations as _plan,
    )
    from app.services.remediators.registry import execute_plans as _exec

    tmpd = _Path(_tf.mkdtemp())
    apply_policy = _Policy(allow_ai_actions=True, require_human_review_for_all=False)

    # 4a. A single heading jump produces exactly ONE violation (duplicate
    # analyzer deregistered — no phantom pendingManual item).
    d = _Docx()
    d.core_properties.title = "Jump Test"
    d.add_heading("Top", level=1)
    d.add_heading("Deep", level=3)
    p = tmpd / "jump.docx"
    d.save(str(p))
    res = _parse(str(p))
    viol = _Engine().detect_violations(res.tree)
    jump_violations = [v for v in viol if v.rule_id in ("HEADING_LEVEL_JUMP", "SKIPPED_HEADING_LEVEL")]
    check("one heading jump -> exactly ONE violation (no duplicate)", len(jump_violations) == 1)

    # 4b. Bare-URL link text: flagged AND the approved fix actually rewrites it.
    d = _Docx()
    d.core_properties.title = "Link Test"
    para = d.add_paragraph("See ")
    rid = d.part.relate_to(
        "https://example.gov/report",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = _Ox("w:hyperlink"); hl.set(_qn("r:id"), rid)
    run_el = _Ox("w:r"); t_el = _Ox("w:t"); t_el.text = "https://example.gov/report"
    run_el.append(t_el); hl.append(run_el); para._p.append(hl)
    p = tmpd / "bareurl.docx"
    d.save(str(p))
    res = _parse(str(p))
    _run(res.tree)
    plans = [pl for pl in _plan(res.tree, apply_policy) if pl.flag.code.value == "LINK_TEXT_NON_DESCRIPTIVE"]
    check("bare-URL link is flagged", len(plans) >= 1)
    execs = _exec(res.tree, plans)
    ok_exec = [e for e in execs if e.action_code.value == "IMPROVE_LINK_TEXT" and e.status.value == "success"]
    check("approved bare-URL fix actually RUNS (no detect/fix disagreement)", len(ok_exec) >= 1)

    # 4c. Derived document title: first heading text, not "Untitled Document".
    d = _Docx()  # no core title
    d.add_heading("Quarterly Financial Review", level=1)
    d.add_paragraph("Body.")
    p = tmpd / "title.docx"
    d.save(str(p))
    res = _parse(str(p))
    _run(res.tree)
    plans = [pl for pl in _plan(res.tree, apply_policy) if pl.flag.code.value == "DOCUMENT_TITLE_MISSING"]
    _exec(res.tree, plans)
    new_title = res.tree.root.metadata.properties.get("title")
    check("title derived from first heading (not 'Untitled Document')", new_title == "Quarterly Financial Review")

    # 4d. PDF_UNTAGGED is a persisted action for pdf (the writer reconstructs
    # the struct tree, so the score may honestly credit it).
    check("TAG_PDF_STRUCTURE persists for pdf", _persists("TAG_PDF_STRUCTURE", "pdf"))

    # 4e. TEXT_STYLED_AS_HEADING — the title-page failure.
    from docx.shared import Pt as _Pt

    def _flag_codes(tree):
        found = set()

        def walk(n):
            for f in n.accessibility_flags:
                found.add(f.code.value)
            for ch in n.children:
                walk(ch)

        walk(tree.root)
        return found

    FAKE_H = "TEXT_STYLED_AS_HEADING"

    d = _Docx()
    d.core_properties.title = "Fake Heading Test"
    # (a) big bold short line typed as a normal paragraph -> flagged
    p_big = d.add_paragraph()
    r = p_big.add_run("Annual Report 2026")
    r.bold = True
    r.font.size = _Pt(24)
    # (b) normal body paragraph -> not flagged
    d.add_paragraph("This is an ordinary sentence of body text in the document.")
    # (c) bold-but-small inline emphasis -> not flagged (no 14pt size)
    p_small = d.add_paragraph()
    rs = p_small.add_run("Important note")
    rs.bold = True
    # (d) real heading style -> becomes a HeadingNode, never flagged
    d.add_heading("Real Section", level=1)
    p = tmpd / "fakeheading.docx"
    d.save(str(p))
    res = _parse(str(p))
    _run(res.tree)
    codes = _flag_codes(res.tree)
    check("big bold non-heading text -> TEXT_STYLED_AS_HEADING fires", FAKE_H in codes)
    flagged_paras = [
        n
        for n in res.tree.root.children[0].children
        if any(f.code.value == FAKE_H for f in n.accessibility_flags)
    ]
    check("exactly the fake-heading paragraph is flagged (no body/emphasis FPs)", len(flagged_paras) == 1)

    # (e) Word Title style -> flagged even without explicit size
    d = _Docx()
    d.core_properties.title = "Title Style Test"
    tp = d.add_paragraph("Quarterly Review")
    tp.style = d.styles["Title"]
    d.add_paragraph("Body text.")
    p = tmpd / "titlestyle.docx"
    d.save(str(p))
    res = _parse(str(p))
    _run(res.tree)
    check("Title-style paragraph -> TEXT_STYLED_AS_HEADING fires", FAKE_H in _flag_codes(res.tree))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
