"""Smoke: /pipeline/remediate delivers ONLY approved fixes, and charges for every one.

The honesty invariant from both sides: a fix is charged if and only if it
persists into the output bytes. Found in adversarial review: the PDF writer
built the full structure tree on EVERY PDF, whatever was approved — so
approved_violations=[] returned the 5-credit tagged PDF, uncharged, and the
paid run's bytes were identical to it. The DOCX writer likewise copied a
w:lang-derived language into dc:language that nobody approved. Pins:

  * approve NOTHING (PDF, DOCX, PPTX, HTML) -> not charged, and the download
    is byte-for-byte the upload (no tagging, no metadata re-assertion);
  * approve only the PDF title -> charged, /Title written, but NOT tagged;
  * approve PDF_UNTAGGED -> tagged (StructTreeRoot + MarkInfo) and charged 5;
  * TAG_PDF_STRUCTURE is writer-confirmed: when the writer builds no tree
    (nothing taggable, or a tagger failure) the fix is reported skipped, is not
    counted or charged, and the file is the source;
  * approve only the DOCX title -> charged, title set, dc:language untouched;
  * approve only the PPTX title -> no a:rPr@lang stamped onto the runs.

Usage:
    python -m app.devtools.smoke_remediate_only_approved
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import zipfile
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="508_smoke_only_approved_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/s.db"
os.environ["MATERIALIZED_ROOT"] = str(Path(_TMP) / "materialized")

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402
from pypdf import PdfReader, PdfWriter  # noqa: E402
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject  # noqa: E402
from sqlalchemy import select  # noqa: E402

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
HTML = "text/html"


def _pdf(pages: int = 3, blank: bool = False) -> bytes:
    w = PdfWriter()
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    fref = w._add_object(font)  # noqa: SLF001
    for i in range(pages):
        page = w.add_blank_page(width=400, height=300)
        if blank:
            continue
        cs = DecodedStreamObject()
        cs.set_data((f"BT /F1 24 Tf 20 260 Td (Section {i} heading) Tj ET\n"
                     f"BT /F1 12 Tf 20 230 Td (Body paragraph one on page {i}.) Tj ET\n"
                     f"BT /F1 12 Tf 20 200 Td (Body paragraph two on page {i}.) Tj ET").encode())
        page[NameObject("/Contents")] = w._add_object(cs)  # noqa: SLF001
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): fref})})
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def _docx() -> bytes:
    # python-docx's template declares w:lang="en-US" in styles.xml and leaves
    # dc:language empty — the exact shape the writer used to "sync".
    d = Document()
    d.add_heading("Quarterly results", level=1)
    d.add_paragraph("Revenue grew in every region this quarter, led by the west.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    prs = Presentation()
    prs.core_properties.language = "en-US"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Regional safety results."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


_HTML = b"<!doctype html><html><head></head><body><h1>Notice</h1><p>Office hours change next week.</p></body></html>"


def _struct(b: bytes) -> dict:
    root = PdfReader(io.BytesIO(b)).trailer["/Root"]
    mark = root.get("/MarkInfo")
    return {"tree": "/StructTreeRoot" in root, "marked": bool(mark.get_object().get("/Marked")) if mark else False}


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.db.models import UserRow
    from app.db.session_sqlalchemy import session_scope
    from app.main import app

    client = TestClient(app, raise_server_exceptions=False)
    email = "only-approved@example.com"
    r = client.post("/auth/sign-in", json={"email": email, "displayName": "OA", "password": "onlyapproved1"})
    auth = {"Authorization": f"Bearer {r.json()['token']}"}
    client.post("/auth/grant-starter", headers=auth)
    with session_scope() as s:
        uid = s.execute(select(UserRow.id).where(UserRow.email == email)).scalar_one()

    def balance() -> int:
        with session_scope() as s:
            return int(s.execute(select(UserRow.credits_balance).where(UserRow.id == uid)).scalar_one() or 0)

    def top_up(n: int = 25) -> None:
        with session_scope() as s:
            s.execute(select(UserRow).where(UserRow.id == uid)).scalar_one().credits_balance = n

    def ids(name: str, data: bytes, mime: str) -> dict:
        a = client.post("/pipeline/analyze", files={"file": (name, data, mime)}, headers=auth)
        assert a.status_code == 200, a.text[:200]
        return {v["ruleId"]: v["id"] for v in a.json()["violations"]}

    def remediate(name: str, data: bytes, mime: str, approved: list) -> tuple:
        rr = client.post("/pipeline/remediate", files={"file": (name, data, mime)},
                         data={"approved_violations": json.dumps(approved), "rejected_violations": "[]"},
                         headers=auth)
        assert rr.status_code == 200, rr.text[:300]
        body = rr.json()
        dl = client.get(body["downloadUrl"])
        assert dl.status_code == 200, dl.text[:200]
        return body, dl.content

    # ---- PDF ---------------------------------------------------------------
    pdf = _pdf()
    pdf_ids = ids("Quarterly_Budget_Review.pdf", pdf, PDF)
    check("fixture PDF is flagged untagged", "PDF_UNTAGGED" in pdf_ids, str(sorted(pdf_ids)))

    b0 = balance()
    body, out = remediate("Quarterly_Budget_Review.pdf", pdf, PDF, [])
    check("PDF approve nothing -> charged=false", body["charged"] is False)
    check("PDF approve nothing -> balance unchanged", balance() == b0, f"{b0} -> {balance()}")
    check("PDF approve nothing -> download is byte-for-byte the upload", out == pdf)
    check("PDF approve nothing -> NOT tagged (no StructTreeRoot)", not _struct(out)["tree"], str(_struct(out)))
    check("PDF approve nothing -> writer reports nothing applied", body["writer"]["applied"] == [], str(body["writer"]))

    b0 = balance()
    body, out = remediate("Quarterly_Budget_Review.pdf", pdf, PDF, [pdf_ids["DOCUMENT_TITLE_MISSING"]])
    title_ok = [e for e in body["executions"] if e["actionCode"] == "SET_DOCUMENT_TITLE" and e["status"] == "success"]
    check("PDF title-only: the title fix ran", len(title_ok) == 1, str(body["executions"]))
    check("PDF title-only -> charged 5", body["charged"] is True and b0 - balance() == 5, f"{b0} -> {balance()}")
    check("PDF title-only -> /Title written", bool((PdfReader(io.BytesIO(out)).metadata or {}).get("/Title")))
    check("PDF title-only -> NOT tagged (tagging was not approved)", _struct(out) == {"tree": False, "marked": False},
          str(_struct(out)))

    b0 = balance()
    body, out = remediate("Quarterly_Budget_Review.pdf", pdf, PDF, [pdf_ids["PDF_UNTAGGED"]])
    tag = [e for e in body["executions"] if e["actionCode"] == "TAG_PDF_STRUCTURE"]
    check("PDF_UNTAGGED approved -> TAG_PDF_STRUCTURE success", [e["status"] for e in tag] == ["success"], str(tag))
    check("PDF_UNTAGGED approved -> tagged (StructTreeRoot + Marked)", _struct(out) == {"tree": True, "marked": True},
          str(_struct(out)))
    check("PDF_UNTAGGED approved -> charged 5", body["charged"] is True and b0 - balance() == 5, f"{b0} -> {balance()}")
    check("PDF_UNTAGGED approved -> persistedFixes counts the tagging", body["persistedFixes"] >= 1)

    # No tree written, nothing charged. TAG_PDF_STRUCTURE is writer-confirmed:
    # the executor only records the request, so counting waits for the
    # writer's own confirmation that a structure tree landed in the bytes.
    from app.api.pipeline import _count_persisted_fixes
    from app.models.accessibility import ActionCode
    from app.parsers import parse_to_tree
    from app.services.remediators.base import ExecutionResult, ExecutionStatus
    from app.writers import pdf_writer

    tag_ok = ExecutionResult(action_code=ActionCode.TAG_PDF_STRUCTURE, target_node_id="doc-1",
                             status=ExecutionStatus.SUCCESS, notes="")
    check("tagging success WITHOUT the writer's confirmation is not counted",
          _count_persisted_fixes([tag_ok], [{"kind": "pdfua_xmp_metadata", "target_id": "document"}], "pdf") == 0)
    check("tagging success WITH the writer's confirmation is counted",
          _count_persisted_fixes([tag_ok], [{"action": "TAG_PDF_STRUCTURE", "target_id": "doc-1"}], "pdf") == 1)

    blank_src = Path(_TMP) / "blank.pdf"
    blank_src.write_bytes(_pdf(pages=1, blank=True))
    res = parse_to_tree(str(blank_src))
    res.tree.root.metadata.properties = dict(res.tree.root.metadata.properties or {})
    res.tree.root.metadata.properties["tag_structure_requested"] = True
    rep = pdf_writer.write_remediated_pdf(blank_src, res.tree, Path(_TMP) / "blank-out.pdf")
    check("writer: nothing taggable -> no TAG_PDF_STRUCTURE confirmation",
          not any(a.get("action") == "TAG_PDF_STRUCTURE" for a in rep["applied"]), str(rep["applied"]))

    # Route level: the tagger writes no tree (patched in-process, as a tagger
    # failure or an all-untaggable document would).
    top_up()
    real_tag_pdf = pdf_writer.tag_pdf
    pdf_writer.tag_pdf = lambda writer, tree: {"applied": [], "structTree": False, "structSkipped": "no_taggable_pages"}
    try:
        b0 = balance()
        body, out = remediate("Quarterly_Budget_Review.pdf", pdf, PDF, [pdf_ids["PDF_UNTAGGED"]])
    finally:
        pdf_writer.tag_pdf = real_tag_pdf
    tag = [e for e in body["executions"] if e["actionCode"] == "TAG_PDF_STRUCTURE"]
    check("no tree written: TAG_PDF_STRUCTURE reported skipped, not success",
          [e["status"] for e in tag] == ["skipped"], str(tag))
    check("no tree written: the note does not promise a structure tree",
          bool(tag) and all("not applied" in (e["notes"] or "") for e in tag), str(tag))
    check("no tree written: not charged", body["charged"] is False and balance() == b0, f"{b0} -> {balance()}")
    check("no tree written: download is the upload", out == pdf)

    # ---- DOCX --------------------------------------------------------------
    top_up()
    dx = _docx()
    dx_ids = ids("Quarterly_Budget_Review.docx", dx, DOCX)
    check("fixture DOCX has DOCUMENT_TITLE_MISSING", "DOCUMENT_TITLE_MISSING" in dx_ids, str(sorted(dx_ids)))
    check("fixture DOCX is NOT flagged for language (w:lang declares en-US)", "DOCUMENT_LANGUAGE_MISSING" not in dx_ids)

    b0 = balance()
    body, out = remediate("Quarterly_Budget_Review.docx", dx, DOCX, [])
    check("DOCX approve nothing -> charged=false, balance unchanged", body["charged"] is False and balance() == b0)
    check("DOCX approve nothing -> download is byte-for-byte the upload", out == dx)

    b0 = balance()
    body, out = remediate("Quarterly_Budget_Review.docx", dx, DOCX, [dx_ids["DOCUMENT_TITLE_MISSING"]])
    core = Document(io.BytesIO(out)).core_properties
    check("DOCX title-only -> charged 3", body["charged"] is True and b0 - balance() == 3, f"{b0} -> {balance()}")
    check("DOCX title-only -> title written", bool((core.title or "").strip()), repr(core.title))
    check("DOCX title-only -> dc:language untouched (nobody approved a language fix)",
          (core.language or "") == "", repr(core.language))
    check("DOCX title-only -> writer did not report a language write",
          not any(str(a.get("kind", "")).startswith("document_language") for a in body["writer"]["applied"]),
          str(body["writer"]["applied"]))

    # ---- PPTX --------------------------------------------------------------
    top_up()
    px = _pptx()
    px_ids = ids("Regional_Safety_Plan.pptx", px, PPTX)
    b0 = balance()
    body, out = remediate("Regional_Safety_Plan.pptx", px, PPTX, [])
    check("PPTX approve nothing -> charged=false, balance unchanged", body["charged"] is False and balance() == b0)
    check("PPTX approve nothing -> download is byte-for-byte the upload", out == px)
    if "DOCUMENT_TITLE_MISSING" in px_ids:
        body, out = remediate("Regional_Safety_Plan.pptx", px, PPTX, [px_ids["DOCUMENT_TITLE_MISSING"]])
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            slides = b"".join(z.read(n) for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        check("PPTX title-only -> no a:rPr@lang stamped onto the runs", b'lang="en-US"' not in slides)
    else:
        check("fixture PPTX has DOCUMENT_TITLE_MISSING", False, str(sorted(px_ids)))

    # ---- HTML --------------------------------------------------------------
    top_up()
    b0 = balance()
    body, out = remediate("notice.html", _HTML, HTML, [])
    check("HTML approve nothing -> charged=false, balance unchanged", body["charged"] is False and balance() == b0)
    check("HTML approve nothing -> download is byte-for-byte the upload", out == _HTML)

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except AssertionError as exc:
        print(f"FAIL only-approved smoke: {exc}")
        sys.exit(1)
