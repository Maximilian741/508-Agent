"""Smoke: the Word shapes real customers upload are read correctly, fixed in
the right place, and never come back as a file Word cannot open.

Each case reproduces something the robustness corpus (content controls,
tracked changes, nested tables, text boxes, 500 pages) showed going wrong:

  * tracked changes — python-docx's ``paragraph.text`` drops text inside
    ``w:ins`` and inline content controls: "work up to <ins>three</ins> days"
    read "work up to  days", and a paragraph that is one tracked insertion
    read as empty and was never audited. Deleted text must stay out.
  * a Heading-styled paragraph with its own Word numbering is a HEADING (it
    was swallowed as a list item, breaking the outline)
  * a table nested in a cell is reported (TABLE_NESTED could never fire on a
    Word file), and a header fix for the INNER table lands on the inner
    table's row — the outer table is untouched
  * a picture in a table cell (a form's logo cell) is found and its alt lands
  * a modern text box is stored twice (mc:Choice + mc:Fallback); its "click
    here" is ONE finding, not two
  * schema order: a header row whose trPr carries a tracked-change marker,
    styles.xml whose docDefaults has only pPrDefault, and a Mac numbering
    part ending in numIdMacAtCleanup all stay in schema order after a fix
  * nothing unapproved is rewritten: a bold first row and a custom heading
    style are left alone when only the title is fixed
  * the output gate: if the saved file does not re-open, the SOURCE bytes are
    restored, nothing is reported applied, and the pipeline counts 0
  * a control byte in generated alt text, link text or a title is dropped —
    it used to abort the whole write (422, no file for the customer)
  * a synthesized header row goes on TOP of a table whose rows all live in a
    repeating-section content control (it was appended after the last row)
  * irregular vertical merges Word opens (a continuation in the first row, or
    under a cell with a different grid span) no longer fail the whole parse
  * no docProps/core.xml: no invented "Word Document" title, in the tree or
    in the output; an unchanged title is not re-"applied"
  * findings carry the PAGE Word last laid them out on (lastRenderedPageBreak
    marks that agree with app.xml's page count) — and no page at all when
    the file does not say

Usage:
    python -m app.devtools.smoke_docx_robustness
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_drb_')}/s.db")
os.environ.pop("OPENAI_API_KEY", None)
os.environ.pop("ANTHROPIC_API_KEY", None)
os.environ["SEMANTIC_PROVIDER"] = "heuristic"

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement, parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402
from docx.shared import Inches  # noqa: E402
from lxml import etree  # noqa: E402
from PIL import Image  # noqa: E402

import app.writers.docx_writer as docx_writer  # noqa: E402
from app.analyzers.registry import run_analyzers  # noqa: E402
from app.api.pipeline import _count_persisted_fixes  # noqa: E402
from app.models.accessibility import (  # noqa: E402
    ContentKind,
    HeadingNode,
    ImageNode,
    LinkNode,
    ListItemNode,
    NodeContent,
    NodeMetadata,
    ParagraphNode,
    TableCellNode,
    TableCellType,
    TableHeaderScope,
    TableNode,
    TableRowNode,
    iter_reading_order,
)
from app.parsers import parse_to_tree  # noqa: E402
from app.parsers.docx_parser import DOCXParser  # noqa: E402
from app.services.remediation_engine import _collect_evidence  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.docx_writer import write_remediated_docx  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
POL = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)


def _png(seed: int) -> io.BytesIO:
    b = io.BytesIO()
    Image.new("RGB", (30 + seed, 20), (seed * 20 % 255, 80, 140)).save(b, "PNG")
    b.seek(0)
    return b


def _ins(text: str, rid: str = "1"):
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), rid)
    ins.set(qn("w:author"), "Reviewer")
    ins.set(qn("w:date"), "2026-09-01T00:00:00Z")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    ins.append(r)
    return ins


def _del(text: str, rid: str = "2"):
    d = OxmlElement("w:del")
    d.set(qn("w:id"), rid)
    d.set(qn("w:author"), "Reviewer")
    d.set(qn("w:date"), "2026-09-01T00:00:00Z")
    r = OxmlElement("w:r")
    t = OxmlElement("w:delText")
    t.text = text
    r.append(t)
    d.append(r)
    return d


def _nodes(tree, kind):
    return [n for n in iter_reading_order(tree.root) if isinstance(n, kind)]


def _codes(tree):
    return [f.code.value for n in iter_reading_order(tree.root) for f in n.accessibility_flags]


def _rewrite_part(path: Path, name: str, fn) -> None:
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    root = etree.fromstring(items[name])
    fn(root)
    items[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in items.items():
            z.writestr(n, b)


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="508_drb_"))

    # ===== 1. Tracked changes =================================================
    d = Document()
    d.core_properties.title = "Remote Work Policy"
    d.add_heading("Remote Work Policy", 1)
    p = d.add_paragraph("Employees may work remotely up to ")
    p._p.append(_ins("three", "11"))
    p._p.append(_del("two", "12"))
    p.add_run(" days per week.")
    whole = d.add_paragraph()
    whole._p.append(_ins("New: equipment stipends are reviewed annually.", "13"))
    gone = d.add_paragraph()
    gone._p.append(_del("This whole sentence was removed by the reviewer.", "14"))
    src = tmp / "tracked.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    texts = [n.content.text for n in _nodes(res.tree, ParagraphNode)]
    check("tracked: inserted word is part of the sentence",
          "Employees may work remotely up to three days per week." in texts, str(texts))
    check("tracked: a paragraph that is one tracked insertion is audited",
          "New: equipment stipends are reviewed annually." in texts, str(texts))
    check("tracked: deleted text is not treated as content",
          not any("removed by the reviewer" in t or "twodays" in t or " two " in t for t in texts), str(texts))

    # ===== 2. Numbered Heading paragraph is a heading ========================
    d = Document()
    d.core_properties.title = "Numbered headings"
    h1 = d.add_heading("Introduction", 1)
    h2 = d.add_heading("Scope", 2)
    for h in (h1, h2):
        numpr = parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>')
        h._p.get_or_add_pPr().append(numpr)
    d.add_paragraph("Body text of the scope section that is long enough to read as prose.")
    src = tmp / "numbered_headings.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    hs = [(n.content.text, n.level) for n in _nodes(res.tree, HeadingNode)]
    check("numbered headings: both are HeadingNodes (not list items)",
          hs == [("Introduction", 1), ("Scope", 2)] and not _nodes(res.tree, ListItemNode), str(hs))
    check("numbered headings: no invented heading jump", "HEADING_LEVEL_JUMP" not in _codes(res.tree))

    # ===== 2b. Outline levels make headings (custom heading styles) ==========
    d = Document()
    d.core_properties.title = "Agency template"
    agency = d.styles.add_style("Agency Heading One", 1)
    agency.base_style = d.styles["Heading 1"]            # inherits outline level 0
    d.add_paragraph("Program Overview", style="Agency Heading One")
    d.add_paragraph("Body text under the overview that is long enough to be prose.")
    direct = d.add_paragraph("Eligibility")               # Normal + direct "Level 2"
    lvl = OxmlElement("w:outlineLvl"); lvl.set(qn("w:val"), "1")
    direct._p.get_or_add_pPr().append(lvl)
    body_lvl = d.add_paragraph("Plain body text explicitly marked as body level text here.")
    lvl9 = OxmlElement("w:outlineLvl"); lvl9.set(qn("w:val"), "9")
    body_lvl._p.get_or_add_pPr().append(lvl9)
    src = tmp / "outline_levels.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    hs = [(n.content.text, n.level) for n in _nodes(res.tree, HeadingNode)]
    check("outline levels: a style based on Heading 1 and a direct Level 2 are headings; Level 9 is body",
          hs == [("Program Overview", 1), ("Eligibility", 2)], str(hs))
    check("outline levels: no false 'document has no headings'", "DOCUMENT_NO_HEADINGS" not in _codes(res.tree))

    # ===== 2c. Internal links (TOC entries) point at bookmarks ================
    d = Document()
    d.core_properties.title = "Handbook"
    toc = d.add_paragraph()
    for anchor, label in (("_Toc100", "Introduction"), ("_Toc999", "Missing section")):
        hl = OxmlElement("w:hyperlink"); hl.set(qn("w:anchor"), anchor)
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = label; r.append(t); hl.append(r)
        toc._p.append(hl)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), ' HYPERLINK \\l "_Toc100" ')
    fr = OxmlElement("w:r"); ft = OxmlElement("w:t"); ft.text = "Back to the introduction"; fr.append(ft); fld.append(fr)
    d.add_paragraph()._p.append(fld)
    h = d.add_heading("Introduction", 1)
    bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), "0"); bs.set(qn("w:name"), "_Toc100")
    be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), "0")
    h._p.insert(1, bs)
    h._p.append(be)
    d.add_paragraph("Introduction body text that is long enough to be prose in this handbook.")
    src = tmp / "toc_links.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    lk = {n.content.text: n for n in _nodes(res.tree, LinkNode)}
    broken = {n.content.text for n in lk.values() if any(f.code.value == "LINK_TARGET_BROKEN" for f in n.accessibility_flags)}
    check("toc: an entry pointing at an existing bookmark is a valid internal link",
          lk.get("Introduction") is not None and lk["Introduction"].target == "#_Toc100"
          and "Introduction" not in broken, str({k: v.target for k, v in lk.items()}))
    check("toc: HYPERLINK \\l field resolves to the bookmark, not to the switch",
          lk.get("Back to the introduction") is not None and lk["Back to the introduction"].target == "#_Toc100")
    check("toc: an entry pointing at a bookmark that does not exist IS reported broken",
          broken == {"Missing section"}, str(broken))

    # ===== 3. Nested table + picture in a cell ================================
    d = Document()
    d.core_properties.title = "Grant Budget"
    d.add_heading("Grant Budget", 1)
    outer = d.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "Personnel"
    outer.cell(0, 0).paragraphs[0].runs[0].bold = True
    outer.cell(0, 1).text = "Breakdown"
    outer.cell(0, 1).paragraphs[0].runs[0].bold = True
    inner = outer.cell(1, 1).add_table(rows=3, cols=3)
    for r, row in enumerate([["Role", "FTE", "Cost"], ["PI", "0.25", "42,000"], ["Analyst", "1.0", "88,000"]]):
        for c, v in enumerate(row):
            inner.cell(r, c).text = v
    outer.cell(1, 0).paragraphs[0].add_run().add_picture(_png(3), width=Inches(0.5))
    src = tmp / "nested.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    tables = _nodes(res.tree, TableNode)
    nested = [t for t in tables if (t.metadata.properties or {}).get("nested_table")]
    check("nested: the inner table is a node and TABLE_NESTED fires once",
          len(tables) == 2 and len(nested) == 1 and _codes(res.tree).count("TABLE_NESTED") == 1,
          str(([tb.id for tb in tables], _codes(res.tree))))
    cell_imgs = [n for n in _nodes(res.tree, ImageNode) if (n.metadata.properties or {}).get("docx_story") == "table"]
    check("cell picture: found and flagged", len(cell_imgs) == 1
          and any(f.code.value == "MISSING_ALT_TEXT" for f in cell_imgs[0].accessibility_flags))
    check("nested: inner data table without a header row is flagged",
          any(f.code.value == "TABLE_MISSING_HEADERS" for f in nested[0].accessibility_flags) if nested else False)
    # Fix: promote the inner table's first row; give the cell picture alt.
    first_row = nested[0].children[0]
    for c in first_row.children:
        c.cell_type = TableCellType.HEADER
    cell_imgs[0].alt_text = "Agency seal"
    out = tmp / "nested_fixed.docx"
    wr = write_remediated_docx(src, res.tree, out)
    with zipfile.ZipFile(out) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    outer_tbl = root.find(f".//{W}body/{W}tbl")
    inner_tbl = outer_tbl.find(f".//{W}tc/{W}tbl")
    inner_hdr = inner_tbl.find(f"{W}tr/{W}trPr/{W}tblHeader") if inner_tbl is not None else None
    outer_hdr = outer_tbl.find(f"{W}tr/{W}trPr/{W}tblHeader")
    check("nested: tblHeader landed on the INNER table's first row", inner_hdr is not None)
    check("nested: the outer (already bold-header) table was NOT touched", outer_hdr is None)
    descrs = [dp.get("descr") for dp in root.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")]
    check("cell picture: alt landed on it", descrs == ["Agency seal"], str(descrs))
    res2 = parse_to_tree(str(out))
    run_analyzers(res2.tree)
    check("nested: re-parse — inner headers fixed, alt fixed, only the manual TABLE_NESTED remains",
          "TABLE_MISSING_HEADERS" not in _codes(res2.tree) and "MISSING_ALT_TEXT" not in _codes(res2.tree)
          and "TABLE_NESTED" in _codes(res2.tree), str(_codes(res2.tree)))

    # ===== 4. mc:AlternateContent text box: one finding, not two ==============
    d = Document()
    d.core_properties.title = "Newsletter"
    d.add_heading("Newsletter", 1)
    host = d.add_paragraph("Sidebar anchor paragraph with enough words to be prose here.")
    rid = d.part.relate_to("https://example.gov/sidebar",
                           "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    tb_p = (f'<w:p><w:r><w:t xml:space="preserve">Sidebar: </w:t></w:r>'
            f'<w:hyperlink r:id="{rid}"><w:r><w:t>click here</w:t></w:r></w:hyperlink></w:p>')
    extra_ns = (
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
    )
    alt = parse_xml(
        f'<w:r {nsdecls("w", "r", "wp", "a")} {extra_ns}><mc:AlternateContent>'
        '<mc:Choice Requires="wps"><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        'relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:extent cx="914400" cy="457200"/><wp:wrapNone/><wp:docPr id="50" name="Text Box 50"/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr/><wps:txbx><w:txbxContent>{tb_p}</w:txbxContent></wps:txbx>'
        '<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
        f'<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>{tb_p}</w:txbxContent></v:textbox></v:shape>'
        '</w:pict></mc:Fallback></mc:AlternateContent></w:r>'
    )
    host._p.append(alt)
    src = tmp / "textbox_alt.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    sidebar = [n for n in _nodes(res.tree, LinkNode) if n.target == "https://example.gov/sidebar"]
    check("text box stored twice (Choice + Fallback): its link is ONE node",
          len(sidebar) == 1, str([n.id for n in sidebar]))
    check("... and ONE finding", _codes(res.tree).count("LINK_TEXT_NON_DESCRIPTIVE") == 1, str(_codes(res.tree)))
    check("host paragraph text excludes the text box copy (no duplicated words)",
          any(n.content.text == "Sidebar anchor paragraph with enough words to be prose here."
              for n in _nodes(res.tree, ParagraphNode)))

    # ===== 5. Schema order is kept ===========================================
    d = Document()
    d.add_heading("Quarterly figures", 1)
    t = d.add_table(rows=4, cols=3)
    for r in range(4):
        for c in range(3):
            t.cell(r, c).text = f"v{r}{c}" if r else ["Region", "Q1", "Q2"][c]
    t.rows[0]._tr.get_or_add_trPr().append(parse_xml(
        f'<w:ins {nsdecls("w")} w:id="90" w:author="R" w:date="2026-09-01T00:00:00Z"/>'))
    d.add_paragraph("- first typed item")
    d.add_paragraph("- second typed item")
    src = tmp / "schema.docx"
    d.save(str(src))

    def _pprdefault_only(root):
        dd = root.find(f"{W}docDefaults")
        for child in list(dd):
            if child.tag == f"{W}rPrDefault":
                dd.remove(child)

    _rewrite_part(src, "word/styles.xml", _pprdefault_only)

    def _mac_cleanup(root):
        mac = etree.SubElement(root, f"{W}numIdMacAtCleanup")
        mac.set(f"{W}val", "3")

    _rewrite_part(src, "word/numbering.xml", _mac_cleanup)
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    res.tree.root.metadata.language = "en-US"            # SET_DOCUMENT_LANGUAGE's effect
    for c in _nodes(res.tree, TableNode)[0].children[0].children:
        c.cell_type = TableCellType.HEADER                # ADD_TABLE_HEADERS' effect
    plans = [p for p in plan_remediations(res.tree, POL) if p.flag.code.value == "LIST_STRUCTURE_INVALID"]
    execute_plans(res.tree, plans)
    out = tmp / "schema_fixed.docx"
    wr = write_remediated_docx(src, res.tree, out)
    with zipfile.ZipFile(out) as z:
        doc_root = etree.fromstring(z.read("word/document.xml"))
        styles_root = etree.fromstring(z.read("word/styles.xml"))
        num_root = etree.fromstring(z.read("word/numbering.xml"))
    trpr = doc_root.find(f".//{W}tbl/{W}tr/{W}trPr")
    order = [etree.QName(c).localname for c in trpr]
    check("schema: tblHeader goes BEFORE the row's tracked-change marker",
          order.index("tblHeader") < order.index("ins"), str(order))
    dd = [etree.QName(c).localname for c in styles_root.find(f"{W}docDefaults")]
    check("schema: rPrDefault precedes pPrDefault", dd == ["rPrDefault", "pPrDefault"], str(dd))
    tail = [etree.QName(c).localname for c in num_root][-1]
    check("schema: numIdMacAtCleanup stays last in numbering.xml", tail == "numIdMacAtCleanup",
          str([etree.QName(c).localname for c in num_root][-4:]))
    check("schema: the fixes were applied", any(a.get("kind") == "list_conversion" for a in wr["applied"])
          and any(a.get("kind") == "table_header_row" for a in wr["applied"]), str(wr["applied"]))
    check("applied list: a 3-column header promotion is ONE entry (per row, not per cell)",
          len([a for a in wr["applied"] if a.get("kind") == "table_header_row"]) == 1, str(wr["applied"]))

    # ===== 6. Nothing unapproved is rewritten ================================
    d = Document()
    custom = d.styles.add_style("Agency Heading", 1)  # paragraph style
    custom.base_style = d.styles["Heading 2"]
    d.add_heading("Report", 1)
    hp = d.add_paragraph("Agency section", style="Heading 2")
    t = d.add_table(rows=2, cols=2)          # small table: row 0 typed header by default
    t.cell(0, 0).text = "Name"
    t.cell(0, 1).text = "Role"
    t.cell(1, 0).text = "Ana"
    t.cell(1, 1).text = "Lead"
    src = tmp / "untouched.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    res.tree.root.metadata.properties["title"] = "Staff Report"  # only the title is fixed
    out = tmp / "untouched_fixed.docx"
    wr = write_remediated_docx(src, res.tree, out)
    kinds = sorted({a.get("kind") for a in wr["applied"]})
    check("unapproved: only the title was written (no heading re-assert, no tblHeader)",
          kinds == ["document_title"], str(wr["applied"]))
    with zipfile.ZipFile(out) as z:
        x = z.read("word/document.xml")
    check("unapproved: no tblHeader appeared", b"tblHeader" not in x)

    # ===== 7. The output gate =================================================
    d = Document()
    d.add_paragraph("Body text long enough to be a real paragraph of prose.")
    src = tmp / "gate.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    plans = [p for p in plan_remediations(res.tree, POL) if p.flag.code.value == "DOCUMENT_TITLE_MISSING"]
    execs = execute_plans(res.tree, plans)
    real = docx_writer._output_problem
    docx_writer._output_problem = lambda _p: "simulated: word/document.xml is not well-formed"
    try:
        out = tmp / "gate_fixed.docx"
        wr = write_remediated_docx(src, res.tree, out)
    finally:
        docx_writer._output_problem = real
    check("gate: a file that does not re-open is NOT delivered (source bytes restored)",
          out.read_bytes() == src.read_bytes())
    check("gate: nothing reported applied", wr["applied"] == [], str(wr["applied"]))
    check("gate: failed_to_save reason present",
          any(str(s.get("reason", "")).startswith("failed_to_save") for s in wr["skipped"]))
    check("gate: pipeline counts 0 (no charge)", _count_persisted_fixes(execs, wr["applied"], "docx", wr["skipped"]) == 0)
    check("gate: the real check passes a sound file", real(src) is None)
    bad = tmp / "broken.docx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(bad, "w") as zout:
        for n in zin.namelist():
            data = zin.read(n)
            zout.writestr(n, data[:-20] if n == "word/document.xml" else data)
    check("gate: the real check rejects a truncated document.xml", real(bad) is not None)

    # ===== 8. One bad character does not cost the customer the whole file ====
    # A control byte in generated alt text / link text / title made lxml raise
    # mid-write; /remediate answered 422 and the customer got NO file for one
    # stray character. It is dropped and every fix still lands.
    d = Document()
    d.add_heading("Control characters", 1)
    d.add_paragraph("Body text long enough to be prose next to the picture below.")
    d.add_picture(_png(8), width=Inches(1))
    lp = d.add_paragraph("Read the report ")
    lrid = d.part.relate_to("https://example.gov/report",
                            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                            is_external=True)
    hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), lrid)
    hr = OxmlElement("w:r"); ht = OxmlElement("w:t"); ht.text = "here"; hr.append(ht); hl.append(hr)
    lp._p.append(hl)
    src = tmp / "ctrl.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    for n in iter_reading_order(res.tree.root):
        if isinstance(n, ImageNode):
            n.alt_text = "Bar chart of revenue\x01 by quarter\x0b"
        if isinstance(n, LinkNode):
            n.content.text = "Annual report\x07 2026"
    res.tree.root.metadata.properties["title"] = "Budget\x0c Summary\ud800"
    out = tmp / "ctrl_fixed.docx"
    try:
        wr = write_remediated_docx(src, res.tree, out)
        crashed = None
    except Exception as exc:  # the regression
        wr, crashed = {"applied": [], "skipped": []}, repr(exc)
    check("control chars: the write completes", crashed is None, str(crashed))
    kinds = sorted(a.get("kind") for a in wr["applied"])
    check("control chars: alt, link and title all applied", kinds == ["document_title", "image_alt_text", "link_text"],
          str(wr))
    if crashed is None:
        od = Document(str(out))
        with zipfile.ZipFile(out) as z:
            droot = etree.fromstring(z.read("word/document.xml"))
        descr = [dp.get("descr") for dp in droot.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")]
        ltext = ["".join(t.text or "" for t in h.iter(f"{W}t")) for h in droot.iter(f"{W}hyperlink")]
        check("control chars: stripped, words kept",
              descr == ["Bar chart of revenue by quarter"] and ltext == ["Annual report 2026"]
              and od.core_properties.title == "Budget Summary", str((descr, ltext, od.core_properties.title)))

    # ===== 9. A synthesized header row goes on TOP, even when rows are in an SDT
    d = Document()
    d.add_heading("Repeating section", 1)
    body = d.element.body
    tbl = parse_xml(
        f'<w:tbl {nsdecls("w")}><w:tblPr/><w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/></w:tblGrid>'
        '<w:sdt><w:sdtPr><w:alias w:val="Rows"/></w:sdtPr><w:sdtContent>'
        + "".join(
            f'<w:tr><w:tc><w:p><w:r><w:t>{a}</w:t></w:r></w:p></w:tc>'
            f'<w:tc><w:p><w:r><w:t>{b}</w:t></w:r></w:p></w:tc></w:tr>'
            for a, b in (("North", "120"), ("South", "95"), ("East", "140"))
        )
        + "</w:sdtContent></w:sdt></w:tbl>"
    )
    body.insert(len(body) - 1, tbl)
    src = tmp / "sdt_rows.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    tnode = _nodes(res.tree, TableNode)[0]
    synth = TableRowNode(
        id=f"{tnode.id}-thead", content=NodeContent(kind=ContentKind.NONE),
        metadata=NodeMetadata(properties={"synthesized": True}), accessibility_flags=[],
        children=[
            TableCellNode(id=f"{tnode.id}-th-{i}", cell_type=TableCellType.HEADER,
                          header_scope=TableHeaderScope.COLUMN,
                          content=NodeContent(kind=ContentKind.TEXT, text=t),
                          metadata=NodeMetadata(properties={"synthesized": True}),
                          children=[], accessibility_flags=[])
            for i, t in enumerate(("Region", "Sales"))
        ],
    )
    tnode.children.insert(0, synth)
    out = tmp / "sdt_rows_fixed.docx"
    wr = write_remediated_docx(src, res.tree, out)
    with zipfile.ZipFile(out) as z:
        droot = etree.fromstring(z.read("word/document.xml"))
    t_el = droot.find(f".//{W}body/{W}tbl")
    order = [etree.QName(c).localname for c in t_el]
    first_tr = t_el.find(f"{W}tr")
    check("sdt rows: the synthesized header row is the table's FIRST row (not after the last)",
          order[:3] == ["tblPr", "tblGrid", "tr"] and first_tr is not None
          and first_tr.find(f"{W}trPr/{W}tblHeader") is not None, str(order))
    check("sdt rows: synthesized cells are not reported as 'not found' noise",
          not [s for s in wr["skipped"] if s.get("reason") == "cell_not_found_in_source"], str(wr["skipped"]))
    res2 = parse_to_tree(str(out))
    t2 = _nodes(res2.tree, TableNode)[0]
    check("sdt rows: re-parse reads Region | Sales as the header row, data rows intact",
          [c.content.text for c in t2.children[0].children] == ["Region", "Sales"] and len(t2.children) == 4,
          str([[c.content.text for c in r.children] for r in t2.children]))

    # ===== 10. Irregular merges Word opens do not fail the whole document =====
    # python-docx's row.cells raises on a vertical-merge continuation in the
    # first row, and on one under a cell spanning a different number of grid
    # columns. One such table made /analyze answer "Failed to parse document"
    # for a file Word opens without complaint.
    def _tc(text, extra=""):
        return f"<w:tc><w:tcPr>{extra}</w:tcPr><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:tc>"

    def _grid(n):
        return "<w:tblGrid>" + '<w:gridCol w:w="1500"/>' * n + "</w:tblGrid>"

    shapes = {
        "vmerge_first_row": _grid(2)
        + "<w:tr>" + _tc("Region", "<w:vMerge/>") + _tc("Sales") + "</w:tr>"
        + "<w:tr>" + _tc("North") + _tc("120") + "</w:tr>"
        + "<w:tr>" + _tc("South") + _tc("95") + "</w:tr>",
        "vmerge_under_span": _grid(3)
        + "<w:tr>" + _tc("Program", '<w:gridSpan w:val="2"/>') + _tc("Budget") + "</w:tr>"
        + "<w:tr>" + _tc("Parks") + _tc("Trails", "<w:vMerge/>") + _tc("40") + "</w:tr>"
        + "<w:tr>" + _tc("Pools") + _tc("Swim") + _tc("25") + "</w:tr>",
    }
    for name, inner in shapes.items():
        d = Document()
        d.core_properties.title = "Irregular table"
        d.add_heading("Irregular table", 1)
        d.element.body.insert(len(d.element.body) - 1,
                              parse_xml(f'<w:tbl {nsdecls("w")}><w:tblPr/>{inner}</w:tbl>'))
        src = tmp / f"{name}.docx"
        d.save(str(src))
        try:
            res = parse_to_tree(str(src))
            legacy = DOCXParser().parse(str(src))
            err = None
        except Exception as exc:  # the regression
            res, legacy, err = None, None, repr(exc)
        check(f"irregular merge ({name}): the document parses (tree + legacy /documents parse)",
              err is None and legacy is not None and legacy.get("tables") == 1, str(err))
        if res is None:
            continue
        run_analyzers(res.tree)
        tbl_nodes = _nodes(res.tree, TableNode)
        check(f"irregular merge ({name}): the table is a node with its 3 rows",
              len(tbl_nodes) == 1 and len(tbl_nodes[0].children) == 3, str([len(t.children) for t in tbl_nodes]))
        for c in tbl_nodes[0].children[0].children:
            c.cell_type = TableCellType.HEADER            # ADD_TABLE_HEADERS' effect
        out = tmp / f"{name}_fixed.docx"
        wr = write_remediated_docx(src, res.tree, out)
        with zipfile.ZipFile(out) as z:
            troot = etree.fromstring(z.read("word/document.xml"))
        rows_x = troot.findall(f".//{W}tbl/{W}tr")
        check(f"irregular merge ({name}): tblHeader lands on the first row only; file re-opens",
              [r.find(f"{W}trPr/{W}tblHeader") is not None for r in rows_x] == [True, False, False]
              and any(a.get("kind") == "table_header_row" for a in wr["applied"])
              and Document(str(out)) is not None, str(wr))

    # ===== 11. No docProps/core.xml: no invented title ========================
    # python-docx CREATES a missing core-properties part pre-filled with title
    # "Word Document" and author "python-docx". Reading through it hid the
    # missing-title finding and the writer saved the invention into the file.
    d = Document()
    d.add_heading("Quarterly report", 1)
    d.add_paragraph("Body text that is long enough to be read as ordinary prose here.")
    full = tmp / "with_core.docx"
    d.save(str(full))
    src = tmp / "no_core.docx"
    with zipfile.ZipFile(full) as zin, zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            data = zin.read(n)
            if n == "docProps/core.xml":
                continue
            if n == "_rels/.rels":
                root_rels = etree.fromstring(data)
                for rel in list(root_rels):
                    if (rel.get("Type") or "").endswith("/core-properties"):
                        root_rels.remove(rel)
                data = etree.tostring(root_rels, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(n, data)
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    check("no core part: no invented title; DOCUMENT_TITLE_MISSING fires",
          not (res.tree.root.metadata.properties or {}).get("title") and "DOCUMENT_TITLE_MISSING" in _codes(res.tree),
          str(((res.tree.root.metadata.properties or {}).get("title"), _codes(res.tree))))
    out = tmp / "no_core_untouched.docx"
    wr = write_remediated_docx(src, res.tree, out)
    with zipfile.ZipFile(out) as z:
        has_core = "docProps/core.xml" in z.namelist()
    check("no core part, nothing approved: nothing applied and no core part invented in the output",
          wr["applied"] == [] and not has_core, str((wr["applied"], has_core)))
    res.tree.root.metadata.properties["title"] = "Quarterly Report 2026"   # SET_DOCUMENT_TITLE's effect
    out = tmp / "no_core_titled.docx"
    wr = write_remediated_docx(src, res.tree, out)
    with zipfile.ZipFile(out) as z:
        core_xml = z.read("docProps/core.xml") if "docProps/core.xml" in z.namelist() else b""
    check("no core part, title approved: the approved title is written, python-docx's defaults are not",
          b"Quarterly Report 2026" in core_xml and b"Word Document" not in core_xml and b"python-docx" not in core_xml
          and [a.get("kind") for a in wr["applied"]] == ["document_title"], core_xml.decode("utf-8", "replace")[:400])
    # An unchanged title is not re-"applied" on every run.
    d = Document()
    d.core_properties.title = "Already titled"
    d.add_paragraph("Body text that is long enough to be read as ordinary prose here.")
    src = tmp / "titled.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    wr = write_remediated_docx(src, res.tree, tmp / "titled_out.docx")
    check("an unchanged title is not reported as applied", wr["applied"] == [], str(wr["applied"]))

    # ===== 12. Findings say which PAGE they are on (as Word last laid it out) ==
    # A DOCX has no pages of its own, so every DOCX finding had page=None and
    # the review screen said "no page available". Word writes a
    # <w:lastRenderedPageBreak/> where each page began at its last save and
    # the page count into docProps/app.xml; when the two agree the parser
    # reports each node's page. When they do not (a file python-docx or
    # another tool wrote: template "Pages 1", no markers) pages stay unknown.
    def _lrpb_run():
        r = OxmlElement("w:r")
        r.append(OxmlElement("w:lastRenderedPageBreak"))
        return r

    def _set_app_pages(path: Path, n: int) -> None:
        def fn(root):
            for el in root.iter():
                if etree.QName(el).localname == "Pages":
                    el.text = str(n)
        _rewrite_part(path, "docProps/app.xml", fn)

    d = Document()
    d.add_heading("Annual Report", 1)
    d.add_paragraph("Page one text that is long enough to be read as ordinary prose.")
    p2 = d.add_paragraph()
    p2._p.append(_lrpb_run())                       # page 2 starts here
    p2.add_run("Page two opens with this sentence, then a chart. ")
    p2.add_run().add_picture(_png(12), width=Inches(1))
    p3 = d.add_paragraph()
    p3._p.append(_lrpb_run())                       # page 3 starts here
    p3.add_run("Page three: see the figures ")
    p3rid = d.part.relate_to("https://example.gov/figures",
                             "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                             is_external=True)
    hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), p3rid)
    hr = OxmlElement("w:r"); ht = OxmlElement("w:t"); ht.text = "here"; hr.append(ht); hl.append(hr)
    p3._p.append(hl)
    t = d.add_table(rows=3, cols=2)
    for r in range(3):
        for c in range(2):
            t.cell(r, c).text = f"v{r}{c}"
    src = tmp / "paged.docx"
    d.save(str(src))
    _set_app_pages(src, 3)
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    by_type = {}
    for n in iter_reading_order(res.tree.root):
        by_type.setdefault(type(n).__name__, []).append(n)
    pages = {
        "heading": [n.metadata.page for n in by_type.get("HeadingNode", [])],
        "image": [n.metadata.page for n in by_type.get("ImageNode", [])],
        "link": [n.metadata.page for n in by_type.get("LinkNode", [])],
        "table": [n.metadata.page for n in by_type.get("TableNode", [])],
    }
    check("pages: heading p1, picture p2, link p3, table p3",
          pages == {"heading": [1], "image": [2], "link": [3], "table": [3]}, str(pages))
    img = by_type["ImageNode"][0]
    alt_flag = next(f for f in img.accessibility_flags if f.code.value == "MISSING_ALT_TEXT")
    check("pages: the finding's evidence carries the page (what the review screen shows)",
          _collect_evidence(img, alt_flag).get("page") == 2)
    # Word's count disagrees with the markers (stale): no pages at all.
    _set_app_pages(src, 7)
    res = parse_to_tree(str(src))
    check("pages: markers that disagree with Word's own page count -> pages unknown, never guessed",
          all(n.metadata.page is None for n in iter_reading_order(res.tree.root)))
    # A python-docx-made file: template "Pages 1" and no markers.
    d = Document()
    for i in range(40):
        d.add_paragraph(f"Paragraph {i} of a long generated file that no word processor paginated.")
    d.add_picture(_png(13), width=Inches(1))
    src = tmp / "unpaged.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    check("pages: a generated file with no render marks gets NO page numbers (not 'page 1' for all)",
          all(n.metadata.page is None for n in iter_reading_order(res.tree.root)))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
