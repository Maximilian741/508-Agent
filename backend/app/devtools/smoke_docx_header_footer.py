"""Smoke: page headers/footers are audited AND fixed, and every writer index
agrees with the parser about which element each node id names.

Before this, ``word/header*.xml`` / ``word/footer*.xml`` were never opened:
the agency logo in the letterhead — on every printed page, the single most
common image in a government DOCX — got no finding and no alt, a low-contrast
grey footer went unmeasured, and a "click here" privacy link in the footer was
invisible. A re-scan of our own output then called the file clean.

Worse, three indexes had drifted apart from the parser without anyone
noticing, because nothing compared them element by element:

  * the writer counted images in EVERY body paragraph, the parser skipped
    images inside headings and list items -> after the first such image,
    alt text written for image N landed on image N+1's picture;
  * links inside list items ("For details, click here" in a bullet) were never
    parsed, so they were never flagged or fixed;
  * low-contrast list items were flagged but the writer had no index for them,
    so the recolour never landed.

Pinned here with distinct bytes/targets per element, so "the right element"
is checked by identity, not by count:

  * header/footer images, links and text become nodes (header, first-page
    header and footer; a header shared by two sections is audited ONCE;
    a letterhead layout table is not reported as a data table)
  * alt written for each image lands on exactly that picture's docPr — body,
    heading, list item, table cell, header — and nowhere else
  * link text written for each link lands on exactly that link
  * a low-contrast header line and a low-contrast list item are recoloured,
    and the writer confirms both
  * re-parsing the output: the alt/links/contrast findings are gone and no
    new finding appears
  * only what Word PRINTS is audited: a first-page header whose "Different
    first page" was switched off, and an even-page header without "Different
    odd & even pages", are not (Word keeps both parts in the file); with the
    option on, the even-page logo is audited and fixed
  * a footer text box's link and words are read, flagged and rewritten in
    place

Usage:
    python -m app.devtools.smoke_docx_header_footer
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_dhf_')}/s.db")
os.environ.pop("OPENAI_API_KEY", None)
os.environ.pop("ANTHROPIC_API_KEY", None)
os.environ["SEMANTIC_PROVIDER"] = "heuristic"

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.oxml import OxmlElement, parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402
from docx.shared import Inches, RGBColor  # noqa: E402
from lxml import etree  # noqa: E402
from PIL import Image  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.models.accessibility import (  # noqa: E402
    ImageNode,
    LinkNode,
    ListItemNode,
    ParagraphNode,
    TableNode,
    iter_reading_order,
)
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.docx_writer import write_remediated_docx  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
POL = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)
_SEQ = [0]


def _png() -> io.BytesIO:
    """A PNG with bytes unique to this call, so Word gives it its own part."""
    _SEQ[0] += 1
    b = io.BytesIO()
    Image.new("RGB", (40 + _SEQ[0], 20), (10 * _SEQ[0] % 255, 90, 160)).save(b, "PNG")
    b.seek(0)
    return b


def _hyperlink(paragraph, part, url: str, text: str) -> None:
    rid = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    hl.append(r)
    paragraph._p.append(hl)


def _grey(run, hex6="AAAAAA"):
    run.font.color.rgb = RGBColor.from_string(hex6)


def _build(path: Path) -> None:
    d = Document()
    sec = d.sections[0]
    sec.different_first_page_header_footer = True
    # Default header: a letterhead LAYOUT table (logo | agency name) + grey text.
    ht = sec.header.add_table(rows=1, cols=2, width=Inches(6))
    ht.cell(0, 0).paragraphs[0].add_run().add_picture(_png(), width=Inches(1))
    ht.cell(0, 1).paragraphs[0].add_run("Department of Examples")
    hp = sec.header.paragraphs[0]
    _grey(hp.add_run("Internal memo - not for distribution"))
    # First-page header: its own logo, inline.
    sec.first_page_header.paragraphs[0].add_run().add_picture(_png(), width=Inches(1.2))
    # Footer: a generic link.
    fp = sec.footer.paragraphs[0]
    fp.add_run("Questions about privacy? ")
    _hyperlink(fp, sec.footer.part, "https://example.gov/privacy", "click here")

    d.add_heading("Records Retention Schedule", 1)
    # An image INSIDE a heading paragraph (a logo set in the title line).
    d.paragraphs[-1].add_run().add_picture(_png(), width=Inches(0.4))
    d.add_paragraph("This memo sets out how long each class of record is kept by the department.")
    d.add_picture(_png(), width=Inches(2))                         # body image
    li = d.add_paragraph("Keep payroll records for seven years.", style="List Bullet")
    _grey(li.runs[0])                                               # low-contrast list item
    li2 = d.add_paragraph("For the full schedule, ", style="List Bullet")
    _hyperlink(li2, d.part, "https://example.gov/schedule", "click here")  # link in a list item
    li3 = d.add_paragraph("Retention icon: ", style="List Bullet")
    li3.add_run().add_picture(_png(), width=Inches(0.3))            # image in a list item
    for item in (li, li2, li3):
        # Word's own bullets: numbering on the paragraph itself (w:numPr).
        numpr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0"); numpr.append(ilvl)
        nid = OxmlElement("w:numId"); nid.set(qn("w:val"), "1"); numpr.append(nid)
        item._p.get_or_add_pPr().append(numpr)
    t = d.add_table(rows=3, cols=2)
    t.cell(0, 0).text = "Record"
    t.cell(0, 1).text = "Years"
    t.cell(1, 0).text = "Payroll"
    t.cell(1, 1).paragraphs[0].add_run().add_picture(_png(), width=Inches(0.3))  # cell image
    t.cell(2, 0).text = "Leave"
    t.cell(2, 1).text = "3"
    d.add_paragraph("Contact the records office with any questions about this schedule.")
    d.add_picture(_png(), width=Inches(1))                          # a later body image
    # A second section that REUSES the first section's header (linked).
    d.add_section(WD_SECTION.NEW_PAGE)
    d.add_paragraph("Annex text in a second section that inherits the same header.")
    d.save(str(path))


def _drawings(xml_bytes: bytes):
    """[(rid, descr)] for every picture in a part, in document order."""
    root = etree.fromstring(xml_bytes)
    out = []
    for wrapper in root.iter(f"{WP}inline", f"{WP}anchor"):
        blip = wrapper.find(f".//{A}blip")
        dp = wrapper.find(f"{WP}docPr")
        out.append((blip.get(f"{R}embed") if blip is not None else None, dp.get("descr") if dp is not None else None))
    return out


def _rels(z, part_name: str):
    rels_name = part_name.replace("word/", "word/_rels/") + ".rels"
    root = etree.fromstring(z.read(rels_name))
    return {r.get("Id"): r.get("Target") for r in root.iter(f"{PKG_REL}Relationship")}


def _links(xml_bytes: bytes, rels):
    root = etree.fromstring(xml_bytes)
    return [
        (rels.get(h.get(f"{R}id")), "".join(t.text or "" for t in h.iter(f"{W}t")))
        for h in root.iter(f"{W}hyperlink")
    ]


def _media_name(z, part_name: str, rid: str) -> str:
    return _rels(z, part_name).get(rid, "")


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="508_dhf_"))
    src = tmp / "memo.docx"
    out = tmp / "memo_fixed.docx"
    _build(src)

    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    nodes = list(iter_reading_order(res.tree.root))
    images = [n for n in nodes if isinstance(n, ImageNode)]
    links = [n for n in nodes if isinstance(n, LinkNode)]

    def story(n):
        return (n.metadata.properties or {}).get("docx_story")

    hdr_imgs = [n for n in images if story(n) in ("header", "footer")]
    check("header images are nodes (letterhead-table logo + first-page logo)", len(hdr_imgs) == 2,
          str([(n.id, story(n)) for n in images]))
    check("body images: heading + body + list item + cell + later body = 5",
          len([n for n in images if story(n) not in ("header", "footer")]) == 5,
          str([(n.id, story(n)) for n in images]))
    check("every image is flagged MISSING_ALT_TEXT (7)",
          sum(1 for n in images for f in n.accessibility_flags if f.code.value == "MISSING_ALT_TEXT") == 7)
    footer_links = [n for n in links if story(n) == "footer"]
    check("footer 'click here' link is a node and is flagged",
          len(footer_links) == 1 and any(f.code.value == "LINK_TEXT_NON_DESCRIPTIVE" for f in footer_links[0].accessibility_flags),
          str([(n.id, n.content.text, story(n)) for n in links]))
    list_links = [n for n in links if n.target == "https://example.gov/schedule"]
    check("link inside a bullet list item is a node and is flagged",
          len(list_links) == 1 and any(f.code.value == "LINK_TEXT_NON_DESCRIPTIVE" for f in list_links[0].accessibility_flags))
    items = [n for n in nodes if isinstance(n, ListItemNode)]
    check("the three bullets are real list items carrying their link / picture",
          len(items) == 3
          and any(isinstance(c, LinkNode) for c in items[1].children)
          and any(isinstance(c, ImageNode) for c in items[2].children),
          str([(n.id, [c.id for c in n.children]) for n in items]))
    check("the link's snippet is its sentence", (list_links[0].metadata.properties or {}).get("snippet")
          == "For the full schedule, click here" if list_links else False)
    grey_hdr = [n for n in nodes if isinstance(n, ParagraphNode) and story(n) == "header"
                and any(f.code.value == "LOW_CONTRAST_TEXT" for f in n.accessibility_flags)]
    check("low-contrast header line is flagged", len(grey_hdr) == 1)
    check("header text is audited once even though two sections share the header",
          sum(1 for n in nodes if (n.content.text or "") == "Internal memo - not for distribution") == 1)
    tables = [n for n in nodes if isinstance(n, TableNode)]
    check("the letterhead layout table is NOT reported as a data table", len(tables) == 1,
          str([t.id for t in tables]))
    check("header/footer nodes say where they are (story + variant) for location",
          all((n.metadata.properties or {}).get("docx_story_variant") in ("default", "first", "even")
              for n in hdr_imgs + footer_links))
    check("the letterhead logo and the footer link carry the surrounding text as a snippet",
          any("Department of Examples" in ((n.metadata.properties or {}).get("snippet") or "") for n in hdr_imgs)
          and "privacy" in ((footer_links[0].metadata.properties or {}).get("snippet") or "").lower())
    first_page = [n for n in hdr_imgs if (n.metadata.properties or {}).get("docx_story_variant") == "first"]
    check("a logo with no words around it gets NO invented snippet",
          len(first_page) == 1 and not (first_page[0].metadata.properties or {}).get("snippet"))

    # Every image gets a DISTINCT alt; every link a distinct name; then write.
    for n in images:
        n.alt_text = f"ALT {n.id}"
    for n in links:
        n.content.text = f"LINK {n.id}"
    # Contrast: run the real deterministic executor on the contrast flags only.
    plans = [p for p in plan_remediations(res.tree, POL) if p.flag.code.value == "LOW_CONTRAST_TEXT"]
    execs = execute_plans(res.tree, plans)
    check("two contrast fixes executed (header line + list item)",
          len([e for e in execs if e.status.value == "success"]) == 2,
          str([(e.target_node_id, e.status.value, e.notes) for e in execs]))
    wr = write_remediated_docx(src, res.tree, out)
    applied = wr["applied"]

    alt_written = {a["target_id"] for a in applied if a.get("kind") == "image_alt_text"}
    check("writer confirms every image alt (7 of 7)", alt_written == {n.id for n in images},
          f"missing={sorted({n.id for n in images} - alt_written)} skipped={wr['skipped']}")
    link_written = {a["target_id"] for a in applied if a.get("kind") == "link_text"}
    check("writer confirms every link rewrite", link_written == {n.id for n in links},
          f"missing={sorted({n.id for n in links} - link_written)}")
    contrast_written = {a["target_id"] for a in applied if a.get("action") == "FIX_CONTRAST"}
    check("writer confirms both recolours", contrast_written == {e.target_node_id for e in execs if e.status.value == "success"},
          f"{contrast_written} vs {[e.target_node_id for e in execs]}")

    # Identity: each alt landed on THE picture whose bytes the node points at.
    by_rid_story = {}
    for n in images:
        props = n.metadata.properties or {}
        by_rid_story[(props.get("docx_part"), props.get("image_rid"))] = f"ALT {n.id}"
    with zipfile.ZipFile(out) as z:
        parts = [n for n in z.namelist() if n.startswith("word/") and n.endswith(".xml")
                 and ("header" in n or "footer" in n or n == "word/document.xml")]
        mism = []
        seen = 0
        for part in parts:
            for rid, descr in _drawings(z.read(part)):
                seen += 1
                want = by_rid_story.get(("/" + part, rid))
                if want != descr:
                    mism.append((part, rid, descr, want))
        check("every picture carries exactly ITS node's alt (no off-by-one)", not mism and seen == 7,
              f"seen={seen} mismatches={mism}")
        doc_links = _links(z.read("word/document.xml"), _rels(z, "word/document.xml"))
        footer_part = next(p for p in parts if "footer" in p)
        ftr_links = _links(z.read(footer_part), _rels(z, footer_part))
    want_links = {n.target: f"LINK {n.id}" for n in links}
    got_links = dict(doc_links + ftr_links)
    check("every link carries exactly ITS node's text", got_links == want_links, f"{got_links} vs {want_links}")

    # Re-parse the OUTPUT: fixed findings gone, no new ones.
    res2 = parse_to_tree(str(out))
    run_analyzers(res2.tree)
    codes2 = [f.code.value for n in iter_reading_order(res2.tree.root) for f in n.accessibility_flags]
    codes1 = [f.code.value for n in nodes for f in n.accessibility_flags]
    check("re-parse: no MISSING_ALT_TEXT left", "MISSING_ALT_TEXT" not in codes2, str(codes2))
    check("re-parse: no LOW_CONTRAST_TEXT left", "LOW_CONTRAST_TEXT" not in codes2, str(codes2))
    check("re-parse: no new finding codes", not (set(codes2) - set(codes1)), str(set(codes2) - set(codes1)))
    reopened = Document(str(out))
    check("output re-opens; body paragraph count unchanged",
          len(reopened.paragraphs) == len(Document(str(src)).paragraphs))

    # The real engine end to end: every GENERATE_ALT_TEXT success must be
    # confirmed by the writer (alt is not writer-confirmed in the pipeline,
    # so index agreement is the only thing standing between a "success" and
    # a charge for a picture that never changed).
    res3 = parse_to_tree(str(src))
    run_analyzers(res3.tree)
    plans3 = plan_remediations(res3.tree, POL)
    execs3 = execute_plans(res3.tree, plans3)
    wr3 = write_remediated_docx(src, res3.tree, tmp / "memo_engine.docx")
    ok_alt = {e.target_node_id for e in execs3 if e.action_code.value == "GENERATE_ALT_TEXT" and e.status.value == "success"}
    wrote = {a["target_id"] for a in wr3["applied"] if a.get("kind") == "image_alt_text"}
    check("engine run: every successful alt execution reached the file", ok_alt <= wrote,
          f"success-but-unwritten={sorted(ok_alt - wrote)}")
    ok_link = {e.target_node_id for e in execs3 if e.action_code.value == "IMPROVE_LINK_TEXT" and e.status.value == "success"}
    wrote_l = {a["target_id"] for a in wr3["applied"] if a.get("kind") == "link_text"}
    check("engine run: every successful link rewrite reached the file", ok_link <= wrote_l,
          f"success-but-unwritten={sorted(ok_link - wrote_l)}")

    # ===== Only what Word PRINTS is audited; footer text boxes are read =====
    # Word keeps a first-page header part (and its reference) after "Different
    # first page" is switched off, and an even-page header after "Different
    # odd & even" is switched off. Neither is printed; a logo in one got a
    # MISSING_ALT_TEXT finding and a charge to fix a picture no reader meets.
    # A footer callout in a text box ("Privacy: click here") was never read.
    src2 = tmp / "variants.docx"
    d = Document()
    s0 = d.sections[0]
    s0.different_first_page_header_footer = True
    s0.first_page_header.paragraphs[0].add_run().add_picture(_png(), width=Inches(1))  # will be unprinted
    s0.even_page_header.paragraphs[0].add_run().add_picture(_png(), width=Inches(1))   # never printed
    s0.header.paragraphs[0].add_run("Department of Examples")
    fp = s0.footer.paragraphs[0]
    fp.add_run("Page footer ")
    rid = s0.footer.part.relate_to(
        "https://example.gov/privacy-policy",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    wps = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    fp._p.append(parse_xml(
        f'<w:r {nsdecls("w", "wp", "a", "r")} xmlns:wps="{wps}"><w:drawing><wp:anchor distT="0" distB="0" '
        'distL="0" distR="0" simplePos="0" relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" '
        'allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>0'
        '</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset>'
        '</wp:positionV><wp:extent cx="914400" cy="457200"/><wp:wrapNone/><wp:docPr id="90" name="Text Box 90"/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr/><wps:txbx><w:txbxContent><w:p><w:r>'
        '<w:t xml:space="preserve">Privacy questions: </w:t></w:r>'
        f'<w:hyperlink r:id="{rid}"><w:r><w:t>click here</w:t></w:r></w:hyperlink></w:p>'
        '<w:p><w:r><w:t>Records Office, Room 12</w:t></w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr/>'
        '</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'))
    d.add_paragraph("Body text of a memo that is long enough to read as ordinary prose.")
    d.save(str(src2))
    d = Document(str(src2))
    d.sections[0].different_first_page_header_footer = False   # the part and its reference stay
    d.save(str(src2))
    with zipfile.ZipFile(src2) as z:
        dx = z.read("word/document.xml")
    check("fixture: first/even header references are still in the file, titlePg is off",
          b'w:type="first"' in dx and b'w:type="even"' in dx and b"titlePg" not in dx)
    r2 = parse_to_tree(str(src2))
    run_analyzers(r2.tree)
    n2 = list(iter_reading_order(r2.tree.root))
    imgs2 = [n for n in n2 if isinstance(n, ImageNode)]
    check("unprinted first-page and even-page headers are NOT audited (no image, no finding)",
          not imgs2, str([(n.id, (n.metadata.properties or {}).get("docx_story_variant")) for n in imgs2]))
    tb_links = [n for n in n2 if isinstance(n, LinkNode) and (n.metadata.properties or {}).get("docx_story") == "footer"]
    check("the footer text box's 'click here' is ONE node, flagged, with its sentence as snippet",
          len(tb_links) == 1 and any(f.code.value == "LINK_TEXT_NON_DESCRIPTIVE" for f in tb_links[0].accessibility_flags)
          and (tb_links[0].metadata.properties or {}).get("snippet") == "Privacy questions: click here",
          str([(n.id, n.content.text, (n.metadata.properties or {}).get("snippet")) for n in tb_links]))
    check("the footer text box's other line is read as text",
          any(isinstance(n, ParagraphNode) and n.content.text == "Records Office, Room 12" for n in n2))
    for n in tb_links:
        n.content.text = "Privacy policy"
    out2 = tmp / "variants_fixed.docx"
    wr2 = write_remediated_docx(src2, r2.tree, out2)
    check("writer confirms the footer text-box link rewrite",
          {a["target_id"] for a in wr2["applied"] if a.get("kind") == "link_text"} == {n.id for n in tb_links},
          str(wr2))
    with zipfile.ZipFile(out2) as z:
        fparts = [p for p in z.namelist() if p.startswith("word/footer")]
        got = [lk for p in fparts for lk in _links(z.read(p), _rels(z, p))]
    check("the rewrite landed on that link in the footer part",
          ("https://example.gov/privacy-policy", "Privacy policy") in got, str(got))
    r3 = parse_to_tree(str(out2))
    run_analyzers(r3.tree)
    codes3 = [f.code.value for n in iter_reading_order(r3.tree.root) for f in n.accessibility_flags]
    check("re-parse: the footer link finding is gone", "LINK_TEXT_NON_DESCRIPTIVE" not in codes3, str(codes3))

    # Different odd & even pages ON: the even-page header IS printed and audited.
    d = Document(str(src2))
    d.settings.odd_and_even_pages_header_footer = True
    src3 = tmp / "even_on.docx"
    d.save(str(src3))
    r4 = parse_to_tree(str(src3))
    run_analyzers(r4.tree)
    imgs4 = [n for n in iter_reading_order(r4.tree.root) if isinstance(n, ImageNode)]
    check("with 'Different odd & even pages' on, the even-page header logo is audited (and only it)",
          [(n.metadata.properties or {}).get("docx_story_variant") for n in imgs4] == ["even"]
          and any(f.code.value == "MISSING_ALT_TEXT" for f in imgs4[0].accessibility_flags),
          str([(n.id, (n.metadata.properties or {}).get("docx_story_variant")) for n in imgs4]))
    if imgs4:
        imgs4[0].alt_text = "Agency seal"
    wr4 = write_remediated_docx(src3, r4.tree, tmp / "even_on_fixed.docx")
    check("... and its alt lands", [a["target_id"] for a in wr4["applied"] if a.get("kind") == "image_alt_text"]
          == [n.id for n in imgs4], str(wr4))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
