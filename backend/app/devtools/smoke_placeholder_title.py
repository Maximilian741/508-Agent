"""Smoke: generic/placeholder document titles are detected AND auto-fixed.

A title like "Document1", "PowerPoint Presentation", "Untitled", or the bare
filename is announced verbatim by a screen reader — as useless as no title.
Previously the engine only flagged a truly EMPTY title. This pins the upgrade:

  analyzer -> DOCUMENT_TITLE_MISSING now fires on placeholder titles too
  executor -> SET_DOCUMENT_TITLE OVERWRITES a placeholder (derives a real title
              from the document's first heading), but only when it can derive
              something genuinely better (never junk-for-junk -> honest)
  writer   -> the new title persists to the file's core properties
  honesty  -> a REAL title is left untouched; re-parsing the OUTPUT shows the
              flag cleared and a real title in the bytes.

Usage:
    python -m app.devtools.smoke_placeholder_title
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_pt_')}/s.db")

from docx import Document  # noqa: E402

from app.analyzers.helpers import is_placeholder_title  # noqa: E402
from app.analyzers.registry import run_analyzers  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.docx_writer import write_remediated_docx  # noqa: E402

FLAG = "DOCUMENT_TITLE_MISSING"


def _has_title_flag(tree) -> bool:
    return any(f.code.value == FLAG for f in tree.root.accessibility_flags)


def _title(tree):
    return (tree.root.metadata.properties or {}).get("title")


def main() -> int:
    failures = 0

    def check(name, cond, extra="") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="placeholder_title_"))
    apply_policy = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)

    # --- helper precision -----------------------------------------------------
    check("helper: 'Document1' is placeholder", is_placeholder_title("Document1"))
    check("helper: 'PowerPoint Presentation' is placeholder", is_placeholder_title("PowerPoint Presentation"))
    check("helper: 'Microsoft Word - X' is placeholder", is_placeholder_title("Microsoft Word - Memo"))
    check("helper: title==filename is placeholder", is_placeholder_title("report.docx", "report.docx"))
    check("helper: real title is NOT placeholder", not is_placeholder_title("Q3 Earnings Summary"))
    check("helper: empty is NOT placeholder (handled as missing)", not is_placeholder_title(""))

    # --- A: placeholder title + a real heading -> detected + overwritten ------
    d = Document()
    d.core_properties.title = "Document1"
    d.add_heading("Annual Sustainability Report", level=1)
    d.add_paragraph("Body.")
    src = tmp / "deck.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    check("placeholder 'Document1' flagged DOCUMENT_TITLE_MISSING", _has_title_flag(res.tree),
          f"title={_title(res.tree)!r}")
    plans = [p for p in plan_remediations(res.tree, apply_policy) if p.flag.code.value == FLAG]
    execs = execute_plans(res.tree, plans)
    ok = [e for e in execs if e.status.value == "success"]
    check("SET_DOCUMENT_TITLE overwrites the placeholder", len(ok) == 1,
          str([(e.status.value, e.notes) for e in execs]))
    check("tree title replaced with the real heading", _title(res.tree) == "Annual Sustainability Report",
          f"got {_title(res.tree)!r}")
    out = tmp / "deck_fixed.docx"
    write_remediated_docx(src, res.tree, out)
    res2 = parse_to_tree(str(out))
    run_analyzers(res2.tree)
    check("output: real title persisted to bytes", _title(res2.tree) == "Annual Sustainability Report",
          f"got {_title(res2.tree)!r}")
    check("output: no title flag (idempotent)", not _has_title_flag(res2.tree))

    # --- B: a REAL title is left alone ---------------------------------------
    d2 = Document()
    d2.core_properties.title = "Q3 Earnings Summary"
    d2.add_heading("Some Heading", level=1)
    src2 = tmp / "real.docx"
    d2.save(str(src2))
    r2 = parse_to_tree(str(src2))
    run_analyzers(r2.tree)
    check("real title: NOT flagged", not _has_title_flag(r2.tree), f"title={_title(r2.tree)!r}")

    # --- C: title equal to the filename -> placeholder ------------------------
    d3 = Document()
    src3 = tmp / "quarterly.docx"
    d3.core_properties.title = "quarterly.docx"
    d3.add_heading("Quarterly Numbers", level=1)
    d3.save(str(src3))
    r3 = parse_to_tree(str(src3))
    run_analyzers(r3.tree)
    check("title==filename: flagged", _has_title_flag(r3.tree), f"title={_title(r3.tree)!r}")

    # --- D: placeholder but nothing better derivable -> honest SKIP -----------
    d4 = Document()
    src4 = tmp / "untitled.docx"
    d4.core_properties.title = "Untitled"
    d4.add_paragraph("Just body text, no heading at all.")
    d4.save(str(src4))
    r4 = parse_to_tree(str(src4))
    run_analyzers(r4.tree)
    check("placeholder 'Untitled' flagged", _has_title_flag(r4.tree))
    pl4 = [p for p in plan_remediations(r4.tree, apply_policy) if p.flag.code.value == FLAG]
    ex4 = execute_plans(r4.tree, pl4)
    # filename stem "untitled" -> derived title would also be "Untitled" -> skip.
    succeeded = [e for e in ex4 if e.status.value == "success"]
    check("no derivable better title -> honest skip (no junk-for-junk)", len(succeeded) == 0,
          str([(e.status.value, e.notes) for e in ex4]))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
