"""Smoke: the legacy /documents fix flow never hands out remediated bytes.

Pre-launch audit finding (critical, reproduced): POST /documents/{id}/apply-fixes,
/finalize and the fixed / rebuilt downloads wrote fixes into the uploaded file
(a title guessed from the filename, a hard-coded en-US language) and served the
result with NO credit check. A zero-balance account got the paid deliverable for
free while /pipeline/remediate answered the same account 402. On top of that,
/documents/{id}/ai-review and POST /remediate were free, unthrottled AI proxies.

Those routes are retired (410). This drives a ZERO-balance account through the
whole legacy click path and pins:

  * every legacy fix / finalize / AI-review / fixed-download route -> 410, and no
    response body is a document;
  * nothing is written: no fixed artifact on disk, no fixedPath on the record;
  * a fixed file left over from BEFORE the retirement is served by no route and
    is not smuggled out inside an evidence bundle;
  * the read-only routes still work (upload, scan, issues, original download);
  * auth and tenant isolation still run first (401 anonymous, 404 other account);
  * the balance stays 0 and the ledger stays empty;
  * neither the paid inference provider nor the direct OpenAI helpers the legacy
    routes used is ever called.

Run: python -m app.devtools.smoke_legacy_remediation_retired
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import time
import zipfile
from pathlib import Path

_TMP = Path(tempfile.mkdtemp(prefix="508_smoke_legacy_retired_"))
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/smoke.db"
os.environ["STORAGE_LOCAL_ROOT"] = str(_TMP / "storage")
os.environ["MATERIALIZED_ROOT"] = str(_TMP / "materialized")
for _key in ("OPENAI_API_KEY", "ANTHROPIC_API_KEY", "SEMANTIC_PROVIDER"):
    os.environ.pop(_key, None)

from fastapi.testclient import TestClient  # noqa: E402

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# (method, path under /documents/{id}/, json body)
_RETIRED_FIX_ROUTES = [
    ("POST", "apply-fixes", None),
    ("POST", "apply-fixes?mode=rebuild", None),
    ("POST", "finalize", None),
    ("POST", "ai-review", {"mode": "propose", "maxItems": 5}),
    ("POST", "ai-review", {"mode": "apply", "maxItems": 5}),
]
_RETIRED_DOWNLOAD_ROUTES = [
    ("GET", "file-fixed", None),
    ("GET", "pdf-fixed", None),
    ("HEAD", "pdf-fixed", None),
    ("GET", "download?variant=fixed", None),
    ("GET", "pdf-rebuilt", None),
    ("HEAD", "pdf-rebuilt", None),
]
_LEFTOVER_MARKER = b"LEFTOVER-REMEDIATED-BYTES"


def _pdf_bytes() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.core_properties.title = ""
    doc.add_heading("Quarterly Report", level=1)
    doc.add_paragraph("Revenue grew in every region this quarter.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _looks_like_document(body: bytes) -> bool:
    return body.startswith(b"%PDF") or body.startswith(b"PK") or _LEFTOVER_MARKER in body


def _install_spies() -> dict:
    """Count (and refuse) every road to a paid model the legacy routes had."""
    import app.ai.alt_text_suggester as alt_text_suggester
    import app.ai.auto_review as auto_review
    import app.ai.semantic_inference as si

    hits = {"paid_provider": 0, "direct_openai": 0}

    class _PaidProviderSpy(si.HeuristicProvider):
        name = "paid-spy"

        def _hit(self, *_args, **_kwargs):
            hits["paid_provider"] += 1
            raise AssertionError("a free legacy path reached the PAID provider")

        alt_text = link_text = document_title = document_language = table_caption = _hit

    def _direct(*_args, **_kwargs):
        hits["direct_openai"] += 1
        raise AssertionError("a legacy route called OpenAI directly")

    si.build_default_provider = lambda *a, **k: _PaidProviderSpy()
    alt_text_suggester._openai_alt_text = _direct
    auto_review._openai_propose = _direct
    return hits


def _balance(client, headers) -> tuple:
    r = client.get("/credits/balance", headers=headers)
    assert r.status_code == 200, r.text
    body = r.json()
    return int(body["balance"]), len(body.get("history") or [])


def _wait_job(docs_api, job_id: str, timeout: float = 60.0) -> str:
    deadline = time.monotonic() + timeout
    status = ""
    while time.monotonic() < deadline:
        status = str((docs_api.JOBS.get(job_id) or {}).get("status") or "")
        if status in {"done", "error"}:
            return status
        time.sleep(0.2)
    return status


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    hits = _install_spies()

    from app.main import app
    from app.api import documents as docs_api

    # The legacy flow's runtime dirs are hard-coded under backend/.runtime.
    docs_api.UPLOADS_DIR = _TMP / "uploads"
    docs_api.FIXED_DIR = _TMP / "fixed"
    docs_api.RESULTS_DIR = _TMP / "results"

    client = TestClient(app)

    def sign_in(email: str) -> dict:
        r = client.post("/auth/sign-in", json={"email": email, "displayName": "LR", "password": "legacyretired1"})
        assert r.status_code == 200, r.text
        return {"Authorization": f"Bearer {r.json()['token']}"}

    auth = sign_in("legacy-retired-smoke@example.com")
    outsider = sign_in("legacy-retired-outsider@example.com")
    check("account never granted credits starts at 0 with an empty ledger", _balance(client, auth) == (0, 0), str(_balance(client, auth)))

    last_doc = None
    for fname, data, mime in (
        ("Budget_2026.pdf", _pdf_bytes(), "application/pdf"),
        ("Quarterly_Report.docx", _docx_bytes(), _DOCX_MIME),
    ):
        up = client.post("/documents/upload", files={"file": (fname, data, mime)}, headers=auth)
        check(f"{fname}: upload still works", up.status_code == 200, up.text[:200])
        if up.status_code != 200:
            continue
        doc_id = up.json()["docId"]
        started = client.post(f"/documents/{doc_id}/scan", json={}, headers=auth)
        check(f"{fname}: scan still starts", started.status_code == 200, started.text[:200])
        job_id = str(started.json().get("jobId") or "")
        status = _wait_job(docs_api, job_id)
        check(f"{fname}: scan completes", status == "done", status)
        issues = client.get(f"/documents/{doc_id}/issues", headers=auth)
        check(f"{fname}: issues still readable", issues.status_code == 200 and isinstance(issues.json(), list), issues.text[:200])

        for method, path, body in _RETIRED_FIX_ROUTES + _RETIRED_DOWNLOAD_ROUTES:
            r = client.request(method, f"/documents/{doc_id}/{path}", json=body, headers=auth)
            # repr of the raw bytes: a regression here returns a PDF/DOCX body,
            # and printing that as text crashes a cp1252 Windows console.
            check(f"{fname}: {method} {path} -> 410", r.status_code == 410, f"got {r.status_code} {r.content[:120]!r}")
            check(f"{fname}: {method} {path} returns no document bytes", not _looks_like_document(r.content))
            if method != "HEAD" and r.status_code == 410:
                check(
                    f"{fname}: {method} {path} points at the paid flow",
                    "/pipeline/remediate" in str(r.json().get("detail")),
                    r.text[:200],
                )

        record = docs_api._get_doc(doc_id) or {}
        check(
            f"{fname}: no fixed/rebuilt artifact recorded on the document",
            not record.get("fixedPath") and not record.get("rebuiltPath") and not record.get("fixReport"),
            str({k: record.get(k) for k in ("fixedPath", "rebuiltPath", "fixReport")}),
        )
        fixed_dir = docs_api.FIXED_DIR / doc_id
        check(f"{fname}: nothing written to the fixed dir", not fixed_dir.exists() or not any(fixed_dir.iterdir()))
        original = client.get(f"/documents/{doc_id}/download", headers=auth)
        check(
            f"{fname}: original download still byte-identical to the upload",
            original.status_code == 200 and original.content == data,
            f"{original.status_code} {len(original.content)} vs {len(data)}",
        )
        check(f"{fname}: anonymous caller still gets 401 first", client.post(f"/documents/{doc_id}/apply-fixes").status_code == 401)
        check(
            f"{fname}: another account still gets 404 (no existence leak)",
            client.post(f"/documents/{doc_id}/apply-fixes", headers=outsider).status_code == 404,
        )
        last_doc = (doc_id, job_id)

    # A file the free flow produced BEFORE the retirement may still be on disk
    # and referenced by the record. No route may serve it, and an evidence
    # bundle must not carry it out either.
    if last_doc:
        doc_id, job_id = last_doc
        leftover = docs_api.FIXED_DIR / doc_id / "fixed.docx"
        leftover.parent.mkdir(parents=True, exist_ok=True)
        leftover.write_bytes(b"PK" + _LEFTOVER_MARKER)
        refs = {"fixedPath": str(leftover), "rebuiltPath": str(leftover)}
        record = dict(docs_api._get_doc(doc_id) or {})
        record.update(refs)
        record["fixReport"] = dict(refs)
        docs_api._save_doc(doc_id, record)
        docs_api.REPO.save_fix_report(doc_id, dict(refs))
        for method, path, _body in _RETIRED_DOWNLOAD_ROUTES:
            r = client.request(method, f"/documents/{doc_id}/{path}", headers=auth)
            check(
                f"leftover artifact: {method} {path} -> 410, not served",
                r.status_code == 410 and _LEFTOVER_MARKER not in r.content,
                f"got {r.status_code}",
            )
        bundle = client.post(
            f"/jobs/{job_id}/evidence-bundle",
            json={"includeFixedIfAvailable": True, "includeRebuiltIfAvailable": True},
            headers=auth,
        )
        check("evidence bundle still builds", bundle.status_code == 200, bundle.text[:200])
        if bundle.status_code == 200:
            zipped = client.get(bundle.json()["downloadUrl"], headers=auth)
            check("evidence bundle downloads", zipped.status_code == 200, zipped.text[:200])
            if zipped.status_code == 200:
                with zipfile.ZipFile(io.BytesIO(zipped.content)) as zf:
                    leaked = [
                        name
                        for name in zf.namelist()
                        if "artifacts/fixed" in name or "artifacts/rebuilt" in name or _LEFTOVER_MARKER in zf.read(name)
                    ]
                check("evidence bundle carries no fixed/rebuilt artifact", not leaked, str(leaked))

    # The legacy JSON pair: /scan is read-only analysis and stays; /remediate is gone.
    scanned = client.post(
        "/scan",
        json={"documentId": "legacy", "sourceFormat": "pdf", "content": json.dumps({"title": "", "images": [{"alt_text": ""}]})},
        headers=auth,
    )
    check("legacy JSON /scan still answers (read-only)", scanned.status_code == 200, scanned.text[:200])
    remediated = client.post("/remediate", json={"targetNodeId": "img-1", "actionCode": "GENERATE_ALT_TEXT"}, headers=auth)
    check("legacy POST /remediate -> 410", remediated.status_code == 410, remediated.text[:200])

    check("balance still 0 and ledger still empty after the whole legacy path", _balance(client, auth) == (0, 0), str(_balance(client, auth)))
    check("paid inference provider never called", hits["paid_provider"] == 0, str(hits))
    check("legacy direct-OpenAI helpers never called", hits["direct_openai"] == 0, str(hits))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
