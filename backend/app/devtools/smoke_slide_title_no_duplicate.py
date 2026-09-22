"""Smoke: SET_SLIDE_TITLE never puts a second visible copy of the title on a slide.

The old writer inserted a NEW title placeholder carrying the derived text
while the text box it came from stayed where it was. The placeholder had no
geometry of its own, so it inherited the layout's title position and the
words appeared twice — on 40 of 200 slides of a real deck — and a data
table's first cell ("North") became one slide's title. Pinned here, through
the real detect -> plan -> execute -> write path:

  * a plain top-level text box holding the title is PROMOTED: it becomes the
    title placeholder (same shape id, same xfrm), its look is frozen onto it
    (explicit size/font/colour/alignment/anchor/insets/no-fill), and the
    slide carries the title text exactly once;
  * a slide whose layout left an EMPTY title placeholder: the empty
    placeholder is dropped and the text box promoted — still one copy;
  * text that can't be promoted (inside a group, first line of a multi-line
    box, a body placeholder) gets an OFF-SLIDE title placeholder: entirely
    above the slide, first in reading order, and the only ON-slide copy of
    the words is the original;
  * a slide with only a table, a date or a page number is REFUSED with a
    plain reason — no "Slide N", nothing written, nothing charged;
  * a running banner repeated on 3+ slides ("ACME Corp") is not a title;
  * the output reopens, re-analysis finds no untitled slide it claimed to
    fix, and the writer's entry names the action so it can be reconciled.

Usage:
    python -m app.devtools.smoke_slide_title_no_duplicate
"""

from __future__ import annotations

import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_stnd_')}/s.db")

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.pptx_writer import write_remediated_pptx  # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
FLAG = "SLIDE_TITLE_MISSING"


def _build(path: Path) -> None:
    prs = Presentation()
    # 1 Blank + a 28pt text box at the top (promote)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    tb.text_frame.text = "Untitled detail slide 2"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    s.shapes.add_textbox(Inches(1), Inches(1.8), Inches(8), Inches(3)).text_frame.text = "- alpha\n- beta\n- gamma"
    # 2 Title Only layout: EMPTY title placeholder + the real title in a text box
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1)).text_frame.text = "Quarterly Revenue Overview"
    # 3 Blank + a data table only (refuse; "North" is not a title)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_table(4, 3, Inches(0.5), Inches(1.0), Inches(9), Inches(3)).table
    for r, row in enumerate([["North", "120", "9%"], ["South", "98", "7%"], ["East", "143", "11%"], ["West", "110", "8%"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    # 4 Blank + the title inside a GROUP (off-slide title)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    g = s.shapes.add_group_shape()
    g.shapes.add_textbox(Inches(1), Inches(0.6), Inches(8), Inches(1)).text_frame.text = "Platform components"
    # 5 Blank + a two-line box (off-slide title "Agenda")
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(2)).text_frame.text = "Agenda\nIntroductions and goals"
    # 6 Blank + bold coloured text box (promote; formatting frozen)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1.5), Inches(3), Inches(7), Inches(1.2))
    tb.text_frame.text = "Strategic Priorities"
    run = tb.text_frame.paragraphs[0].runs[0]
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    # 7 only a page number and a date (refuse)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(8), Inches(7), Inches(1), Inches(0.4)).text_frame.text = "7"
    s.shapes.add_textbox(Inches(1), Inches(7), Inches(3), Inches(0.4)).text_frame.text = "June 16, 2026"
    # 8-10 a running banner on every slide above the real heading
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(3), Inches(0.4)).text_frame.text = "ACME Corp"
        s.shapes.add_textbox(Inches(1), Inches(1.0), Inches(8), Inches(1)).text_frame.text = f"Programme update {i + 1}"
    # 11 control: already titled
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Already Titled"
    prs.save(str(path))


def _slide_xml(path: Path, n: int):
    with zipfile.ZipFile(path) as z:
        return etree.fromstring(z.read(f"ppt/slides/slide{n}.xml"))


def _shapes_with_text(sld, text: str):
    """(sp element, is_title_ph, y, cy) for every p:sp whose text is ``text``."""
    out = []
    for sp in sld.iter(P + "sp"):
        t = "".join(x.text or "" for x in sp.iter(A + "t")).strip()
        if t != text:
            continue
        ph = sp.find(f"{P}nvSpPr/{P}nvPr/{P}ph")
        off = sp.find(f"{P}spPr/{A}xfrm/{A}off")
        ext = sp.find(f"{P}spPr/{A}xfrm/{A}ext")
        y = int(off.get("y")) if off is not None else None
        cy = int(ext.get("cy")) if ext is not None else None
        out.append((sp, ph is not None and ph.get("type") in ("title", "ctrTitle"), y, cy))
    return out


def _on_slide(entry) -> bool:
    _sp, _is_title, y, cy = entry
    if y is None:
        return True  # inherits the layout position: on the slide
    return y + (cy or 0) > 0


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="508_stnd_"))
    src = tmp / "deck.pptx"
    _build(src)
    before = {n: etree.tostring(_slide_xml(src, n)) for n in (3, 7, 11)}

    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    policy = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)
    plans = [p for p in plan_remediations(res.tree, policy) if p.flag.code.value == FLAG]
    execs = {e.target_node_id: e for e in execute_plans(res.tree, plans)}
    status = {int(k.split("-")[1]): (e.status.value, e.notes or "") for k, e in execs.items()}

    check("10 untitled slides were planned", len(plans) == 10, str(len(plans)))
    for n in (3, 7):
        st, note = status.get(n, ("?", ""))
        check(f"slide {n}: refused (nothing on it reads as a title)", st == "skipped" and "did not make one up" in note, f"{st} {note[:120]}")
    for n in (1, 2, 4, 5, 6, 8, 9, 10):
        check(f"slide {n}: title set", status.get(n, ("?",))[0] == "success", str(status.get(n)))

    out = tmp / "deck_fixed.pptx"
    rep = write_remediated_pptx(src, res.tree, out)
    applied = {a["target_id"]: a for a in rep["applied"] if a.get("kind") == "slide_title"}
    check("every applied title names its action (reconcilable)", all(a.get("action") == "SET_SLIDE_TITLE" for a in applied.values()), str(applied))
    check("writer applied exactly the 8 approved titles", len(applied) == 8, str(sorted(applied)))

    prs = Presentation(str(out))  # reopens
    titles = [((s.shapes.title.text if s.shapes.title is not None else None)) for s in prs.slides]
    want = ["Untitled detail slide 2", "Quarterly Revenue Overview", None, "Platform components", "Agenda",
            "Strategic Priorities", None, "Programme update 1", "Programme update 2", "Programme update 3", "Already Titled"]
    check("titles are the slides' own words (never 'North', never 'ACME Corp', never 'Slide N')", titles == want, str(titles))

    # ---- promoted: one copy, same shape, same place, look frozen ----------
    src_prs = Presentation(str(src))
    for n, text in ((1, "Untitled detail slide 2"), (2, "Quarterly Revenue Overview"), (6, "Strategic Priorities"),
                    (8, "Programme update 1")):
        hits = _shapes_with_text(_slide_xml(out, n), text)
        check(f"slide {n}: the title text is on the slide exactly once", len(hits) == 1, str(len(hits)))
        check(f"slide {n}: that one shape IS the title placeholder", bool(hits) and hits[0][1])
        orig = next(sh for sh in src_prs.slides[n - 1].shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text == text)
        new = prs.slides[n - 1].shapes.title
        check(f"slide {n}: promoted in place (same shape id and geometry)",
              new.shape_id == orig.shape_id and (new.left, new.top, new.width, new.height) == (orig.left, orig.top, orig.width, orig.height),
              f"{(orig.shape_id, orig.left, orig.top, orig.width, orig.height)} -> {(new.shape_id, new.left, new.top, new.width, new.height)}")
    check("slide 2: the empty layout title placeholder was dropped (not left beside the promoted box)",
          sum(1 for sp in _slide_xml(out, 2).iter(P + "sp") if sp.find(f"{P}nvSpPr/{P}nvPr/{P}ph") is not None) == 1)

    sp6 = _shapes_with_text(_slide_xml(out, 6), "Strategic Priorities")[0][0]
    r_pr = sp6.find(f".//{A}r/{A}rPr")
    body = sp6.find(f"{P}txBody/{A}bodyPr")
    p_pr = sp6.find(f".//{A}p/{A}pPr")
    check("slide 6: size/bold/colour/font pinned on the run (the title style can't restyle it)",
          r_pr.get("sz") == "1800" and r_pr.get("b") == "1" and r_pr.find(f"{A}solidFill/{A}srgbClr").get("val") == "1F4E79"
          and r_pr.find(f"{A}latin").get("typeface") == "+mn-lt", etree.tostring(r_pr)[:300])
    check("slide 6: alignment + anchor + insets pinned (title style centres/bottom-anchors)",
          p_pr.get("algn") == "l" and body.get("anchor") == "t" and body.get("lIns") == "91440", etree.tostring(body)[:200])
    check("slide 6: shape has explicit no-fill / no-line (a layout title band can't paint behind it)",
          sp6.find(f"{P}spPr/{A}noFill") is not None and sp6.find(f"{P}spPr/{A}ln/{A}noFill") is not None)
    check("slide 6: the promoted shape is no longer marked txBox",
          sp6.find(f"{P}nvSpPr/{P}cNvSpPr").get("txBox") is None)
    check("slide 1: the 28pt size set by the author is kept", _shapes_with_text(_slide_xml(out, 1), "Untitled detail slide 2")[0][0].find(f".//{A}rPr").get("sz") == "2800")

    # ---- off-slide: hidden title, original stays the only visible copy ----
    for n, text in ((4, "Platform components"), (5, "Agenda")):
        sld = _slide_xml(out, n)
        titles_sp = [e for e in _shapes_with_text(sld, text) if e[1]]
        check(f"slide {n}: an off-slide title placeholder carries the text", len(titles_sp) == 1 and not _on_slide(titles_sp[0]),
              str([(e[2], e[3]) for e in titles_sp]))
        tree_children = [c for c in sld.find(f"{P}cSld/{P}spTree") if c.tag in (P + "sp", P + "grpSp", P + "graphicFrame", P + "pic")]
        check(f"slide {n}: the hidden title is first in reading order", bool(titles_sp) and tree_children and tree_children[0] is titles_sp[0][0])
    visible4 = [e for e in _shapes_with_text(_slide_xml(out, 4), "Platform components") if _on_slide(e) and not e[1]]
    check("slide 4: the grouped original is untouched and is the only on-slide copy", len(visible4) == 1)
    body5 = [sp for sp in _slide_xml(out, 5).iter(P + "sp") if "Introductions" in "".join(t.text or "" for t in sp.iter(A + "t"))]
    check("slide 5: the two-line box is untouched (still both lines, still not a placeholder)",
          len(body5) == 1 and body5[0].find(f"{P}nvSpPr/{P}nvPr/{P}ph") is None)

    # ---- refused slides are byte-identical ---------------------------------
    for n in (3, 7, 11):
        check(f"slide {n}: slide XML unchanged", etree.tostring(_slide_xml(out, n)) == before[n])

    # ---- words: nothing lost, nothing duplicated on-slide ------------------
    def words(path: Path) -> int:
        n = 0
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                    root = etree.fromstring(z.read(name))
                    for sp in root.iter(P + "sp"):
                        off = sp.find(f"{P}spPr/{A}xfrm/{A}off")
                        ext = sp.find(f"{P}spPr/{A}xfrm/{A}ext")
                        if off is not None and ext is not None and int(off.get("y")) + int(ext.get("cy")) <= 0:
                            continue  # off-slide: never rendered
                        n += sum(len((t.text or "").split()) for t in sp.iter(A + "t"))
                    for gf in root.iter(P + "graphicFrame"):
                        n += sum(len((t.text or "").split()) for t in gf.iter(A + "t"))
        return n

    check("rendered (on-slide) words unchanged: no title was added as a second visible copy",
          words(out) == words(src), f"{words(src)} -> {words(out)}")

    # ---- re-analysis -------------------------------------------------------
    res2 = parse_to_tree(str(out))
    run_analyzers(res2.tree)
    still = sorted(int(n.id.split("-")[1]) for n in res2.tree.root.children
                   if any(f.code.value == FLAG for f in n.accessibility_flags))
    check("re-analysis: only the two refused slides are still untitled", still == [3, 7], str(still))

    # ---- the deck-wide boilerplate tally is computed once per job ----------
    # It used to be recounted over the whole deck for every untitled slide
    # (O(slides^2): 5 s of a 500-slide deck's remediation).
    from app.services.remediators import set_slide_title_executor as sste

    calls = {"n": 0}
    real = sste._slide_text_counts

    def counting(tree):
        calls["n"] += 1
        return real(tree)

    sste._slide_text_counts = counting
    try:
        res3 = parse_to_tree(str(src))
        run_analyzers(res3.tree)
        plans3 = [p for p in plan_remediations(res3.tree, policy) if p.flag.code.value == FLAG]
        execs3 = execute_plans(res3.tree, plans3)
    finally:
        sste._slide_text_counts = real
    check("the boilerplate tally is computed once for all 10 untitled slides", calls["n"] == 1, str(calls["n"]))
    check("and the outcome is the same as before",
          sorted((e.target_node_id, e.status.value) for e in execs3) == sorted((k, e.status.value) for k, e in execs.items()))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
