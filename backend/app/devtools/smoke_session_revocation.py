"""Smoke: a session token dies when the account says so.

Before the fix a stolen 7-day token survived password reset, set-password and
sign-out (a no-op 204): nothing server-side could revoke a JWT. Now every token
carries the user's ``token_version`` as ``ver``. Pins:

1. mint/verify helpers: ``ver`` embedded; legacy (no ``ver``) == version 0.
2. A legacy token minted before ``ver`` existed works until the first bump.
3. Password RESET revokes every outstanding token, and a revoked token is
   refused on every identity path (require_user_id, require_admin, optional_user,
   require_user_id_or_api_key).
4. Recovery kills API KEYS too. A stolen session can mint one (POST /api-keys),
   and bumping token_version would not touch it, so reset-password and
   set-password revoke every active key. The owner sees the reason in
   GET /api-keys and can create a replacement that works.
5. set-password revokes the caller's and other devices' tokens and returns a
   fresh one.
6. Sign-out revokes all devices; missing / stale tokens still get 204 and a
   stale token can't sign out the current session.
7. An email change revokes and returns a fresh token; a name-only PATCH doesn't.
8. A deleted account's token is refused.

Usage:
    python -m app.devtools.smoke_session_revocation
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import time

_TMP = tempfile.mkdtemp(prefix="508_smoke_sessions_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/s.db"
os.environ.pop("SMTP_HOST", None)

import jwt  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.main import app  # noqa: E402
from app.config import get_settings  # noqa: E402
from app.db.models import EmailVerifyTokenRow  # noqa: E402
from app.db.session_sqlalchemy import session_scope  # noqa: E402
from app.security.sessions import mint_session, session_is_current, verify_session  # noqa: E402

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _legacy_token(user_id: str) -> str:
    """A token exactly as minted before token_version existed: no ``ver``."""
    now = int(time.time())
    tok = jwt.encode(
        {"sub": user_id, "iss": "508-agent", "aud": "508-agent-session", "iat": now, "exp": now + 3600},
        get_settings().app_secret,
        algorithm="HS256",
    )
    return tok.decode("utf-8") if isinstance(tok, bytes) else tok


def _docx_bytes() -> bytes:
    """A small valid .docx, so a live API key gets 200 from /pipeline/analyze."""
    from docx import Document

    d = Document()
    d.add_heading("Report", level=1)
    d.add_paragraph("Body text.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def main() -> int:
    c = TestClient(app)
    failures = 0

    def check(name: str, cond: bool, detail: object = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, "" if cond else f"  [{detail}]")
        if not cond:
            failures += 1

    def h(tok: str) -> dict:
        return {"Authorization": f"Bearer {tok}"}

    def me(tok: str) -> int:
        return c.get("/auth/me", headers=h(tok)).status_code

    def signin(pw: str) -> str:
        r = c.post("/auth/sign-in", json={"email": EMAIL, "password": pw})
        assert r.status_code == 200, r.text
        return r.json()["token"]

    # --- 1. helpers -----------------------------------------------------------
    claims = verify_session(mint_session("unit-user", version=3))
    check("mint_session embeds ver", claims is not None and claims.get("ver") == 3, claims)
    check("session_is_current: same version", session_is_current(claims, 3))
    check("session_is_current: bumped version is stale", not session_is_current(claims, 4))
    check(
        "legacy claims (no ver) count as version 0 only",
        session_is_current({"sub": "x"}, 0) and not session_is_current({"sub": "x"}, 1),
    )
    check(
        "malformed ver never matches",
        not session_is_current({"sub": "x", "ver": "0"}, 0) and not session_is_current({"sub": "x", "ver": True}, 1),
    )

    # --- 2. stolen + legacy tokens live before any bump -------------------------
    EMAIL = "victim@example.com"
    r = c.post("/auth/sign-in", json={"email": EMAIL, "password": "victim-pass-1"})
    stolen, victim_id = r.json()["token"], r.json()["user"]["id"]
    legacy = _legacy_token(victim_id)
    check("fresh token works", me(stolen) == 200)
    check("legacy token (no ver claim) still works before the first bump", me(legacy) == 200)

    # The attacker mints an API key from the stolen session: a second
    # credential, which bumping token_version would not touch.
    r = c.post("/api-keys", headers=h(stolen), json={"name": "attacker persistence"})
    stolen_key = r.json().get("key", "")
    check("API key minted from the stolen session", r.status_code == 200 and stolen_key.startswith("ak_"), r.text)
    scan = c.post("/pipeline/analyze", headers={"X-API-Key": stolen_key}, files={"file": ("r.docx", _docx_bytes(), DOCX_MIME)})
    check("that key scans before the reset", scan.status_code == 200, scan.status_code)

    # --- 3. password reset -------------------------------------------------------
    c.post("/auth/request-password-reset", json={"email": EMAIL})
    with session_scope() as s:
        reset_token = s.execute(
            select(EmailVerifyTokenRow.token).where(
                EmailVerifyTokenRow.user_id == victim_id, EmailVerifyTokenRow.token.like("pr_%")
            )
        ).scalars().first()
    r = c.post("/auth/reset-password", json={"token": reset_token, "password": "victim-pass-2"})
    check("reset-password 200", r.status_code == 200, r.text)
    check("reset reports the API keys it revoked", r.json().get("apiKeysRevoked") == 1, r.text)
    check("stolen token rejected after password reset", me(stolen) == 401)
    check("legacy token rejected after the first bump", me(legacy) == 401)
    check("revoked token: /credits/balance 401", c.get("/credits/balance", headers=h(stolen)).status_code == 401)
    check("revoked token: /admin/metrics 401 (not 403)", c.get("/admin/metrics", headers=h(stolen)).status_code == 401)
    check(
        "revoked token: whoami answers as anonymous",
        c.get("/api/admin/whoami", headers=h(stolen)).json() == {"email": None, "isAdmin": False},
    )
    r = c.post("/pipeline/analyze", headers=h(stolen), files={"file": ("x.docx", b"PK", DOCX_MIME)})
    check("revoked token: /pipeline/analyze (session-or-API-key route) 401", r.status_code == 401, r.status_code)
    scan = c.post("/pipeline/analyze", headers={"X-API-Key": stolen_key}, files={"file": ("r.docx", _docx_bytes(), DOCX_MIME)})
    check("the attacker's API key is dead after recovery", scan.status_code == 401, scan.status_code)

    # --- 4. the owner can see it happened, and start over -----------------------
    laptop = signin("victim-pass-2")
    keys = c.get("/api-keys", headers=h(laptop)).json()
    check(
        "owner sees the key revoked, when, and why",
        len(keys) == 1
        and keys[0]["revoked"] is True
        and keys[0]["revokedReason"] == "password_reset"
        and bool(keys[0]["revokedAt"]),
        keys,
    )
    r = c.post("/api-keys", headers=h(laptop), json={"name": "replacement"})
    new_key = r.json().get("key", "")
    check("owner can create a replacement key", r.status_code == 200 and bool(new_key), r.text)
    scan = c.post("/pipeline/analyze", headers={"X-API-Key": new_key}, files={"file": ("r.docx", _docx_bytes(), DOCX_MIME)})
    check("the replacement key scans", scan.status_code == 200, scan.status_code)

    # --- 5. set-password --------------------------------------------------------
    phone = signin("victim-pass-2")
    r = c.post(
        "/auth/set-password",
        headers=h(laptop),
        json={"password": "victim-pass-3", "currentPassword": "victim-pass-2"},
    )
    fresh = r.json().get("token")
    check("set-password 200 returns a fresh token", r.status_code == 200 and bool(fresh) and fresh != laptop, r.text)
    check("set-password revokes the caller's old token", me(laptop) == 401)
    check("set-password revokes other devices", me(phone) == 401)
    check("the returned token works", me(fresh) == 200)
    check("set-password revokes API keys as well", r.json().get("apiKeysRevoked") == 1, r.text)
    scan = c.post("/pipeline/analyze", headers={"X-API-Key": new_key}, files={"file": ("r.docx", _docx_bytes(), DOCX_MIME)})
    check("the replacement key dies with the password change", scan.status_code == 401, scan.status_code)

    # --- 6. sign-out ------------------------------------------------------------
    other = signin("victim-pass-3")
    check("sign-out with a valid token -> 204", c.post("/auth/sign-out", headers=h(fresh)).status_code == 204)
    check("signed-out token rejected", me(fresh) == 401)
    check("sign-out ends every device", me(other) == 401)
    check("sign-out with an already-revoked token -> 204", c.post("/auth/sign-out", headers=h(fresh)).status_code == 204)
    check("sign-out with no token -> 204", c.post("/auth/sign-out").status_code == 204)
    current = signin("victim-pass-3")
    check("sign-in works after sign-out", me(current) == 200)
    c.post("/auth/sign-out", headers=h(other))
    check("a stale token can't sign out the current session", me(current) == 200)

    # --- 7. email change ----------------------------------------------------------
    r = c.patch(
        "/auth/me",
        headers=h(current),
        json={"email": "victim-moved@example.com", "currentPassword": "victim-pass-3"},
    )
    moved = r.json().get("token")
    check("email change -> 200 with a fresh token", r.status_code == 200 and bool(moved), r.text)
    check("email change revokes the previous token", me(current) == 401)
    check("fresh token after email change works", me(moved) == 200)
    r = c.patch("/auth/me", headers=h(moved), json={"displayName": "Just A Rename"})
    check(
        "name-only PATCH doesn't rotate (token null, old token fine)",
        r.status_code == 200 and r.json().get("token") is None and me(moved) == 200,
        r.text,
    )

    # --- 8. deletion ---------------------------------------------------------------
    check("delete account 204", c.delete("/auth/me", headers=h(moved)).status_code == 204)
    check("deleted account's token rejected", me(moved) == 401)

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
