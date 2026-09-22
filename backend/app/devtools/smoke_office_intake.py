"""Smoke: legacy and OpenDocument files convert with LibreOffice, safely and honestly.

.doc .rtf .odt -> .docx, .xls .ods -> .xlsx, .ppt .odp -> .pptx, through
``soffice --headless --convert-to``. LibreOffice is not installed on dev
boxes or CI, so a FAKE soffice (a stdlib-only script, run as a real
subprocess through the real argv/env/timeout path) stands in for it:

  * each legacy type is accepted, converted, and analysed as its modern
    format, with an ``intake`` record saying so;
  * the converter is sandboxed: it sees a neutral input name (never the
    customer's filename), a throwaway profile, and an environment WITHOUT the
    app's secrets; its working directory is gone afterwards;
  * a fix on a converted .doc is priced as a .docx (3) and delivered as a
    .docx; a run where nothing persisted is free AND hands back the original
    .doc bytes (the conversion alone is not given away);
  * no LibreOffice -> 415 telling the person exactly how to save the modern
    format themselves (per format), and nothing charged;
  * a converter that fails, produces junk, or hangs -> 422 with a sentence,
    killed at the timeout rather than waited out;
  * an OpenDocument file whose ``mimetype`` says it is something else is
    refused naming what it really is.

Usage:
    python -m app.devtools.smoke_office_intake
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import time
import zipfile
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="508_smoke_office_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/s.db"
os.environ["MATERIALIZED_ROOT"] = str(Path(_TMP) / "materialized")
os.environ.setdefault("APP_SECRET", "x" * 64)
os.environ["SEMANTIC_PROVIDER"] = "heuristic"
os.environ["OFFICE_CONVERT_TIMEOUT_SECONDS"] = "5"
for _key in ("OPENAI_API_KEY", "ANTHROPIC_API_KEY"):
    os.environ.pop(_key, None)

from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import select  # noqa: E402

OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

SHIM = r'''
import json, os, shutil, sys, time
from pathlib import Path
here = Path(__file__).resolve().parent
args = sys.argv[1:]
with open(here / "fake_soffice.log", "a", encoding="utf-8") as fh:
    fh.write(json.dumps({"argv": args, "env": sorted(os.environ.keys())}) + "\n")
target = args[args.index("--convert-to") + 1].split(":")[0]
outdir = Path(args[args.index("--outdir") + 1])
inp = Path(args[-1])
data = inp.read_bytes()
if b"SLEEP" in data:
    time.sleep(60)
if b"FAIL" in data:
    sys.exit(1)
out = outdir / (inp.stem + "." + target)
if b"JUNK" in data:
    out.write_bytes(b"this is not a zip")
    sys.exit(0)
cfg = json.loads((here / "fake_soffice.json").read_text(encoding="utf-8"))
shutil.copyfile(cfg[target], out)
'''


def _docx(path: Path) -> None:
    from docx import Document

    d = Document()
    d.add_heading("Annual Budget Memo", level=1)
    d.add_paragraph("The committee reviewed spending for every department this year.")
    d.core_properties.title = ""
    d.save(str(path))


def _xlsx(path: Path) -> None:
    import xlsxwriter

    wb = xlsxwriter.Workbook(str(path))
    ws = wb.add_worksheet("Totals")
    bold = wb.add_format({"bold": True})
    ws.write_row(0, 0, ["Region", "Total"], bold)
    ws.write_row(1, 0, ["North", 10])
    ws.write_row(2, 0, ["South", 20])
    wb.close()


def _pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Quarterly review"
    prs.save(str(path))


def _odf(kind: str, marker: bytes = b"") -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", f"application/vnd.oasis.opendocument.{kind}", compress_type=zipfile.ZIP_STORED)
        z.writestr("content.xml", b"<office:document-content/>" + marker)
    return buf.getvalue()


def main() -> int:  # noqa: PLR0915
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.db.models import UserRow
    from app.db.session_sqlalchemy import session_scope
    from app.intake.office import set_soffice_command_for_testing
    from app.main import app

    tmp = Path(_TMP)
    shim_dir = tmp / "shim"
    shim_dir.mkdir()
    (shim_dir / "fake_soffice.py").write_text(SHIM, encoding="utf-8")
    fixtures = {"docx": shim_dir / "out.docx", "xlsx": shim_dir / "out.xlsx", "pptx": shim_dir / "out.pptx"}
    _docx(fixtures["docx"])
    _xlsx(fixtures["xlsx"])
    _pptx(fixtures["pptx"])
    (shim_dir / "fake_soffice.json").write_text(json.dumps({k: str(v) for k, v in fixtures.items()}), encoding="utf-8")
    log = shim_dir / "fake_soffice.log"

    def last_call() -> dict:
        lines = log.read_text(encoding="utf-8").strip().splitlines() if log.exists() else []
        return json.loads(lines[-1]) if lines else {}

    set_soffice_command_for_testing([sys.executable, str(shim_dir / "fake_soffice.py")])

    client = TestClient(app, raise_server_exceptions=False)
    email = "office-smoke@example.com"
    r = client.post("/auth/sign-in", json={"email": email, "displayName": "O", "password": "officesmokepass1"})
    assert r.status_code == 200, r.text
    headers = {"Authorization": f"Bearer {r.json()['token']}"}

    def set_balance(n: int) -> None:
        with session_scope() as s:
            row = s.execute(select(UserRow).where(UserRow.email == email)).scalars().first()
            row.credits_balance = n

    def balance() -> int:
        return int(client.get("/credits/balance", headers=headers).json()["balance"])

    def post(path: str, name: str, data: bytes, ids=None):
        kw = {"files": {"file": (name, data, "application/octet-stream")}, "headers": headers}
        if ids is not None:
            kw["data"] = {"approved_violations": json.dumps(ids), "rejected_violations": "[]"}
        return client.post(path, **kw)

    set_balance(40)
    legacy_doc = OLE2 + b"\x00" * 504 + b"legacy word body"

    # ---- 1. every legacy type converts and is analysed as its modern format
    cases = {
        "report.doc": (legacy_doc, "docx"),
        "memo.rtf": (b"{\\rtf1\\ansi Budget memo}", "docx"),
        "notes.odt": (_odf("text"), "docx"),
        "data.xls": (OLE2 + b"\x00" * 504, "xlsx"),
        "sheet.ods": (_odf("spreadsheet"), "xlsx"),
        "deck.ppt": (OLE2 + b"\x00" * 504, "pptx"),
        "slides.odp": (_odf("presentation"), "pptx"),
    }
    for name, (data, target) in cases.items():
        rr = post("/pipeline/analyze", name, data)
        body = rr.json() if rr.status_code == 200 else {}
        intake = body.get("intake") or {}
        check(
            f"{name}: converted and analysed as .{target}",
            rr.status_code == 200 and body["summary"]["sourceFormat"] == target and intake.get("converter") == "libreoffice"
            and intake.get("originalFormat") == name.rsplit(".", 1)[1] and intake.get("convertedTo") == target,
            rr.text[:240],
        )

    # ---- 2. the sandbox ------------------------------------------------------
    rr = post("/pipeline/analyze", "Quarterly Report (FINAL; rm -rf).doc", legacy_doc)
    call = last_call()
    argv = call.get("argv") or []
    check("sandbox: converter sees a neutral input name, not the customer's",
          rr.status_code == 200 and argv and Path(argv[-1]).name == "input.doc" and not any("Quarterly" in a for a in argv), str(argv)[:300])
    check("sandbox: headless, throwaway profile, docx filter",
          "--headless" in argv and any(a.startswith("-env:UserInstallation=file:") for a in argv) and "docx:MS Word 2007 XML" in argv)
    env_keys = set(call.get("env") or [])
    leaked = {k for k in ("APP_SECRET", "DATABASE_URL", "MATERIALIZED_ROOT", "SEMANTIC_PROVIDER") if k in env_keys}
    check("sandbox: the app's secrets are not in the converter's environment", not leaked, str(leaked))
    outdir = Path(argv[argv.index("--outdir") + 1]) if "--outdir" in argv else None
    check("sandbox: the working directory is removed afterwards", outdir is not None and not outdir.parent.exists(), str(outdir))

    # ---- 3. priced and delivered as the modern format ----------------------
    rep = post("/pipeline/analyze", "report.doc", legacy_doc).json()
    ids = [v["id"] for v in rep["violations"]]
    b0 = balance()
    rr = post("/pipeline/remediate", "report.doc", legacy_doc, ids)
    body = rr.json()
    check("remediate .doc: charged at the docx price (3)", rr.status_code == 200 and body["charged"] is True and b0 - balance() == 3,
          f"{rr.status_code} {b0}->{balance()} {rr.text[:200]}")
    check("remediate .doc: a .docx comes back", body.get("filename") == "report-remediated.docx", str(body.get("filename")))
    check("remediate .doc: priced/recorded as docx", body.get("sourceFormat") == "docx", str(body.get("sourceFormat")))
    dl = client.get(body["downloadUrl"], headers=headers)
    from docx import Document

    try:
        doc = Document(io.BytesIO(dl.content))
        title = doc.core_properties.title
    except Exception as exc:  # pragma: no cover - reported below
        title = f"<unreadable: {exc}>"
    check("remediate .doc: the .docx opens and carries the fix", title == "Annual Budget Memo", str(title))

    b0 = balance()
    rr = post("/pipeline/remediate", "report.doc", legacy_doc, [])
    body = rr.json()
    check("nothing approved: free", rr.status_code == 200 and body["charged"] is False and balance() == b0)
    check("nothing approved: the ORIGINAL .doc comes back, not a free conversion",
          body.get("filename") == "report-remediated.doc" and client.get(body["downloadUrl"], headers=headers).content == legacy_doc,
          str(body.get("filename")))

    # ---- 4. no LibreOffice: say exactly what to do --------------------------
    set_soffice_command_for_testing([])
    try:
        for name, data, choice in (
            ("report.doc", legacy_doc, "Word Document (.docx)"),
            ("data.xls", OLE2 + b"\x00" * 504, "Excel Workbook (.xlsx)"),
            ("deck.ppt", OLE2 + b"\x00" * 504, "PowerPoint Presentation (.pptx)"),
        ):
            rr = post("/pipeline/analyze", name, data)
            detail = rr.json().get("detail") or ""
            check(f"no soffice: {name} -> 415 'Save As {choice}'", rr.status_code == 415 and "Save As" in detail and choice in detail, rr.text[:200])
        b0 = balance()
        rr = post("/pipeline/remediate", "report.doc", legacy_doc, [])
        check("no soffice: remediate refused, not charged", rr.status_code == 415 and balance() == b0)
    finally:
        set_soffice_command_for_testing([sys.executable, str(shim_dir / "fake_soffice.py")])

    # ---- 5. converter failures are sentences --------------------------------
    rr = post("/pipeline/analyze", "broken.doc", legacy_doc + b"FAIL")
    check("converter fails -> 422 'couldn't convert'", rr.status_code == 422 and "couldn't convert" in rr.json()["detail"], rr.text[:200])
    rr = post("/pipeline/analyze", "junk.doc", legacy_doc + b"JUNK")
    check("converter output unusable -> 422", rr.status_code == 422 and "usable" in rr.json()["detail"], rr.text[:200])
    t0 = time.monotonic()
    rr = post("/pipeline/analyze", "hang.doc", legacy_doc + b"SLEEP")
    took = time.monotonic() - t0
    check("converter hangs -> 422 'took too long', killed at the timeout", rr.status_code == 422 and "too long" in rr.json()["detail"] and took < 30,
          f"{rr.status_code} after {took:.1f}s {rr.text[:160]}")

    # ---- 6. OpenDocument kind is checked ------------------------------------
    rr = post("/pipeline/analyze", "notes.odt", _odf("spreadsheet"))
    check("an .odt that is really a spreadsheet -> 400 naming it", rr.status_code == 400 and "OpenDocument spreadsheet" in rr.json()["detail"], rr.text[:200])

    set_soffice_command_for_testing(None)
    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
