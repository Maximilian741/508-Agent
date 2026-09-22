"""Smoke: "we can fix N of these" is exactly what fixing then does.

/pipeline/analyze marks each finding ``autoFixable`` and sums them into
``summary.autoFixable`` — the number the one-step flow promises BEFORE the
customer pays. It used to be the format's capability only (an action for the
flag persists into this format), so a PDF with an uncaptioned picture and an
unsure language guess was promised 3 fixes and delivered 1: the executors
refuse placeholder alt text and low-confidence languages, correctly — but
after the promise.

With no paid AI provider configured, /remediate runs the offline executors,
which are deterministic, so /analyze now dry-runs them on a copy of the tree
(same apply policy, same plan selection) and promises only what they make.

Pinned here, per format (DOCX, HTML, PPTX, PDF):
  1. promise == outcome: the ids analyze marks autoFixable are EXACTLY the ids
     remediate (approving everything) reports ``fixed``;
  2. a refusal shrinks the promise: at least one finding whose format
     capability says "fixable" is not promised, and is indeed not fixed;
  3. never wider than the capability;
  4. summary.cost is 0 when nothing is promised (there is nothing to buy);
  5. the account-free scan promises the same thing as a signed-in scan;
  6. with a paid provider configured, analyze does NOT predict (that would
     mean calling it) — the capability stands, and no client is built.

Run: python -m app.devtools.smoke_fix_promise
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
from pathlib import Path

_TMP = Path(tempfile.mkdtemp(prefix="508_smoke_promise_"))
os.environ["DATABASE_URL"] = f"sqlite:///{(_TMP / 'promise.db').as_posix()}"
os.environ["STORAGE_LOCAL_ROOT"] = str(_TMP / "storage")
os.environ["MATERIALIZED_ROOT"] = str(_TMP / "materialized")
for _key in ("OPENAI_API_KEY", "ANTHROPIC_API_KEY", "SEMANTIC_PROVIDER", "SMTP_HOST"):
    os.environ.pop(_key, None)

from fastapi.testclient import TestClient  # noqa: E402
from PIL import Image  # noqa: E402

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _png(w: int, h: int, color=(30, 110, 200)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "PNG")
    return buf.getvalue()


def _docx() -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = Document()
    d.add_heading("Quarterly Results", level=1)
    d.add_heading("Regional detail", level=3)
    for item in ("- first typed item", "- second typed item", "- third typed item"):
        d.add_paragraph(item)
    d.add_paragraph().add_run().add_picture(io.BytesIO(_png(300, 200)))
    para = d.add_paragraph("For the annual numbers ")
    rid = d.part.relate_to(
        "https://example.com/annual.pdf",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "click here"
    r.append(t)
    hl.append(r)
    para._p.append(hl)
    table = d.add_table(rows=3, cols=3)
    for ri, row in enumerate([("Region", "Q1", "Q2"), ("North", "120", "130"), ("South", "90", "95")]):
        for ci, val in enumerate(row):
            table.cell(ri, ci).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _html() -> bytes:
    return (
        "<!doctype html><html><head></head><body>"
        "<h1>Report</h1><h3>Details</h3>"
        '<p>For the full report, <a href="https://example.com/r.pdf">click here</a> today.</p>'
        '<img src="chart.png">'
        "<table><tr><td>Region</td><td>Q1</td></tr><tr><td>North</td><td>12</td></tr></table>"
        '<form><p>Search: <input type="text" name="q"></p></form>'
        '<p><a href="/next" tabindex="3">Skip ahead</a></p>'
        "</body></html>"
    ).encode("utf-8")


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(_png(400, 250, (40, 160, 90))), Inches(1), Inches(2))
    tb = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1))
    p = tb.text_frame.paragraphs[0]
    p.add_run().text = "Read the report: "
    link = p.add_run()
    link.text = "click here"
    link.hyperlink.address = "https://example.com/deck-report"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pdf() -> bytes:
    """Untagged, no title/language, one uncaptioned picture, one link."""
    from pypdf import PdfWriter
    from pypdf.annotations import Link
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

    w = PdfWriter()
    font = w._add_object(DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
        NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
    }))
    img = DecodedStreamObject()
    img.set_data(Image.new("RGB", (40, 30), (200, 30, 30)).tobytes())
    img.update({
        NameObject("/Type"): NameObject("/XObject"),
        NameObject("/Subtype"): NameObject("/Image"),
        NameObject("/Width"): NumberObject(40),
        NameObject("/Height"): NumberObject(30),
        NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
        NameObject("/BitsPerComponent"): NumberObject(8),
    })
    pg = w.add_blank_page(612, 792)
    pg[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        NameObject("/XObject"): DictionaryObject({NameObject("/Im1"): w._add_object(img)}),
    })
    body = (
        "BT /F1 24 Tf 72 700 Td (Annual Report 2025) Tj ET\n"
        "BT /F1 12 Tf 72 660 Td (Revenue grew in every region this year, led by the north.) Tj ET\n"
        "BT /F1 12 Tf 72 640 Td (For the details, click here today. Costs held flat.) Tj ET\n"
        "q 200 0 0 150 300 400 cm /Im1 Do Q\n"
    ).encode("latin-1")
    cs = DecodedStreamObject()
    cs.set_data(body)
    pg[NameObject("/Contents")] = w._add_object(cs)
    w.add_annotation(0, Link(rect=(164, 635, 224, 652), url="https://example.com/details"))
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, detail: object = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, "" if cond else f"  [{str(detail)[:500]}]")
        if not cond:
            failures += 1

    import app.api.pipeline as pipeline
    from app.api.credits import DOC_FORMAT_COSTS
    from app.db.models import UserRow
    from app.db.session_sqlalchemy import session_scope
    from app.main import app

    client = TestClient(app)
    r = client.post("/auth/sign-in", json={"email": "promise@example.com", "password": "promisepass1"})
    assert r.status_code == 200, r.text
    auth = {"Authorization": f"Bearer {r.json()['token']}"}
    with session_scope() as s:
        s.get(UserRow, r.json()["user"]["id"]).credits_balance = 1000
    anon = TestClient(app)

    refusals = []
    fixtures = [
        ("promise.docx", _docx(), DOCX, "docx"),
        ("promise.html", _html(), "text/html", "html"),
        ("promise.pptx", _pptx(), PPTX, "pptx"),
        ("promise.pdf", _pdf(), "application/pdf", "pdf"),
    ]
    for name, data, mime, fmt in fixtures:
        a = client.post("/pipeline/analyze", files={"file": (name, data, mime)}, headers=auth)
        check(f"{fmt}: analyze -> 200", a.status_code == 200, a.text[:300])
        body = a.json() if a.status_code == 200 else {"violations": [], "summary": {}}
        vs = body["violations"]
        promised = {v["id"] for v in vs if v["autoFixable"]}
        capable = {v["id"] for v in vs if pipeline._auto_fixable(v["ruleId"], fmt)}
        check(f"{fmt}: the promise is never wider than the format capability", promised <= capable, promised - capable)
        refusals += [(fmt, v["ruleId"], v["id"]) for v in vs if v["id"] in capable - promised]
        s = body["summary"]
        check(
            f"{fmt}: summary.autoFixable counts the promise; cost is the price only when something is promised",
            s.get("autoFixable") == len(promised) and s.get("needsYou") == len(vs) - len(promised)
            and s.get("cost") == (DOC_FORMAT_COSTS[fmt] if promised else 0),
            s,
        )

        an = anon.post("/pipeline/analyze", files={"file": (name, data, mime)})
        anon_promised = {v["id"] for v in an.json().get("violations", []) if v["autoFixable"]} if an.status_code == 200 else None
        check(f"{fmt}: the account-free scan promises the same fixes", anon_promised == promised, (anon_promised, promised))

        rr = client.post(
            "/pipeline/remediate",
            files={"file": (name, data, mime)},
            data={"approved_violations": json.dumps([v["id"] for v in vs]), "rejected_violations": "[]"},
            headers=auth,
        )
        check(f"{fmt}: remediate (approve everything) -> 200", rr.status_code == 200, rr.text[:300])
        rb = rr.json() if rr.status_code == 200 else {}
        fixed = {v["id"] for v in rb.get("violations") or [] if v.get("fixed")}
        by_id = {v["id"]: v["ruleId"] for v in vs}
        check(
            f"{fmt}: promise == outcome (promised {len(promised)}, fixed {len(fixed)})",
            promised == fixed,
            {"promised_not_fixed": sorted(by_id.get(i, i) for i in promised - fixed),
             "fixed_not_promised": sorted(by_id.get(i, i) for i in fixed - promised)},
        )
        check(f"{fmt}: charged exactly when something was promised", bool(rb.get("charged")) == bool(promised), (rb.get("charged"), len(promised)))

    check(
        "a fix the executors refuse is not promised (at least one capable finding held back, e.g. an uncaptioned PDF picture)",
        any(fmt == "pdf" and rule == "MISSING_ALT_TEXT" for fmt, rule, _ in refusals),
        refusals,
    )

    # --- a paid provider configured: no prediction, and nothing is built -----------
    import app.ai.semantic_inference as si

    built = {"n": 0}
    real_build = si.build_default_provider

    def _counting_build(*a, **k):  # pragma: no cover - only hit on a regression
        built["n"] += 1
        return real_build(*a, **k)

    si.build_default_provider = _counting_build
    real_name = pipeline._configured_provider_name
    pipeline._configured_provider_name = lambda: "claude"
    try:
        name, data, mime, fmt = fixtures[3]
        a = client.post("/pipeline/analyze", files={"file": (name, data, mime)}, headers=auth)
        vs = a.json()["violations"] if a.status_code == 200 else []
        check(
            "paid provider configured: autoFixable is the format capability (predicting would mean paying)",
            bool(vs) and all(v["autoFixable"] == pipeline._auto_fixable(v["ruleId"], fmt) for v in vs),
            [(v["ruleId"], v["autoFixable"]) for v in vs],
        )
        check("paid provider configured: analyze built no provider", built["n"] == 0, built)
    finally:
        pipeline._configured_provider_name = real_name
        si.build_default_provider = real_build

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
