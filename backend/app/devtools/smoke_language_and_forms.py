"""Smoke: real language tagging (w:lang / a:rPr@lang) + PDF form-field labels.

- DOCX language is written to ``w:lang`` on the document defaults (where AT and
  the Word Accessibility Checker read it), not only ``docProps/core.xml``.
- PPTX language is written to each run's ``a:rPr@lang``.
- PDF AcroForm fields with no ``/TU`` accessible label are flagged
  (WCAG 3.3.2 / 4.1.2); labeled fields are not.

Usage:
    python -m app.devtools.smoke_language_and_forms
"""

from __future__ import annotations

import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_langforms_')}/s.db"

from docx import Document  # noqa: E402
from pptx import Presentation  # noqa: E402
from pypdf import PdfWriter  # noqa: E402
from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_engine import RemediationEngine  # noqa: E402
from app.writers import write_remediated  # noqa: E402


def main() -> int:
    failures = 0

    def check(name: str, cond: bool) -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name)
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp())

    # --- DOCX w:lang ---
    dp = tmp / "in.docx"
    d = Document(); d.add_paragraph("Hello world."); d.save(str(dp))
    res = parse_to_tree(str(dp)); res.tree.root.metadata.language = "en-US"
    dout = tmp / "out.docx"; write_remediated(dp, res.tree, dout, source_format=res.format)
    with zipfile.ZipFile(dout) as z:
        styles = z.read("word/styles.xml").decode("utf-8", "ignore") if "word/styles.xml" in z.namelist() else ""
    check("DOCX w:lang written to styles.xml", "w:lang" in styles and "en-US" in styles)
    Document(str(dout))  # not corrupt
    check("DOCX reopens cleanly", True)

    # --- PPTX a:rPr@lang ---
    pp = tmp / "in.pptx"
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(0, 0, 100, 100); tb.text_frame.text = "Hello slide."
    prs.save(str(pp))
    res2 = parse_to_tree(str(pp)); res2.tree.root.metadata.language = "en-US"
    pout = tmp / "out.pptx"; write_remediated(pp, res2.tree, pout, source_format=res2.format)
    xml = b""
    with zipfile.ZipFile(pout) as z:
        for nm in z.namelist():
            if nm.startswith("ppt/slides/slide") and nm.endswith(".xml"):
                xml += z.read(nm)
    check("PPTX a:rPr@lang written on runs", b'lang="en-US"' in xml)
    Presentation(str(pout))
    check("PPTX reopens cleanly", True)

    # --- PDF form-field labels ---
    def build_pdf(path: Path, with_tu: bool) -> None:
        w = PdfWriter(); w.add_blank_page(width=300, height=200)
        f = DictionaryObject({NameObject("/FT"): NameObject("/Tx"), NameObject("/T"): TextStringObject("field1")})
        if with_tu:
            f[NameObject("/TU")] = TextStringObject("Your full name")
        fr = w._add_object(f)  # noqa: SLF001
        w._root_object[NameObject("/AcroForm")] = w._add_object(DictionaryObject({NameObject("/Fields"): ArrayObject([fr])}))  # noqa: SLF001
        with open(path, "wb") as fh:
            w.write(fh)

    def has_form_flag(path: Path) -> bool:
        res = parse_to_tree(str(path)); run_analyzers(res.tree)
        return any(v.rule_id == "FORM_FIELD_UNLABELED" for v in RemediationEngine().detect_violations(res.tree))

    build_pdf(tmp / "unlabeled.pdf", with_tu=False)
    build_pdf(tmp / "labeled.pdf", with_tu=True)
    check("PDF unlabeled form field -> FORM_FIELD_UNLABELED", has_form_flag(tmp / "unlabeled.pdf"))
    check("PDF labeled form field -> not flagged", not has_form_flag(tmp / "labeled.pdf"))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
