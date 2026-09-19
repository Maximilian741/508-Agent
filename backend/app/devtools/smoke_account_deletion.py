"""Smoke: deleting an account ends every credential and leaves nothing dangling.

Deletion used to remove two tables — ``users`` and ``credit_ledger`` — and walk
away, so an erased account kept handing out authority. Pins:

1. The account's API keys die with it. A key minted before the delete is
   refused as ``X-API-Key`` and as ``Authorization: Bearer ak_…``, writes no
   further rows, and its ``api_keys`` row is gone. Re-registering the same
   address mints a NEW user id, so the ghost key can't follow the mailbox.
2. ``deps._api_key_user_id`` fails closed on its own: an orphaned key row
   pointing at a missing user authenticates nobody, even if a future writer
   forgets to clean the table up.
3. Outstanding verify/reset links for the account are gone.
4. A team OWNER with other members in it cannot delete (409 team_has_members) —
   deleting would silently take away other people's shared wallet, and nothing
   could ever disband the team afterwards (that needs owner_id == caller).
   Remove the members and the same call succeeds.
5. A SOLE owner is never stuck: their team is disbanded as part of the delete,
   and pending invites for it are revoked.
6. A plain MEMBER deleting their account frees their seat and leaves the team
   standing for everyone else.
7. The subscription stops reading as active, so the dead account grants no
   benefits.
8. ``teams.resolve_credit_user_id`` falls back to the caller when the resolved
   owner row is missing, so a dangling owner_id can never point a live user's
   wallet at a nonexistent account.

Usage:
    python -m app.devtools.smoke_account_deletion
"""

from __future__ import annotations

import io
import os
import sys
import tempfile

_TMP = tempfile.mkdtemp(prefix="508_smoke_acctdel_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/d.db"
os.environ.pop("SMTP_HOST", None)
# One rate-limit bucket per section (/auth allows 60 req/min per IP); the app
# only reads the header when told it sits behind a proxy, which now defaults
# to off. See security/rate_limit.py.
os.environ["TRUST_PROXY_HEADERS"] = "true"

from datetime import datetime  # noqa: E402

from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.api.deps import _api_key_user_id, hash_api_key  # noqa: E402
from app.api.teams import resolve_credit_user_id  # noqa: E402
from app.db.models import (  # noqa: E402
    ApiKeyRow,
    EmailVerifyTokenRow,
    SubscriptionRow,
    TeamInviteRow,
    TeamMemberRow,
    TeamRow,
    UserRow,
)
from app.db.session_sqlalchemy import session_scope  # noqa: E402
from app.main import app  # noqa: E402

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_bytes() -> bytes:
    """A small valid .docx, so a LIVE key gets 200 from /pipeline/analyze."""
    from docx import Document

    d = Document()
    d.add_heading("Report", level=1)
    d.add_paragraph("Body text.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _subscribe(user_id: str, sub_id: str) -> None:
    with session_scope() as s:
        s.add(SubscriptionRow(
            id=sub_id, user_id=user_id, plan="team", status="active",
            stripe_customer_id="cus_del", current_period_end=None,
            overage_enabled=False, created_at=datetime.utcnow(),
            updated_at=datetime.utcnow(),
        ))


def _verify(client: TestClient, auth: dict) -> None:
    """Prove an address the way clicking the emailed link does."""
    uid = client.get("/auth/me", headers=auth).json()["id"]
    assert client.post("/auth/request-verify-email", headers=auth).status_code == 200
    with session_scope() as s:
        link = s.execute(
            select(EmailVerifyTokenRow.token).where(EmailVerifyTokenRow.user_id == uid)
        ).scalars().first()
    assert client.get("/auth/verify-email", params={"token": link}).status_code == 200


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, detail: object = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, "" if cond else f"  [{detail}]")
        if not cond:
            failures += 1

    def client(ip: str) -> TestClient:
        return TestClient(app, headers={"X-Forwarded-For": ip})

    def signin(c: TestClient, email: str, password: str = "deletepass123"):
        r = c.post("/auth/sign-in", json={"email": email, "password": password})
        assert r.status_code == 200, r.text
        return {"Authorization": f"Bearer {r.json()['token']}"}, r.json()["user"]["id"]

    body = _docx_bytes()

    # --- 1. the ghost API key ------------------------------------------------
    c = client("10.9.0.1")
    auth, uid = signin(c, "ghost@example.com")
    key = c.post("/api-keys", headers=auth, json={"name": "cli"}).json()["key"]
    r = c.post("/pipeline/analyze", headers={"X-API-Key": key}, files={"file": ("a.docx", body, DOCX_MIME)})
    check("the key scans while the account is alive", r.status_code == 200, r.status_code)
    c.post("/auth/request-verify-email", headers=auth)  # an outstanding link

    check("delete account -> 204", c.delete("/auth/me", headers=auth).status_code == 204)
    check("the account's session token is refused", c.get("/auth/me", headers=auth).status_code == 401)

    r = c.post("/pipeline/analyze", headers={"X-API-Key": key}, files={"file": ("a.docx", body, DOCX_MIME)})
    check("the DELETED account's API key is refused (X-API-Key)", r.status_code == 401, r.status_code)
    r = c.post("/pipeline/analyze", headers={"Authorization": f"Bearer {key}"}, files={"file": ("a.docx", body, DOCX_MIME)})
    check("...and as Authorization: Bearer ak_ prefixed", r.status_code == 401, r.status_code)

    with session_scope() as s:
        keys_left = s.execute(select(ApiKeyRow).where(ApiKeyRow.user_id == uid)).scalars().all()
        tokens_left = s.execute(
            select(EmailVerifyTokenRow).where(EmailVerifyTokenRow.user_id == uid)
        ).scalars().all()
    check("no api_keys rows survive the account", keys_left == [], keys_left)
    check("no email_verify_tokens rows survive the account", tokens_left == [], tokens_left)

    # The ex-owner cannot revoke through the product either, so the key must be
    # dead on its own: re-registering the address mints a different identity.
    auth2, uid2 = signin(c, "ghost@example.com")
    check("re-registering the address mints a NEW user id", uid2 != uid, (uid, uid2))
    check("the new owner sees none of the old keys", c.get("/api-keys", headers=auth2).json() == [])
    r = c.post("/pipeline/analyze", headers={"X-API-Key": key}, files={"file": ("a.docx", body, DOCX_MIME)})
    check("the ghost key still can't follow the mailbox", r.status_code == 401, r.status_code)

    # --- 2. the resolver fails closed even on a hand-orphaned row ------------
    stray = "ak_live_" + "f" * 32
    with session_scope() as s:
        s.add(ApiKeyRow(
            id="orphan01", user_id="no-such-user", name="orphan",
            key_hash=hash_api_key(stray), key_prefix="ak_live_", created_at=datetime.utcnow(),
        ))
    check("an api key whose owner row is missing resolves to nobody", _api_key_user_id(stray) is None)

    # --- 3. a team owner with members cannot just vanish ---------------------
    c = client("10.9.0.2")
    owner_auth, owner_id = signin(c, "boss@example.com")
    _subscribe(owner_id, "sub_del_boss")
    check("owner creates a team", c.post("/teams", headers=owner_auth, json={"name": "Acme"}).status_code == 200)

    member_auth, member_id = signin(c, "colleague@example.com")
    _verify(c, member_auth)
    inv = c.post("/teams/invite", headers=owner_auth, json={"email": "colleague@example.com"}).json()
    check(
        "member accepts",
        c.post("/teams/accept", headers=member_auth, json={"token": inv["acceptUrl"].split("token=")[-1]}).status_code == 200,
    )

    r = c.delete("/auth/me", headers=owner_auth)
    check("owner of a populated team -> 409 team_has_members", r.status_code == 409 and "team_has_members" in r.text, r.text)
    check("the refusal changed nothing: owner still signed in", c.get("/auth/me", headers=owner_auth).status_code == 200)
    check("...and the team is still there", c.get("/teams/me", headers=owner_auth).json().get("team") is not None)

    # --- 4. a plain MEMBER deleting frees their seat, team survives ----------
    check("member deletes their own account -> 204", c.delete("/auth/me", headers=member_auth).status_code == 204)
    with session_scope() as s:
        seats = s.execute(select(TeamMemberRow.user_id)).scalars().all()
    check("the member's seat is released", member_id not in seats, seats)
    team = c.get("/teams/me", headers=owner_auth).json().get("team")
    check("the owner's team survives a member leaving that way", team is not None and team["seatsUsed"] == 1, team)

    # --- 5. a SOLE owner can still delete; the team goes with them -----------
    pending = c.post("/teams/invite", headers=owner_auth, json={"email": "never-joins@example.com"}).json()
    check("a pending invite exists before the delete", pending.get("status") == "pending", pending)
    check("sole owner deletes -> 204", c.delete("/auth/me", headers=owner_auth).status_code == 204)
    with session_scope() as s:
        teams_left = s.execute(select(TeamRow.id)).scalars().all()
        members_left = s.execute(select(TeamMemberRow.user_id)).scalars().all()
        invite_status = s.execute(
            select(TeamInviteRow.status).where(TeamInviteRow.id == pending["id"])
        ).scalars().one()
        sub_status = s.execute(
            select(SubscriptionRow.status).where(SubscriptionRow.user_id == owner_id)
        ).scalars().one()
        user_left = s.execute(select(UserRow.id).where(UserRow.id == owner_id)).scalars().first()
    check("the team is disbanded", teams_left == [], teams_left)
    check("no member rows are left behind", members_left == [], members_left)
    check("the pending invite is revoked, not left live", invite_status == "revoked", invite_status)
    check("the subscription no longer reads as active", sub_status == "canceled", sub_status)
    check("the user row is gone", user_left is None, user_left)

    # --- 6. a dangling owner_id never redirects a live user's wallet ---------
    c = client("10.9.0.3")
    solo_auth, solo_id = signin(c, "solo@example.com")
    _subscribe(solo_id, "sub_del_solo")
    c.post("/teams", headers=solo_auth, json={"name": "Orphanage"})
    stranded_auth, stranded_id = signin(c, "stranded@example.com")
    _verify(c, stranded_auth)
    inv = c.post("/teams/invite", headers=solo_auth, json={"email": "stranded@example.com"}).json()
    c.post("/teams/accept", headers=stranded_auth, json={"token": inv["acceptUrl"].split("token=")[-1]})
    check("a member resolves to the owner while the owner exists", resolve_credit_user_id(stranded_id) == solo_id)
    # Hand-orphan the team the way the old delete_me did, bypassing the route.
    with session_scope() as s:
        s.delete(s.execute(select(UserRow).where(UserRow.id == solo_id)).scalar_one())
    check(
        "with the owner row gone the member resolves to THEMSELVES",
        resolve_credit_user_id(stranded_id) == stranded_id,
        resolve_credit_user_id(stranded_id),
    )
    check("...and their balance read still works", c.get("/credits/balance", headers=stranded_auth).status_code == 200)

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
