"""Smoke: unlabeled DOCX content controls get an accessible name (auto-fix).

A content control (``w:sdt``) with no ``w:alias`` has no accessible name — a
screen reader announces "edit" with no idea what to type. Previously detect-only.
This pins the new CONSERVATIVE auto-fix:

  parser   -> counts unlabeled controls AND how many have a confident nearby
              label (form_fields_derivable); shared _derive_sdt_label
  analyzer -> FORM_FIELD_UNLABELED fires on the document root
  executor -> FILL_FORM_FIELD_LABELS succeeds only when >=1 is derivable
              (else SKIP -> manual review), reports "labeled N of M"
  writer   -> w:alias written on exactly the derivable controls (inline
              "Name: [__]" + table label cell); ambiguous ones left untouched
  honesty  -> FILL_FORM_FIELD_LABELS persists for docx; re-parsing the OUTPUT
              shows the unlabeled count dropped by exactly the derivable count,
              and the ambiguous control still needs manual work (idempotent).

Usage:
    python -m app.devtools.smoke_form_field_labels
"""

from __future__ import annotations

import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_ff_')}/s.db")

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.api.pipeline import _action_persists  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.docx_writer import write_remediated_docx  # noqa: E402

FLAG = "FORM_FIELD_UNLABELED"


def _new_sdt(placeholder: str = "Enter text", alias: str | None = None):
    sdt = OxmlElement("w:sdt")
    sdtPr = OxmlElement("w:sdtPr")
    if alias is not None:
        a = OxmlElement("w:alias")
        a.set(qn("w:val"), alias)
        sdtPr.append(a)
    sdt.append(sdtPr)
    content = OxmlElement("w:sdtContent")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = placeholder
    r.append(t)
    content.append(r)
    sdt.append(content)
    return sdt


def _add_inline_sdt(paragraph, label_text: str, alias: str | None = None):
    """A 'Label: [content control]' inline form field."""
    paragraph.add_run(label_text)
    paragraph._p.append(_new_sdt(alias=alias))


def _root_props(tree) -> dict:
    return tree.root.metadata.properties or {}


def _alias_vals(path) -> list[list[str]]:
    """Per content control (in document order), the list of its w:alias values."""
    doc = Document(str(path))
    out: list[list[str]] = []
    for sdt in doc.element.body.iter(qn("w:sdt")):
        sdtPr = sdt.find(qn("w:sdtPr"))
        aliases = sdtPr.findall(qn("w:alias")) if sdtPr is not None else []
        out.append([(a.get(qn("w:val")) or "") for a in aliases])
    return out


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="form_fields_"))
    apply_policy = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)

    d = Document()
    d.core_properties.title = "Form"
    # 1. Inline labeled-by-preceding-text control -> derivable "Full Name".
    _add_inline_sdt(d.add_paragraph(), "Full Name: ")
    # 2. Table: left cell "Email Address" | right cell with a bare control -> derivable.
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Email Address"
    table.rows[0].cells[1].paragraphs[0]._p.append(_new_sdt())
    # 3. Ambiguous: a control alone in a paragraph, no nearby label -> NOT derivable.
    d.add_paragraph()._p.append(_new_sdt())
    # 4. Already labeled control -> not counted unlabeled, never touched.
    _add_inline_sdt(d.add_paragraph(), "Phone: ", alias="Phone Number")
    src = tmp / "form.docx"
    d.save(str(src))

    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    props = _root_props(res.tree)
    check("FORM_FIELD_UNLABELED flag fires", any(
        f.code.value == FLAG for f in res.tree.root.accessibility_flags))
    check("4 total controls, 3 unlabeled, 2 derivable",
          props.get("form_fields_total") == 4
          and props.get("form_fields_unlabeled") == 3
          and props.get("form_fields_derivable") == 2,
          f"total={props.get('form_fields_total')} unlabeled={props.get('form_fields_unlabeled')} derivable={props.get('form_fields_derivable')}")
    check("FILL_FORM_FIELD_LABELS persists for docx (honesty matrix)",
          _action_persists("FILL_FORM_FIELD_LABELS", "docx"))

    plans = [p for p in plan_remediations(res.tree, apply_policy) if p.flag.code.value == FLAG]
    fill = [p for p in plans if any(a.action_code.value == "FILL_FORM_FIELD_LABELS" for a in p.actions)]
    check("a FILL_FORM_FIELD_LABELS plan exists", len(fill) >= 1, str([a.action_code.value for p in plans for a in p.actions]))
    execs = execute_plans(res.tree, plans)
    ok = [e for e in execs if e.status.value == "success" and e.action_code.value == "FILL_FORM_FIELD_LABELS"]
    check("fill executes successfully (labeled 2 of 3)", len(ok) == 1, str([(e.action_code.value, e.status.value, e.notes) for e in execs]))

    out = tmp / "form_fixed.docx"
    result = write_remediated_docx(src, res.tree, out)
    labels = [a for a in result["applied"] if a.get("kind") == "form_field_label"]
    check("writer applied exactly 2 form-field labels", len(labels) == 2, str(result["applied"]))

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8", "replace")
    check("derived aliases present in bytes",
          'w:val="Full Name"' in doc_xml and 'w:val="Email Address"' in doc_xml, "alias missing")
    check("pre-existing alias preserved", 'w:val="Phone Number"' in doc_xml)

    # Re-parse OUTPUT: the 2 derivable are labeled now, only the ambiguous remains.
    res2 = parse_to_tree(str(out))
    run_analyzers(res2.tree)
    p2 = _root_props(res2.tree)
    check("output: unlabeled dropped 3 -> 1 (only the ambiguous control remains)",
          p2.get("form_fields_unlabeled") == 1, f"got {p2.get('form_fields_unlabeled')}")
    check("output: 0 derivable left (the remaining one has no nearby label)",
          p2.get("form_fields_derivable") == 0, f"got {p2.get('form_fields_derivable')}")

    # Idempotent: re-running labels nothing more.
    plans2 = [p for p in plan_remediations(res2.tree, apply_policy) if p.flag.code.value == FLAG]
    execs2 = execute_plans(res2.tree, plans2)
    fill_ok2 = [e for e in execs2 if e.status.value == "success" and e.action_code.value == "FILL_FORM_FIELD_LABELS"]
    check("idempotent: no further successful fill (nothing derivable)", len(fill_ok2) == 0,
          str([(e.action_code.value, e.status.value) for e in execs2]))

    # =========================================================================
    # Fixture B — the two HIGH bugs the adversarial review found:
    #   (1) multiple controls in ONE paragraph: control #2 must derive "Date",
    #       NOT "Name <ctrl1 value> Date".
    #   (2) an existing EMPTY <w:alias>: must be OVERWRITTEN (one alias), not
    #       have a second <w:alias> appended (schema-invalid; LibreOffice drops
    #       both).
    # =========================================================================
    d2 = Document()
    d2.core_properties.title = "Form2"
    p = d2.add_paragraph()
    p.add_run("Name: ")
    p._p.append(_new_sdt(placeholder="John Doe"))   # control #1
    p.add_run(" Date: ")
    p._p.append(_new_sdt())                          # control #2
    p3 = d2.add_paragraph()
    p3.add_run("Country: ")
    p3._p.append(_new_sdt(alias=""))                 # existing EMPTY alias
    src2 = tmp / "form2.docx"
    d2.save(str(src2))

    r2 = parse_to_tree(str(src2))
    run_analyzers(r2.tree)
    pr2 = _root_props(r2.tree)
    check("multi-control + empty-alias: 3 unlabeled, 3 derivable",
          pr2.get("form_fields_unlabeled") == 3 and pr2.get("form_fields_derivable") == 3,
          f"unlabeled={pr2.get('form_fields_unlabeled')} derivable={pr2.get('form_fields_derivable')}")
    execute_plans(r2.tree, [p for p in plan_remediations(r2.tree, apply_policy) if p.flag.code.value == FLAG])
    out2 = tmp / "form2_fixed.docx"
    write_remediated_docx(src2, r2.tree, out2)
    vals = _alias_vals(out2)
    # Three controls, each with EXACTLY ONE alias.
    check("every control has exactly one <w:alias> (no double-alias corruption)",
          all(len(v) == 1 for v in vals), str(vals))
    flat = [v[0] for v in vals if v]
    check("control #1 = 'Name', control #2 = 'Date' (not the concatenation)",
          flat[0] == "Name" and flat[1] == "Date", str(flat))
    check("empty-alias control overwritten to 'Country'", "Country" in flat, str(flat))
    r2b = parse_to_tree(str(out2))
    run_analyzers(r2b.tree)
    check("multi/empty: all labeled now (unlabeled 3 -> 0)",
          (_root_props(r2b.tree).get("form_fields_unlabeled") or 0) == 0,
          f"got {_root_props(r2b.tree).get('form_fields_unlabeled')}")

    # =========================================================================
    # Fixture C — derivation must REJECT non-labels (sentences, questions,
    # instructions, headings, artifacts) and leave them for manual review.
    # =========================================================================
    d3 = Document()
    d3.core_properties.title = "Form3"
    for prompt in (
        "What is your name? ",                       # question
        "Please enter your full mailing address. ",  # sentence
        "See page 3 for detailed instructions on completing this ",  # instruction (>6 words)
        "Section 4 Employment ",                     # heading
        "* ",                                        # artifact only
    ):
        pp = d3.add_paragraph()
        pp.add_run(prompt)
        pp._p.append(_new_sdt())
    src3 = tmp / "form3.docx"
    d3.save(str(src3))
    r3 = parse_to_tree(str(src3))
    run_analyzers(r3.tree)
    pr3 = _root_props(r3.tree)
    check("non-label prompts are NOT auto-derived (all 5 left manual)",
          pr3.get("form_fields_unlabeled") == 5 and pr3.get("form_fields_derivable") == 0,
          f"unlabeled={pr3.get('form_fields_unlabeled')} derivable={pr3.get('form_fields_derivable')}")

    # Clean doc: a properly labeled control -> no flag, no action.
    c = Document()
    c.core_properties.title = "Clean"
    _add_inline_sdt(c.add_paragraph(), "City: ", alias="City")
    csrc = tmp / "clean.docx"
    c.save(str(csrc))
    rc = parse_to_tree(str(csrc))
    run_analyzers(rc.tree)
    check("clean doc: no form-field flag",
          not any(f.code.value == FLAG for f in rc.tree.root.accessibility_flags))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
