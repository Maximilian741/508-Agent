"""Smoke: DOCX text-box (w:txbxContent) content is analyzed AND remediable.

Text boxes are invisible to ``doc.paragraphs`` — before this, sidebar and
callout content (everywhere in government documents) was never analyzed at
all. Pins:

  - text-box paragraph text becomes a ParagraphNode (marked in_text_box)
  - text-box links (w:hyperlink AND fldSimple) become LinkNodes, get flagged
    ("click here"), and IMPROVE_LINK_TEXT genuinely rewrites them in the
    OUTPUT BYTES (docx-tblink ids mint identically in parser and writer)
  - body link ids are NOT disturbed by text-box links (separate id space) —
    rewrites land on the right elements when both exist
  - plain documents emit no tb nodes (no regression)

Usage:
    python -m app.devtools.smoke_textbox_content
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_tb_')}/s.db")

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from lxml import etree  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.models.accessibility import ContentKind, LinkNode, NodeContent, ParagraphNode, iter_reading_order  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.writers import write_remediated  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"


def _add_text_box(doc, inner_xml: str, shape_id: int = 99) -> None:
    """Append a paragraph hosting a wps text box whose w:txbxContent is
    ``inner_xml`` (one or more <w:p> elements)."""
    xml = (
        f'<w:drawing xmlns:w="{W}"><wp:inline xmlns:wp="{WP}" distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="3000000" cy="1000000"/><wp:docPr id="{shape_id}" name="TextBox {shape_id}"/>'
        f'<a:graphic xmlns:a="{A}"><a:graphicData uri="{WPS}">'
        f'<wps:wsp xmlns:wps="{WPS}"><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="3000000" cy="1000000"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>'
        f"<wps:txbx><w:txbxContent>{inner_xml}</w:txbxContent></wps:txbx>"
        f"<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing>"
    )
    drawing = etree.fromstring(xml)
    host = doc.add_paragraph()
    run = OxmlElement("w:r")
    run.append(drawing)
    host._p.append(run)


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="tb_"))

    # --- fixture: body link + text box (text, hyperlink, fldSimple link) -----
    d = Document()
    d.core_properties.title = "TB Test"
    d.add_paragraph("Body paragraph outside any box.")
    body_para = d.add_paragraph("Body link: ")
    rid_body = d.part.relate_to(
        "https://example.com/body",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), rid_body)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "read the methodology"
    r.append(t); hl.append(r); body_para._p.append(hl)

    rid_tb = d.part.relate_to(
        "https://example.com/sidebar",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    inner = (
        "<w:p><w:r><w:t>Important sidebar callout text lives here.</w:t></w:r></w:p>"
        f'<w:p><w:hyperlink r:id="{rid_tb}" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        "<w:r><w:t>click here</w:t></w:r></w:hyperlink></w:p>"
        '<w:p><w:fldSimple w:instr=" HYPERLINK &quot;https://example.gov/field&quot; ">'
        "<w:r><w:t>https://example.gov/field</w:t></w:r></w:fldSimple></w:p>"
    )
    _add_text_box(d, inner)
    src = tmp / "tb.docx"
    d.save(str(src))

    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    nodes = list(iter_reading_order(res.tree.root))
    tb_paras = [n for n in nodes if isinstance(n, ParagraphNode) and (n.metadata.properties or {}).get("in_text_box")]
    check("text-box paragraph text visible", any("sidebar callout" in (n.content.text or "") for n in tb_paras), str([n.content.text for n in tb_paras]))
    links = [n for n in nodes if isinstance(n, LinkNode)]
    tb_links = [n for n in links if (n.metadata.properties or {}).get("in_text_box")]
    check("2 text-box links emitted (hyperlink + fldSimple)", len(tb_links) == 2, str([(n.id, n.content.text) for n in links]))
    check("tb link ids use the docx-tblink space", all(n.id.startswith("docx-tblink") for n in tb_links))
    check("body link unaffected (docx-link-1)", any(n.id == "docx-link-1" and n.content.text == "read the methodology" for n in links))
    flagged = [n for n in tb_links if any(f.code.value == "LINK_TEXT_NON_DESCRIPTIVE" for f in n.accessibility_flags)]
    check("both tb links flagged non-descriptive", len(flagged) == 2, str([(n.content.text, [f.code.value for f in n.accessibility_flags]) for n in tb_links]))

    # --- remediation round-trip: rewrite ALL links, body + tb ----------------
    for i, ln in enumerate(links):
        ln.content = NodeContent(kind=ContentKind.TEXT, text=f"FIXED-{i}: {ln.id}")
    out = tmp / "tb_out.docx"
    write_remediated(src, res.tree, out, source_format=res.format)
    res2 = parse_to_tree(str(out))
    links2 = [n for n in iter_reading_order(res2.tree.root) if isinstance(n, LinkNode)]
    texts2 = [n.content.text for n in links2]
    check(
        "every link rewritten in OUTPUT BYTES with no cross-assignment",
        all(t and t.startswith("FIXED-") and links2[i].id in t for i, t in enumerate(texts2)),
        str(list(zip([n.id for n in links2], texts2))),
    )
    check("tb links rewritten specifically", sum(1 for n in links2 if (n.metadata.properties or {}).get("in_text_box") and n.content.text.startswith("FIXED-")) == 2)

    # --- plain doc: no tb nodes ----------------------------------------------
    c = Document()
    c.core_properties.title = "Plain"
    c.add_paragraph("Nothing fancy here.")
    plain = tmp / "plain.docx"
    c.save(str(plain))
    res3 = parse_to_tree(str(plain))
    n_tb = sum(1 for n in iter_reading_order(res3.tree.root) if (n.metadata.properties or {}).get("in_text_box"))
    check("plain doc emits no text-box nodes", n_tb == 0)

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
