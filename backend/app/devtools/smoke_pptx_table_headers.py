"""Smoke: PPTX table headers reach parity with DOCX (detect + persist).

Before this, PPTX tables had the same gaps DOCX had pre-#15/#17:
* the parser typed row 0 as a header unconditionally, so TABLE_MISSING_HEADERS
  never fired; and
* ADD_TABLE_HEADERS wasn't credited as persisted for pptx.

Now:
* Detection is firstRow-aware — a data grid (>=3 rows, >=2 cols) whose
  PowerPoint "Header Row" band is OFF (Table.first_row is False) is flagged.
* Promote path: the writer sets <a:tblPr firstRow="1"> on the existing first row.
* First row not header-like (a sentence, numbers): REFUSED. Header text is
  never invented — a "Column 1 | Column 2" row used to be inserted and charged,
  which tells a screen-reader user nothing. The table stays flagged for a
  person, and the file is not touched.

Usage:
    python -m app.devtools.smoke_pptx_table_headers
"""

from __future__ import annotations

import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_pptxtbl_')}/s.db"

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.api.pipeline import _action_persists  # noqa: E402
from app.models.accessibility import AccessibilityFlagCode as F  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers import write_remediated  # noqa: E402

_POLICY = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)


def _build(path: Path, rows: int, cols: int, first_row_text, header_band: bool):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    gf = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(0.5), Inches(9), Inches(4))
    table = gf.table
    table.first_row = header_band
    for r in range(rows):
        for c in range(cols):
            table.cell(r, c).text = (first_row_text if r == 0 else f"v{r}{c}")
    prs.save(str(path))


def _codes(tree) -> set:
    found = set()

    def walk(n):
        for f in n.accessibility_flags:
            found.add(f.code.value)
        for ch in n.children:
            walk(ch)

    walk(tree.root)
    return found


def _remediate(path: Path, out: Path):
    res = parse_to_tree(str(path))
    run_analyzers(res.tree)
    plans = [p for p in plan_remediations(res.tree, _POLICY) if p.flag.code == F.TABLE_MISSING_HEADERS]
    execs = execute_plans(res.tree, plans) if plans else []
    write_remediated(path, res.tree, out, source_format=res.format)
    slide_xml = b""
    with zipfile.ZipFile(out) as z:
        for nm in z.namelist():
            if nm.startswith("ppt/slides/slide") and nm.endswith(".xml"):
                slide_xml += z.read(nm)
    return execs, slide_xml.decode("utf-8", "ignore")


def main() -> int:
    failures = 0

    def check(name, cond):
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name)
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp())
    TMH = F.TABLE_MISSING_HEADERS.value

    check("ADD_TABLE_HEADERS now persists for pptx (score honesty)", _action_persists("ADD_TABLE_HEADERS", "pptx"))

    # --- Detection ---
    p = tmp / "noheader.pptx"
    _build(p, 3, 3, "Region", header_band=False)
    check("pptx data grid with header band OFF -> TABLE_MISSING_HEADERS", TMH in _codes(run_analyzers(parse_to_tree(str(p)).tree)))

    p = tmp / "withband.pptx"
    _build(p, 3, 3, "Region", header_band=True)
    check("pptx data grid with header band ON -> not flagged", TMH not in _codes(run_analyzers(parse_to_tree(str(p)).tree)))

    p = tmp / "small.pptx"
    _build(p, 2, 2, "A", header_band=False)
    check("pptx 2x2 table -> not flagged (avoids small/layout FP)", TMH not in _codes(run_analyzers(parse_to_tree(str(p)).tree)))

    # --- Promote: label-like first row, band off -> firstRow=1 persists ---
    p = tmp / "promote.pptx"
    _build(p, 3, 3, "Region", header_band=False)
    execs, xml = _remediate(p, tmp / "promote_out.pptx")
    check("promote: ADD_TABLE_HEADERS ran", any(e.action_code.value == "ADD_TABLE_HEADERS" and e.status.value == "success" for e in execs))
    check("promote: firstRow band set on output", 'firstRow="1"' in xml)
    check("promote: no row inserted (still 3 <a:tr>)", xml.count("<a:tr ") + xml.count("<a:tr>") == 3)
    Presentation(str(tmp / "promote_out.pptx"))
    check("promote: output reopens cleanly", True)

    # --- Sentence first row -> refused: no invented header row, no band ---
    p = tmp / "synth.pptx"
    _build(p, 3, 3, "This whole row is a sentence.", header_band=False)
    execs2, xml2 = _remediate(p, tmp / "synth_out.pptx")
    refused = [e for e in execs2 if e.action_code.value == "ADD_TABLE_HEADERS"]
    check("no header row: ADD_TABLE_HEADERS refused with a reason",
          bool(refused) and all(e.status.value == "skipped" and "not charged" in (e.notes or "") for e in refused))
    check("no header row: firstRow band NOT set on output", 'firstRow="1"' not in xml2)
    check("no header row: no row inserted (still 3 <a:tr>)", xml2.count("<a:tr ") + xml2.count("<a:tr>") == 3)
    check("no header row: no placeholder header text written", "Column 1" not in xml2)
    out = Presentation(str(tmp / "synth_out.pptx"))
    tbl = None
    for shp in out.slides[0].shapes:
        if shp.has_table:
            tbl = shp.table
            break
    check("no header row: reopened table still has 3 rows, 3 cols", tbl is not None and len(tbl.rows) == 3 and len(tbl.columns) == 3)

    # --- Defense in depth: a too-narrow header is clamped to the grid width so
    # the table always reopens (expert-found robustness gap). ---
    from app.writers.pptx_writer import _insert_pptx_header_row

    p = tmp / "narrow.pptx"
    _build(p, 3, 3, "x", header_band=False)
    prs = Presentation(str(p))
    narrow_tbl = next(s.table for s in prs.slides[0].shapes if s.has_table)
    _insert_pptx_header_row(narrow_tbl, ["only", "two"])  # fewer than 3 grid cols
    p_out = tmp / "narrow_out.pptx"
    prs.save(str(p_out))
    reopened = Presentation(str(p_out))
    rtbl = next(s.table for s in reopened.slides[0].shapes if s.has_table)
    check("narrow header clamped to grid width (reopens, row has 3 cells)",
          len(rtbl.rows[0].cells) == 3 and len(rtbl.columns) == 3)

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
