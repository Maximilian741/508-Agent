"""Smoke: a signed artifact URL dies with the session that minted it.

``GET /pipeline/files`` has no auth dependency on purpose — the deferred-charge
path exists precisely because the client was already gone, so the signature IS
the credential. But the signature used to cover only job/filename/expiry, which
made it an UNREVOCABLE BEARER TOKEN: a URL lifted from a hijacked session (or
from ``GET /pipeline/jobs``, which re-mints one per listing with a 24h TTL) kept
serving the remediated document after sign-out, after a password reset, and
after the account was deleted — whose job directory was never purged either.

The signature is now bound to the owner's current ``token_version``, the same
counter that already invalidates session JWTs and API keys. The binding is not
in the URL: the download route re-derives it from the job manifest's userId plus
the live user row, so nothing the holder of the URL can see or forge is
involved. Pinned here:

  * the owner's ordinary download works, with and without a session header;
  * sign-out kills every outstanding URL (410), and /pipeline/jobs re-mints one
    that works — so the legitimate flow is never stuck;
  * a password change does the same;
  * deleting the account makes its URLs stop working AND removes the bytes;
  * a tampered signature is still a plain 403, and an expired one still a 410.

Usage:
    python -m app.devtools.smoke_artifact_url_revocation
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="508_smoke_revoke_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/s.db"
os.environ["MATERIALIZED_ROOT"] = str(Path(_TMP) / "materialized")

from PIL import Image  # noqa: E402
from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import select  # noqa: E402

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PASSWORD = "revokepass12345"


def _docx() -> bytes:
    buf = io.BytesIO()
    png = io.BytesIO()
    Image.new("RGB", (24, 24), (10, 120, 200)).save(png, format="PNG")
    d = Document()
    d.add_heading("Section", level=1)
    d.add_paragraph("The chart below shows quarterly revenue by region for the year.")
    d.add_paragraph().add_run().add_picture(io.BytesIO(png.getvalue()))
    d.save(buf)
    return buf.getvalue()


def main() -> int:  # noqa: PLR0915
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.config import get_settings
    from app.db.models import UserRow
    from app.db.session_sqlalchemy import session_scope
    from app.main import app

    client = TestClient(app, raise_server_exceptions=False)
    data = _docx()

    def sign_in(email: str) -> dict:
        r = client.post("/auth/sign-in", json={"email": email, "displayName": "V", "password": PASSWORD})
        assert r.status_code == 200, r.text
        return {"Authorization": f"Bearer {r.json()['token']}"}

    def set_balance(email: str, n: int) -> None:
        with session_scope() as s:
            s.execute(select(UserRow).where(UserRow.email == email)).scalars().first().credits_balance = n

    def make_job(headers: dict, name: str = "book.docx") -> dict:
        a = client.post("/pipeline/analyze", files={"file": (name, data, DOCX)}, headers=headers)
        alt = [v["id"] for v in a.json()["violations"] if v["ruleId"] == "MISSING_ALT_TEXT"]
        assert alt, str([v["ruleId"] for v in a.json()["violations"]])
        r = client.post(
            "/pipeline/remediate",
            files={"file": (name, data, DOCX)},
            data={"approved_violations": json.dumps(alt), "rejected_violations": "[]"},
            headers=headers,
        )
        assert r.status_code == 200, r.text
        return r.json()

    # ---- 1. the owner's ordinary download, session or not ----------------
    h = sign_in("victim@example.com")
    set_balance("victim@example.com", 100)
    job = make_job(h)
    url = job["downloadUrl"]
    r = client.get(url, headers=h)
    check("owner downloads their artifact -> 200", r.status_code == 200 and len(r.content) > 0, str(r.status_code))
    r = client.get(url)
    check("...and the signed URL still works with NO session header (deferred-charge path)",
          r.status_code == 200 and len(r.content) > 0, str(r.status_code))

    # ---- 2. sign-out revokes it, and /pipeline/jobs re-mints -------------
    check("sign-out -> 204", client.post("/auth/sign-out", headers=h).status_code == 204)
    check("the old session token is dead", client.get("/auth/me", headers=h).status_code == 401)
    r = client.get(url)
    check("the artifact URL minted under that session is GONE (410)", r.status_code == 410, str(r.status_code))

    h2 = sign_in("victim@example.com")
    jobs = client.get("/pipeline/jobs", headers=h2).json()["jobs"]
    fresh = next((j for j in jobs if j["jobId"] == job["jobId"]), None)
    check("the owner can still find the job after signing back in", fresh is not None, str(jobs)[:200])
    if fresh:
        check("...and its re-minted URL works", client.get(fresh["downloadUrl"]).status_code == 200)
        check("...while the pre-sign-out URL stays dead", client.get(url).status_code == 410)
        url = fresh["downloadUrl"]

    # ---- 3. a password change revokes it too ----------------------------
    r = client.post(
        "/auth/set-password",
        json={"password": "brandnewpass98765", "currentPassword": PASSWORD},
        headers=h2,
    )
    check("set-password -> 200", r.status_code == 200, r.text[:160])
    check("the artifact URL minted before the password change is GONE (410)",
          client.get(url).status_code == 410, str(client.get(url).status_code))
    h3 = {"Authorization": f"Bearer {r.json()['token']}"} if r.json().get("token") else None
    if h3:
        jobs = client.get("/pipeline/jobs", headers=h3).json()["jobs"]
        fresh = next((j for j in jobs if j["jobId"] == job["jobId"]), None)
        check("...and the owner's fresh URL works again",
              fresh is not None and client.get(fresh["downloadUrl"]).status_code == 200)

    # ---- 4. deleting the account revokes AND purges ----------------------
    h4 = sign_in("deleter@example.com")
    set_balance("deleter@example.com", 100)
    job2 = make_job(h4, "report.docx")
    url2 = job2["downloadUrl"]
    check("the second account's artifact downloads", client.get(url2).status_code == 200)
    job_dir = get_settings().materialized_root / "pipeline" / job2["jobId"]
    check("DELETE /auth/me -> 204", client.delete("/auth/me", headers=h4).status_code == 204)
    check("the deleted account's artifact URL no longer serves the document",
          client.get(url2).status_code in (403, 404), str(client.get(url2).status_code))
    check("...and its bytes are gone from disk", not job_dir.exists(), str(job_dir))

    # ---- 5. the ordinary signature failures still behave -----------------
    h5 = sign_in("tamper@example.com")
    set_balance("tamper@example.com", 100)
    job3 = make_job(h5, "third.docx")
    good = job3["downloadUrl"]
    bad = good[:-1] + ("0" if good[-1] != "0" else "1")
    check("a tampered signature -> 403", client.get(bad).status_code == 403, str(client.get(bad).status_code))
    stale = good.split("?")[0] + "?exp=1&sig=" + good.split("sig=")[1]
    check("an expired URL -> 410", client.get(stale).status_code == 410, str(client.get(stale).status_code))
    check("another tenant cannot list the job", not any(
        j["jobId"] == job3["jobId"] for j in client.get("/pipeline/jobs", headers=sign_in("bystander@example.com")).json()["jobs"]
    ))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except AssertionError as exc:
        print(f"FAIL artifact-url-revocation smoke: {exc}")
        sys.exit(1)
    except Exception as exc:  # pragma: no cover
        print(f"FAIL artifact-url-revocation smoke: {exc.__class__.__name__}: {exc}")
        sys.exit(1)
