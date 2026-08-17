"""Smoke: the fast paragraph-style lookup answers exactly like python-docx.

``paragraph.style`` resolves the DEFAULT paragraph style — the case for most
body text — by walking every style in styles.xml per call. Profiling a
200-page document put that one property at 89% of the DOCX writer's wall
time; the parser paid it again. A 500-page document took ~44s end to end,
most of the way to a proxy timeout, for a lookup whose answer is the same for
every body paragraph in the file.

:func:`paragraph_style_name` reads the ``w:pStyle`` id off the XML and
resolves names through a per-document cache, falling back to python-docx for
anything it cannot resolve. The parser and writer BOTH use it, and they must
agree on ids — so the one thing that matters is that it never disagrees with
what python-docx would have said. Pinned here on a document that mixes every
case: Title, Heading 1-6, explicit named styles, list styles, and empty
paragraphs with no pStyle at all (the implicit-default path).

Also pins the speedup, loosely (>=10x), so a well-meaning "simplification"
back to ``paragraph.style.name`` fails loudly instead of quietly costing 40s.

Usage:
    python -m app.devtools.smoke_docx_style_lookup
"""

from __future__ import annotations

import os
import sys
import tempfile
import time

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_dsl_')}/s.db")

import docx  # noqa: E402

from app.parsers.docx_parser import paragraph_style_name  # noqa: E402


def main() -> int:
    failures = 0

    def check(name, cond, extra=""):
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    d = docx.Document()
    d.add_heading("Title", 0)
    for i in range(300):
        d.add_heading(f"Heading {i}", level=(i % 6) + 1)
        d.add_paragraph("body text " * 5)                       # implicit default
        d.add_paragraph("quote", style="Intense Quote" if i % 7 == 0 else "Normal")
        d.add_paragraph("item", style="List Bullet")
        d.add_paragraph("")                                     # empty, no pStyle
    path = tempfile.mkdtemp(prefix="508_dsl_") + "/mixed.docx"
    d.save(path)
    doc = docx.Document(path)

    t0 = time.monotonic()
    slow = [(p.style.name or "") if p.style else "" for p in doc.paragraphs]
    t1 = time.monotonic()
    cache: dict = {}
    fast = [paragraph_style_name(p, cache) for p in doc.paragraphs]
    t2 = time.monotonic()

    mismatches = [(i, a, b) for i, (a, b) in enumerate(zip(slow, fast)) if a != b]
    check(f"fast lookup agrees with python-docx on all {len(slow)} paragraphs",
          not mismatches, str(mismatches[:5]))
    check("every heading level 0-6, named, list and empty paragraph was covered",
          len(set(slow)) >= 9, str(sorted(set(slow))))
    check("cache stays small (one entry per DISTINCT style id, not per paragraph)",
          len(cache) <= 16, str(len(cache)))
    slow_s, fast_s = t1 - t0, max(t2 - t1, 1e-9)
    check("fast path is at least 10x faster than paragraph.style",
          slow_s / fast_s >= 10, f"slow={slow_s:.3f}s fast={fast_s:.3f}s")

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
