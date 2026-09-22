"""Smoke: link-text remediation now PERSISTS to the output (DOCX + PPTX).

Previously IMPROVE_LINK_TEXT mutated only the in-memory tree (no writer), so the
fix never reached the downloaded file. Now the DOCX and PPTX writers rewrite the
hyperlink's display text — preserving the link target and formatting — so the
rewrite round-trips.

Usage:
    python -m app.devtools.smoke_link_text_writer
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_link_')}/s.db"

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.models.accessibility import ContentKind, LinkNode, NodeContent, iter_reading_order  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.writers import write_remediated  # noqa: E402

_NEW = "Visit example.com"
_HREF = "https://example.com/very/long/path"


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp())

    def _links(path: Path):
        res = parse_to_tree(str(path))
        return res, [n for n in iter_reading_order(res.tree.root) if isinstance(n, LinkNode)]

    def _rewrite_and_check(src: Path, fmt: str) -> None:
        res, links = _links(src)
        check(f"{fmt}: hyperlink 'click here' detected", bool(links) and links[0].content.text == "click here")
        if not links:
            return
        links[0].content = NodeContent(kind=ContentKind.TEXT, text=_NEW)
        out = src.with_name(src.stem + "_out" + src.suffix)
        write_remediated(src, res.tree, out, source_format=res.format)
        _, links2 = _links(out)
        check(f"{fmt}: link text rewritten in output", bool(links2) and links2[0].content.text == _NEW)
        check(f"{fmt}: link target preserved", bool(links2) and bool(links2[0].target) and "example.com" in links2[0].target)
        # not corrupt
        (Document if fmt == "docx" else Presentation)(str(out))
        check(f"{fmt}: output reopens cleanly", True)

    # DOCX
    dp = tmp / "in.docx"
    d = Document(); para = d.add_paragraph("See ")
    rid = d.part.relate_to(_HREF, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "click here"; r.append(t); hl.append(r)
    para._p.append(hl); d.save(str(dp))
    _rewrite_and_check(dp, "docx")

    # PPTX
    pp = tmp / "in.pptx"
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = tb.text_frame.paragraphs[0].add_run(); run.text = "click here"
    run.hyperlink.address = _HREF
    prs.save(str(pp))
    _rewrite_and_check(pp, "pptx")

    # --- DOCX multi-link id alignment (3 links get distinct text, no cross-assignment) ---
    md = tmp / "multi.docx"
    d2 = Document()
    for i in range(3):
        para = d2.add_paragraph(f"Item {i}: ")
        rid = d2.part.relate_to(f"https://example.com/{i}", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), rid)
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "click here"; r.append(t); hl.append(r)
        para._p.append(hl)
    d2.save(str(md))
    res, links = _links(md)
    for i, ln in enumerate(links):
        ln.content = NodeContent(kind=ContentKind.TEXT, text=f"LINK-{i}")
    mo = tmp / "multi_out.docx"; write_remediated(md, res.tree, mo, source_format=res.format)
    _, links2 = _links(mo)
    check("docx: 3 links each get their own text (no cross-assignment)", [l.content.text for l in links2] == ["LINK-0", "LINK-1", "LINK-2"])

    # --- PPTX split-run hyperlink: one visual link across 2 same-address runs -> no doubled text ---
    sp = tmp / "split.pptx"
    prs2 = Presentation(); slide2 = prs2.slides.add_slide(prs2.slide_layouts[6])
    tb2 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    pp2 = tb2.text_frame.paragraphs[0]
    a = pp2.add_run(); a.text = "click "; a.hyperlink.address = "https://example.com/same"
    b = pp2.add_run(); b.text = "here"; b.hyperlink.address = "https://example.com/same"
    prs2.save(str(sp))
    res, links = _links(sp)
    check("pptx: split-run link coalesced into one node", len(links) == 1 and links[0].content.text == "click here")
    if links:
        links[0].content = NodeContent(kind=ContentKind.TEXT, text=_NEW)
    so = tmp / "split_out.pptx"; write_remediated(sp, res.tree, so, source_format=res.format)
    import zipfile
    sx = b""
    with zipfile.ZipFile(so) as z:
        for nm in z.namelist():
            if nm.startswith("ppt/slides/slide") and nm.endswith(".xml"):
                sx += z.read(nm)
    check("pptx: split-run rewrite not duplicated", sx.count(_NEW.encode()) == 1)

    # --- DOCX fldSimple HYPERLINK field: detected + rewritten -----------------
    fp = tmp / "fld.docx"
    d3 = Document()
    para = d3.add_paragraph("Field link: ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), ' HYPERLINK "https://example.gov/report" ')
    fr = OxmlElement("w:r"); ft = OxmlElement("w:t"); ft.text = "click here"
    fr.append(ft); fld.append(fr); para._p.append(fld)
    d3.save(str(fp))
    res, links = _links(fp)
    check("docx: fldSimple link detected", len(links) == 1 and links[0].content.text == "click here")
    check("docx: fldSimple target parsed from instr", bool(links) and links[0].target == "https://example.gov/report")
    if links:
        links[0].content = NodeContent(kind=ContentKind.TEXT, text=_NEW)
    fo = tmp / "fld_out.docx"; write_remediated(fp, res.tree, fo, source_format=res.format)
    _, links2 = _links(fo)
    check("docx: fldSimple text rewritten in output", bool(links2) and links2[0].content.text == _NEW)
    check("docx: fldSimple instr target preserved", bool(links2) and links2[0].target == "https://example.gov/report")

    # --- DOCX in-cell hyperlink: detected + rewritten -------------------------
    cp = tmp / "cell.docx"
    d4 = Document()
    d4.add_paragraph("Intro paragraph.")
    tbl = d4.add_table(rows=2, cols=2)
    cell = tbl.cell(1, 1)
    cell_para = cell.paragraphs[0]
    cell_para.add_run("See ")
    rid = d4.part.relate_to(
        "https://example.com/cell-target",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "click here"; r.append(t); hl.append(r)
    cell_para._p.append(hl)
    d4.save(str(cp))
    res, links = _links(cp)
    check("docx: in-cell hyperlink detected", len(links) == 1 and links[0].content.text == "click here")
    if links:
        links[0].content = NodeContent(kind=ContentKind.TEXT, text=_NEW)
    co = tmp / "cell_out.docx"; write_remediated(cp, res.tree, co, source_format=res.format)
    _, links2 = _links(co)
    check("docx: in-cell link text rewritten in output", bool(links2) and links2[0].content.text == _NEW)
    check(
        "docx: in-cell link target preserved",
        bool(links2) and links2[0].target == "https://example.com/cell-target",
    )

    # --- DOCX merged cell: link emitted exactly ONCE ---------------------------
    mp2 = tmp / "merged.docx"
    d5 = Document()
    tbl2 = d5.add_table(rows=2, cols=2)
    merged = tbl2.cell(0, 0).merge(tbl2.cell(0, 1))
    mpar = merged.paragraphs[0]
    rid = d5.part.relate_to(
        "https://example.com/merged",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "merged link"; r.append(t); hl.append(r)
    mpar._p.append(hl)
    d5.save(str(mp2))
    _, links = _links(mp2)
    check("docx: merged-cell link emitted exactly once", len(links) == 1)

    # --- DOCX id-alignment torture: anchor-only link + heading link + mix -----
    # A text-less anchor used to desync the writer's paragraph index (it
    # skipped the paragraph; the parser minted docx-p). A hyperlink inside a
    # Heading paragraph is a real link a screen-reader user tabs to: it used
    # to be skipped by the parser (heading branch first) — now it is emitted
    # AND rewritten, in the same id order by parser and writer (the writer
    # runs the parser's own walk), so nothing drifts.
    tp = tmp / "torture.docx"
    d6 = Document()
    p_anchor = d6.add_paragraph("Anchor paragraph keeps its docx-p id ")
    bare = OxmlElement("w:hyperlink"); bare.set(qn("w:anchor"), "top")
    p_anchor._p.append(bare)  # no runs, no text
    h = d6.add_heading("Heading with ", level=1)
    rid = d6.part.relate_to(
        "https://example.com/heading",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hh = OxmlElement("w:hyperlink"); hh.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "a link inside"; r.append(t); hh.append(r)
    h._p.append(hh)
    para = d6.add_paragraph("Body: ")
    rid = d6.part.relate_to(
        "https://example.com/body",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hb = OxmlElement("w:hyperlink"); hb.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "click here"; r.append(t); hb.append(r)
    para._p.append(hb)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), ' HYPERLINK "https://example.com/fld" ')
    fr = OxmlElement("w:r"); ft = OxmlElement("w:t"); ft.text = "field link"
    fr.append(ft); fld.append(fr)
    d6.add_paragraph("Field: ")._p.append(fld)
    d6.save(str(tp))
    res, links = _links(tp)
    check(
        "docx: torture doc emits heading+body+fld links (text-less anchor excluded)",
        [l.content.text for l in links] == ["a link inside", "click here", "field link"],
        str([l.content.text for l in links]),
    )
    for i, ln in enumerate(links):
        ln.content = NodeContent(kind=ContentKind.TEXT, text=f"FIXED-{i}")
    to = tmp / "torture_out.docx"; write_remediated(tp, res.tree, to, source_format=res.format)
    _, links2 = _links(to)
    check(
        "docx: torture rewrite lands on the right links (no drift)",
        [(l.content.text, l.target) for l in links2]
        == [("FIXED-0", "https://example.com/heading"), ("FIXED-1", "https://example.com/body"),
            ("FIXED-2", "https://example.com/fld")],
        str([(l.content.text, l.target) for l in links2]),
    )

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
