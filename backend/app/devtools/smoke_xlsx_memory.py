"""Smoke: a big workbook is analyzed and fixed in bounded memory, and nothing
it did not read is passed off as checked.

The deploy docs size the whole stack at 2 GB of RAM. A normal data export used
to blow through that: /pipeline/remediate on a 10 MB workbook of four dense
sheets peaked at 2.9 GB (every sheet's scan window alive at once, a second
full re-scan in verification, and each touched sheet parsed into an lxml DOM
to append one element); a 4 MB one-sheet ledger of 100k rows peaked at 950 MB.

  * both workbooks go through /pipeline/analyze and /pipeline/remediate over
    HTTP in a FRESH interpreter each, whose own peak memory (the OS's peak
    working set / max RSS, which sees lxml's C allocations that tracemalloc
    cannot) must stay under PEAK_LIMIT_MB, and every clear header row is
    still found and fixed;
  * the per-workbook cell budget: a sheet whose share is spent still follows
    its data block to the end and gets the fix; data nothing continues is
    disclosed (ANALYSIS_TRUNCATED); a sheet with no budget left is not
    mistaken for an empty one; the scan keeps only the cells read later;
  * hidden sheets are not read cell by cell;
  * <tableParts> is spliced into a sheet that already has an Excel table and
    an <extLst>, past a comment in its data that reads "</sheetData>": the
    old table, the new one and the sparkline all survive, and every byte
    before the real </sheetData> is the customer's;
  * a busy server says so (503, "not charged") instead of blaming the file.

Usage:
    python -m app.devtools.smoke_xlsx_memory
"""

from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import tempfile
import time
import warnings
import zipfile
from pathlib import Path
from typing import Callable, Iterable, List, Optional, Tuple

_TMP = tempfile.mkdtemp(prefix="508_smoke_xlsx_mem_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/s.db"
os.environ["MATERIALIZED_ROOT"] = str(Path(_TMP) / "materialized")
os.environ.setdefault("APP_SECRET", "x" * 64)
os.environ["SEMANTIC_PROVIDER"] = "heuristic"
for _key in ("OPENAI_API_KEY", "ANTHROPIC_API_KEY"):
    os.environ.pop(_key, None)
warnings.simplefilter("ignore")

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Measured after the fix: ~230 MB (four sheets) and ~170 MB (ledger), of which
# ~140 MB is the interpreter with the app imported. Before: 2.9 GB and 950 MB.
PEAK_LIMIT_MB = 450
GROWTH_LIMIT_MB = 300

_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def _peak_mb() -> float:
    """This process's peak memory so far, in MB, as the OS counts it."""
    if sys.platform == "win32":
        import ctypes
        from ctypes import wintypes

        class _Counters(ctypes.Structure):
            _fields_ = [
                ("cb", wintypes.DWORD), ("PageFaultCount", wintypes.DWORD),
                ("PeakWorkingSetSize", ctypes.c_size_t), ("WorkingSetSize", ctypes.c_size_t),
                ("QuotaPeakPagedPoolUsage", ctypes.c_size_t), ("QuotaPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t), ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
                ("PagefileUsage", ctypes.c_size_t), ("PeakPagefileUsage", ctypes.c_size_t),
            ]

        kernel32 = ctypes.WinDLL("kernel32")
        psapi = ctypes.WinDLL("psapi")
        kernel32.GetCurrentProcess.restype = wintypes.HANDLE
        psapi.GetProcessMemoryInfo.argtypes = [wintypes.HANDLE, ctypes.POINTER(_Counters), wintypes.DWORD]
        counters = _Counters()
        counters.cb = ctypes.sizeof(_Counters)
        if not psapi.GetProcessMemoryInfo(kernel32.GetCurrentProcess(), ctypes.byref(counters), counters.cb):
            raise OSError("GetProcessMemoryInfo failed")
        return counters.PeakWorkingSetSize / 1e6
    import resource

    peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return peak / 1e6 if sys.platform == "darwin" else peak / 1e3  # bytes on macOS, KiB on Linux


# ---------------------------------------------------------------------------
# Workbooks, written as raw SpreadsheetML (XlsxWriter needs ~9 s for the
# four-sheet one; this needs ~2)
# ---------------------------------------------------------------------------

Row = Tuple[int, List[object]]


def _letters(n: int) -> str:
    out = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


_COLS = [_letters(i) for i in range(1, 401)]


def _sheet_xml(rows: Iterable[Row]) -> bytes:
    """Rows are (row number, values): text is an inline string, a str
    starting with '*' is bold text (style 1), anything else is a number."""
    parts = [f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<worksheet xmlns="{_MAIN}" xmlns:r="{_REL}"><sheetData>']
    for r, values in rows:
        cells = []
        for i, v in enumerate(values):
            if v is None:
                continue
            ref = f"{_COLS[i]}{r}"
            if isinstance(v, str) and v.startswith("*"):
                cells.append(f'<c r="{ref}" s="1" t="inlineStr"><is><t>{v[1:]}</t></is></c>')
            elif isinstance(v, str):
                cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{v}</t></is></c>')
            else:
                cells.append(f'<c r="{ref}"><v>{v}</v></c>')
        parts.append(f'<row r="{r}">{"".join(cells)}</row>')
    parts.append("</sheetData></worksheet>")
    return "".join(parts).encode("utf-8")


def write_workbook(path: Path, sheets: List[Tuple[str, Callable[[], Iterable[Row]], str]]) -> bytes:
    """``sheets`` is [(name, rows factory, state)]; returns the file's bytes."""
    overrides = "".join(
        f'<Override PartName="/xl/worksheets/sheet{i}.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        for i in range(1, len(sheets) + 1)
    )
    ctypes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        f"{overrides}</Types>"
    )
    sheet_els = "".join(
        f'<sheet name="{name}" sheetId="{i}"' + ("" if state == "visible" else f' state="{state}"') + f' r:id="rId{i}"/>'
        for i, (name, _rows, state) in enumerate(sheets, start=1)
    )
    wb_rels = "".join(
        f'<Relationship Id="rId{i}" Type="{_REL}/worksheet" Target="worksheets/sheet{i}.xml"/>'
        for i in range(1, len(sheets) + 1)
    ) + f'<Relationship Id="rId{len(sheets) + 1}" Type="{_REL}/styles" Target="styles.xml"/>'
    styles = (
        f'<styleSheet xmlns="{_MAIN}"><fonts count="2"><font><sz val="11"/><name val="Calibri"/></font>'
        '<font><b/><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        '<cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        '<xf numFmtId="0" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1"/></cellXfs></styleSheet>'
    )
    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ctypes_xml)
        z.writestr("_rels/.rels", f'<Relationships xmlns="{_PKG}"><Relationship Id="rId1" '
                   f'Type="{_REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr("xl/workbook.xml", f'<workbook xmlns="{_MAIN}" xmlns:r="{_REL}"><sheets>{sheet_els}</sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", f'<Relationships xmlns="{_PKG}">{wb_rels}</Relationships>')
        z.writestr("xl/styles.xml", styles)
        cache: dict = {}
        for i, (_name, rows, _state) in enumerate(sheets, start=1):
            if rows not in cache:
                cache[rows] = _sheet_xml(rows())
            z.writestr(f"xl/worksheets/sheet{i}.xml", cache[rows])
    return path.read_bytes()


def dense_rows(rows: int = 5000, cols: int = 200) -> Callable[[], Iterable[Row]]:
    def gen() -> Iterable[Row]:
        yield 1, [f"*Metric {c}" for c in range(cols)]
        for r in range(2, rows + 1):
            yield r, [(r + c) % 10 for c in range(cols)]
    return gen


def ledger_rows(rows: int) -> Callable[[], Iterable[Row]]:
    def gen() -> Iterable[Row]:
        yield 1, ["*Date", "*Account", "*Amount", "*Memo", "*Region", "*Owner", "*Status", "*Code"]
        for r in range(2, rows + 1):
            yield r, [45000 + r % 900, f"ACC{r % 97}", r * 1.25, f"memo text {r}", f"R{r % 5}", f"owner{r % 13}", "ok", r]
    return gen


# ---------------------------------------------------------------------------
# Child: one workbook through HTTP analyze + remediate, then its own peak
# ---------------------------------------------------------------------------


def _child(path: str) -> int:
    from fastapi.testclient import TestClient
    from sqlalchemy import select

    from app.db.models import UserRow
    from app.db.session_sqlalchemy import session_scope
    from app.main import app

    client = TestClient(app, raise_server_exceptions=False)
    email = "xlsx-mem@example.com"
    r = client.post("/auth/sign-in", json={"email": email, "displayName": "M", "password": "xlsxmempass12"})
    assert r.status_code == 200, r.text
    headers = {"Authorization": f"Bearer {r.json()['token']}"}
    with session_scope() as s:
        s.execute(select(UserRow).where(UserRow.email == email)).scalars().first().credits_balance = 50
    data = Path(path).read_bytes()
    name = Path(path).name
    base = _peak_mb()
    t0 = time.monotonic()
    a = client.post("/pipeline/analyze", files={"file": (name, data, XLSX)}, headers=headers)
    analyze_s = time.monotonic() - t0
    analyze_peak = _peak_mb()
    rules = [v["ruleId"] for v in a.json().get("violations", [])] if a.status_code == 200 else []
    ids = [v["id"] for v in a.json()["violations"] if v["ruleId"] == "TABLE_MISSING_HEADERS"] if a.status_code == 200 else []
    t0 = time.monotonic()
    rr = client.post(
        "/pipeline/remediate",
        files={"file": (name, data, XLSX)},
        data={"approved_violations": json.dumps(ids), "rejected_violations": "[]"},
        headers=headers,
    )
    remediate_s = time.monotonic() - t0
    body = rr.json() if rr.headers.get("content-type", "").startswith("application/json") else {}
    print("RESULT_JSON " + json.dumps({
        "analyze": a.status_code, "remediate": rr.status_code, "candidates": len(ids),
        "persisted": body.get("persistedFixes"), "truncated": "ANALYSIS_TRUNCATED" in rules,
        "base": round(base), "analyze_peak": round(analyze_peak), "peak": round(_peak_mb()),
        "analyze_s": round(analyze_s, 1), "remediate_s": round(remediate_s, 1),
    }))
    return 0


def _run_child(path: Path) -> dict:
    here = Path(__file__).resolve().parents[2]
    env = dict(os.environ)
    child_tmp = tempfile.mkdtemp(prefix="508_smoke_xlsx_mem_child_")
    env["DATABASE_URL"] = f"sqlite:///{child_tmp}/s.db"
    env["MATERIALIZED_ROOT"] = str(Path(child_tmp) / "materialized")
    proc = subprocess.run(
        [sys.executable, "-m", "app.devtools.smoke_xlsx_memory", "--child", str(path)],
        cwd=str(here), env=env, capture_output=True, text=True, timeout=540,
    )
    for line in (proc.stdout or "").splitlines():
        if line.startswith("RESULT_JSON "):
            return json.loads(line[len("RESULT_JSON "):])
    return {"error": (proc.stdout or "")[-500:] + (proc.stderr or "")[-1500:]}


# ---------------------------------------------------------------------------
# Parent
# ---------------------------------------------------------------------------


def main() -> int:  # noqa: PLR0915
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(_TMP)

    # ---- 1. peak memory, each in a fresh interpreter ------------------------
    multi = tmp / "export.xlsx"
    write_workbook(multi, [(f"Region {i}", dense_rows(), "visible") for i in range(1, 5)])
    ledger = tmp / "ledger.xlsx"
    write_workbook(ledger, [("Ledger", ledger_rows(100_000), "visible")])
    for label, path, want in (("4 dense sheets", multi, 4), ("100k-row ledger", ledger, 1)):
        res = _run_child(path)
        size = path.stat().st_size / 1e6
        print(f"  {label}: {size:.1f} MB upload -> {res}")
        check(f"{label}: analyze and remediate 200", res.get("analyze") == 200 and res.get("remediate") == 200, str(res))
        check(f"{label}: every clear header row found and fixed",
              res.get("candidates") == want and res.get("persisted") == want, str(res))
        check(f"{label}: nothing reported as unread", res.get("truncated") is False, str(res))
        peak = float(res.get("peak") or 1e9)
        check(f"{label}: peak memory under {PEAK_LIMIT_MB} MB (was 2.9 GB / 950 MB)", peak <= PEAK_LIMIT_MB, str(res))
        check(f"{label}: grows under {GROWTH_LIMIT_MB} MB over the idle app",
              peak - float(res.get("base") or 0) <= GROWTH_LIMIT_MB, str(res))

    from app.parsers import xlsx_parser
    from app.services.remediation_engine import RemediationEngine

    # ---- 2. the cell budget -------------------------------------------------
    def south() -> Iterable[Row]:
        yield 1, [f"*Field {c}" for c in range(60)]
        for r in range(2, 61):
            yield r, [r * c for c in range(60)]
        # rows 61-79 blank; a second block nothing continues into
        yield 80, [f"*Other {c}" for c in range(10)]
        for r in range(81, 101):
            yield r, [r + c for c in range(10)]

    budget_path = tmp / "budget.xlsx"
    write_workbook(budget_path, [
        ("North", dense_rows(120, 150), "visible"),
        ("South", south, "visible"),
        ("Sheet3", dense_rows(30, 20), "visible"),
    ])
    saved = xlsx_parser.CELL_BUDGET
    xlsx_parser.CELL_BUDGET = 10_000
    try:
        scan = xlsx_parser.scan_workbook(budget_path)
        result = xlsx_parser.build_tree(budget_path, scan)
    finally:
        xlsx_parser.CELL_BUDGET = saved
    north, south_s, third = scan.sheets
    n_blk = north.blocks[0] if north.blocks else None
    check("budget: the window ends early once a sheet's share is spent", 0 < north.window_rows < 120, str(north.window_rows))
    check("budget: ...but the block is followed to its end and still offered the fix",
          n_blk is not None and n_blk.ref == "A1:ET120" and n_blk.state == "candidate",
          str(n_blk and (n_blk.ref, n_blk.state, n_blk.reason)))
    check("budget: cells held across the workbook stay near the budget",
          sum(s.window_cells for s in scan.sheets) <= 10_000 + xlsx_parser.MIN_WINDOW_ROWS * 150,
          str([s.window_cells for s in scan.sheets]))
    check("budget: a block past the window that nothing continues is disclosed", south_s.truncated is True)
    check("budget: ...while the block that does continue keeps its fix",
          any(b.ref == "A1:BH60" and b.state == "candidate" for b in south_s.blocks), str([(b.ref, b.state) for b in south_s.blocks]))
    check("budget: a sheet with nothing left is not read, and says so", third.window_cells == 0 and third.truncated is True)
    viols = RemediationEngine().detect_violations(result.tree)
    found = {(v.rule_id, v.location.node_id) for v in viols}
    check("budget: the report says part of the workbook was not checked (ANALYSIS_TRUNCATED)",
          any(rule == "ANALYSIS_TRUNCATED" for rule, _ in found), str(sorted(found)))
    check("budget: the unread sheet is not mistaken for an empty one (its default tab name is still flagged)",
          ("SHEET_NAME_DEFAULT", "xlsx-s3") in found, str(sorted(found)))
    kept = sum(len(s.grid) for s in scan.sheets)
    check("budget: after judging, the scan keeps only the cells read later", kept <= 600, str(kept))

    # ---- 3. hidden sheets are not read cell by cell ------------------------
    hidden_path = tmp / "hidden.xlsx"
    write_workbook(hidden_path, [("Summary", dense_rows(10, 5), "visible"), ("Raw", dense_rows(400, 50), "hidden")])
    hscan = xlsx_parser.scan_workbook(hidden_path)
    raw = next(s for s in hscan.sheets if s.name == "Raw")
    check("hidden: a hidden sheet's cells are not read", raw.window_cells == 0 and not raw.blocks and not raw.grid)
    htree = xlsx_parser.build_tree(hidden_path, hscan).tree
    check("hidden: ...and it is still counted", htree.root.metadata.properties.get("hidden_sheet_count") == 1)

    # ---- 4. splicing into a sheet with a table, an extLst and a trap --------
    import openpyxl
    import xlsxwriter
    from fastapi.testclient import TestClient
    from sqlalchemy import select

    from app.db.models import UserRow
    from app.db.session_sqlalchemy import session_scope
    from app.main import app

    trap_path = tmp / "stock.xlsx"
    wb = xlsxwriter.Workbook(str(trap_path))
    bold = wb.add_format({"bold": True})
    ws = wb.add_worksheet("Stock")
    ws.add_table("A1:C4", {"columns": [{"header": "Item"}, {"header": "Qty"}, {"header": "Cost"}],
                           "data": [["a", 1, 2], ["b", 3, 4], ["c", 5, 6]]})
    ws.write_row("E1", ["Month", "Sales", "Returns"], bold)
    for i in range(4):
        ws.write_row(1 + i, 4, [f"M{i + 1}", 10 + i, 1 + i])
    ws.add_sparkline("H2", {"range": "Stock!F2:G2"})   # lives in the worksheet's <extLst>
    wb.close()
    with zipfile.ZipFile(str(trap_path)) as z:
        entries = [(i, z.read(i)) for i in z.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for info, body in entries:
            if info.filename == "xl/worksheets/sheet1.xml":
                assert b'<row r="2"' in body and b"<extLst>" in body and b"<tableParts" in body
                body = body.replace(b'<row r="2"', b'<!-- totals end at </sheetData> --><row r="2"', 1)
            z.writestr(info, body)
    trap = buf.getvalue()
    src_sheet = zipfile.ZipFile(io.BytesIO(trap)).read("xl/worksheets/sheet1.xml")

    client = TestClient(app, raise_server_exceptions=False)
    email = "xlsx-mem-parent@example.com"
    r = client.post("/auth/sign-in", json={"email": email, "displayName": "P", "password": "xlsxmempass12"})
    assert r.status_code == 200, r.text
    headers = {"Authorization": f"Bearer {r.json()['token']}"}

    def set_balance(n: int) -> None:
        with session_scope() as s:
            s.execute(select(UserRow).where(UserRow.email == email)).scalars().first().credits_balance = n

    def balance() -> int:
        return int(client.get("/credits/balance", headers=headers).json()["balance"])

    def analyze(name: str, data: bytes):
        return client.post("/pipeline/analyze", files={"file": (name, data, XLSX)}, headers=headers)

    def remediate(name: str, data: bytes, ids: list):
        return client.post("/pipeline/remediate", files={"file": (name, data, XLSX)},
                           data={"approved_violations": json.dumps(ids), "rejected_violations": "[]"}, headers=headers)

    set_balance(50)
    a = analyze("stock.xlsx", trap)
    ids = [v["id"] for v in a.json()["violations"] if v["ruleId"] == "TABLE_MISSING_HEADERS"] if a.status_code == 200 else []
    check("splice: the plain range beside the existing table is offered the fix", len(ids) == 1, a.text[:300])
    rr = remediate("stock.xlsx", trap, ids)
    body = rr.json() if rr.status_code == 200 else {}
    check("splice: remediate 200 and the fix is counted", rr.status_code == 200 and body.get("persistedFixes") == 1, rr.text[:300])
    out: Optional[bytes] = None
    if rr.status_code == 200:
        dl = client.get(body["downloadUrl"], headers=headers)
        out = dl.content if dl.status_code == 200 else None
    if out is not None:
        out_sheet = zipfile.ZipFile(io.BytesIO(out)).read("xl/worksheets/sheet1.xml")
        cut = src_sheet.rindex(b"</sheetData>") + len(b"</sheetData>")
        check("splice: every byte through the real </sheetData> is the customer's", out_sheet[:cut] == src_sheet[:cut])
        check("splice: the comment in the data is untouched", b"<!-- totals end at </sheetData> -->" in out_sheet)
        check("splice: the sparkline's <extLst> is still last", out_sheet.rstrip().endswith(b"</extLst></worksheet>"))
        owb = openpyxl.load_workbook(io.BytesIO(out))
        tables = dict(owb["Stock"].tables.items())
        check("splice: openpyxl sees the old table and the new one", tables.get("Table1") == "A1:C4"
              and sorted(tables.values()) == ["A1:C4", "E1:G5"], str(tables))
        check("splice: the cell values read back", [c.value for c in owb["Stock"]["E"][:3]] == ["Month", "M1", "M2"])

    # ---- 4b. a sheet the splicer cannot walk through --------------------------
    # A UTF-16 sheet part: our scanner (libxml2) reads it, the byte splicer
    # does not try to. The table is left as it was, with a reason, and the
    # workbook's other fixes still land and are the only thing charged.
    u16_path = tmp / "u16.xlsx"
    wb = xlsxwriter.Workbook(str(u16_path))
    ws = wb.add_worksheet("Sales")
    ws.write("A1", "Quarterly sales")
    ws.write_row("A3", ["Region", "Q1", "Q2"], wb.add_format({"bold": True}))
    for i in range(4):
        ws.write_row(3 + i, 0, [f"R{i}", 10 + i, 20 + i])
    wb.close()
    with zipfile.ZipFile(str(u16_path)) as z:
        entries = [(i, z.read(i)) for i in z.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for info, body in entries:
            if info.filename == "xl/worksheets/sheet1.xml":
                text = body.decode("utf-8").replace('encoding="UTF-8"', 'encoding="UTF-16"', 1)
                body = text.encode("utf-16")
            z.writestr(info, body)
    u16 = buf.getvalue()
    a = analyze("u16.xlsx", u16)
    rep = a.json() if a.status_code == 200 else {"violations": []}
    rules = {v["ruleId"]: v["id"] for v in rep["violations"]}
    check("utf-16 sheet: read, and its clear range is offered the fix", "TABLE_MISSING_HEADERS" in rules, a.text[:300])
    approve = [v["id"] for v in rep["violations"] if v["ruleId"] in ("TABLE_MISSING_HEADERS", "DOCUMENT_TITLE_MISSING")]
    set_balance(20)
    b0 = balance()
    rr = remediate("u16.xlsx", u16, approve)
    body = rr.json() if rr.status_code == 200 else {}
    writer = body.get("writer") or {}
    actions = [a_.get("action") for a_ in writer.get("applied") or [] if isinstance(a_, dict)]
    skipped = [s for s in writer.get("skipped") or [] if isinstance(s, dict) and s.get("action") == "ADD_TABLE_HEADERS"]
    check("utf-16 sheet: remediate still 200, the title lands", rr.status_code == 200 and "SET_DOCUMENT_TITLE" in actions,
          rr.text[:400])
    check("utf-16 sheet: the table is not claimed; it is skipped with a reason",
          "ADD_TABLE_HEADERS" not in actions and bool(skipped) and "left as it was" in str(skipped[0].get("reason")), str(body.get("skipped"))[:400])
    check("utf-16 sheet: charged for the fix that persisted, once", body.get("persistedFixes") == 1 and b0 - balance() == 3,
          f"{body.get('persistedFixes')} {b0}->{balance()}")

    # ---- 5. busy: a 503 that says so, and no charge -------------------------
    from app.writers import xlsx_writer

    saved_wait = xlsx_parser._XLSX_SLOT_WAIT_SECONDS
    xlsx_parser._XLSX_SLOT_WAIT_SECONDS = 0.2
    held = 0
    while xlsx_parser._XLSX_SLOTS.acquire(blocking=False):
        held += 1
    try:
        set_balance(20)
        b0 = balance()
        a = analyze("stock.xlsx", trap)
        check("busy: analyze is a 503 that says so", a.status_code == 503 and "not charged" in a.text, f"{a.status_code} {a.text[:200]}")
        rr = remediate("stock.xlsx", trap, ids)
        check("busy: remediate is a 503, not 'we couldn't read this file'",
              rr.status_code == 503 and "not charged" in rr.text, f"{rr.status_code} {rr.text[:200]}")
        check("busy: nothing charged", balance() == b0)
    finally:
        for _ in range(held):
            xlsx_parser._XLSX_SLOTS.release()
        xlsx_parser._XLSX_SLOT_WAIT_SECONDS = saved_wait

    real_slot = xlsx_writer.xlsx_work_slot

    def _busy_slot():
        from fastapi import HTTPException

        raise HTTPException(status_code=503, detail="busy (simulated). You were not charged.")

    xlsx_writer.xlsx_work_slot = _busy_slot
    try:
        b0 = balance()
        rr = remediate("stock.xlsx", trap, ids)
        check("busy writer: a 503 after a good parse, and nothing charged",
              rr.status_code == 503 and balance() == b0, f"{rr.status_code} {rr.text[:200]}")
    finally:
        xlsx_writer.xlsx_work_slot = real_slot
    a = analyze("stock.xlsx", trap)
    check("busy: once a slot frees up, the same upload works", a.status_code == 200, a.text[:200])

    print()
    print("RESULT:", "all passed" if failures == 0 else f"{failures} failed")
    return 0 if failures == 0 else 1


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "--child":
        sys.exit(_child(sys.argv[2]))
    sys.exit(main())
