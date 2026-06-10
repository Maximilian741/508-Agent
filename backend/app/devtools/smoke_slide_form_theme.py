"""Smoke: PPTX slide titles, DOCX content-control labels, theme-colour contrast.

- A PowerPoint slide with no title is flagged (SLIDE_TITLE_MISSING); a titled
  deck is not.
- A DOCX content control (w:sdt) with no alias is flagged (FORM_FIELD_UNLABELED);
  one with an alias is not.
- A DOCX run using a heavily-tinted (light) THEME colour is now caught by the
  contrast check (theme colours are resolved, not skipped).

Usage:
    python -m app.devtools.smoke_slide_form_theme
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_sft_')}/s.db"

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from pptx import Presentation  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.parsers.docx_parser import _apply_tint_shade  # noqa: E402
from app.services.remediation_engine import RemediationEngine  # noqa: E402


def _has(path: Path, rule: str) -> bool:
    res = parse_to_tree(str(path))
    run_analyzers(res.tree)
    return any(v.rule_id == rule for v in RemediationEngine().detect_violations(res.tree))


def main() -> int:
    failures = 0

    def check(name: str, cond: bool) -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name)
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp())

    # --- PPTX slide titles ---
    titled = Presentation()
    for _ in range(2):
        s = titled.slides.add_slide(titled.slide_layouts[5]); s.shapes.title.text = "A title"
    pt = tmp / "titled.pptx"; titled.save(str(pt))
    missing = Presentation()
    s1 = missing.slides.add_slide(missing.slide_layouts[5]); s1.shapes.title.text = "Has title"
    missing.slides.add_slide(missing.slide_layouts[6])  # blank, no title
    pm = tmp / "missing.pptx"; missing.save(str(pm))
    check("PPTX all-titled deck not flagged", not _has(pt, "SLIDE_TITLE_MISSING"))
    check("PPTX missing-title deck flagged", _has(pm, "SLIDE_TITLE_MISSING"))

    # --- Magnitude: one issue PER untitled slide (not one per deck) ---
    def _count(path: Path, rule: str) -> int:
        res = parse_to_tree(str(path))
        run_analyzers(res.tree)
        return sum(1 for v in RemediationEngine().detect_violations(res.tree) if v.rule_id == rule)

    multi = Presentation()
    sm = multi.slides.add_slide(multi.slide_layouts[5]); sm.shapes.title.text = "Only titled slide"
    multi.slides.add_slide(multi.slide_layouts[6])  # blank
    multi.slides.add_slide(multi.slide_layouts[6])  # blank
    multi.slides.add_slide(multi.slide_layouts[6])  # blank
    pmu = tmp / "multi_missing.pptx"; multi.save(str(pmu))
    check("3 untitled slides -> exactly 3 SLIDE_TITLE_MISSING issues", _count(pmu, "SLIDE_TITLE_MISSING") == 3)
    check("1 untitled slide -> exactly 1 issue", _count(pm, "SLIDE_TITLE_MISSING") == 1)

    # --- Title double-emit: title text appears in exactly ONE node ---
    res_t = parse_to_tree(str(pt))

    def _texts(node, acc):
        if node.content and node.content.text:
            acc.append(node.content.text.strip())
        for ch in node.children:
            _texts(ch, acc)
        return acc

    all_texts = _texts(res_t.tree.root, [])
    # Each titled slide contributes its SectionNode label + ONE HeadingNode —
    # the title placeholder must not ALSO surface as a ParagraphNode.
    check(
        "slide title not double-emitted as paragraph",
        all_texts.count("A title") == 2 * 2,  # 2 slides x (section label + heading)
    )
    from app.models.accessibility import ParagraphNode as _PN

    def _paras(node, acc):
        if isinstance(node, _PN):
            acc.append(node.content.text if node.content else "")
        for ch in node.children:
            _paras(ch, acc)
        return acc

    check(
        "no ParagraphNode carries the title text",
        all("A title" not in (t or "") for t in _paras(res_t.tree.root, [])),
    )

    # --- DOCX content controls ---
    def build_cc(path: Path, labeled: bool) -> None:
        d = Document(); d.add_paragraph("Body.")
        sdt = OxmlElement("w:sdt"); pr = OxmlElement("w:sdtPr")
        if labeled:
            alias = OxmlElement("w:alias"); alias.set(qn("w:val"), "Full name"); pr.append(alias)
        sdt.append(pr); sdt.append(OxmlElement("w:sdtContent"))
        d.element.body.append(sdt); d.save(str(path))

    build_cc(tmp / "cc_unlabeled.docx", labeled=False)
    build_cc(tmp / "cc_labeled.docx", labeled=True)
    check("DOCX unlabeled content control flagged", _has(tmp / "cc_unlabeled.docx", "FORM_FIELD_UNLABELED"))
    check("DOCX labeled content control not flagged", not _has(tmp / "cc_labeled.docx", "FORM_FIELD_UNLABELED"))

    # --- Theme-colour contrast ---
    check("tint=FF keeps colour", _apply_tint_shade("4472C4", "FF", None) == "4472C4")
    check("tint=33 lightens past #C0", int(_apply_tint_shade("4472C4", "33", None)[0:2], 16) > 0xC0)
    th = tmp / "theme.docx"
    d = Document(); para = d.add_paragraph(); run = para.add_run("Light themed text")
    rpr = run._element.get_or_add_rPr()
    col = OxmlElement("w:color"); col.set(qn("w:themeColor"), "accent1"); col.set(qn("w:themeTint"), "33")
    rpr.append(col)
    d.save(str(th))
    check("light themed text caught by contrast", _has(th, "LOW_CONTRAST_TEXT"))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
