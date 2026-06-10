"""Smoke: DOCX footnote/endnote content is analyzed AND remediable.

Notes live in separate package parts (footnotes.xml / endnotes.xml) that
``doc.paragraphs`` never opens — before this, footnote text and the
"click here" links that government documents love to bury in footnotes
were never analyzed. Pins:

  - footnote paragraph text becomes a ParagraphNode (marked in_footnote)
  - footnote + endnote links become LinkNodes (docx-fnlink space), get
    flagged, and IMPROVE_LINK_TEXT genuinely rewrites them in the OUTPUT
    BYTES (note parts re-serialized; targets resolve via the NOTE part's
    own rels)
  - separator/continuation stubs never emit nodes
  - body link ids are untouched; plain docs emit nothing

Usage:
    python -m app.devtools.smoke_footnotes
"""

from __future__ import annotations

import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_fn_')}/s.db")

from docx import Document  # noqa: E402
from docx.opc.constants import CONTENT_TYPE as CT  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.opc.packuri import PackURI  # noqa: E402
from docx.opc.part import Part  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.models.accessibility import ContentKind, LinkNode, NodeContent, ParagraphNode, iter_reading_order  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.writers import write_remediated  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="fn_"))

    # --- fixture: body link + footnotes part + endnotes part -----------------
    d = Document()
    d.core_properties.title = "FN Test"
    body_para = d.add_paragraph("Body link: ")
    rid_body = d.part.relate_to(
        "https://example.com/body", f"{R}/hyperlink", is_external=True
    )
    hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), rid_body)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "annual data tables"
    r.append(t); hl.append(r); body_para._p.append(hl)

    fn_xml = (
        f'<w:footnotes xmlns:w="{W}" xmlns:r="{R}">'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:t>Source data published at </w:t></w:r>'
        '<w:hyperlink r:id="rIdF1"><w:r><w:t>click here</w:t></w:r></w:hyperlink></w:p></w:footnote>'
        '<w:footnote w:id="2"><w:p><w:r><w:t>Methodology detailed in the technical annex.</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    ).encode()
    fn_part = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, fn_xml, d.part.package)
    d.part.relate_to(fn_part, RT.FOOTNOTES)
    fn_part.rels.add_relationship(RT.HYPERLINK, "https://example.com/fn-source", "rIdF1", is_external=True)

    en_xml = (
        f'<w:endnotes xmlns:w="{W}" xmlns:r="{R}">'
        '<w:endnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:endnote>'
        '<w:endnote w:id="1"><w:p><w:hyperlink r:id="rIdE1"><w:r><w:t>https://example.gov/endnote</w:t></w:r></w:hyperlink></w:p></w:endnote>'
        "</w:endnotes>"
    ).encode()
    en_part = Part(PackURI("/word/endnotes.xml"), CT.WML_ENDNOTES, en_xml, d.part.package)
    d.part.relate_to(en_part, RT.ENDNOTES)
    en_part.rels.add_relationship(RT.HYPERLINK, "https://example.gov/endnote", "rIdE1", is_external=True)

    src = tmp / "fn.docx"
    d.save(str(src))

    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    nodes = list(iter_reading_order(res.tree.root))
    fn_paras = [n for n in nodes if isinstance(n, ParagraphNode) and (n.metadata.properties or {}).get("in_footnote")]
    check("footnote text visible", any("technical annex" in (n.content.text or "") for n in fn_paras), str([n.content.text for n in fn_paras]))
    links = [n for n in nodes if isinstance(n, LinkNode)]
    fn_links = [n for n in links if (n.metadata.properties or {}).get("in_footnote")]
    check("2 note links emitted (footnote + endnote)", len(fn_links) == 2, str([(n.id, n.content.text) for n in links]))
    check("note link ids use the docx-fnlink space", all(n.id.startswith("docx-fnlink") for n in fn_links))
    check("footnote link target from NOTE part rels", any(n.target == "https://example.com/fn-source" for n in fn_links), str([n.target for n in fn_links]))
    check("body link unaffected (docx-link-1)", any(n.id == "docx-link-1" and n.content.text == "annual data tables" for n in links))
    flagged = [n for n in fn_links if any(f.code.value == "LINK_TEXT_NON_DESCRIPTIVE" for f in n.accessibility_flags)]
    check("both note links flagged non-descriptive", len(flagged) == 2, str([(n.content.text, [f.code.value for f in n.accessibility_flags]) for n in fn_links]))

    # --- remediation round-trip ------------------------------------------------
    for i, ln in enumerate(links):
        ln.content = NodeContent(kind=ContentKind.TEXT, text=f"FIXED-{i}: {ln.id}")
    out = tmp / "fn_out.docx"
    write_remediated(src, res.tree, out, source_format=res.format)
    res2 = parse_to_tree(str(out))
    links2 = [n for n in iter_reading_order(res2.tree.root) if isinstance(n, LinkNode)]
    check(
        "every link rewritten in OUTPUT with no cross-assignment",
        all((n.content.text or "").startswith("FIXED-") and n.id in n.content.text for n in links2),
        str([(n.id, n.content.text) for n in links2]),
    )
    with zipfile.ZipFile(out) as z:
        fn_out_xml = z.read("word/footnotes.xml").decode("utf-8", "replace")
        en_out_xml = z.read("word/endnotes.xml").decode("utf-8", "replace")
    check("footnotes.xml bytes carry the rewrite", "FIXED-" in fn_out_xml and "click here" not in fn_out_xml)
    check("endnotes.xml bytes carry the rewrite", "FIXED-" in en_out_xml)
    check("separator stubs emitted no nodes", not any("separator" in (n.content.text or "").lower() for n in nodes))

    # --- plain doc -------------------------------------------------------------
    c = Document()
    c.core_properties.title = "Plain"
    c.add_paragraph("No notes here.")
    plain = tmp / "plain.docx"
    c.save(str(plain))
    res3 = parse_to_tree(str(plain))
    n_fn = sum(1 for n in iter_reading_order(res3.tree.root) if (n.metadata.properties or {}).get("in_footnote"))
    check("plain doc emits no note nodes", n_fn == 0)

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
