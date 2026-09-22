"""Smoke: /pipeline/remediate never charges for a file it didn't change.

Billing-honesty guard. The remediate endpoint charges credits only AFTER the
writer persists at least one fix. If every approved item is manual-only (or the
caller approves nothing), the downloaded file is byte-for-byte the source — so
the user must NOT be billed. This drives the in-process app through TestClient
and asserts:

  * Remediate approving a real fix (alt text) -> charged; balance drops by the
    docx cost; response ``charged`` is true.
  * Remediate approving NOTHING -> not charged; balance unchanged; response
    ``charged`` is false; a file is still returned.

Run: python -m app.devtools.smoke_no_charge_no_fix
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
from pathlib import Path

_SMOKE_DB_DIR = tempfile.mkdtemp(prefix="508_smoke_nocharge_")
os.environ["DATABASE_URL"] = f"sqlite:///{_SMOKE_DB_DIR}/smoke.db"

from fastapi.testclient import TestClient  # noqa: E402

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _valid_png() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (24, 24), (10, 120, 200)).save(buf, format="PNG")
    return buf.getvalue()


def _build_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Section", level=1)
    # A real sentence before the picture: the alt executor now derives its
    # description from nearby text and REFUSES to write a placeholder when
    # there is none. "Filler paragraph." (two words) was below the bar, so the
    # approved alt fix honestly skipped and — correctly — nothing was charged,
    # which broke this smoke's "a real fix charges" assertion.
    doc.add_paragraph("The chart below shows quarterly revenue by region for the year.")
    doc.add_paragraph().add_run().add_picture(io.BytesIO(_valid_png()))
    doc.save(str(path))


def _balance(client, headers) -> int:
    r = client.get("/credits/balance", headers=headers)
    assert r.status_code == 200, r.text
    return int(r.json()["balance"])


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.main import app

    client = TestClient(app)

    signin = client.post(
        "/auth/sign-in",
        json={"email": "nocharge-smoke@example.com", "displayName": "NC", "password": "nochargepass1"},
    )
    assert signin.status_code == 200, signin.text
    headers = {"Authorization": f"Bearer {signin.json()['token']}"}
    assert client.post("/auth/grant-starter", headers=headers).status_code == 200

    with tempfile.TemporaryDirectory() as tmp_dir:
        src = Path(tmp_dir) / "book.docx"
        _build_docx(src)

        # Analyze to get violation ids.
        with src.open("rb") as fh:
            an = client.post("/pipeline/analyze", files={"file": ("book.docx", fh, _DOCX_MIME)}, headers=headers)
        assert an.status_code == 200, an.text
        violations = an.json()["violations"]
        alt_ids = [v["id"] for v in violations if v["ruleId"] == "MISSING_ALT_TEXT"]
        all_ids = [v["id"] for v in violations]
        check("doc has a fixable MISSING_ALT_TEXT finding", len(alt_ids) >= 1, str([v["ruleId"] for v in violations]))

        # ---- Case 1: approve the alt fix -> CHARGED ----
        bal0 = _balance(client, headers)
        with src.open("rb") as fh:
            rem = client.post(
                "/pipeline/remediate",
                files={"file": ("book.docx", fh, _DOCX_MIME)},
                data={"approved_violations": json.dumps(alt_ids), "rejected_violations": json.dumps([])},
                headers=headers,
            )
        assert rem.status_code == 200, rem.text
        body = rem.json()
        bal1 = _balance(client, headers)
        check("applying a real fix reports charged=true", body.get("charged") is True, str(body.get("charged")))
        check("applying a real fix decrements the balance by the docx cost (3)",
              bal0 - bal1 == 3, f"bal0={bal0} bal1={bal1}")

        # ---- Case 2: approve NOTHING -> NOT charged ----
        bal_before = _balance(client, headers)
        with src.open("rb") as fh:
            rem2 = client.post(
                "/pipeline/remediate",
                files={"file": ("book.docx", fh, _DOCX_MIME)},
                data={"approved_violations": json.dumps([]), "rejected_violations": json.dumps(all_ids)},
                headers=headers,
            )
        assert rem2.status_code == 200, rem2.text
        body2 = rem2.json()
        bal_after = _balance(client, headers)
        check("approving nothing reports charged=false", body2.get("charged") is False, str(body2.get("charged")))
        check("approving nothing does NOT change the balance",
              bal_after == bal_before, f"before={bal_before} after={bal_after}")
        check("approving nothing still returns a downloadable file", bool(body2.get("downloadUrl")))
        # The honest signal is that no APPROVED fix ran (executions empty) — the
        # writer may still re-assert existing structure into ``applied`` (e.g. an
        # already-correct heading style), which is why the charge keys off
        # successful persisted executions, not the writer's applied list.
        check("approving nothing ran zero fix executions",
              isinstance(body2.get("executions"), list) and len(body2["executions"]) == 0,
              str(body2.get("executions")))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except AssertionError as exc:
        print(f"FAIL no-charge smoke: {exc}")
        sys.exit(1)
    except Exception as exc:  # pragma: no cover
        print(f"FAIL no-charge smoke: {exc.__class__.__name__}: {exc}")
        sys.exit(1)
