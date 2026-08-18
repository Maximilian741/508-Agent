"""Smoke: DOCX language is read from where screen readers read it, and never downgraded.

The parser read the language only from the Dublin Core dc:language core
property, which is rarely set. Word writes the document language into
styles.xml docDefaults <w:lang w:val="en-US"/> on save — and, per the
writer's own docstring, THAT is what screen readers and Word's Accessibility
Checker read. So essentially every Word document was flagged
DOCUMENT_LANGUAGE_MISSING, and the "fix" then overwrote en-US with a less
specific "en" — a downgrade credited as a fix.

Pinned here through the real engine + writer:
  * a default Word document (docDefaults en-US) reads en-US, is NOT flagged,
    and its docDefaults still say en-US after remediation
  * a document with NO w:lang anywhere IS flagged, gets a real detected
    language, and the writer sets it

Usage:
    python -m app.devtools.smoke_docx_language_source
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import zipfile

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_dls_')}/s.db")
os.environ.pop("OPENAI_API_KEY", None)
os.environ.pop("ANTHROPIC_API_KEY", None)
os.environ["SEMANTIC_PROVIDER"] = "heuristic"

from pathlib import Path  # noqa: E402

import docx  # noqa: E402
from lxml import etree  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_engine import RemediationEngine  # noqa: E402
from app.services.remediation_planner import RemediationPolicy  # noqa: E402
from app.writers.docx_writer import write_remediated_docx  # noqa: E402

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_POL = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)


def _default_lang(path: Path):
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/styles.xml"))
    el = root.find(f".//{_W}docDefaults/{_W}rPrDefault/{_W}rPr/{_W}lang")
    return el.get(f"{_W}val") if el is not None else None


def _strip_lang(path: Path) -> None:
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    root = etree.fromstring(items["word/styles.xml"])
    for el in list(root.iter(f"{_W}lang")):
        el.getparent().remove(el)
    items["word/styles.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in items.items():
            z.writestr(n, b)
    path.write_bytes(buf.getvalue())


def _run(src: Path, out: Path):
    tree = parse_to_tree(str(src)).tree
    run_analyzers(tree)
    eng = RemediationEngine(policy=_POL)
    eng.detect_violations(tree)
    eng.execute(tree)
    write_remediated_docx(src, tree, out)
    flagged = any(f.code.value == "DOCUMENT_LANGUAGE_MISSING" for f in tree.root.accessibility_flags)
    return tree, flagged


def main() -> int:
    failures = 0

    def check(name, cond, extra=""):
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    d = Path(tempfile.mkdtemp(prefix="508_dls_"))
    text = "The annual report presents the results of the year for the board and the committee."

    # A: default Word template carries docDefaults en-US
    a = d / "a.docx"
    doc = docx.Document(); doc.add_paragraph(text); doc.save(str(a))
    check("fixture A: source docDefaults say en-US", _default_lang(a) == "en-US", str(_default_lang(a)))
    tree, flagged = _run(a, d / "a-out.docx")
    check("A: parser reads the language from w:lang (en-US)", tree.root.metadata.language == "en-US", repr(tree.root.metadata.language))
    check("A: NOT flagged DOCUMENT_LANGUAGE_MISSING (was a false positive on every Word doc)", not flagged)
    check("A: output docDefaults STILL say en-US (never downgraded to 'en')", _default_lang(d / "a-out.docx") == "en-US", str(_default_lang(d / "a-out.docx")))

    # B: no w:lang anywhere -> a real finding, a real fix
    b = d / "b.docx"
    doc = docx.Document(); doc.add_paragraph(text); doc.save(str(b)); _strip_lang(b)
    check("fixture B: source has NO w:lang", _default_lang(b) is None, str(_default_lang(b)))
    tree, flagged = _run(b, d / "b-out.docx")
    check("B: IS flagged DOCUMENT_LANGUAGE_MISSING", flagged)
    check("B: a real language was detected (en)", tree.root.metadata.language == "en", repr(tree.root.metadata.language))
    check("B: writer set docDefaults w:lang", _default_lang(d / "b-out.docx") == "en", str(_default_lang(d / "b-out.docx")))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
