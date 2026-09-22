"""Smoke: DOCX contrast is measured against the background the text is REALLY on.

A verifier found white letterhead text on a navy header-table cell reported as
1.0:1 "white on white", recoloured to grey (2.49:1 against the navy it really
sits on), charged for, and the re-scan called the file clean. The parser read
only paragraph shading and assumed white for everything else.

What this pins (real python-docx files, real parse -> analyze -> execute ->
write -> re-parse):

  header / footer
   1. white text on a navy header CELL: no finding, nothing recoloured
   2. grey text on that navy cell: flagged against the navy, recoloured to a
      shade that passes against the navy (not against white), re-scan clean
   3. white text on a THEME-coloured cell fill: resolved from the theme
   4. white text on TABLE-level shading: no finding
   5. grey text in a table whose STYLE shades cells: not measured (the
      style's banding/header shading depends on position — unknown)
   6. white text on PATTERN shading (pct50): not measured
   7. white text in a header TEXT BOX with a navy fill: no finding; grey text
      in a see-through text box: not measured
   8. a shape BEHIND the header text: header text not measured
   9. control: grey text in a plain header paragraph is still flagged + fixed
  body
  10. white text on a dark HIGHLIGHT: not measured; in the same paragraph,
      grey unhighlighted text is still flagged, and only it is recoloured
  11. a highlighted run sharing its colour with a measured run: the paragraph
      is not measured (the writer recolours by colour value)
  12. white text in a paragraph STYLE with navy shading: no finding
  13. a painted PAGE colour: body text not measured
  14. a shape behind the body text on a page the file does not name: body
      text not measured
  15. Word's page WATERMARK does not switch measuring off
  16. ... but a dark shape that only borrows the watermark's name does

Usage:
    python -m app.devtools.smoke_docx_contrast_background
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_dcbg_')}/s.db")

from docx import Document  # noqa: E402
from docx.enum.text import WD_COLOR_INDEX  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.oxml import OxmlElement, parse_xml  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402
from lxml import etree  # noqa: E402

from app.analyzers.contrast import contrast_ratio  # noqa: E402
from app.analyzers.registry import run_analyzers  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.docx_writer import write_remediated_docx  # noqa: E402

CONTRAST = "LOW_CONTRAST_TEXT"
NAVY = "1F3864"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
POLICY = RemediationPolicy(allow_ai_actions=False, require_human_review_for_all=False)

_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office"'
)


# ---------------------------------------------------------------- builders --

def _run(par, text: str, hex6: str, size: float = 12.0):
    r = par.add_run(text)
    r.font.color.rgb = RGBColor.from_string(hex6)
    r.font.size = Pt(size)
    return r


def _shd(parent_pr, fill: str, val: str = "clear", **extra):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), val)
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    for k, v in extra.items():
        shd.set(qn(f"w:{k}"), v)
    parent_pr.append(shd)
    return shd


def _cell_para(container, fill=None, style=None, **shd_extra):
    tbl = container.add_table(rows=1, cols=1, width=Inches(6))
    if style:
        tbl.style = style
    cell = tbl.cell(0, 0)
    if fill:
        _shd(cell._tc.get_or_add_tcPr(), fill, **shd_extra)
    return tbl, cell.paragraphs[0]


def _anchor(inner: str, behind: bool, shape_id: int) -> str:
    return (
        f'<w:r {_NS}><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" '
        f'simplePos="0" relativeHeight="{shape_id}" behindDoc="{1 if behind else 0}" locked="0" '
        'layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:extent cx="7772400" cy="1371600"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/><wp:docPr id="{shape_id}" name="Shape {shape_id}"/><wp:cNvGraphicFramePr/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'{inner}</a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'
    )


def _panel(fill_xml: str) -> str:
    """A plain rectangle (no text)."""
    return (
        '<wps:wsp><wps:cNvSpPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="7772400" cy="1371600"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>{fill_xml}</wps:spPr><wps:bodyPr/></wps:wsp>'
    )


def _text_box(fill_xml: str, text: str, hex6: str) -> str:
    return (
        '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2743200" cy="457200"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>{fill_xml}</wps:spPr>'
        '<wps:txbx><w:txbxContent><w:p><w:r><w:rPr>'
        f'<w:color w:val="{hex6}"/><w:sz w:val="24"/></w:rPr><w:t>{text}</w:t></w:r></w:p></w:txbxContent></wps:txbx>'
        '<wps:bodyPr/></wps:wsp>'
    )


def _append_run_xml(par, xml: str) -> None:
    par._p.append(parse_xml(xml))


def _save(doc, path: Path) -> Path:
    doc.core_properties.title = "Contrast background smoke"
    doc.save(str(path))
    return path


def build_letterhead(path: Path) -> Path:
    """Header cases 1-7 and 9 in one header (no shape behind any text)."""
    doc = Document()
    hdr = doc.sections[0].header
    _run(hdr.paragraphs[0], "Plain grey header line", "999999")                       # 9
    _, p = _cell_para(hdr, NAVY)
    _run(p, "White on navy cell", "FFFFFF")                                            # 1
    _, p = _cell_para(hdr, NAVY)
    _run(p, "Grey on navy cell", "747474")                                             # 2
    _, p = _cell_para(hdr, NAVY, themeFill="accent1", themeFillShade="80")
    _run(p, "White on themed cell", "FFFFFF")                                          # 3
    tbl, p = _cell_para(hdr)
    _shd(tbl._tbl.tblPr, NAVY)
    _run(p, "White on table shading", "FFFFFF")                                        # 4
    _, p = _cell_para(hdr, style="Light Shading Accent 1")
    _run(p, "Grey in a style-shaded table", "999999")                                  # 5
    _, p = _cell_para(hdr, NAVY, val="pct50")
    _run(p, "White on pattern shading", "FFFFFF")                                      # 6
    anchor_p = hdr.add_paragraph()
    _append_run_xml(anchor_p, _anchor(_text_box(
        f'<a:solidFill><a:srgbClr val="{NAVY}"/></a:solidFill>', "White in a navy text box", "FFFFFF"),
        behind=False, shape_id=101))                                                   # 7a
    anchor_p2 = hdr.add_paragraph()
    _append_run_xml(anchor_p2, _anchor(_text_box(
        "<a:noFill/>", "Grey in a see-through text box", "999999"), behind=False, shape_id=102))  # 7b
    doc.add_paragraph("Body text of the memo, which is long enough to be ordinary prose.")
    return _save(doc, path)


def build_banner(path: Path) -> Path:
    """Case 8: a navy panel behind the header text."""
    doc = Document()
    hdr = doc.sections[0].header
    p = hdr.paragraphs[0]
    _append_run_xml(p, _anchor(_panel(f'<a:solidFill><a:srgbClr val="{NAVY}"/></a:solidFill>'), True, 201))
    _run(p, "White on a banner behind the text", "FFFFFF")
    doc.add_paragraph("Body text.")
    return _save(doc, path)


def build_body(path: Path) -> Path:
    """Cases 10-12 (and the control for 15)."""
    doc = Document()
    p = doc.add_paragraph()
    hl = _run(p, "White on a dark-blue highlight. ", "FFFFFF")
    hl.font.highlight_color = WD_COLOR_INDEX.DARK_BLUE
    _run(p, "Grey without a highlight.", "999999")                                     # 10
    p = doc.add_paragraph()
    hl = _run(p, "Grey on a black highlight. ", "999999")
    hl.font.highlight_color = WD_COLOR_INDEX.BLACK
    _run(p, "Grey on the page.", "999999")                                             # 11
    st = doc.styles.add_style("Navy Band", WD_STYLE_TYPE.PARAGRAPH)
    _shd(st.element.get_or_add_pPr(), NAVY)
    p = doc.add_paragraph(style="Navy Band")
    _run(p, "White in a navy-shaded style", "FFFFFF")                                  # 12
    doc.add_paragraph().add_run("Plain grey body control").font.color.rgb = RGBColor.from_string("999999")
    return _save(doc, path)


def build_page_colour(path: Path) -> Path:
    """Case 13."""
    doc = Document()
    bg = parse_xml(f'<w:background {_NS} w:color="{NAVY}"/>')
    doc.element.insert(0, bg)
    _run(doc.add_paragraph(), "White on a navy page", "FFFFFF")
    return _save(doc, path)


def build_body_panel(path: Path) -> Path:
    """Case 14: a cover-page panel behind body text (no page information)."""
    doc = Document()
    p = doc.add_paragraph()
    _append_run_xml(p, _anchor(_panel(f'<a:solidFill><a:srgbClr val="{NAVY}"/></a:solidFill>'), True, 301))
    _run(p, "Annual Report cover title", "FFFFFF", 28)
    _run(doc.add_paragraph(), "Grey text later in the report", "999999")
    return _save(doc, path)


def build_watermark(path: Path, real: bool = True) -> Path:
    """Case 15: Word's VML DRAFT watermark in the header. ``real=False``: a
    dark full-page rectangle that merely carries the watermark's name."""
    doc = Document()
    hdr = doc.sections[0].header
    body = ('<v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="DRAFT"/>'
            if real else "")
    fill = "silver" if real else f"#{NAVY}"
    _append_run_xml(hdr.paragraphs[0], (
        f'<w:r {_NS}><w:pict><v:shape id="PowerPlusWaterMarkObject357831064" o:spid="_x0000_s2049" '
        'type="#_x0000_t136" style="position:absolute;margin-left:0;margin-top:0;width:412pt;height:164pt;'
        'z-index:-251657216;mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
        f'mso-position-vertical:center;mso-position-vertical-relative:margin" fillcolor="{fill}" stroked="f">'
        f'{body}</v:shape></w:pict></w:r>'
    ))
    _run(doc.add_paragraph(), "Grey text under a watermark", "999999")
    return _save(doc, path)


# ---------------------------------------------------------------- helpers --

def _nodes(tree):
    out = []

    def walk(n):
        out.append(n)
        for c in n.children:
            walk(c)

    walk(tree.root)
    return out


def _analyze(path: Path):
    res = parse_to_tree(str(path))
    run_analyzers(res.tree)
    return res


def _by_text(res):
    out = {}
    for n in _nodes(res.tree):
        t = getattr(n.content, "text", None)
        if t:
            out[" ".join(t.split())] = n
    return out


def _flagged(node) -> bool:
    return node is not None and any(f.code.value == CONTRAST for f in node.accessibility_flags)


def _remediate(res, src: Path, out: Path):
    plans = plan_remediations(res.tree, POLICY)
    execute_plans(res.tree, plans)
    return write_remediated_docx(src, res.tree, out)


def _run_colors(path: Path, part: str):
    """{text: color} for every visible run in ``part``."""
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(path.read_bytes())).read(part))
    out = {}
    for r in root.iter(f"{W}r"):
        t = "".join(x.text or "" for x in r.iter(f"{W}t"))
        c = r.find(f"{W}rPr/{W}color")
        if t.strip():
            out[t] = c.get(f"{W}val") if c is not None else None
    return out


def main() -> int:
    failures = 0

    def check(name, cond, extra=""):
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, "" if cond else extra)
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="dcbg_smoke_"))

    # ---- header / footer -------------------------------------------------
    src = build_letterhead(tmp / "letterhead.docx")
    res = _analyze(src)
    t = _by_text(res)
    white = t.get("White on navy cell")
    check("1. white on a navy header cell: found, not flagged",
          white is not None and not _flagged(white), str(white and white.metadata.properties))
    grey = t.get("Grey on navy cell")
    finding = (grey.metadata.properties or {}).get("contrast_finding") if grey else None
    check("2. grey on the navy cell: flagged against the NAVY, not white",
          _flagged(grey) and finding and finding.get("bg") == NAVY, str(finding))
    check("3. white on a theme-coloured cell: not flagged", not _flagged(t.get("White on themed cell")),
          str(t.get("White on themed cell") and t["White on themed cell"].metadata.properties))
    themed = (t.get("White on themed cell").metadata.properties or {}) if t.get("White on themed cell") else {}
    check("3. the themed fill was resolved from the theme (not assumed white)",
          themed.get("bg_color") not in (None, "FFFFFF"), str(themed))
    check("4. white on table-level shading: not flagged", not _flagged(t.get("White on table shading")))
    styled = t.get("Grey in a style-shaded table")
    check("5. grey in a table whose style shades cells: not measured",
          styled is not None and "explicit_text_colors" not in (styled.metadata.properties or {}),
          str(styled and styled.metadata.properties))
    patt = t.get("White on pattern shading")
    check("6. white on pattern shading: not measured",
          patt is not None and "explicit_text_colors" not in (patt.metadata.properties or {}))
    tb = t.get("White in a navy text box")
    check("7. white in a navy-filled header text box: found, not flagged",
          tb is not None and not _flagged(tb), str(tb and tb.metadata.properties))
    clear_tb = t.get("Grey in a see-through text box")
    check("7. grey in a see-through text box: not measured",
          clear_tb is not None and "explicit_text_colors" not in (clear_tb.metadata.properties or {}))
    check("9. control: grey plain header line still flagged", _flagged(t.get("Plain grey header line")))

    out = tmp / "letterhead.fixed.docx"
    result = _remediate(res, src, out)
    fixed_ids = {a.get("target_id") for a in result["applied"] if a.get("action") == "FIX_CONTRAST"}
    check("writer recoloured exactly the two real failures",
          fixed_ids == {grey.id, t["Plain grey header line"].id}, str(fixed_ids))
    colours = _run_colors(out, "word/header1.xml")
    for text in ("White on navy cell", "White on themed cell", "White on table shading",
                 "White on pattern shading", "White in a navy text box"):
        check(f"   '{text}' still white in the output", colours.get(text) == "FFFFFF", str(colours.get(text)))
    for text in ("Grey in a style-shaded table", "Grey in a see-through text box"):
        check(f"   '{text}' not recoloured", colours.get(text) == "999999", str(colours.get(text)))
    new = colours.get("Grey on navy cell")
    ratio = contrast_ratio(new, NAVY) if new else None
    check("2. the recolour passes AA against the navy it sits on",
          ratio is not None and ratio >= 4.5, f"{new} -> {ratio}")
    res2 = _analyze(out)
    left = [n.content.text for n in _nodes(res2.tree) if _flagged(n)]
    check("re-scan of the output: no contrast finding left", not left, str(left))

    src = build_banner(tmp / "banner.docx")
    res = _analyze(src)
    banner = _by_text(res).get("White on a banner behind the text")
    check("8. white on a shape behind the header text: not measured",
          banner is not None and "explicit_text_colors" not in (banner.metadata.properties or {}),
          str(banner and banner.metadata.properties))
    out = tmp / "banner.fixed.docx"
    result = _remediate(res, src, out)
    check("8. nothing recoloured",
          not [a for a in result["applied"] if a.get("action") == "FIX_CONTRAST"]
          and _run_colors(out, "word/header1.xml").get("White on a banner behind the text") == "FFFFFF")

    # ---- body --------------------------------------------------------------
    src = build_body(tmp / "body.docx")
    res = _analyze(src)
    nodes = [n for n in _nodes(res.tree) if (n.metadata.properties or {}).get("explicit_text_colors")]
    by_start = {" ".join((n.content.text or "").split())[:20]: n for n in nodes}
    mixed = next((n for n in _nodes(res.tree) if (n.content.text or "").startswith("White on a dark-blue")), None)
    mprops = (mixed.metadata.properties or {}) if mixed else {}
    check("10. highlighted white run left out, grey run measured",
          [c["c"] for c in mprops.get("explicit_text_colors", [])] == ["999999"] and _flagged(mixed), str(mprops))
    shared = next((n for n in _nodes(res.tree) if (n.content.text or "").startswith("Grey on a black")), None)
    check("11. a highlighted run sharing a measured run's colour: paragraph not measured",
          shared is not None and "explicit_text_colors" not in (shared.metadata.properties or {}),
          str(shared and shared.metadata.properties))
    band = next((n for n in _nodes(res.tree) if (n.content.text or "") == "White in a navy-shaded style"), None)
    check("12. white in a navy-shaded paragraph style: not flagged",
          band is not None and not _flagged(band) and (band.metadata.properties or {}).get("bg_color") == NAVY,
          str(band and band.metadata.properties))
    check("control: plain grey body text still flagged",
          _flagged(next((n for n in nodes if n.content.text == "Plain grey body control"), None)), str(by_start))
    out = tmp / "body.fixed.docx"
    _remediate(res, src, out)
    colours = _run_colors(out, "word/document.xml")
    check("10. the highlighted white run is untouched", colours.get("White on a dark-blue highlight. ") == "FFFFFF",
          str(colours))
    check("10. the unhighlighted grey run was recoloured", colours.get("Grey without a highlight.") not in (None, "999999"),
          str(colours))
    check("11. both runs of the shared-colour paragraph untouched",
          colours.get("Grey on a black highlight. ") == "999999" and colours.get("Grey on the page.") == "999999",
          str(colours))

    res = _analyze(build_page_colour(tmp / "page.docx"))
    page = next((n for n in _nodes(res.tree) if (n.content.text or "") == "White on a navy page"), None)
    check("13. a painted page colour: body text not measured",
          page is not None and "explicit_text_colors" not in (page.metadata.properties or {}),
          str(page and page.metadata.properties))

    res = _analyze(build_body_panel(tmp / "panel.docx"))
    measured = [n.content.text for n in _nodes(res.tree) if (n.metadata.properties or {}).get("explicit_text_colors")]
    check("14. a shape behind body text on an unknown page: body not measured", not measured, str(measured))

    res = _analyze(build_watermark(tmp / "watermark.docx"))
    wm = next((n for n in _nodes(res.tree) if (n.content.text or "") == "Grey text under a watermark"), None)
    check("15. Word's page watermark does not switch measuring off", _flagged(wm),
          str(wm and wm.metadata.properties))
    res = _analyze(build_watermark(tmp / "fake_watermark.docx", real=False))
    fake = next((n for n in _nodes(res.tree) if (n.content.text or "") == "Grey text under a watermark"), None)
    check("16. a dark panel that only borrows the watermark's name still counts as behind the text",
          fake is not None and "explicit_text_colors" not in (fake.metadata.properties or {}),
          str(fake and fake.metadata.properties))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
