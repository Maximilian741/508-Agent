"""Smoke: every PPTX finding can be drawn where it is on the slide.

The shared location contract (violations carry ``location.page``,
``location.bbox``, ``location.pageSize``) needs the parser to record WHERE each
node is. Pinned:
  * every shape-derived node — heading (title), paragraph, image, chart,
    table, link — carries ``bbox`` = [x0, y0, x1, y1] in POINTS with the
    origin at the slide's bottom-left (the PDF convention, so one overlay
    serves both), and ``page_size`` = the slide size in points;
  * the slide's section carries ``page_size``, and every node's page is its
    1-based slide number;
  * a shape inside a group that was MOVED and SCALED is placed in slide
    coordinates, not the group's private child space — python-pptx reports
    the latter, so the "topmost text" the slide-title fix picks was compared
    across two different coordinate systems.

Usage:
    python -m app.devtools.smoke_pptx_locations
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_pptx_loc_')}/s.db")

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.models.accessibility import (  # noqa: E402
    HeadingNode,
    ImageNode,
    LinkNode,
    ParagraphNode,
    SectionNode,
    TableNode,
    iter_reading_order,
)
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360f80f000001010100187dd3a50000000049454e44ae426082"
)
EMU = 914400


def _set_group_xfrm(grp, off, ext, ch_off, ch_ext) -> None:
    xfrm = grp._element.find(f"{P}grpSpPr/{A}xfrm")
    xfrm.find(f"{A}off").set("x", str(off[0]))
    xfrm.find(f"{A}off").set("y", str(off[1]))
    xfrm.find(f"{A}ext").set("cx", str(ext[0]))
    xfrm.find(f"{A}ext").set("cy", str(ext[1]))
    xfrm.find(f"{A}chOff").set("x", str(ch_off[0]))
    xfrm.find(f"{A}chOff").set("y", str(ch_off[1]))
    xfrm.find(f"{A}chExt").set("cx", str(ch_ext[0]))
    xfrm.find(f"{A}chExt").set("cy", str(ch_ext[1]))


def _build(path: Path) -> None:
    prs = Presentation()  # 10in x 7.5in = 720 x 540 pt
    # Slide 1: title, text box, picture, table, a link.
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Permit volumes"
    tb = s.shapes.add_textbox(Inches(1), Inches(1.5), Inches(6), Inches(1))
    tb.text = "North leads the year"
    s.shapes.add_picture(io.BytesIO(PNG), Inches(4), Inches(2), Inches(2), Inches(2))
    s.shapes.add_table(2, 2, Inches(0.5), Inches(5), Inches(4), Inches(1))
    ltb = s.shapes.add_textbox(Inches(6), Inches(6.5), Inches(3), Inches(0.5))
    run = ltb.text_frame.paragraphs[0].add_run()
    run.text = "click here"
    run.hyperlink.address = "https://example.com/permits"
    # Slide 2 (untitled): a group moved to (5in, 4in) and scaled 50%, holding a
    # picture and a text box at child (0, 0). A plain text box sits at 2in.
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    grp = s2.shapes.add_group_shape()
    grp.shapes.add_picture(io.BytesIO(PNG), 0, 0, Inches(2), Inches(2))
    gtb = grp.shapes.add_textbox(0, Inches(2), Inches(2), Inches(0.5))
    gtb.text = "Grouped footnote"
    _set_group_xfrm(grp, (5 * EMU, 4 * EMU), (1 * EMU, int(1.25 * EMU)), (0, 0), (2 * EMU, int(2.5 * EMU)))
    top = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    top.text = "Budget outlook"
    prs.save(str(path))


def _close(a, b, tol=0.6) -> bool:
    return a is not None and b is not None and len(a) == len(b) and all(abs(x - y) <= tol for x, y in zip(a, b))


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="508_pptx_loc_"))
    src = tmp / "deck.pptx"
    _build(src)
    res = parse_to_tree(str(src))
    slides = [n for n in res.tree.root.children if isinstance(n, SectionNode)]

    def props(n):
        return n.metadata.properties or {}

    check("each slide section carries the slide size in points", all(props(s).get("page_size") == [720.0, 540.0] for s in slides),
          str([props(s).get("page_size") for s in slides]))
    s1 = list(iter_reading_order(slides[0]))
    heading = next(n for n in s1 if isinstance(n, HeadingNode))
    para = next(n for n in s1 if isinstance(n, ParagraphNode) and "North" in (n.content.text or ""))
    image = next(n for n in s1 if isinstance(n, ImageNode))
    table = next(n for n in s1 if isinstance(n, TableNode))
    link = next(n for n in s1 if isinstance(n, LinkNode))
    check("every slide-1 node is on page 1", all(n.metadata.page == 1 for n in (heading, para, image, table, link)))
    check("text box bbox: (1in, 1.5in) 6x1in -> [72, 360, 504, 432] pt, bottom-left origin",
          _close(props(para).get("bbox"), [72, 360, 504, 432]), str(props(para).get("bbox")))
    check("picture bbox: (4in, 2in) 2x2in -> [288, 252, 432, 396]", _close(props(image).get("bbox"), [288, 252, 432, 396]),
          str(props(image).get("bbox")))
    check("table bbox: (0.5in, 5in) 4in wide -> x 36..324, top edge at y=180",
          props(table).get("bbox") is not None and _close(props(table)["bbox"][0::2], [36, 324]) and abs(props(table)["bbox"][3] - 180) < 0.6,
          str(props(table).get("bbox")))
    check("link bbox is its text box (6in, 6.5in) 3x0.5in: [432, 36, 648, 72]", _close(props(link).get("bbox"), [432, 36, 648, 72]), str(props(link).get("bbox")))
    check("title (inherited placeholder position) has a bbox inside the slide",
          props(heading).get("bbox") is not None and 0 <= props(heading)["bbox"][0] < props(heading)["bbox"][2] <= 720
          and 0 <= props(heading)["bbox"][1] < props(heading)["bbox"][3] <= 540, str(props(heading).get("bbox")))
    check("every bbox-bearing node also names the page size", all(props(n).get("page_size") == [720.0, 540.0] for n in (heading, para, image, table, link)))

    s2 = list(iter_reading_order(slides[1]))
    gimg = next(n for n in s2 if isinstance(n, ImageNode))
    gpara = next(n for n in s2 if isinstance(n, ParagraphNode) and "footnote" in (n.content.text or ""))
    check("grouped picture is placed in SLIDE space (moved to 5in,4in; scaled 50%): [360, 180, 432, 252]",
          _close(props(gimg).get("bbox"), [360, 180, 432, 252]), str(props(gimg).get("bbox")))
    check("grouped text box: child y=2in -> slide y=5in, so it is BELOW the 2in box",
          props(gpara).get("order_hint") is not None and props(gpara)["order_hint"][0] > 4.9 * EMU, str(props(gpara).get("order_hint")))

    run_analyzers(res.tree)
    policy = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)
    plans = [p for p in plan_remediations(res.tree, policy) if p.flag.code.value == "SLIDE_TITLE_MISSING"]
    execute_plans(res.tree, plans)
    chosen = props(slides[1]).get("set_slide_title")
    check("slide 2's title is the text that is really topmost ('Budget outlook'), not the grouped footnote",
          chosen == "Budget outlook", repr(chosen))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
