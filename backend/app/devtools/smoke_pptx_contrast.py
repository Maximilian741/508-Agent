"""Smoke: FIX_CONTRAST for PPTX — recolour low-contrast runs to nearest AA shade.

Real python-pptx round trip proving the auto-fix + honesty invariant: the
failing run's colour is rewritten to an AA-passing shade in the output bytes,
re-parsing clears LOW_CONTRAST_TEXT, a run that already passes (same shape) is
untouched, and an approved recolour whose shape can't be resolved is recorded
in skipped (never credited/charged).

PPTX contrast detection only fires when the shape has an explicit solid fill
(we never assume white on slides), so the fixtures set one.

Usage:
    python -m app.devtools.smoke_pptx_contrast
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_pcx_')}/s.db")

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.api.pipeline import _action_persists  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.pptx_writer import write_remediated_pptx  # noqa: E402

CONTRAST = "LOW_CONTRAST_TEXT"


def _white_filled_box(slide, top_in):
    tb = slide.shapes.add_textbox(Inches(1), Inches(top_in), Inches(8), Inches(1))
    tb.fill.solid()
    tb.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")  # explicit bg
    return tb


def _build_pptx(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Shape 1: single grey run (#999999 ~2.8:1 on white) — must be fixed.
    tb1 = _white_filled_box(slide, 1)
    r = tb1.text_frame.paragraphs[0].add_run()
    r.text = "Grey text fails AA on white and must be auto-darkened."
    r.font.color.rgb = RGBColor.from_string("999999")
    # Shape 2: failing grey run + passing near-black run (mixed FP guard).
    tb2 = _white_filled_box(slide, 3)
    p2 = tb2.text_frame.paragraphs[0]
    bad = p2.add_run(); bad.text = "Light grey fails. "
    bad.font.color.rgb = RGBColor.from_string("AAAAAA")
    good = p2.add_run(); good.text = "Dark already passes."
    good.font.color.rgb = RGBColor.from_string("111111")
    prs.save(str(path))


def _flag_count(tree, code):
    n = 0

    def walk(node):
        nonlocal n
        n += sum(1 for f in node.accessibility_flags if f.code.value == code)
        for ch in node.children:
            walk(ch)

    walk(tree.root)
    return n


def _run_hexes(path: Path):
    prs = Presentation(str(path))
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    try:
                        rgb = run.font.color.rgb
                    except Exception:
                        rgb = None
                    if rgb is not None:
                        out.append(str(rgb).upper())
    return out


def main() -> int:
    failures = 0

    def check(name, cond, extra=""):
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="pcx_smoke_"))
    policy = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)

    check("honesty matrix: FIX_CONTRAST persists for pptx", _action_persists("FIX_CONTRAST", "pptx"))

    src = tmp / "contrast.pptx"
    _build_pptx(src)

    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    n_flags = _flag_count(res.tree, CONTRAST)
    check("two low-contrast shapes flagged", n_flags == 2, f"got {n_flags}")

    plans = plan_remediations(res.tree, policy)
    execs = execute_plans(res.tree, plans)
    fix_ok = [e for e in execs if e.action_code.value == "FIX_CONTRAST" and e.status.value == "success"]
    check("FIX_CONTRAST executed for both shapes", len(fix_ok) == 2, f"got {len(fix_ok)}")

    out = tmp / "contrast.fixed.pptx"
    result = write_remediated_pptx(src, res.tree, out)
    applied_contrast = [a for a in result["applied"] if a.get("action") == "FIX_CONTRAST"]
    check("writer applied FIX_CONTRAST to both shapes", len(applied_contrast) == 2, str(result["applied"]))

    hexes = _run_hexes(out)
    check("failing greys (#999999/#AAAAAA) gone from output", "999999" not in hexes and "AAAAAA" not in hexes, str(hexes))
    check("passing dark run (#111111) left untouched", "111111" in hexes, str(hexes))

    res2 = parse_to_tree(str(out))
    run_analyzers(res2.tree)
    check("re-parse: ALL low-contrast flags cleared", _flag_count(res2.tree, CONTRAST) == 0,
          f"got {_flag_count(res2.tree, CONTRAST)}")

    # honesty: unresolved shape not credited
    res3 = parse_to_tree(str(src))
    run_analyzers(res3.tree)
    flagged = None

    def find(node):
        nonlocal flagged
        if flagged is None and any(f.code.value == CONTRAST for f in node.accessibility_flags):
            flagged = node
        for ch in node.children:
            find(ch)

    find(res3.tree.root)
    plans3 = plan_remediations(res3.tree, policy)
    execute_plans(res3.tree, plans3)
    flagged.id = "slide-1-p-DOESNOTEXIST"
    out3 = tmp / "contrast.broken.pptx"
    res_broken = write_remediated_pptx(src, res3.tree, out3)
    applied_ids = {a.get("target_id") for a in res_broken["applied"] if a.get("action") == "FIX_CONTRAST"}
    skipped_reasons = {s.get("reason") for s in res_broken["skipped"]}
    check("unresolved shape NOT credited as FIX_CONTRAST", "slide-1-p-DOESNOTEXIST" not in applied_ids)
    check("unresolved shape recorded in skipped",
          any("contrast" in str(r) for r in skipped_reasons), str(skipped_reasons))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
