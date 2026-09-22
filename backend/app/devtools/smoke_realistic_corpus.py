"""Smoke: realistic-corpus end-to-end — analyze -> remediate ALL -> re-analyze.

Three invariants a buyer implicitly relies on, pinned across realistic
documents in all three formats plus a clean control:

  1. The pipeline never crashes on mixed real-world content.
  2. Remediation makes documents BETTER: total violations strictly decrease,
     and the headline codes the fixers target disappear from the output.
  3. Remediation never makes documents WORSE: no violation code appears in
     the remediated output that wasn't already present before (no regression,
     no fix-induced damage), and a clean document stays at zero.

Usage:
    python -m app.devtools.smoke_realistic_corpus
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_rc_')}/s.db")

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402
from pypdf import PdfWriter  # noqa: E402
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_engine import RemediationEngine  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers import write_remediated  # noqa: E402

APPLY = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)


def _tiny_png(path: Path) -> Path:
    Image.new("RGB", (24, 24), (200, 60, 20)).save(str(path))
    return path


def _violations(path: Path):
    res = parse_to_tree(str(path))
    run_analyzers(res.tree)
    return res, RemediationEngine().detect_violations(res.tree)


def _remediate_all(src: Path, out: Path):
    """Analyze, execute EVERY plan, write. Returns (pre_codes, post_codes)."""
    res, pre = _violations(src)
    plans = plan_remediations(res.tree, APPLY)
    execute_plans(res.tree, plans)
    write_remediated(src, res.tree, out, source_format=res.format)
    _, post = _violations(out)
    return [v.rule_id for v in pre], [v.rule_id for v in post]


def main() -> int:  # noqa: PLR0915
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="corpus_"))
    png = _tiny_png(tmp / "chart.png")

    # ===== 1. Realistic dirty DOCX report ====================================
    d = Document()  # NO core title
    h = d.add_paragraph()
    r = h.add_run("Annual Operations Report")  # fake heading: big bold text
    r.bold = True
    r.font.size = Pt(22)
    d.add_paragraph(
        "This report summarises operational metrics for the fiscal year and "
        "presents the consolidated results of all regional programmes."
    )
    d.add_heading("Background", level=1)
    d.add_heading("Detailed Findings", level=3)  # heading JUMP 1 -> 3
    d.add_picture(str(png))  # image with no alt text
    d.add_paragraph("Key priorities for next quarter:")
    d.add_paragraph("- expand the access audit programme")  # typed fake list
    d.add_paragraph("- replace legacy intake forms")
    d.add_paragraph("- publish quarterly metrics")
    para = d.add_paragraph("Full data is available at ")
    rid = d.part.relate_to(
        "https://data.example.gov/ops",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), rid)
    lr = OxmlElement("w:r"); lt = OxmlElement("w:t"); lt.text = "click here"
    lr.append(lt); hl.append(lr); para._p.append(hl)
    t = d.add_table(rows=3, cols=3)  # data grid, row 0 NOT header-marked
    for i in range(3):
        for j in range(3):
            t.cell(i, j).text = f"{i*10+j}"
    report = tmp / "report.docx"
    d.save(str(report))

    pre, post = _remediate_all(report, tmp / "report_fixed.docx")
    check("docx report: dirty as designed (>=6 issues)", len(pre) >= 6, f"pre={sorted(set(pre))}")
    for code in (
        "DOCUMENT_TITLE_MISSING",
        "MISSING_ALT_TEXT",
        "LIST_STRUCTURE_INVALID",
        "LINK_TEXT_NON_DESCRIPTIVE",
        "TABLE_MISSING_HEADERS",
        "TEXT_STYLED_AS_HEADING",
    ):
        check(f"docx report: detects {code}", code in pre)
    check(
        "docx report: violations strictly decrease",
        len(post) < len(pre),
        f"pre={len(pre)} post={len(post)}",
    )
    new_codes = set(post) - set(pre)
    check("docx report: NO new violation codes after fixing", not new_codes, str(new_codes))
    for fixed_code in ("DOCUMENT_TITLE_MISSING", "MISSING_ALT_TEXT", "LIST_STRUCTURE_INVALID", "TABLE_MISSING_HEADERS"):
        check(f"docx report: {fixed_code} GONE from output", fixed_code not in post)

    # ===== 2. Realistic dirty PPTX deck ======================================
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = "Programme Review"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])  # untitled
    pic = s2.shapes.add_picture(str(png), Inches(1), Inches(1))
    # python-pptx stamps descr=<filename>; PowerPoint itself leaves descr
    # empty on a no-alt picture. Clear it to simulate the real authoring tool.
    # (The filename-descr case is itself caught as ALT_TEXT_NOT_DESCRIPTIVE —
    # covered by smoke_alt_quality_headings.)
    for nv in pic._element.iter():
        if isinstance(nv.tag, str) and nv.tag.endswith("}cNvPr") and nv.get("descr"):
            del nv.attrib["descr"]
    tb = s2.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(2))
    tf = tb.text_frame
    tf.text = "- launch beta cohort"
    tf.add_paragraph().text = "- expand pilot regions"
    s3 = prs.slides.add_slide(prs.slide_layouts[6])  # untitled
    gw = s3.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(6), Inches(2))
    gw.table.first_row = False  # data grid, band off -> missing headers
    for i in range(3):
        for j in range(2):
            gw.table.cell(i, j).text = f"{i}-{j}"
    deck = tmp / "deck.pptx"
    prs.save(str(deck))

    pre, post = _remediate_all(deck, tmp / "deck_fixed.pptx")
    for code in ("SLIDE_TITLE_MISSING", "MISSING_ALT_TEXT", "LIST_STRUCTURE_INVALID", "TABLE_MISSING_HEADERS"):
        check(f"pptx deck: detects {code}", code in pre)
    check("pptx deck: 2 untitled slides -> 2 slide-title issues", pre.count("SLIDE_TITLE_MISSING") == 2)
    check("pptx deck: violations strictly decrease", len(post) < len(pre), f"pre={len(pre)} post={len(post)}")
    new_codes = set(post) - set(pre)
    check("pptx deck: NO new violation codes after fixing", not new_codes, str(new_codes))
    for fixed_code in ("LIST_STRUCTURE_INVALID", "TABLE_MISSING_HEADERS"):
        check(f"pptx deck: {fixed_code} GONE from output", fixed_code not in post)
    # MISSING_ALT_TEXT is deliberately NOT in that list. This picture sits
    # alone on an untitled slide — no caption, no nearby text — and this
    # smoke runs without an AI key, so there is genuinely nothing to describe
    # it from. The old behaviour "fixed" it by writing "Image slide-2-img1
    # shown in slide 2." — a location, not a description — and this smoke
    # asserted that overclaim. Now the executor refuses to write a placeholder
    # and the image honestly stays flagged. (The DOCX picture above DOES get
    # fixed: it has nearby text, so the heuristic derives a real alt from it.)
    check("pptx deck: caption-less image is left HONESTLY unfixed (no placeholder written)",
          "MISSING_ALT_TEXT" in post)

    # ===== 3. Untagged PDF with structure-worthy content =====================
    w = PdfWriter()
    pg = w.add_blank_page(width=612, height=792)
    body = (
        b"BT /F1 24 Tf 72 720 Td (Service Accessibility Review) Tj ET "
        b"BT /F1 11 Tf 72 690 Td (This review covers intake, triage and publication services) Tj ET "
        b"BT /F1 11 Tf 72 672 Td (offered by the agency during the reporting period in detail.) Tj ET "
        b"BT /F1 16 Tf 72 640 Td (Findings) Tj ET "
        b"BT /F1 11 Tf 90 620 Td (- intake forms lack labels) Tj ET "
        b"BT /F1 11 Tf 90 602 Td (- triage queue uses colour alone) Tj ET "
        b"BT /F1 11 Tf 90 584 Td (- publication PDFs are untagged) Tj ET "
        b"BT /F1 11 Tf 72 550 Td (The remainder of this page describes remediation milestones and the) Tj ET "
        b"BT /F1 11 Tf 72 532 Td (owners responsible for each stream across the next two quarters.) Tj ET "
    )
    cs = DecodedStreamObject(); cs.set_data(body)
    pg[NameObject("/Contents")] = w._add_object(cs)  # noqa: SLF001
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    pg[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): w._add_object(font)})
    })
    pdf = tmp / "review.pdf"
    with open(pdf, "wb") as fh:
        w.write(fh)

    pre, post = _remediate_all(pdf, tmp / "review_fixed.pdf")
    check("pdf: detects PDF_UNTAGGED", "PDF_UNTAGGED" in pre)
    check("pdf: PDF_UNTAGGED GONE after tagging", "PDF_UNTAGGED" not in post)
    check("pdf: violations strictly decrease", len(post) < len(pre), f"pre={len(pre)} post={len(post)}")
    new_codes = set(post) - set(pre)
    check("pdf: NO new violation codes after fixing", not new_codes, str(new_codes))

    # ===== 4. Clean control DOCX: zero before, zero after ====================
    c = Document()
    c.core_properties.title = "Quarterly Newsletter"
    c.core_properties.language = "en-US"
    c.add_heading("Quarterly Newsletter", level=1)
    c.add_paragraph("Welcome to the spring edition of our accessible newsletter.")
    c.add_heading("Community News", level=2)
    c.add_paragraph("first real item", style="List Bullet")
    c.add_paragraph("second real item", style="List Bullet")
    c.add_picture(str(png))
    # give the image descriptive alt text directly
    last_inline = c.inline_shapes[-1]
    last_inline._inline.docPr.set("descr", "Bar chart of programme participation by region")
    clean = tmp / "clean.docx"
    c.save(str(clean))

    res, pre_v = _violations(clean)
    check("clean docx: ZERO violations before", len(pre_v) == 0, str([v.rule_id for v in pre_v]))
    plans = plan_remediations(res.tree, APPLY)
    execute_plans(res.tree, plans)
    out_clean = tmp / "clean_fixed.docx"
    write_remediated(clean, res.tree, out_clean, source_format=res.format)
    _, post_v = _violations(out_clean)
    check("clean docx: ZERO violations after", len(post_v) == 0, str([v.rule_id for v in post_v]))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
