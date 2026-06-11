"""Smoke: /pipeline/batch-zip bundles remediated files into one ZIP.

The bulk/agency workflow: remediate several documents, then download all the
fixed files as a single ZIP. Pins:

  - a valid set of jobs -> 200, a real ZIP containing each remediated file
  - de-duplicates identical archive names
  - auth required (401 without a token)
  - all-invalid job ids -> 404 (nothing to bundle)
  - empty job list -> 400

Usage:
    python -m app.devtools.smoke_batch_zip
"""

from __future__ import annotations

import io
import json as _json
import os
import sys
import tempfile
import zipfile

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_bz_')}/s.db"

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.main import app

    client = TestClient(app, raise_server_exceptions=False)

    r = client.post(
        "/auth/sign-in",
        json={"email": "batchzip@example.com", "displayName": "BZ", "password": "batchzip123"},
    )
    auth = {"Authorization": f"Bearer {r.json()['token']}"}
    client.post("/auth/grant-starter", headers=auth)

    def _docx(title_missing: bool = True) -> bytes:
        d = Document()
        if not title_missing:
            d.core_properties.title = "Has title"
        d.add_heading("Report", level=1)
        d.add_paragraph("Body text.")
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    # Remediate two documents -> two jobs.
    jobs = []
    for i in range(2):
        b = _docx()
        ar = client.post("/pipeline/analyze", files={"file": (f"doc{i}.docx", b, DOCX_MIME)}, headers=auth)
        approved = [v["id"] for v in ar.json().get("violations", [])][:1]
        rr = client.post(
            "/pipeline/remediate",
            files={"file": (f"doc{i}.docx", b, DOCX_MIME)},
            data={"approved_violations": _json.dumps(approved), "rejected_violations": "[]"},
            headers=auth,
        )
        check(f"remediate doc{i} -> 200", rr.status_code == 200, rr.text[:160])
        jr = rr.json()
        jobs.append({"jobId": jr["jobId"], "filename": jr["filename"]})

    # --- 1. Valid batch zip --------------------------------------------------
    r = client.post("/pipeline/batch-zip", json={"jobs": jobs}, headers=auth)
    check("batch-zip -> 200", r.status_code == 200, f"got {r.status_code} {r.text[:160]}")
    check("content-type is zip", "zip" in r.headers.get("content-type", ""))
    try:
        zf = zipfile.ZipFile(io.BytesIO(r.content))
        names = zf.namelist()
    except Exception as exc:
        names = []
        check("response is a valid zip", False, str(exc))
    check("zip contains both remediated files", len(names) == 2, str(names))
    check("zipped files are non-empty docx", all(zf.read(n)[:2] == b"PK" and len(zf.read(n)) > 500 for n in names) if names else False)

    # --- 2. Duplicate archive names de-duplicated ----------------------------
    dup = [jobs[0], {"jobId": jobs[0]["jobId"], "filename": jobs[0]["filename"]}]
    r = client.post("/pipeline/batch-zip", json={"jobs": dup}, headers=auth)
    # Same job twice -> the single real file; names must not collide/overwrite.
    if r.status_code == 200:
        z = zipfile.ZipFile(io.BytesIO(r.content))
        check("duplicate job de-dups to one archive entry", len(z.namelist()) == 1, str(z.namelist()))
    else:
        check("duplicate job de-dups (200)", False, f"got {r.status_code}")

    # --- 3. Auth required ----------------------------------------------------
    r = client.post("/pipeline/batch-zip", json={"jobs": jobs})
    check("no token -> 401", r.status_code == 401, f"got {r.status_code}")

    # --- 4. All-invalid jobs -> 404 -----------------------------------------
    r = client.post("/pipeline/batch-zip", json={"jobs": [{"jobId": "deadbeef00", "filename": "nope.docx"}]}, headers=auth)
    check("all-invalid jobs -> 404", r.status_code == 404, f"got {r.status_code}")

    # --- 5. Empty list -> 400 ------------------------------------------------
    r = client.post("/pipeline/batch-zip", json={"jobs": []}, headers=auth)
    check("empty job list -> 400", r.status_code == 400, f"got {r.status_code}")

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
