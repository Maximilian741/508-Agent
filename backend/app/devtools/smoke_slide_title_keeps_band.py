"""Smoke: giving a slide a title never erases the band its EMPTY title box draws.

Found by rendering in PowerPoint: a "Title Only" layout whose title
placeholder has a dark-blue fill, a slide that keeps that placeholder EMPTY,
and the real title typed on top of it in a white 40pt text box. The writer
used to delete the empty placeholder (promote path) or move it off the slide
to hold the title (off-slide path) — the band went with it and the white
title sat on white. It was charged.

Pinned here, through the real detect -> plan -> execute -> write path and the
pipeline's honesty gate:

  * promote path: the band stays — same shape id, same z-order (behind the
    title), the layout's position and fill written onto it, no longer a
    placeholder, marked decorative; the promoted box is the only title;
  * off-slide path (two-line box): the band stays on the slide and a
    separate off-slide title placeholder carries the title;
  * an outline inherited from the MASTER (the layout only sets a width) is
    carried over complete;
  * an empty title placeholder that draws NOTHING is still dropped (no
    stray empty shape left behind);
  * a band that can't be carried over faithfully (a picture fill that lives in
    the layout's own part) is REFUSED: nothing written, the execution turns
    SKIPPED with a plain reason, and it is not counted or charged;
  * re-analysis of every output finds the slide titled and nothing new.

Usage:
    python -m app.devtools.smoke_slide_title_keeps_band
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_stkb_')}/s.db")

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.api.pipeline import _count_persisted_fixes, _reconcile_executions  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.pptx_writer import write_remediated_pptx  # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
DECORATIVE_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
BAND = "1F4E79"


def _title_only_deck(style: str, two_lines: bool = False) -> bytes:
    """A one-slide deck on the Title Only layout, the title typed in a white
    text box over the (empty) title placeholder. ``style`` decorates the
    layout/master title placeholder."""
    prs = Presentation()
    prs.core_properties.title = "Quarterly review"
    prs.core_properties.language = "en-US"
    layout = prs.slide_layouts[5]  # Title Only
    lt = next(p for p in layout.placeholders if p.placeholder_format.idx == 0)
    sp_pr = lt._element.spPr  # noqa: SLF001
    if style == "fill":
        fill = etree.SubElement(sp_pr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", BAND)
    elif style == "master_outline":
        # The layout only sets a width; the colour lives on the master.
        etree.SubElement(sp_pr, qn("a:ln")).set("w", "38100")
        mt = next(p for p in prs.slide_master.placeholders if p.placeholder_format.type == 1)  # TITLE
        m_ln = etree.SubElement(mt._element.spPr, qn("a:ln"))  # noqa: SLF001
        m_fill = etree.SubElement(m_ln, qn("a:solidFill"))
        etree.SubElement(m_fill, qn("a:srgbClr")).set("val", "C00000")
    elif style == "picture":
        blip_fill = etree.SubElement(sp_pr, qn("a:blipFill"))
        etree.SubElement(blip_fill, qn("a:blip")).set(qn("r:embed"), "rId99")
        etree.SubElement(etree.SubElement(blip_fill, qn("a:stretch")), qn("a:fillRect"))
    slide = prs.slides.add_slide(layout)
    tph = slide.shapes.title
    box = slide.shapes.add_textbox(tph.left, tph.top, tph.width, tph.height)
    box.text_frame.text = "Quarterly Results\nFiscal year 2026" if two_lines else "Quarterly Results"
    for r in box.text_frame.paragraphs[0].runs:
        r.font.size = Pt(40)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(2)).text_frame.text = (
        "Revenue grew in every region and margins improved across the business this quarter."
    )
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


def _slide1(path: Path):
    with zipfile.ZipFile(path) as z:
        return etree.fromstring(z.read("ppt/slides/slide1.xml"))


def _layout_title_xfrm(path: Path):
    prs = Presentation(str(path))
    lt = next(p for p in prs.slides[0].slide_layout.placeholders if p.placeholder_format.idx == 0)
    return (lt.left, lt.top, lt.width, lt.height)


def _run(src: Path, out: Path):
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    before = sorted(f.code.value for n in _walk(res.tree.root) for f in n.accessibility_flags)
    policy = RemediationPolicy(allow_ai_actions=False, require_human_review_for_all=False)
    plans = [p for p in plan_remediations(res.tree, policy) if p.flag.code.value == "SLIDE_TITLE_MISSING"]
    execs = execute_plans(res.tree, plans)
    rep = write_remediated_pptx(src, res.tree, out)
    return before, plans, execs, rep


def _walk(node):
    yield node
    for c in node.children or []:
        yield from _walk(c)


def _flags_after(out: Path):
    res = parse_to_tree(str(out))
    run_analyzers(res.tree)
    return sorted(f.code.value for n in _walk(res.tree.root) for f in n.accessibility_flags)


def _title_phs(sld):
    return [
        sp for sp in sld.iter(P + "sp")
        if (ph := sp.find(f"{P}nvSpPr/{P}nvPr/{P}ph")) is not None and ph.get("type") in ("title", "ctrTitle")
    ]


def _by_id(sld, shape_id: str):
    return next((sp for sp in sld.iter(P + "sp") if sp.find(f"{P}nvSpPr/{P}cNvPr").get("id") == shape_id), None)


def _is_decorative(sp) -> bool:
    ext_lst = sp.find(f"{P}nvSpPr/{P}cNvPr/{A}extLst")
    return ext_lst is not None and any(e.get("uri") == DECORATIVE_URI for e in ext_lst)


def _xfrm(sp):
    off = sp.find(f"{P}spPr/{A}xfrm/{A}off")
    ext = sp.find(f"{P}spPr/{A}xfrm/{A}ext")
    if off is None or ext is None:
        return None
    return (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="508_stkb_"))

    # ---- 1. promote path: band kept behind the promoted title --------------
    src = tmp / "band.pptx"
    src.write_bytes(_title_only_deck("fill"))
    band_id = _title_phs(_slide1(src))[0].find(f"{P}nvSpPr/{P}cNvPr").get("id")
    out = tmp / "band_out.pptx"
    before, plans, execs, rep = _run(src, out)
    applied = [a for a in rep["applied"] if a.get("action") == "SET_SLIDE_TITLE"]
    check("promote: the slide was flagged and the title applied", len(plans) == 1 and len(applied) == 1, f"{len(plans)} {rep}")
    sld = _slide1(out)
    band = _by_id(sld, band_id)
    check("promote: the empty title box is still on the slide (same shape id)", band is not None)
    if band is not None:
        check("promote: it is no longer a placeholder", band.find(f"{P}nvSpPr/{P}nvPr/{P}ph") is None)
        fill = band.find(f"{P}spPr/{A}solidFill/{A}srgbClr")
        check("promote: it still paints the layout's band colour", fill is not None and fill.get("val") == BAND,
              etree.tostring(band.find(f"{P}spPr"))[:300])
        check("promote: at the layout title's position and size", _xfrm(band) == _layout_title_xfrm(src),
              f"{_xfrm(band)} vs {_layout_title_xfrm(src)}")
        check("promote: marked decorative (the slide's design, not content)", _is_decorative(band))
        shapes = [sp for sp in sld.find(f"{P}cSld/{P}spTree") if sp.tag == P + "sp"]
        titles = _title_phs(sld)
        check("promote: exactly one title placeholder, carrying the slide's own words",
              len(titles) == 1 and "".join(t.text or "" for t in titles[0].iter(A + "t")) == "Quarterly Results")
        check("promote: the band is still drawn BEHIND the title (z-order kept)",
              bool(titles) and shapes.index(band) < shapes.index(titles[0]))
    prs = Presentation(str(out))
    check("promote: the deck reopens and PowerPoint's title is the slide's text",
          prs.slides[0].shapes.title is not None and prs.slides[0].shapes.title.text == "Quarterly Results")
    after = _flags_after(out)
    check("promote: re-analysis finds the slide titled and nothing new",
          "SLIDE_TITLE_MISSING" not in after and set(after) <= set(before), f"{before} -> {after}")
    check("promote: counted and charged (the writer confirmed it)",
          _count_persisted_fixes(execs, rep["applied"], "pptx", rep["skipped"]) == 1)

    # ---- 2. off-slide path (two-line box): band kept, title hidden above ----
    src = tmp / "band2.pptx"
    src.write_bytes(_title_only_deck("fill", two_lines=True))
    out = tmp / "band2_out.pptx"
    _before, _plans, execs, rep = _run(src, out)
    sld = _slide1(out)
    band = _by_id(sld, band_id)
    titles = _title_phs(sld)
    check("off-slide: the band is still on the slide with its colour",
          band is not None and band.find(f"{P}nvSpPr/{P}nvPr/{P}ph") is None
          and band.find(f"{P}spPr/{A}solidFill/{A}srgbClr") is not None
          and _xfrm(band) == _layout_title_xfrm(src))
    check("off-slide: a separate title placeholder holds the title, entirely above the slide",
          len(titles) == 1 and titles[0] is not band and _xfrm(titles[0]) is not None
          and _xfrm(titles[0])[1] + _xfrm(titles[0])[3] < 0)
    check("off-slide: re-analysis finds the slide titled", "SLIDE_TITLE_MISSING" not in _flags_after(out))

    # ---- 3. outline inherited from the master is carried over complete -----
    src = tmp / "outline.pptx"
    src.write_bytes(_title_only_deck("master_outline"))
    out = tmp / "outline_out.pptx"
    _before, _plans, execs, rep = _run(src, out)
    band = _by_id(_slide1(out), band_id)
    ln = band.find(f"{P}spPr/{A}ln") if band is not None else None
    check("outline: the empty title box is kept (it draws a red outline)",
          band is not None and band.find(f"{P}nvSpPr/{P}nvPr/{P}ph") is None)
    check("outline: width from the layout AND colour from the master",
          ln is not None and ln.get("w") == "38100" and ln.find(f"{A}solidFill/{A}srgbClr") is not None
          and ln.find(f"{A}solidFill/{A}srgbClr").get("val") == "C00000",
          etree.tostring(ln)[:300] if ln is not None else "no ln")
    check("outline: no fill invented (the chain has none)", band is not None and band.find(f"{P}spPr/{A}noFill") is not None)

    # ---- 4. an empty title box that draws nothing is still dropped ---------
    src = tmp / "plain.pptx"
    src.write_bytes(_title_only_deck("none"))
    out = tmp / "plain_out.pptx"
    _before, _plans, execs, rep = _run(src, out)
    sld = _slide1(out)
    check("plain: the invisible empty placeholder is dropped, not kept as a stray shape",
          _by_id(sld, band_id) is None and len(_title_phs(sld)) == 1, str(rep["applied"]))

    # ---- 5. a picture fill living in the layout's part: refused, not charged
    src = tmp / "picture.pptx"
    src.write_bytes(_title_only_deck("picture"))
    out = tmp / "picture_out.pptx"
    src_slide = etree.tostring(_slide1(src))
    _before, plans, execs, rep = _run(src, out)
    check("picture: the executor picked the text (the writer is what refuses)",
          len(plans) == 1 and [e.status.value for e in execs] == ["success"], str([(e.status.value, e.notes) for e in execs]))
    check("picture: the writer refused with a reason and applied nothing",
          not [a for a in rep["applied"] if a.get("action") == "SET_SLIDE_TITLE"]
          and any(s.get("reason") == "empty_title_box_draws_slide_design" for s in rep["skipped"]), str(rep["skipped"]))
    check("picture: the slide is byte-for-byte untouched", etree.tostring(_slide1(out)) == src_slide)
    check("picture: not counted, so not charged", _count_persisted_fixes(execs, rep["applied"], "pptx", rep["skipped"]) == 0)
    _reconcile_executions(execs, rep["applied"], "pptx")
    e = execs[0]
    check("picture: the execution the UI shows says SKIPPED, in plain words, not charged",
          e.status.value == "skipped" and "not charged" in (e.notes or "") and "title box" in (e.notes or ""), f"{e.status} {e.notes}")

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 0 if failures == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
