"""Smoke: three new detection refinements (batch 5).

1. ALT_TEXT_NOT_DESCRIPTIVE — a meaningful image whose alt text is just a
   filename ("image1.png") or a generic placeholder ("Picture 1") passes the
   missing-alt check but says nothing to a screen-reader user (WCAG 1.1.1).
2. DOCUMENT_NO_HEADINGS — a long DOCX/PDF with substantial running text but zero
   headings is hard to navigate (WCAG 2.4.6 / 1.3.1). Scoped to page-flow formats
   so slide decks (which organise by slide title) are never falsely flagged.
3. LINK_TEXT_NON_DESCRIPTIVE (strengthened) — now also catches link text that is
   a bare URL ("https://example.com/a") plus more generic phrases. This is the
   high-value half: such links are remediable and the fix now PERSISTS (batch 4).

Usage:
    python -m app.devtools.smoke_alt_quality_headings
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_altq_')}/s.db"

from docx import Document  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.analyzers.image_analyzer import is_nondescriptive_alt  # noqa: E402
from app.analyzers.link_analyzer import _looks_like_url  # noqa: E402
from app.analyzers.registry import run_analyzers  # noqa: E402
from app.services.remediation_planner import plan_remediations, RemediationPolicy  # noqa: E402
from app.services.remediators.generate_alt_text_executor import GenerateAltTextExecutor  # noqa: E402
from app.models.accessibility import (  # noqa: E402
    AccessibilityFlagCode,
    AccessibilityTree,
    ContentKind,
    DocumentNode,
    ImageNode,
    LinkNode,
    NodeContent,
    NodeMetadata,
)
from app.parsers import parse_to_tree  # noqa: E402


def _codes(tree: AccessibilityTree) -> set[str]:
    found: set[str] = set()

    def walk(node) -> None:
        for f in node.accessibility_flags:
            found.add(f.code.value)
        for ch in node.children:
            walk(ch)

    walk(tree.root)
    return found


def _img_tree(alt: str, caption: str | None = None) -> AccessibilityTree:
    # A Word Caption-styled paragraph is text written FOR the picture; the
    # DOCX parser records that as caption_source="caption_style".
    props = {"caption": caption, "caption_source": "caption_style"} if caption else {}
    img = ImageNode(
        id="img-1",
        content=NodeContent(kind=ContentKind.NONE),
        metadata=NodeMetadata(source_format="docx", properties=props),
        is_decorative=False,
        alt_text=alt,
    )
    root = DocumentNode(
        id="doc-1",
        content=NodeContent(kind=ContentKind.NONE),
        metadata=NodeMetadata(source_format="docx"),
        children=[img],
    )
    return AccessibilityTree(root=root)


def _link_tree(text: str) -> AccessibilityTree:
    link = LinkNode(
        id="link-1",
        content=NodeContent(kind=ContentKind.TEXT, text=text),
        metadata=NodeMetadata(source_format="docx"),
        target="https://example.com/destination",
    )
    root = DocumentNode(
        id="doc-1",
        content=NodeContent(kind=ContentKind.NONE),
        metadata=NodeMetadata(source_format="docx"),
        children=[link],
    )
    return AccessibilityTree(root=root)


def main() -> int:
    failures = 0

    def check(name: str, cond: bool) -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name)
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp())

    # --- 1a. is_nondescriptive_alt: pure-function patterns ---
    for bad in ["image1.png", "DSC_0042.JPG", "logo.svg", "Picture 1", "image",
                "img_2", "Graphic", "Screenshot 2024", "screen shot 5", "Untitled",
                "Content Placeholder 3", "Chart"]:
        check(f"alt {bad!r} flagged non-descriptive", is_nondescriptive_alt(bad))
    for good in ["A red car on a wet street", "Q3 revenue grew 12 percent",
                 "Company logo: blue circle with the letter A", "Bar chart of sales by region",
                 "Photograph of the new campus entrance", "Diagram of the request lifecycle"]:
        check(f"alt {good!r} NOT flagged", not is_nondescriptive_alt(good))

    # --- 1b. analyzer end-to-end on a constructed tree ---
    check("filename alt -> ALT_TEXT_NOT_DESCRIPTIVE",
          AccessibilityFlagCode.ALT_TEXT_NOT_DESCRIPTIVE.value in _codes(run_analyzers(_img_tree("image1.png"))))
    check("good alt -> no ALT_TEXT_NOT_DESCRIPTIVE flag",
          AccessibilityFlagCode.ALT_TEXT_NOT_DESCRIPTIVE.value not in _codes(run_analyzers(_img_tree("A red car on a wet street"))))
    # an empty-alt image is the MissingAltText case, not this one
    check("empty alt -> not the non-descriptive flag (missing instead)",
          AccessibilityFlagCode.ALT_TEXT_NOT_DESCRIPTIVE.value not in _codes(run_analyzers(_img_tree(""))))

    # --- 1c. non-descriptive alt is AUTO-FIXABLE: the executor REPLACES it ---
    # ...when there is something REAL to replace it with. The image carries a
    # caption so the heuristic provider (no AI key in smokes) can derive a
    # genuine description. Without any context the executor now refuses to
    # write the "Image ... shown in document." placeholder — which is what
    # this fixture used to receive, and which the "no loop" check below only
    # passed because the classifier could not see it. See 1d.
    policy = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)
    ex = GenerateAltTextExecutor()
    t_bad = _img_tree("image1.png", caption="Bar chart of quarterly sales by region")
    run_analyzers(t_bad)
    bad_img = t_bad.root.children[0]
    plans = [p for p in plan_remediations(t_bad, policy)
             if p.flag.code == AccessibilityFlagCode.ALT_TEXT_NOT_DESCRIPTIVE]
    check("non-descriptive alt -> a GENERATE_ALT_TEXT plan exists", len(plans) == 1)
    if plans:
        res = ex.execute(plans[0], t_bad)
        check("executor replaces filename alt (status success)", res.status.value == "success")
        check("filename alt was overwritten with a description",
              bad_img.alt_text and bad_img.alt_text != "image1.png")
        check("replacement does not itself re-trigger the flag (no loop)",
              not is_nondescriptive_alt(bad_img.alt_text or ""))
        check("replacement marked pending human review",
              bad_img.metadata.properties.get("alt_text_pending_review") is True)

    # --- 1d. no context at all -> the executor REFUSES the placeholder ---
    t_bare = _img_tree("image1.png")
    run_analyzers(t_bare)
    bare_img = t_bare.root.children[0]
    bare_plans = [p for p in plan_remediations(t_bare, policy)
                  if p.flag.code == AccessibilityFlagCode.ALT_TEXT_NOT_DESCRIPTIVE]
    if bare_plans:
        res_b = ex.execute(bare_plans[0], t_bare)
        check("no context: executor SKIPS rather than writing a location placeholder",
              res_b.status.value == "skipped")
        check("no context: the filename alt is left as-is (still flagged, honestly)",
              bare_img.alt_text == "image1.png")
        check("no context: the note names the reason",
              "caption" in (res_b.notes or "").lower() and "not charged" in (res_b.notes or "").lower())

    # good alt must NEVER be overwritten (it carries no non-descriptive flag)
    t_good = _img_tree("A red car on a wet street at night")
    run_analyzers(t_good)
    check("good alt is not flagged non-descriptive (so never regenerated)",
          AccessibilityFlagCode.ALT_TEXT_NOT_DESCRIPTIVE.value not in _codes(t_good))

    # --- 2. DOCUMENT_NO_HEADINGS ---
    NO_H = AccessibilityFlagCode.DOCUMENT_NO_HEADINGS.value
    body = ("This paragraph carries enough running prose to count as real reading "
            "content for the heading-density heuristic under test. ")

    # long docx, zero headings -> flagged
    d_long = tmp / "long.docx"
    doc = Document()
    for _ in range(15):
        doc.add_paragraph(body)
    doc.save(str(d_long))
    check("long docx with no headings -> DOCUMENT_NO_HEADINGS",
          NO_H in _codes(run_analyzers(parse_to_tree(str(d_long)).tree)))

    # long docx WITH a heading -> not flagged
    d_head = tmp / "head.docx"
    doc = Document()
    doc.add_heading("Overview", level=1)
    for _ in range(15):
        doc.add_paragraph(body)
    doc.save(str(d_head))
    check("docx with a heading -> no DOCUMENT_NO_HEADINGS",
          NO_H not in _codes(run_analyzers(parse_to_tree(str(d_head)).tree)))

    # short docx -> not flagged (below threshold)
    d_short = tmp / "short.docx"
    doc = Document()
    for _ in range(3):
        doc.add_paragraph(body)
    doc.save(str(d_short))
    check("short docx -> no DOCUMENT_NO_HEADINGS (below threshold)",
          NO_H not in _codes(run_analyzers(parse_to_tree(str(d_short)).tree)))

    # pptx with lots of text, no headings -> NOT flagged (format-scoped out)
    p_deck = tmp / "deck.pptx"
    prs = Presentation()
    for _ in range(4):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5))
        tf = tb.text_frame
        for _ in range(6):
            tf.add_paragraph().text = body
    prs.save(str(p_deck))
    check("pptx with text but no headings -> NOT DOCUMENT_NO_HEADINGS (scoped to docx/pdf)",
          NO_H not in _codes(run_analyzers(parse_to_tree(str(p_deck)).tree)))

    # --- 3. LINK_TEXT_NON_DESCRIPTIVE strengthened (URL + more phrases) ---
    for bad_url in ["https://example.com/a/b", "http://foo.org", "www.example.com",
                    "example.com", "example.com/path?x=1", "SUB.example.co.uk/page"]:
        check(f"link text {bad_url!r} looks like a URL", _looks_like_url(bad_url))
    for ok_text in ["Visit example.com today", "Our annual report", "fig.5",
                    "Acme Inc.", "see section 3.2", "email us at hello"]:
        check(f"link text {ok_text!r} not a bare URL", not _looks_like_url(ok_text))

    LT = AccessibilityFlagCode.LINK_TEXT_NON_DESCRIPTIVE.value
    check("bare-URL link text -> LINK_TEXT_NON_DESCRIPTIVE",
          LT in _codes(run_analyzers(_link_tree("https://example.com/pricing"))))
    check("new generic phrase 'website' -> LINK_TEXT_NON_DESCRIPTIVE",
          LT in _codes(run_analyzers(_link_tree("website"))))
    check("descriptive link text -> no LINK_TEXT_NON_DESCRIPTIVE",
          LT not in _codes(run_analyzers(_link_tree("Download the 2024 annual report"))))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
