"""Smoke: the PPTX writer reports only what it actually changed.

The writer's ``applied`` list is what the UI shows as "fixes made". It used to
report, on EVERY run of EVERY deck:
  * each picture's existing alt text as a fresh ``image_alt_text`` fix —
    including PowerPoint's own ``descr="image.png"`` that nobody approved
    (the alt-text executor had just refused to write a placeholder for it);
  * ``table_first_row_header`` for every header cell of a table whose header
    band was already on ("tblPr/@firstRow already 1");
  * ``document_title`` for a deck whose title was already set.
Pinned: a parsed deck written back UNCHANGED yields an empty ``applied``
list and slide XML identical to the source.

Also pinned: promoting a hyperlinked text box to the slide title freezes its
look WITHOUT pinning the text colour or underline on the link run — those
come from the theme's hyperlink styling, and pinning tx1 / u="none" turned a
blue underlined link into plain black text.

Usage:
    python -m app.devtools.smoke_pptx_writer_noop_honesty
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_pptx_noop_')}/s.db")

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.pptx_writer import write_remediated_pptx  # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360f80f000001010100187dd3a50000000049454e44ae426082"
)


def _build(path: Path) -> None:
    prs = Presentation()
    prs.core_properties.title = "Quarterly review"
    prs.core_properties.language = "en-US"
    s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    s.shapes.title.text = "Overview"
    pic = s.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(2), Inches(2), Inches(2))
    pic._element.find(f"{P}nvPicPr/{P}cNvPr").set("descr", "Team photo outside City Hall")
    s.shapes.add_picture(io.BytesIO(PNG), Inches(4), Inches(2), Inches(2), Inches(2))  # descr="image.png"
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(4.5), Inches(4), Inches(1)).table  # firstRow="1"
    for c, (h, v) in enumerate((("Region", "North"), ("Permits", "120"))):
        tbl.cell(0, c).text = h
        tbl.cell(1, c).text = v
    # Slide 2: the only title candidate is a hyperlinked text box.
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s2.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "Permit application portal"
    run.hyperlink.address = "https://example.com/permits"
    prs.save(str(path))


def _slide(path: Path, n: int) -> bytes:
    with zipfile.ZipFile(path) as z:
        return z.read(f"ppt/slides/slide{n}.xml")


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="508_pptx_noop_"))
    src = tmp / "deck.pptx"
    _build(src)

    # ---- 1. untouched tree -> nothing reported, nothing changed -------------
    res = parse_to_tree(str(src))
    out = tmp / "untouched.pptx"
    rep = write_remediated_pptx(src, res.tree, out)
    check("untouched deck: the writer reports no applied fixes", rep["applied"] == [], str(rep["applied"]))
    check("untouched deck: slide 1 XML is unchanged", etree.tostring(etree.fromstring(_slide(out, 1)))
          == etree.tostring(etree.fromstring(_slide(src, 1))))
    reopened = Presentation(str(out))
    check("untouched deck: title and language still set",
          reopened.core_properties.title == "Quarterly review" and reopened.core_properties.language == "en-US")

    # ---- 2. the real path: only the approved slide title is reported --------
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    policy = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)
    plans = plan_remediations(res.tree, policy)
    execs = execute_plans(res.tree, plans)
    by_action = {}
    for e in execs:
        by_action.setdefault(e.action_code.value, []).append(e.status.value)
    check("the placeholder 'image.png' alt was refused, not written", by_action.get("GENERATE_ALT_TEXT") == ["skipped"], str(by_action))
    out = tmp / "fixed.pptx"
    rep = write_remediated_pptx(src, res.tree, out)
    kinds = sorted(a.get("kind") for a in rep["applied"])
    check("fixed deck: the ONLY applied entry is the slide title it set", kinds == ["slide_title"], str(rep["applied"]))
    check("fixed deck: the unapproved 'image.png' descr is not reported as a fix",
          not any(a.get("kind") == "image_alt_text" for a in rep["applied"]))

    s2 = etree.fromstring(_slide(out, 2))
    sps = [sp for sp in s2.iter(f"{P}sp") if "Permit application portal" in "".join(t.text or "" for t in sp.iter(f"{A}t"))]
    check("hyperlinked box: exactly one shape carries the words", len(sps) == 1, str(len(sps)))
    sp = sps[0] if sps else None
    is_title = sp is not None and sp.find(f"{P}nvSpPr/{P}nvPr/{P}ph") is not None
    check("hyperlinked box: it was promoted to the title placeholder", is_title)
    r_pr = sp.find(f".//{A}r/{A}rPr") if sp is not None else None
    check("hyperlinked box: the link is kept", r_pr is not None and r_pr.find(f"{A}hlinkClick") is not None)
    check("hyperlinked box: no text colour pinned on the link run (theme hlink colour stays)",
          r_pr is not None and not any(c.tag in (f"{A}solidFill", f"{A}schemeClr") for c in r_pr), etree.tostring(r_pr) if r_pr is not None else "")
    check("hyperlinked box: no underline override pinned on the link run", r_pr is not None and r_pr.get("u") is None)
    check("hyperlinked box: size and font are still pinned (title style can't restyle it)",
          r_pr is not None and r_pr.get("sz") is not None and r_pr.find(f"{A}latin") is not None)
    again = parse_to_tree(str(out))
    run_analyzers(again.tree)
    untitled = [n for n in again.tree.root.children if (n.metadata.properties or {}).get("missing_title")]
    check("re-analysis: no untitled slide left", untitled == [], str([n.id for n in untitled]))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
