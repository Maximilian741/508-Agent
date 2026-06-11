"""Smoke: remediation never charges for a file it didn't produce (the C1 fix).

The cardinal billing bug the audit found: credits were spent BEFORE the
document was parsed/written, so a parse failure (exactly what corrupt/exotic
real-world PDFs cause) charged the user and returned a 422 with no file and
no refund. This pins the corrected behaviour over real HTTP:

  - a corrupt upload -> 422 AND the balance is UNCHANGED (no charge)
  - a valid upload  -> 200 AND the balance drops by exactly the format cost
  - a broke user    -> 402 up front, balance untouched, no work done

Usage:
    python -m app.devtools.smoke_billing_safety
"""

from __future__ import annotations

import io
import os
import sys
import tempfile

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_bill_')}/s.db"
os.environ.pop("SMTP_HOST", None)
os.environ.pop("STRIPE_SECRET_KEY", None)

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.main import app

    client = TestClient(app, raise_server_exceptions=False)

    EMAIL = "billing-safety@example.com"
    r = client.post(
        "/auth/sign-in",
        json={"email": EMAIL, "displayName": "Bill", "password": "billsafe123"},
    )
    token = r.json()["token"]
    auth = {"Authorization": f"Bearer {token}"}
    client.post("/auth/grant-starter", headers=auth)  # +25 credits

    def balance() -> int:
        b = client.get("/credits/balance", headers=auth)
        return int(b.json().get("balance", 0))

    start = balance()
    check("starter credits granted", start >= 25, f"balance={start}")

    # A valid docx (no title -> at least one fixable finding).
    d = Document()
    d.add_heading("Quarterly Report", level=1)
    d.add_paragraph("Body text for the report.")
    valid = io.BytesIO()
    d.save(valid)
    valid_bytes = valid.getvalue()

    # --- 1. Corrupt upload: 422 and NO charge --------------------------------
    before = balance()
    r = client.post(
        "/pipeline/remediate",
        files={"file": ("broken.docx", b"PK\x03\x04 not really a docx at all" * 20, DOCX_MIME)},
        data={"approved_violations": "[]", "rejected_violations": "[]"},
        headers=auth,
    )
    # Either 400 (rejected by the upload/OOXML sniff) or 422 (parse failure) —
    # both are correct "no file, no charge" outcomes.
    check("corrupt upload -> 4xx", 400 <= r.status_code < 500, f"got {r.status_code}")
    after = balance()
    check("corrupt upload did NOT charge (balance unchanged)", after == before, f"{before} -> {after}")

    # --- 2. Truncated (parseable magic, unparseable body): 422, no charge ----
    before = balance()
    r = client.post(
        "/pipeline/remediate",
        files={"file": ("trunc.docx", valid_bytes[:120], DOCX_MIME)},
        data={"approved_violations": "[]", "rejected_violations": "[]"},
        headers=auth,
    )
    check("truncated upload -> 4xx", 400 <= r.status_code < 500, f"got {r.status_code}")
    after = balance()
    check("truncated upload did NOT charge", after == before, f"{before} -> {after}")

    # --- 3. Valid upload: 200 and charged EXACTLY the docx cost (3) -----------
    before = balance()
    # analyze first to get a real violation id to approve.
    ar = client.post("/pipeline/analyze", files={"file": ("ok.docx", valid_bytes, DOCX_MIME)}, headers=auth)
    viols = ar.json().get("violations", [])
    approved = [v["id"] for v in viols][:1]
    import json as _json

    r = client.post(
        "/pipeline/remediate",
        files={"file": ("ok.docx", valid_bytes, DOCX_MIME)},
        data={"approved_violations": _json.dumps(approved), "rejected_violations": "[]"},
        headers=auth,
    )
    check("valid upload -> 200", r.status_code == 200, f"got {r.status_code} {r.text[:160]}")
    check("valid upload returns a downloadUrl", bool(r.json().get("downloadUrl")))
    after = balance()
    check("valid upload charged exactly the docx cost (3)", before - after == 3, f"{before} -> {after}")

    # --- 4. Broke user: 402 up front, no charge ------------------------------
    # Drain the wallet, then attempt a remediate.
    drain = balance()
    # Spend down to < 3 by issuing remediates until close, or just check the
    # 402 path by draining via repeated valid remediates would cost more; use
    # a fresh user with 0 credits instead.
    r2 = client.post(
        "/auth/sign-in",
        json={"email": "broke@example.com", "displayName": "Broke", "password": "brokepass123"},
    )
    broke_auth = {"Authorization": f"Bearer {r2.json()['token']}"}
    # No grant-starter -> 0 credits.
    bb = client.get("/credits/balance", headers=broke_auth)
    broke_bal = int(bb.json().get("balance", 0))
    r = client.post(
        "/pipeline/remediate",
        files={"file": ("ok.docx", valid_bytes, DOCX_MIME)},
        data={"approved_violations": "[]", "rejected_violations": "[]"},
        headers=broke_auth,
    )
    check("broke user (0 credits) -> 402", r.status_code == 402, f"got {r.status_code} (bal {broke_bal})")
    bb2 = client.get("/credits/balance", headers=broke_auth)
    check("broke user balance untouched", int(bb2.json().get("balance", 0)) == broke_bal)

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
