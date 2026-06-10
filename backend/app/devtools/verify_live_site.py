"""Post-deploy smoke — verify a LIVE 508 Agent deployment serves the core journey.

Unlike the in-process smoke suite (which uses a TestClient), this hits a REAL
running deployment over HTTP. Run it right after `docker compose up` on the host,
or against your public API hostname once the tunnel is live.

    python -m app.devtools.verify_live_site --api-url https://api.yourdomain.com

What it checks (uses a throwaway account; does NOT spend money):
  1. GET /healthz            -> process is up
  2. GET /readyz             -> database is reachable (readiness)
  3. POST /auth/sign-in      -> creates a throwaway account, returns a JWT
  4. POST /pipeline/analyze  -> a generated DOCX with known issues is analysed
                                and the expected violations come back

Billing/checkout is intentionally NOT exercised (it needs a real card). The
script prints the one manual billing check to run by hand at the end.

Exit code is 0 on success, 1 on any failure — safe to wire into a deploy script.
"""

from __future__ import annotations

import argparse
import io
import os
import sys
import time


def _build_fixture_docx() -> bytes:
    """A tiny .docx with deliberate issues: an image with no alt + no title."""
    from docx import Document
    from PIL import Image

    img = io.BytesIO()
    Image.new("RGB", (24, 24), (200, 40, 40)).save(img, format="PNG")

    doc = Document()  # deliberately no core_properties.title
    doc.add_heading("Quarterly Report", level=1)
    doc.add_paragraph("Body text so the document has content to analyse.")
    doc.add_paragraph().add_run().add_picture(io.BytesIO(img.getvalue()))
    # Typed fake list -> the engine's signature LIST_STRUCTURE_INVALID check.
    doc.add_paragraph("- first typed item in the fake list")
    doc.add_paragraph("- second typed item in the fake list")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def main() -> int:
    parser = argparse.ArgumentParser(description="Smoke-test a live 508 Agent deployment.")
    parser.add_argument(
        "--api-url",
        default=os.environ.get("API_URL", "http://127.0.0.1:8000"),
        help="Base URL of the backend API (e.g. https://api.yourdomain.com).",
    )
    parser.add_argument("--timeout", type=float, default=30.0)
    args = parser.parse_args()
    base = args.api_url.rstrip("/")

    import httpx

    DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    failures = 0

    def check(name: str, cond: bool, detail: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, (f"  ({detail})" if detail and not cond else ""))
        if not cond:
            failures += 1

    print(f"Verifying live deployment at {base}\n")

    with httpx.Client(base_url=base, timeout=args.timeout, follow_redirects=True) as c:
        # 1. liveness
        try:
            r = c.get("/healthz")
            check("GET /healthz -> 200", r.status_code == 200, f"got {r.status_code}")
        except Exception as exc:  # noqa: BLE001
            check("GET /healthz reachable", False, str(exc))
            print("\nCannot reach the API at all — is the URL right and the stack up?")
            return 1

        # 2. readiness (DB)
        r = c.get("/readyz")
        ready = r.status_code == 200 and (r.json() or {}).get("ready") is True
        check("GET /readyz -> ready (database reachable)", ready, f"got {r.status_code} {r.text[:120]}")

        # 3. throwaway sign-in
        email = f"deploy-smoke+{int(time.time())}@example.com"
        r = c.post("/auth/sign-in", json={"email": email, "displayName": "Deploy Smoke", "password": "deploysmoke123"})
        token = (r.json() or {}).get("token") if r.status_code == 200 else None
        check("POST /auth/sign-in -> token", bool(token), f"got {r.status_code} {r.text[:160]}")
        if not token:
            print("\nAuth failed — check APP_SECRET is set and the DB migrated.")
            return 1
        headers = {"Authorization": f"Bearer {token}"}

        # 4. analyze a generated document
        try:
            fixture = _build_fixture_docx()
        except Exception as exc:  # noqa: BLE001
            check("build local fixture (needs python-docx + Pillow)", False, str(exc))
            return 1
        r = c.post("/pipeline/analyze", files={"file": ("report.docx", fixture, DOCX_MIME)}, headers=headers)
        ok = r.status_code == 200
        check("POST /pipeline/analyze -> 200", ok, f"got {r.status_code} {r.text[:200]}")
        if ok:
            rule_ids = [v.get("ruleId") for v in (r.json() or {}).get("violations", [])]
            check(
                "analyze surfaces the expected issues (missing alt / no title)",
                "MISSING_ALT_TEXT" in rule_ids or "DOCUMENT_TITLE_MISSING" in rule_ids,
                f"got {rule_ids}",
            )
            check(
                "analyze detects the typed fake list (signature check)",
                "LIST_STRUCTURE_INVALID" in rule_ids,
                f"got {rule_ids}",
            )

    print("\n" + ("ALL CHECKS PASSED" if failures == 0 else f"{failures} CHECK(S) FAILED"))
    print(
        "\nManual billing check (needs a card): open the app, buy a Starter pack with the\n"
        "Stripe test card 4242 4242 4242 4242, and confirm your credit balance updates\n"
        "and the Stripe webhook recorded the grant."
    )
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
