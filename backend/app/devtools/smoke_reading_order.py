"""Smoke: PPTX reading-order detection (WCAG 1.3.2).

Screen readers follow the slide's shape-tree order, not visual position. A
slide whose substantial text shapes are stacked vertically but appear in
REVERSE tree order reads bottom-then-top — flagged READING_ORDER_AMBIGUOUS
on the slide. Precision guards pinned here:

  - correct top-to-bottom tree order -> no flag
  - side-by-side columns (vertical overlap) in any tree order -> no flag
    (column order is a legitimate authoring choice)
  - tiny labels (<12 chars, e.g. page numbers) never participate
  - title + single body -> no flag

Usage:
    python -m app.devtools.smoke_reading_order
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_ro_')}/s.db")

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_engine import RemediationEngine  # noqa: E402

FLAG = "READING_ORDER_AMBIGUOUS"


def _count(path: Path) -> int:
    res = parse_to_tree(str(path))
    run_analyzers(res.tree)
    return sum(1 for v in RemediationEngine().detect_violations(res.tree) if v.rule_id == FLAG)


def _box(slide, text, top_in, left_in=1.0, w=6.0, h=1.0):
    tb = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(w), Inches(h))
    tb.text_frame.text = text
    return tb


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="ro_"))

    # 1. Inverted stack: bottom box FIRST in the tree -> flagged once.
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _box(s, "This paragraph belongs at the bottom of the slide.", top_in=4.5)
    _box(s, "This paragraph belongs at the top of the slide.", top_in=1.0)
    p = tmp / "inverted.pptx"; prs.save(str(p))
    check("inverted stack flagged exactly once", _count(p) == 1, f"got {_count(p)}")

    # 2. Same boxes in correct visual order -> no flag.
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _box(s, "This paragraph belongs at the top of the slide.", top_in=1.0)
    _box(s, "This paragraph belongs at the bottom of the slide.", top_in=4.5)
    p = tmp / "correct.pptx"; prs.save(str(p))
    check("correct order not flagged", _count(p) == 0)

    # 3. Side-by-side columns, right column first in the tree -> no flag
    #    (vertical overlap means column order is an authoring choice).
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _box(s, "Right-hand column body text for the layout.", top_in=1.5, left_in=5.0, w=3.5)
    _box(s, "Left-hand column body text for the layout.", top_in=1.5, left_in=0.8, w=3.5)
    p = tmp / "columns.pptx"; prs.save(str(p))
    check("side-by-side columns not flagged", _count(p) == 0)

    # 4. Tiny label (a page number) below, added first -> ignored, no flag.
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _box(s, "7", top_in=6.5, w=0.5, h=0.4)  # page number, <12 chars
    _box(s, "Main body content paragraph for this slide.", top_in=1.0)
    p = tmp / "pagenum.pptx"; prs.save(str(p))
    check("tiny label below never flags", _count(p) == 0)

    # 5. Titled slide + one body box in normal order -> no flag.
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Quarterly Update"
    _box(s, "Body paragraph below the title, as expected.", top_in=3.0)
    p = tmp / "titled.pptx"; prs.save(str(p))
    check("titled slide in normal order not flagged", _count(p) == 0)

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
