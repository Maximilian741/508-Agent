"""Smoke: with no AI key, "Fix everything" writes only what it can stand behind.

The launch state has no AI key. Before this fix the offline heuristic wrote
nonsense into customers' files and charged for it (heuristic-quality scout,
reproduced end to end through /pipeline/remediate):

  * ``<a href="#section-2">Read more about #section-2</a>``, "Read more about
    click here", "Read more about hr@example.com", "Read more about
    +18005551212" — link text that says less than "click here" did;
  * ``alt="Image html-img-1 — Home About Contact Login"`` — an internal node
    id plus the nav bar that happened to sit above the picture; a byline
    ("Posted on … by admin") REPLACED the author's alt "DSC_0042.jpg";
  * ``<html lang="fr">`` on a Spanish page (French and Spanish shared "de"/"la"
    and the tie went to whichever came first in a dict);
  * a Caption paragraph "Table: 2023, 410, 12%" built from a DATA row, and a
    visible "Column 1 | Column 2" header row inserted into a label/value table;
  * ``<dc:title>Doc 20240912 Wa0003</dc:title>`` from a WhatsApp export name;
  * a scanned PDF reporting one MISSING_ALT_TEXT per page on top of
    SCANNED_DOCUMENT_NO_TEXT;
  * the free /tools/alt-text answering "Uploaded image shown in image." as
    alt text to copy.

Every one of those is now REFUSED: nothing written, nothing counted, nothing
charged, and a plain reason (no settings or provider names) is left for the
customer. The good cases still go through and are still charged: a real
``<figcaption>Figure 2: Revenue by region</figcaption>`` becomes the alt
"Revenue by region", "/docs/benefits-guide.pdf" names its link "Benefits
guide (PDF)", a real header row is promoted, and Spanish/Portuguese pages get
``es``/``pt``. A fake AI provider that answers with the same junk is refused by
the same gates, so the rule holds whichever provider answered.

Checks are made on the DOWNLOADED BYTES and on the wallet.

Run: python -m app.devtools.smoke_semantic_refusals
"""

from __future__ import annotations

import io
import json
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path

_TMP = Path(tempfile.mkdtemp(prefix="508_smoke_semantic_"))
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/s.db"
os.environ["STORAGE_LOCAL_ROOT"] = str(_TMP / "storage")
os.environ["MATERIALIZED_ROOT"] = str(_TMP / "materialized")
for _key in ("OPENAI_API_KEY", "ANTHROPIC_API_KEY", "SEMANTIC_PROVIDER"):
    os.environ.pop(_key, None)

from fastapi.testclient import TestClient  # noqa: E402
from pypdf import PdfReader, PdfWriter  # noqa: E402
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject  # noqa: E402
from sqlalchemy import select  # noqa: E402

HTML = "text/html"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF = "application/pdf"

_ENV_WORDS = ("API_KEY", "ANTHROPIC", "OPENAI", "SEMANTIC_PROVIDER", "env var", "environment variable")

def _png() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 8), (40, 90, 160)).save(buf, format="PNG")
    return buf.getvalue()


_PNG = _png()

_ES_PROSE = (
    "<p>Los empleados que no elijan un plan durante el periodo de inscripción mantendrán su "
    "cobertura actual. Si tiene preguntas, comuníquese con Recursos Humanos.</p>"
    "<p>El Ayuntamiento de la ciudad informa a los vecinos de que las obras en la calle Mayor "
    "comenzarán el lunes y se prolongarán durante tres semanas. Se ruega disculpen las molestias.</p>"
    "<p>La biblioteca estará cerrada el jueves por mantenimiento. Los libros prestados se pueden "
    "devolver en el buzón que está junto a la puerta principal, y no habrá multas durante esos días.</p>"
)

# Everything the old heuristic got wrong, next to the cases it may still fix.
MIXED_HTML = (
    "<!doctype html><html><head><title>Guía de beneficios para empleados</title></head><body>"
    "<h1>Guía de beneficios para empleados</h1>"
    "<p>Home About Contact Login</p>"
    '<img src="IMG_2041.png">'
    "<p>Posted on September 12, 2026 by admin</p>"
    '<img src="photo.jpg" alt="DSC_0042.jpg">'
    '<figure><img src="chart.png"><figcaption>Figure 2: Revenue by region</figcaption></figure>'
    + _ES_PROSE
    + '<p>Guía completa: <a href="/docs/benefits-guide.pdf">click here</a>.</p>'
    '<p>Informe anual: <a href="https://example.com/reports/annual-report-2025.pdf">read more</a>.</p>'
    '<p>Sección dos: <a href="#section-2">click here</a>.</p>'
    '<p>Correo: <a href="mailto:hr@example.com">click here</a>.</p>'
    '<p>Teléfono: <a href="tel:+18005551212">click here</a>.</p>'
    '<p>Directo: <a href="https://example.com/reports/q3.pdf">https://example.com/reports/q3.pdf</a>.</p>'
    '<p>Portal: <a href="https://example.com/">read more</a>.</p>'
    '<table id="t-data"><tr><td>2023</td><td>410</td><td>12%</td></tr>'
    "<tr><td>2024</td><td>455</td><td>11%</td></tr><tr><td>2025</td><td>470</td><td>9%</td></tr></table>"
    '<table id="t-good"><tr><td>Región</td><td>Q1</td><td>Q2</td></tr>'
    "<tr><td>Norte</td><td>10</td><td>20</td></tr><tr><td>Sur</td><td>8</td><td>9</td></tr></table>"
    # A roster: every row is words, so its first row is as likely to be data.
    '<table id="t-words"><tr><td>Alicia</td><td>Ingeniería</td><td>Denver</td></tr>'
    "<tr><td>Bruno</td><td>Ventas</td><td>Austin</td></tr><tr><td>Carla</td><td>Legal</td><td>Boston</td></tr></table>"
    "</body></html>"
)

# Nothing here can be fixed honestly without a person: every approved fix must
# be refused, and the run must be free and hand back the upload unchanged.
ALL_BAD_HTML = (
    "<!doctype html><html><head></head><body>"
    "<h1>1</h1>"
    "<p>The report is attached. Le rapport est joint. Please review by Friday. Veuillez le lire avant vendredi.</p>"
    "<p>Home About Contact Login</p>"
    '<img src="IMG_2041.png">'
    '<p>Jump: <a href="#section-2">click here</a>.</p>'
    '<p>Mail: <a href="mailto:hr@example.com">click here</a>.</p>'
    "<table><tr><td>2023</td><td>410</td><td>12%</td></tr><tr><td>2024</td><td>455</td><td>11%</td></tr></table>"
    "<table><tr><td>Emergency contact</td><td></td></tr><tr><td>Phone</td><td></td></tr></table>"
    "</body></html>"
)


def _docx_bytes() -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    def hyperlink(paragraph, text: str, *, url: str | None = None, anchor: str | None = None) -> None:
        h = OxmlElement("w:hyperlink")
        if url:
            rid = paragraph.part.relate_to(
                url,
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            h.set(qn("r:id"), rid)
        if anchor:
            h.set(qn("w:anchor"), anchor)
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        h.append(run)
        paragraph._p.append(h)  # noqa: SLF001

    img = _TMP / "px.png"
    img.write_bytes(_PNG)
    doc = Document()
    doc.core_properties.title = ""
    doc.add_paragraph(
        "This handbook explains how employees enroll in benefits, which plans are available, "
        "and who to contact with questions about coverage or eligibility."
    )
    doc.add_paragraph("Home About Contact Login")
    doc.add_picture(str(img), width=Inches(1))  # nav-preceded: must NOT get alt
    doc.add_paragraph("Manager email: Click or tap here to enter text.")
    doc.add_picture(str(img), width=Inches(1))  # form-label-preceded: must NOT get alt
    doc.add_picture(str(img), width=Inches(1))  # captioned: the good case
    doc.add_paragraph("Figure 1: Enrollment steps for new employees", style="Caption")
    p = doc.add_paragraph("The full benefits guide: ")
    hyperlink(p, "click here", url="https://example.com/docs/benefits-guide.pdf")
    p = doc.add_paragraph("Section two: ")
    hyperlink(p, "click here", anchor="_Toc12345")
    t = doc.add_table(rows=3, cols=2)  # label/value form: no header row to invent
    for r, label in enumerate(("Emergency contact", "Phone", "Relationship")):
        t.cell(r, 0).text = label
    t = doc.add_table(rows=3, cols=3)  # first row is DATA
    for r, row in enumerate((("2023", "410", "12%"), ("2024", "455", "11%"), ("2025", "470", "9%"))):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    t = doc.add_table(rows=3, cols=3)  # a real (unmarked) header row: the good case
    for r, row in enumerate((("Region", "Q1", "Q2"), ("North", "10", "20"), ("South", "8", "9"))):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _text_pdf(paragraphs) -> bytes:
    """One page of Helvetica/WinAnsi text (so accented Latin-1 extracts right)."""
    w = PdfWriter()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
        NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
    })
    fonts = DictionaryObject({NameObject("/F1"): w._add_object(font)})  # noqa: SLF001
    page = w.add_blank_page(width=612, height=792)
    ops = []
    y = 740
    for para in paragraphs:
        words = para.split()
        line = ""
        for word in words + [None]:
            if word is not None and len(line) + len(word) < 80:
                line = f"{line} {word}".strip()
                continue
            safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(b"BT /F1 11 Tf 72 %d Td (%s) Tj ET" % (y, safe.encode("cp1252")))
            y -= 15
            line = word or ""
        y -= 10
    cs = DecodedStreamObject()
    cs.set_data(b"\n".join(ops))
    page[NameObject("/Contents")] = w._add_object(cs)  # noqa: SLF001
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): fonts})
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def _scanned_pdf(n_pages: int = 3) -> bytes:
    """Every page is just a stretched image: no text operators at all."""
    w = PdfWriter()
    for _ in range(n_pages):
        xo = DecodedStreamObject()
        xo.set_data(bytes([128, 128, 128, 128]))
        xo[NameObject("/Type")] = NameObject("/XObject")
        xo[NameObject("/Subtype")] = NameObject("/Image")
        xo[NameObject("/Width")] = NumberObject(2)
        xo[NameObject("/Height")] = NumberObject(2)
        xo[NameObject("/BitsPerComponent")] = NumberObject(8)
        xo[NameObject("/ColorSpace")] = NameObject("/DeviceGray")
        ref = w._add_object(xo)  # noqa: SLF001
        page = w.add_blank_page(width=300, height=400)
        cs = DecodedStreamObject()
        cs.set_data(b"q 300 0 0 400 0 0 cm /Im0 Do Q")
        page[NameObject("/Contents")] = w._add_object(cs)  # noqa: SLF001
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): ref}),
        })
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


_ES_PDF = [
    "El Ayuntamiento de la ciudad informa a los vecinos de que las obras en la calle Mayor comenzarán el "
    "lunes y se prolongarán durante tres semanas. Se ruega disculpen las molestias.",
    "Los empleados que no elijan un plan durante el periodo de inscripción mantendrán su cobertura actual. "
    "Si tiene preguntas, comuníquese con Recursos Humanos.",
    "La biblioteca estará cerrada el jueves por mantenimiento. Los libros prestados se pueden devolver en el "
    "buzón que está junto a la puerta principal.",
]
_PT_PDF = [
    "A Câmara Municipal informa os moradores de que as obras na rua principal começarão na segunda-feira e "
    "durarão três semanas. Pedimos desculpa pelo incómodo causado a todos os cidadãos.",
    "Os funcionários que não escolherem um plano durante o período de inscrição manterão a sua cobertura "
    "atual. Em caso de dúvidas, entre em contato com os Recursos Humanos.",
    "A biblioteca estará fechada na quinta-feira para manutenção. Os livros emprestados podem ser devolvidos "
    "na caixa ao lado da entrada principal.",
]
_MIXED_PDF = [
    "The quarterly report is attached for your review. Le rapport trimestriel est joint pour votre examen.",
    "Please send comments to the finance team by Friday. Veuillez envoyer vos commentaires avant vendredi.",
]


def main() -> int:  # noqa: PLR0915
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.api.credits import DOC_FORMAT_COSTS
    from app.db.models import UserRow
    from app.db.session_sqlalchemy import session_scope
    from app.main import app

    client = TestClient(app, raise_server_exceptions=False)
    email = "semantic-refusals@example.com"
    r = client.post("/auth/sign-in", json={"email": email, "displayName": "S", "password": "semanticpass1"})
    assert r.status_code == 200, r.text
    headers = {"Authorization": f"Bearer {r.json()['token']}"}

    def set_balance(n: int) -> None:
        with session_scope() as s:
            s.execute(select(UserRow).where(UserRow.email == email)).scalars().first().credits_balance = n

    def balance() -> int:
        return int(client.get("/credits/balance", headers=headers).json()["balance"])

    def analyze(name: str, data: bytes, mime: str) -> dict:
        rr = client.post("/pipeline/analyze", files={"file": (name, data, mime)}, headers=headers)
        assert rr.status_code == 200, rr.text
        return rr.json()

    def fix(name: str, data: bytes, mime: str, rules=None):
        """Approve every finding (what "Fix everything" does) or only ``rules``."""
        found = analyze(name, data, mime)
        ids = [v["id"] for v in found["violations"] if rules is None or v["ruleId"] in rules]
        b0 = balance()
        rr = client.post(
            "/pipeline/remediate",
            files={"file": (name, data, mime)},
            data={"approved_violations": json.dumps(ids), "rejected_violations": "[]"},
            headers=headers,
        )
        assert rr.status_code == 200, rr.text
        body = rr.json()
        out = client.get(body["downloadUrl"], headers=headers)
        assert out.status_code == 200, out.text[:200]
        return found, body, out.content, b0 - balance()

    def by_action(body: dict, code: str, status: str) -> list:
        return [e for e in body["executions"] if e["actionCode"] == code and e["status"] == status]

    def notes_are_plain(body: dict, label: str) -> None:
        skipped = [e for e in body["executions"] if e["status"] == "skipped"]
        leaks = [e["notes"] for e in skipped if any(w.lower() in (e["notes"] or "").lower() for w in _ENV_WORDS)]
        check(f"{label}: no refusal note names a setting or provider key", not leaks, str(leaks[:2]))

    set_balance(500)

    # ---- 1. HTML: bad cases refused, good cases written and charged ---------
    found, body, out_b, debit = fix("benefits.html", MIXED_HTML.encode("utf-8"), HTML)
    out = out_b.decode("utf-8")
    check("html: the analysis reports the offline provider (sanity)", found.get("aiProvider") == "heuristic", str(found.get("aiProvider")))

    check("html: Spanish page tagged es (not fr)", re.search(r'<html[^>]*\blang="es"', out) is not None, out[:120])
    check("html: figcaption 'Figure 2: Revenue by region' became the alt, label stripped",
          'alt="Revenue by region"' in out)
    check("html: the nav-preceded image got NO alt", re.search(r'<img src="IMG_2041\.png"\s*/?>', out) is not None,
          re.findall(r"<img[^>]*IMG_2041[^>]*>", out)[:1])
    check("html: the author's alt was not replaced by a byline", 'alt="DSC_0042.jpg"' in out)
    check("html: no internal node id in any alt", "html-img-" not in out)
    check("html: no nav / byline text written as alt",
          "alt=\"Home" not in out and "Posted on" not in re.sub(r"<p>Posted on[^<]*</p>", "", out))
    check("html: '/docs/benefits-guide.pdf' link named from its address",
          '<a href="/docs/benefits-guide.pdf">Benefits guide (PDF)</a>' in out)
    check("html: annual report link named from its address",
          ">Annual report 2025 (PDF)</a>" in out)
    check("html: no 'Read more about' anywhere", "Read more about" not in out)
    for href in ("#section-2", "mailto:hr@example.com", "tel:+18005551212"):
        check(f"html: {href} link left as the author wrote it", f'<a href="{href}">click here</a>' in out)
    check("html: bare-URL link with a code-like slug left alone",
          '<a href="https://example.com/reports/q3.pdf">https://example.com/reports/q3.pdf</a>' in out)
    check("html: home-page link not renamed to 'Visit example.com'",
          '<a href="https://example.com/">read more</a>' in out)
    data_tbl = re.search(r'<table id="t-data">.*?</table>', out, re.S).group(0)
    good_tbl = re.search(r'<table id="t-good">.*?</table>', out, re.S).group(0)
    words_tbl = re.search(r'<table id="t-words">.*?</table>', out, re.S).group(0)
    check("html: numeric first row NOT promoted to headers", "<th" not in data_tbl, data_tbl[:120])
    check("html: a roster's first row of names NOT promoted to headers", "<th" not in words_tbl, words_tbl[:120])
    header_refusals = by_action(body, "ADD_TABLE_HEADERS", "skipped")
    check("html: both the numeric table and the roster were REFUSED (not merely unflagged)",
          len(header_refusals) == 2 and any("every row of this table is words" in (e["notes"] or "") for e in header_refusals),
          str([e["notes"] for e in header_refusals]))
    check("html: real header row promoted", '<th scope="col">Región</th>' in good_tbl, good_tbl[:160])
    check("html: no placeholder 'Column 1' header row anywhere", "Column 1" not in out)
    check("html: no caption synthesized from headers or values", "<caption" not in out)

    ok_actions = sorted((e["actionCode"], e["targetNodeId"]) for e in body["executions"] if e["status"] == "success")
    check("html: exactly five fixes succeeded (alt, 2 links, headers, language)",
          [a for a, _ in ok_actions] == sorted(
              ["GENERATE_ALT_TEXT", "IMPROVE_LINK_TEXT", "IMPROVE_LINK_TEXT", "ADD_TABLE_HEADERS", "SET_DOCUMENT_LANGUAGE"]),
          str(ok_actions))
    check("html: persistedFixes counts exactly those five", body.get("persistedFixes") == 5, str(body.get("persistedFixes")))
    check("html: charged the HTML price once", body.get("charged") is True and debit == DOC_FORMAT_COSTS["html"],
          f"charged={body.get('charged')} debit={debit}")
    refused = by_action(body, "IMPROVE_LINK_TEXT", "skipped") + by_action(body, "GENERATE_ALT_TEXT", "skipped")
    check("html: every refused link/alt says it was left for a person and not charged",
          refused and all("not charged" in (e["notes"] or "") for e in refused), str([e["notes"] for e in refused][:2]))
    cap_refusals = by_action(body, "GENERATE_TABLE_CAPTION", "skipped")
    check("html: table captions refused with a reason", len(cap_refusals) >= 1 and all(
        "caption" in (e["notes"] or "").lower() for e in cap_refusals), str([e["notes"] for e in cap_refusals]))
    check("html: no refusal note carries a provider chip ('via …')",
          not any(" via " in (e["notes"] or "") for e in body["executions"] if e["status"] == "skipped"))
    notes_are_plain(body, "html")

    # ---- 2. HTML with nothing fixable: free, and the upload comes back ------
    raw = ALL_BAD_HTML.encode("utf-8")
    found, body, out_b, debit = fix("IMG_2041.html", raw, HTML)
    rules = {v["ruleId"] for v in found["violations"]}
    check("all-bad html: the findings are there to refuse (sanity)",
          {"MISSING_ALT_TEXT", "LINK_TEXT_NON_DESCRIPTIVE", "TABLE_MISSING_HEADERS", "DOCUMENT_TITLE_MISSING",
           "DOCUMENT_LANGUAGE_MISSING"} <= rules, str(sorted(rules)))
    check("all-bad html: no fix succeeded", not [e for e in body["executions"] if e["status"] == "success"],
          str([(e["actionCode"], e["notes"]) for e in body["executions"] if e["status"] == "success"]))
    check("all-bad html: persistedFixes 0, not charged, wallet untouched",
          body.get("persistedFixes") == 0 and body.get("charged") is False and debit == 0,
          f"{body.get('persistedFixes')} {body.get('charged')} {debit}")
    check("all-bad html: the delivered bytes are the upload", out_b == raw)
    title_note = (by_action(body, "SET_DOCUMENT_TITLE", "skipped") or [{"notes": ""}])[0]["notes"]
    check("all-bad html: title refusal explains itself ('1' / camera file name)",
          "'1'" in title_note and "IMG_2041" in title_note, title_note)
    lang_note = (by_action(body, "SET_DOCUMENT_LANGUAGE", "skipped") or [{"notes": ""}])[0]["notes"]
    check("all-bad html: a half-English, half-French page gets no language guess",
          "didn't guess" in lang_note, lang_note)
    notes_are_plain(body, "all-bad html")

    # ---- 3. DOCX -------------------------------------------------------------
    docx_raw = _docx_bytes()
    found, body, out_b, debit = fix("DOC-20240912-WA0003.docx", docx_raw, DOCX)
    with zipfile.ZipFile(io.BytesIO(out_b)) as z:
        docxml = z.read("word/document.xml").decode("utf-8")
        core = z.read("docProps/core.xml").decode("utf-8")
    descrs = re.findall(r'<wp:docPr[^>]*?descr="([^"]*)"', docxml)
    check("docx: the Word Caption became the captioned picture's alt (label stripped)",
          "Enrollment steps for new employees" in descrs, str(descrs))
    check("docx: exactly one picture got alt text (nav- and form-label-preceded ones refused)",
          len([d for d in descrs if d.strip()]) == 1, str(descrs))
    check("docx: no node id / nav / form-placeholder text in any alt",
          not any(("docx-img" in d or "Home About" in d or "Click or tap" in d) for d in descrs), str(descrs))
    links = ["".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", m)) for m in
             re.findall(r"<w:hyperlink[^>]*>(.*?)</w:hyperlink>", docxml, re.S)]
    check("docx: external link named from its address", "Benefits guide (PDF)" in links, str(links))
    check("docx: bookmark link (_Toc12345) left as written", links.count("click here") == 1, str(links))
    check("docx: no 'Read more about' in the document", "Read more about" not in docxml)
    check("docx: no placeholder 'Column 1' header row inserted", "Column 1" not in docxml)
    check("docx: no 'Table: …' caption paragraph synthesized", "Table:" not in docxml)
    tbls = re.findall(r"<w:tbl>.*?</w:tbl>", docxml, re.S)
    check("docx: three tables survive (sanity)", len(tbls) == 3, str(len(tbls)))
    if len(tbls) == 3:
        check("docx: label/value table got no header row", "tblHeader" not in tbls[0])
        check("docx: numeric first row NOT marked as a header", "tblHeader" not in tbls[1])
        check("docx: real header row marked as a repeating header", "tblHeader" in tbls[2])
    check("docx: no title made from the WhatsApp export name",
          not re.search(r"<dc:title>[^<]*(Doc|20240912|Wa0003)", core, re.I), core[:300])
    langs = set(re.findall(r'<dc:language>([^<]*)</dc:language>', core))
    check("docx: any language written is English (the text is English)", langs <= {"en"}, str(langs))
    ok = sorted(e["actionCode"] for e in body["executions"] if e["status"] == "success")
    allowed = {"GENERATE_ALT_TEXT", "IMPROVE_LINK_TEXT", "ADD_TABLE_HEADERS", "SET_DOCUMENT_LANGUAGE"}
    check("docx: only the good cases succeeded", set(ok) <= allowed and ok.count("GENERATE_ALT_TEXT") == 1
          and ok.count("IMPROVE_LINK_TEXT") == 1 and ok.count("ADD_TABLE_HEADERS") == 1, str(ok))
    check("docx: charged once, for persisted fixes",
          body.get("charged") is True and debit == DOC_FORMAT_COSTS["docx"] and body.get("persistedFixes") == len(ok),
          f"charged={body.get('charged')} debit={debit} pf={body.get('persistedFixes')} ok={ok}")
    notes_are_plain(body, "docx")

    # ---- 4. PDF language: Spanish -> es, Portuguese -> pt, mixed -> nothing --
    for name, paras, want in (("aviso.pdf", _ES_PDF, "es"), ("aviso-pt.pdf", _PT_PDF, "pt")):
        pdf = _text_pdf(paras)
        found, body, out_b, debit = fix(name, pdf, PDF, rules={"DOCUMENT_LANGUAGE_MISSING"})
        lang = str(PdfReader(io.BytesIO(out_b)).trailer["/Root"].get("/Lang") or "")
        check(f"pdf {want}: /Lang written as {want!r}", lang == want, repr(lang))
        check(f"pdf {want}: charged the PDF price for it",
              body.get("charged") is True and debit == DOC_FORMAT_COSTS["pdf"], f"{body.get('charged')} {debit}")
    mixed = _text_pdf(_MIXED_PDF)
    found, body, out_b, debit = fix("bilingual.pdf", mixed, PDF, rules={"DOCUMENT_LANGUAGE_MISSING"})
    check("pdf mixed en/fr: the language finding exists (sanity)",
          any(v["ruleId"] == "DOCUMENT_LANGUAGE_MISSING" for v in found["violations"]))
    check("pdf mixed en/fr: no /Lang guessed, nothing charged, upload returned",
          "/Lang" not in PdfReader(io.BytesIO(out_b)).trailer["/Root"] and debit == 0 and out_b == mixed,
          f"debit={debit}")

    # ---- 5. scanned PDF: one finding for one root cause ---------------------
    found = analyze("scan.pdf", _scanned_pdf(3), PDF)
    rules = [v["ruleId"] for v in found["violations"]]
    check("scanned pdf: SCANNED_DOCUMENT_NO_TEXT raised", "SCANNED_DOCUMENT_NO_TEXT" in rules, str(rules))
    check("scanned pdf: page scans do NOT each raise MISSING_ALT_TEXT", "MISSING_ALT_TEXT" not in rules, str(rules))

    # ---- 6. free /tools/alt-text: no placeholder to copy --------------------
    rr = client.post("/tools/alt-text", files={"file": ("photo.png", _PNG, "image/png")}, headers=headers)
    tool = rr.json() if rr.status_code == 200 else {}
    check("tools/alt-text: answers 200", rr.status_code == 200, rr.text[:200])
    check("tools/alt-text: no alt text handed back without vision AI", tool.get("altText") == "", str(tool))
    check("tools/alt-text: says AI is not configured", tool.get("aiConfigured") is False, str(tool))
    msg = str(tool.get("message") or "")
    check("tools/alt-text: explains in customer words (no setting names)",
          bool(msg) and not any(w.lower() in msg.lower() for w in _ENV_WORDS), msg)
    # The page can say so BEFORE anyone signs in and uploads (no auth needed).
    av = client.get("/tools/alt-text/availability")
    avj = av.json() if av.status_code == 200 else {}
    check("tools/alt-text/availability: says unavailable, without sign-in",
          av.status_code == 200 and avj.get("available") is False and bool(avj.get("message")), av.text[:200])
    check("tools/alt-text/availability: no provider or setting names",
          not any(w.lower() in av.text.lower() for w in _ENV_WORDS + ("heuristic", "claude", "openai")), av.text)

    # ---- 7. the same gates hold for an AI provider that answers with junk ---
    import app.ai.semantic_inference as si

    class _JunkProvider(si.HeuristicProvider):
        name = "fake-junk"

        def alt_text(self, payload):
            return si.InferenceResult(text=f"Image {payload.get('node_id')} — Home About Contact Login", confidence=0.9, provider=self.name)

        def link_text(self, payload):
            return si.InferenceResult(text=f"Read more about {payload.get('text')}", confidence=0.9, provider=self.name)

        def table_caption(self, payload):
            return si.InferenceResult(text="Data table", confidence=0.9, provider=self.name)

        def document_language(self, payload):
            return si.InferenceResult(text="French (probably)", confidence=0.9, provider=self.name)

    class _GoodProvider(si.HeuristicProvider):
        name = "fake-good"

        def alt_text(self, payload):
            return si.InferenceResult(text="Bar chart of quarterly revenue by region", confidence=0.9, provider=self.name)

        def table_caption(self, payload):
            return si.InferenceResult(text="Quarterly permit volume by district", confidence=0.9, provider=self.name)

    real_build = si.build_default_provider
    try:
        si.build_default_provider = lambda *a, **k: _JunkProvider()
        # Same page, minus the one real header row (a junk provider has no say
        # in ADD_TABLE_HEADERS, which is rule-based).
        # The captioned figure is dropped too: its alt comes from the author's
        # caption without asking any provider, so it is not the junk's to spoil.
        raw = MIXED_HTML.replace(
            '<tr><td>Región</td><td>Q1</td><td>Q2</td></tr>', '<tr><td>2023</td><td>1</td><td>2</td></tr>'
        ).replace(
            '<figure><img src="chart.png"><figcaption>Figure 2: Revenue by region</figcaption></figure>', ''
        ).encode("utf-8")
        found, body, out_b, debit = fix("junk.html", raw, HTML)
        out = out_b.decode("utf-8")
        check("junk AI: nothing it said was written, nothing charged",
              not [e for e in body["executions"] if e["status"] == "success"] and debit == 0 and out_b == raw,
              str([(e["actionCode"], e["notes"]) for e in body["executions"] if e["status"] == "success"]))
        check("junk AI: refusal notes are plain", all("not charged" in (e["notes"] or "") for e in
              by_action(body, "IMPROVE_LINK_TEXT", "skipped") + by_action(body, "GENERATE_TABLE_CAPTION", "skipped")))

        si.build_default_provider = lambda *a, **k: _GoodProvider()
        found, body, out_b, debit = fix("good.html", MIXED_HTML.encode("utf-8"), HTML)
        out = out_b.decode("utf-8")
        check("good AI: a grounded caption is added to the real table",
              "<caption>Quarterly permit volume by district</caption>" in out)
        check("good AI: an image with no caption and no pixels is still refused (nothing to look at)",
              re.search(r'<img src="IMG_2041\.png"\s*/?>', out) is not None)
        rr = client.post("/tools/alt-text", files={"file": ("photo.png", _PNG, "image/png")}, headers=headers)
        check("good AI: tools/alt-text returns the model's description",
              rr.status_code == 200 and rr.json().get("altText") == "Bar chart of quarterly revenue by region"
              and rr.json().get("aiConfigured") is True, rr.text[:200])
        av = client.get("/tools/alt-text/availability")
        check("good AI: tools/alt-text/availability says available",
              av.status_code == 200 and av.json().get("available") is True and av.json().get("message") is None,
              av.text[:200])
    finally:
        si.build_default_provider = real_build

    print()
    print(f"{'OK' if failures == 0 else 'FAILED'}: smoke_semantic_refusals ({failures} failure(s))")
    return 0 if failures == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
