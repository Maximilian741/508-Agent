"""Smoke: a plain team member's remediation never charges the owner's card.

stripe_billing.ensure_balance_for / _overage_eligible take an ``actor_id`` and
let only the wallet owner or a team admin trigger an off-session overage
charge. /pipeline omitted it at both call sites (the affordability pre-check
and the charge), so a plain member's remediate — or a download of their
deferred-charge job — was treated as the owner acting. Stripe HTTP is stubbed
and every PaymentIntent that would be sent is counted. Pins:

  - member remediate on an empty owner wallet (overage on) -> 402, zero intents
  - member's deferred-charge job: signed download and batch-zip -> 402, zero intents
  - control: a team admin's remediate -> 200 via exactly one overage pack

Usage:
    python -m app.devtools.smoke_overage_pipeline_actor
"""

from __future__ import annotations

import hashlib
import hmac
import io
import json
import os
import sys
import tempfile
import time
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="508_smoke_ovactor_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/oa.db"
os.environ["MATERIALIZED_ROOT"] = str(Path(_TMP) / "materialized")
os.environ["STRIPE_SECRET_KEY"] = "sk_test_dummy"
os.environ["STRIPE_WEBHOOK_SECRET"] = "whsec_test_ovactor"
os.environ.pop("OVERAGE_TEST_MODE", None)
os.environ.pop("OVERAGE_MAX_PACKS_PER_DAY", None)

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from starlette.requests import Request  # noqa: E402

from app.api import stripe_billing  # noqa: E402
from app.main import app  # noqa: E402

_WEBHOOK_SECRET = "whsec_test_ovactor"
_INTENTS: list = []
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _fake_stripe_post(path: str, form: dict, secret: str) -> dict:
    _INTENTS.append((path, dict(form)))
    return {"id": f"pi_actor_{len(_INTENTS)}", "status": "succeeded"}


stripe_billing._stripe_post = _fake_stripe_post


def _signin(client: TestClient, email: str):
    r = client.post("/auth/sign-in", json={"email": email, "displayName": "A", "password": "actorpass123"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}, r.json()["user"]["id"]


def _docx() -> bytes:
    d = Document()
    d.add_heading("Quarterly results", level=1)
    d.add_paragraph("Revenue grew in every region this quarter, led by the west.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def main() -> int:
    client = TestClient(app, raise_server_exceptions=False)
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.db.models import UserRow
    from app.db.session_sqlalchemy import session_scope

    def bal(uid: str) -> int:
        with session_scope() as s:
            return int(s.get(UserRow, uid).credits_balance or 0)

    # Owner subscribes through a real signed webhook (overage defaults on).
    owner_auth, owner_id = _signin(client, "actor-owner@example.com")
    body = json.dumps({"type": "checkout.session.completed", "data": {"object": {
        "id": "cs_actor", "mode": "subscription", "subscription": "sub_actor", "customer": "cus_actor",
        "payment_status": "paid", "client_reference_id": owner_id,
        "metadata": {"plan": "team", "user_id": owner_id}}}}).encode("utf-8")
    ts = int(time.time())
    sig = hmac.new(_WEBHOOK_SECRET.encode("utf-8"), f"{ts}.".encode("utf-8") + body, hashlib.sha256).hexdigest()
    client.post("/billing/webhook", content=body, headers={"Stripe-Signature": f"t={ts},v1={sig}"})
    check("owner funded by team checkout (1000)", bal(owner_id) == 1000, f"bal={bal(owner_id)}")

    client.post("/teams", headers=owner_auth, json={"name": "Actors"})
    member_auth, _ = _signin(client, "actor-member@example.com")
    admin_auth, _ = _signin(client, "actor-admin@example.com")
    for email, auth, role in (("actor-member@example.com", member_auth, "member"),
                              ("actor-admin@example.com", admin_auth, "admin")):
        inv = client.post("/teams/invite", headers=owner_auth, json={"email": email, "role": role}).json()
        r = client.post("/teams/accept", headers=auth, json={"token": inv["acceptUrl"].split("token=")[-1]})
        check(f"{role} joined the team", r.status_code == 200)

    dx = _docx()
    a = client.post("/pipeline/analyze", files={"file": ("Quarterly_Budget_Review.docx", dx, DOCX)}, headers=member_auth)
    title_id = {v["ruleId"]: v["id"] for v in a.json()["violations"]}.get("DOCUMENT_TITLE_MISSING")
    check("fixture has a fixable DOCUMENT_TITLE_MISSING", bool(title_id))
    form = {"approved_violations": json.dumps([title_id]), "rejected_violations": "[]"}

    # The member's job whose charge is deferred (client left) while the wallet is funded.
    real_is_disconnected = Request.is_disconnected

    async def _gone(self):  # noqa: ANN001
        return True

    Request.is_disconnected = _gone
    try:
        r = client.post("/pipeline/remediate", files={"file": ("Annual_Program_Report.docx", dx, DOCX)},
                        data=form, headers=member_auth)
    finally:
        Request.is_disconnected = real_is_disconnected
    pending = r.json()
    check("member's job is chargePending", r.status_code == 200 and pending.get("chargePending") is True,
          f"{r.status_code} {str(pending)[:160]}")

    r = client.post("/credits/spend", headers=owner_auth, json={"amount": 1000, "description": "drain"})
    check("owner drains the shared wallet", r.status_code == 200 and bal(owner_id) == 0)

    # 1. The pre-check and the charge both act as the member.
    r = client.post("/pipeline/remediate", files={"file": ("Regional_Safety_Plan.docx", dx, DOCX)},
                    data=form, headers=member_auth)
    check("member remediate on an empty wallet -> 402", r.status_code == 402, f"got {r.status_code}")
    check("...and no PaymentIntent on the owner's card", len(_INTENTS) == 0, str(_INTENTS))

    # 2. The deferred charge acts as the job's owner — here, the member.
    r = client.get(pending["downloadUrl"])
    check("member's deferred-charge download on an empty wallet -> 402", r.status_code == 402, f"got {r.status_code}")
    r = client.post("/pipeline/batch-zip", headers=member_auth,
                    json={"jobs": [{"jobId": pending["jobId"], "filename": pending["filename"]}]})
    check("member's deferred-charge batch-zip on an empty wallet -> 402", r.status_code == 402, f"got {r.status_code}")
    check("...still no PaymentIntent", len(_INTENTS) == 0, str(_INTENTS))
    check("wallet untouched", bal(owner_id) == 0)

    # 3. Control: an admin may top up the owner's wallet.
    r = client.post("/pipeline/remediate", files={"file": ("Board_Meeting_Minutes.docx", dx, DOCX)},
                    data=form, headers=admin_auth)
    check("admin remediate on an empty wallet -> 200, charged", r.status_code == 200 and r.json().get("charged") is True,
          f"{r.status_code} {r.text[:160]}")
    check("admin triggered exactly one PaymentIntent", len(_INTENTS) == 1, str(len(_INTENTS)))
    check("owner wallet = 200 pack - 3", bal(owner_id) == 197, f"bal={bal(owner_id)}")

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
