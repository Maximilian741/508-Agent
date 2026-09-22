"""Smoke: DOCX alt text is written per IMAGE INSTANCE, never per shared rId.

Word dedupes identical image bytes to ONE relationship id, so a logo pasted
four times is four <w:drawing> instances sharing rId10 — each with its own
<wp:docPr> and its own alt text. The parser mints four ImageNodes. The
writer keyed docPr by rId and wrote each node's alt to EVERY docPr sharing
it, so all four instances ended up with the LAST node's text — including two
a human author had already described — and all four were reported as fixed.

Pinned here on exactly that document (one PNG x4: two with human alt, two
without), driving the real engine and writer:
  * the two human-written descriptions survive byte-for-byte
  * only the two blank instances receive generated alt
  * the writer reports at most 2 alt applications (one per changed node),
    each naming its own instance, none claiming "x4"
  * an unchanged node produces no write at all

Usage:
    python -m app.devtools.smoke_docx_shared_image_alt
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import zipfile

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_dsi_')}/s.db")
os.environ.pop("OPENAI_API_KEY", None)
os.environ.pop("ANTHROPIC_API_KEY", None)
os.environ["SEMANTIC_PROVIDER"] = "heuristic"

from pathlib import Path  # noqa: E402

import docx  # noqa: E402
from docx.shared import Inches  # noqa: E402
from lxml import etree  # noqa: E402

from app.models.accessibility import ImageNode, iter_reading_order  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_engine import RemediationEngine  # noqa: E402
from app.services.remediation_planner import RemediationPolicy  # noqa: E402
from app.writers.docx_writer import write_remediated_docx  # noqa: E402

_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _tiny_png() -> bytes:
    # 1x1 white PNG
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde"
        b"\x00\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02\xfe\xa7V\xbd\xfa\x00\x00\x00\x00IEND\xaeB`\x82"
    )


def _docx_with_shared_logo(path: Path) -> None:
    png = path.parent / "logo.png"
    png.write_bytes(_tiny_png())
    d = docx.Document()
    d.add_paragraph("Company report with a repeated logo, and a caption for context.")
    for i in range(4):
        d.add_paragraph(f"Section {i} text near the logo")
        d.add_picture(str(png), width=Inches(0.5))
        # A Word Caption under each picture: text written FOR that instance,
        # which is the only thing the offline alt path may turn into alt.
        d.add_paragraph(f"Figure {i + 1}: Company logo in section {i}", style="Caption")
    d.save(str(path))
    # Give instances 0 and 2 a HUMAN-written alt; leave 1 and 3 blank.
    _set_descr(path, {0: "Company logo variant 0", 2: "Company logo variant 2"})


def _docprs(path: Path):
    """[(rid, descr)] for every drawing instance in document order."""
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    out = []
    for wrapper in root.iter(f"{_WP}inline", f"{_WP}anchor"):
        blip = wrapper.find(f".//{_A}blip")
        rid = blip.get(f"{_R}embed") if blip is not None else None
        dp = wrapper.find(f".//{_WP}docPr")
        out.append((rid, (dp.get("descr") if dp is not None else None)))
    return out


def _set_descr(path: Path, by_index: dict) -> None:
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    root = etree.fromstring(items["word/document.xml"])
    for i, wrapper in enumerate(root.iter(f"{_WP}inline", f"{_WP}anchor")):
        if i in by_index:
            dp = wrapper.find(f".//{_WP}docPr")
            dp.set("descr", by_index[i])
    items["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in items.items():
            z.writestr(n, b)
    path.write_bytes(buf.getvalue())


def main() -> int:
    failures = 0

    def check(name, cond, extra=""):
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="508_dsi_"))
    src = tmp / "logo_x4.docx"
    out = tmp / "logo_x4-fixed.docx"
    _docx_with_shared_logo(src)

    before = _docprs(src)
    check("fixture: 4 image instances", len(before) == 4, str(before))
    check("fixture: all 4 share ONE rId (Word dedupes identical bytes)", len({r for r, _ in before}) == 1, str(before))
    check("fixture: instances 0 and 2 carry human-written alt; 1 and 3 are blank",
          before[0][1] == "Company logo variant 0" and before[2][1] == "Company logo variant 2"
          and not before[1][1] and not before[3][1], str(before))

    res = parse_to_tree(str(src))
    imgs = [n for n in iter_reading_order(res.tree.root) if isinstance(n, ImageNode)]
    check("parser mints one node per INSTANCE (4)", len(imgs) == 4, str(len(imgs)))

    # Real engine, production apply policy: generate alt where missing.
    eng = RemediationEngine(policy=RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False))
    eng.detect_violations(res.tree)
    eng.execute(res.tree)
    rep = write_remediated_docx(src, res.tree, out)

    after = _docprs(out)
    check("human alt on instance 0 survives byte-for-byte", after[0][1] == "Company logo variant 0", repr(after[0][1]))
    check("human alt on instance 2 survives byte-for-byte", after[2][1] == "Company logo variant 2", repr(after[2][1]))
    check("blank instance 1 received generated alt", bool(after[1][1]), repr(after[1][1]))
    check("blank instance 3 received generated alt", bool(after[3][1]), repr(after[3][1]))
    check("the generated alts did NOT overwrite the human ones (no single value on all four)",
          len({d for _, d in after}) >= 3, str(after))

    alt_apps = [a for a in rep.get("applied", []) if a.get("kind") == "image_alt_text"]
    check("writer reports at most 2 alt writes (one per CHANGED node)", len(alt_apps) <= 2, str(alt_apps))
    check("no summary claims to have written multiple instances at once ('x4')",
          all("x4" not in str(a.get("summary")) and "x2" not in str(a.get("summary")) for a in alt_apps), str(alt_apps))
    changed_ids = {a["target_id"] for a in alt_apps}
    untouched = [n.id for n in imgs if (n.alt_text or "").startswith("Company logo variant")]
    check("unchanged (human-described) nodes produced NO write", not (set(untouched) & changed_ids),
          f"untouched={untouched} written={sorted(changed_ids)}")

    reopened = docx.Document(str(out))
    check("output re-opens with the same paragraph count", len(reopened.paragraphs) == len(docx.Document(str(src)).paragraphs))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
