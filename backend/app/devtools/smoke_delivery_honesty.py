"""Smoke: we charge for bytes, report bytes, and never hand out either for free.

Six adversarial-review findings that all reduce to the same product rule —
*charged if and only if the fix persisted into the output bytes*, in both
directions, and *report only what the file actually carries*:

  * A failed FINAL SAVE charged full price. pdf_writer and pptx_writer caught
    the save exception, appended a ``failed_to_save_*`` skip and returned their
    in-memory ``applied`` list; the pipeline's hard-failure guard matched only
    ``failed_to_open``/``copy_failed`` and was additionally gated on
    ``not applied``, which can never help because ``applied`` records intent.
    Result: 5 credits for a ZERO-BYTE PDF (``open(..., "wb")`` had already
    truncated the copy) and 4 for a byte-identical PPTX. Now: writers save to a
    temp file and ``os.replace`` it into place, report ``applied=[]`` on
    failure, and /remediate 422s with no charge — what docx_writer always did.
  * ADD_OCR_TEXT_LAYER was charged on OCR *availability*. The executor returns
    SUCCESS as soon as a provider exists; a scan the engine can't read (blank,
    handwritten, low-DPI, no language pack) recognised nothing, the writer
    added no overlay, and the user paid 5 credits for the same image-only PDF.
    Now writer-confirmed: no overlay, no count, no charge, and the execution is
    reported SKIPPED instead of "success".
  * /pipeline/remediate reported ``status="success"`` with a before/after
    rewrite for fixes that don't persist for the format, while shipping the
    upload unchanged. The honesty suffix existed only on the free /analyze
    path. Now one reconciliation pass covers every action on both paths.
  * A free ``/pipeline/analyze?execute=true`` minted a paid, verifiable
    certificate claiming "N issue(s) were automatically remediated" — the
    executors ran against a tree that was thrown away, no file was ever
    written and no credit spent. /analyze now persists the document AS FOUND;
    only /remediate writes a fix count, and only the persisted one.
  * POST /pipeline/analyze-url raised NameError on EVERY request (500) —
    ``fingerprint_violations`` was never imported. The public free-scan funnel
    was 100% broken and no smoke exercised the route over HTTP.
  * batch-zip silently dropped a job it couldn't charge whenever another file
    in the batch was payable: 200, a short ZIP, nothing anywhere saying so.

Usage:
    python -m app.devtools.smoke_delivery_honesty
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import zipfile
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="508_smoke_delivery_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/s.db"
os.environ["MATERIALIZED_ROOT"] = str(Path(_TMP) / "materialized")

from PIL import Image  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pypdf import PdfReader, PdfWriter  # noqa: E402
from pypdf.generic import (  # noqa: E402
    ArrayObject,
    DecodedStreamObject,
    DictionaryObject,
    FloatObject,
    NameObject,
    NumberObject,
    StreamObject,
    TextStringObject,
)
from sqlalchemy import select  # noqa: E402

PDF = "application/pdf"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
HTML = "text/html"


def _font(w: PdfWriter) -> DictionaryObject:
    f = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    return DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): w._add_object(f)})})  # noqa: SLF001


def _text_pdf() -> bytes:
    """A plain page missing /Title and /Lang — two fixes the PDF writer persists."""
    w = PdfWriter()
    res = _font(w)
    page = w.add_blank_page(width=320, height=320)
    cs = DecodedStreamObject()
    cs.set_data(
        b"BT /F1 20 Tf 20 280 Td (Quarterly Board Report) Tj ET\n"
        b"BT /F1 11 Tf 20 240 Td (Revenue rose across every region this quarter.) Tj ET\n"
        b"BT /F1 11 Tf 20 220 Td (The board reviewed the accessibility programme.) Tj ET"
    )
    page[NameObject("/Contents")] = w._add_object(cs)  # noqa: SLF001
    page[NameObject("/Resources")] = res
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def _link_pdf() -> bytes:
    """Bare-URL link annotations: IMPROVE_LINK_TEXT succeeds in the tree, and
    the PDF writer cannot persist it — the exact shape that was reported as a
    success while the delivered bytes were the upload."""
    w = PdfWriter()
    res = _font(w)
    page = w.add_blank_page(width=400, height=400)
    cs = DecodedStreamObject()
    cs.set_data(
        b"BT /F1 18 Tf 20 360 Td (Regional Revenue) Tj ET\n"
        b"BT /F1 11 Tf 20 320 Td (Read the regional revenue report for details.) Tj ET\n"
        b"BT /F1 11 Tf 20 280 Td (https://example.com/reports/region-0.pdf) Tj ET"
    )
    page[NameObject("/Contents")] = w._add_object(cs)  # noqa: SLF001
    page[NameObject("/Resources")] = res
    annots = ArrayObject()
    for i, rect in enumerate([(20, 295, 120, 310), (20, 275, 320, 292)]):
        a = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Annot"),
                NameObject("/Subtype"): NameObject("/Link"),
                NameObject("/Rect"): ArrayObject([FloatObject(v) for v in rect]),
                NameObject("/A"): DictionaryObject(
                    {
                        NameObject("/S"): NameObject("/URI"),
                        NameObject("/URI"): TextStringObject(f"https://example.com/reports/region-{i}.pdf"),
                    }
                ),
            }
        )
        annots.append(w._add_object(a))  # noqa: SLF001
    page[NameObject("/Annots")] = annots
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def _scanned_pdf() -> bytes:
    """One page whose only content is a full-page JPEG (a classic scan)."""
    img = Image.new("RGB", (1000, 1294), (245, 242, 235))
    jpeg = io.BytesIO()
    img.save(jpeg, format="JPEG")
    w = PdfWriter()
    page = w.add_blank_page(width=612, height=792)
    xobj = StreamObject()
    xobj._data = jpeg.getvalue()  # noqa: SLF001 — raw DCT bytes, no re-encode
    xobj.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(1000),
            NameObject("/Height"): NumberObject(1294),
            NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
            NameObject("/BitsPerComponent"): NumberObject(8),
            NameObject("/Filter"): NameObject("/DCTDecode"),
        }
    )
    ref = w._add_object(xobj)  # noqa: SLF001
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): ref})}
    )
    cs = DecodedStreamObject()
    cs.set_data(b"q 612 0 0 792 0 0 cm /Im0 Do Q")
    page[NameObject("/Contents")] = w._add_object(cs)  # noqa: SLF001
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    tb.text_frame.text = "Regional revenue overview for the year"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class BlindOcr:
    """available() True, recognize() None — what TesseractOcrProvider actually
    returns for a blank, handwritten, low-DPI or unsupported-script page."""

    name = "tesseract"

    def available(self) -> bool:
        return True

    def recognize(self, image_bytes: bytes):
        return None


class ReadableOcr:
    """The control: an engine that DOES read the page, so the writer-confirmed
    gate is shown to pass a real fix through rather than refusing everything."""

    name = "stub"

    def available(self) -> bool:
        return True

    def recognize(self, image_bytes: bytes):
        from app.services.ocr import OcrPageResult, OcrWord

        words = [OcrWord("ANNUAL", 80, 90, 220, 48), OcrWord("REPORT", 320, 90, 230, 48)]
        x = 80
        for token in "Operations summary for the reporting period scanned original".split():
            words.append(OcrWord(token, x, 220, 16 * len(token), 28))
            x += 16 * len(token) + 12
        return OcrPageResult(width_px=1000, height_px=1294, words=words)


def main() -> int:  # noqa: PLR0915
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.db.models import UserRow
    from app.db.session_sqlalchemy import session_scope
    from app.main import app
    from app.services.ocr import set_ocr_provider_for_testing

    client = TestClient(app, raise_server_exceptions=False)
    r = client.post(
        "/auth/sign-in",
        json={"email": "delivery-smoke@example.com", "displayName": "D", "password": "deliverypass1"},
    )
    assert r.status_code == 200, r.text
    headers = {"Authorization": f"Bearer {r.json()['token']}"}

    def set_balance(n: int) -> None:
        with session_scope() as s:
            row = (
                s.execute(select(UserRow).where(UserRow.email == "delivery-smoke@example.com"))
                .scalars()
                .first()
            )
            row.credits_balance = n

    def balance() -> int:
        return int(client.get("/credits/balance", headers=headers).json()["balance"])

    def analyze(name: str, data: bytes, mime: str, execute: bool = False) -> dict:
        url = "/pipeline/analyze" + ("?execute=true" if execute else "")
        rr = client.post(url, files={"file": (name, data, mime)}, headers=headers)
        assert rr.status_code == 200, rr.text
        return rr.json()

    def remediate(name: str, data: bytes, mime: str, ids: list):
        return client.post(
            "/pipeline/remediate",
            files={"file": (name, data, mime)},
            data={"approved_violations": json.dumps(ids), "rejected_violations": "[]"},
            headers=headers,
        )

    set_balance(100)

    # ---- 1. a failed final save charges nothing and delivers nothing ------
    pdf_bytes = _text_pdf()
    ids = [v["id"] for v in analyze("board.pdf", pdf_bytes, PDF)["violations"]]
    check("PDF fixture carries fixable findings", len(ids) >= 2, str(ids))

    import pypdf

    _real_write = pypdf.PdfWriter.write

    def _enospc(self, stream, *a, **k):
        raise OSError(28, "No space left on device")

    b0 = balance()
    pypdf.PdfWriter.write = _enospc
    try:
        rr = remediate("board.pdf", pdf_bytes, PDF, ids)
    finally:
        pypdf.PdfWriter.write = _real_write
    check("PDF save fails -> 422, not a 200 with a 0-byte file", rr.status_code == 422, f"got {rr.status_code}")
    check("PDF save fails -> nothing charged", balance() == b0, f"{b0} -> {balance()}")
    check(
        "PDF save fails -> the failure is explained, not blamed on the document",
        "not charged" in (rr.json().get("detail") or ""),
        str(rr.json())[:160],
    )

    pptx_bytes = _pptx()
    pids = [v["id"] for v in analyze("deck.pptx", pptx_bytes, PPTX)["violations"]]
    import pptx.presentation as _pp

    _real_save = _pp.Presentation.save

    def _enospc_save(self, f):
        raise OSError(28, "No space left on device")

    b0 = balance()
    _pp.Presentation.save = _enospc_save
    try:
        rr = remediate("deck.pptx", pptx_bytes, PPTX, pids)
    finally:
        _pp.Presentation.save = _real_save
    check("PPTX save fails -> 422, not the untouched upload at full price", rr.status_code == 422, f"got {rr.status_code}")
    check("PPTX save fails -> nothing charged", balance() == b0, f"{b0} -> {balance()}")

    # The same documents remediate normally once the save works, so the guard
    # is not just refusing everything.
    b0 = balance()
    rr = remediate("board.pdf", pdf_bytes, PDF, ids)
    body = rr.json()
    check(
        "a working save still delivers and charges (pdf = 5)",
        rr.status_code == 200 and body["charged"] is True and b0 - balance() == 5,
        f"{rr.status_code} charged={body.get('charged')} {b0} -> {balance()}",
    )
    dl = client.get(body["downloadUrl"], headers=headers)
    check("...and the delivered file is non-empty and changed", dl.status_code == 200 and 0 < len(dl.content) and dl.content != pdf_bytes,
          f"{dl.status_code} {len(dl.content)} bytes")

    # ---- 2. OCR that recognises nothing is not charged --------------------
    scan = _scanned_pdf()
    set_ocr_provider_for_testing(BlindOcr())
    try:
        vs = analyze("scan.pdf", scan, PDF)["violations"]
        scanned_ids = [v["id"] for v in vs if v["ruleId"] == "SCANNED_DOCUMENT_NO_TEXT"]
        check("scanned fixture raises SCANNED_DOCUMENT_NO_TEXT", len(scanned_ids) == 1, str([v["ruleId"] for v in vs]))
        b0 = balance()
        rr = remediate("scan.pdf", scan, PDF, scanned_ids)
        body = rr.json()
        check("unreadable scan -> 200 (the file still comes back)", rr.status_code == 200, f"{rr.status_code}")
        check("unreadable scan -> persistedFixes 0, not charged",
              body.get("persistedFixes") == 0 and body.get("charged") is False,
              f"persisted={body.get('persistedFixes')} charged={body.get('charged')}")
        check("unreadable scan -> balance untouched", balance() == b0, f"{b0} -> {balance()}")
        ocr_execs = [e for e in body["executions"] if e["actionCode"] == "ADD_OCR_TEXT_LAYER"]
        check("unreadable scan -> the OCR execution is reported SKIPPED, not success",
              len(ocr_execs) == 1 and ocr_execs[0]["status"] == "skipped", str(ocr_execs))
        check("...with a note that says why and that nothing was charged",
              "recognized no text" in (ocr_execs[0]["notes"] or "") and "not charged" in (ocr_execs[0]["notes"] or ""),
              str(ocr_execs[0].get("notes"))[:160])
        dl = client.get(body["downloadUrl"], headers=headers)
        out = Path(_TMP) / "ocr_out.pdf"
        out.write_bytes(dl.content)
        check("unreadable scan -> the output is still image-only (nothing claimed)",
              (PdfReader(str(out)).pages[0].extract_text() or "").strip() == "")

        # Control: an engine that reads the page is still counted and charged,
        # so the writer-confirmed gate discriminates rather than blocking.
        set_ocr_provider_for_testing(ReadableOcr())
        b0 = balance()
        rr = remediate("scan.pdf", scan, PDF, scanned_ids)
        body = rr.json()
        check("readable scan -> persistedFixes 1 and charged (pdf = 5)",
              rr.status_code == 200 and body.get("persistedFixes") == 1
              and body.get("charged") is True and b0 - balance() == 5,
              f"{rr.status_code} persisted={body.get('persistedFixes')} {b0} -> {balance()}")
        dl = client.get(body["downloadUrl"], headers=headers)
        out2 = Path(_TMP) / "ocr_ok.pdf"
        out2.write_bytes(dl.content)
        text = PdfReader(str(out2)).pages[0].extract_text() or ""
        check("...and the recognized text is really in the delivered bytes",
              "ANNUAL" in text and "REPORT" in text, text[:120])
    finally:
        set_ocr_provider_for_testing(None)

    # ---- 3. remediate reports only fixes the file carries -----------------
    links = _link_pdf()
    vs = analyze("links.pdf", links, PDF)["violations"]
    link_ids = [v["id"] for v in vs if v["ruleId"] in ("LINK_TEXT_NON_DESCRIPTIVE", "LINK_NAME_MISSING")]
    check("link fixture raises a link-text finding", len(link_ids) >= 1, str([v["ruleId"] for v in vs]))
    b0 = balance()
    rr = remediate("links.pdf", links, PDF, link_ids)
    body = rr.json()
    check("non-persisting fix -> not charged", body.get("charged") is False and balance() == b0,
          f"charged={body.get('charged')} {b0} -> {balance()}")
    link_execs = [e for e in body["executions"] if e["actionCode"] == "IMPROVE_LINK_TEXT"]
    check("non-persisting fix -> reported SKIPPED, never success",
          bool(link_execs) and all(e["status"] == "skipped" for e in link_execs),
          str([(e["actionCode"], e["status"]) for e in body["executions"]]))
    check("...and every note says it was not applied to this format",
          all("Not applied to the PDF file" in (e["notes"] or "") for e in link_execs),
          str([e.get("notes") for e in link_execs])[:200])
    dl = client.get(body["downloadUrl"], headers=headers)
    check("...and the delivered bytes are the upload, unchanged", dl.content == links)

    # ---- 4. a certificate cannot be minted from a free analyze ------------
    page = (
        b'<!DOCTYPE html><html lang="en"><head></head><body>'
        b"<h1>Annual Accessibility Statement</h1>"
        b"<p>This page describes our commitment to accessible digital services.</p>"
        b"</body></html>"
    )
    b0 = balance()
    jobs0 = len(client.get("/pipeline/jobs", headers=headers).json()["jobs"])
    an = analyze("statement.html", page, HTML, execute=True)
    doc_id = an["summary"]["documentId"]
    check("the free analyze preview still shows what COULD be fixed",
          an["score"]["fixedAutomatically"] >= 1, str(an["score"]))
    jobs1 = len(client.get("/pipeline/jobs", headers=headers).json()["jobs"])
    check("...and costs nothing and writes no artifact",
          balance() == b0 and jobs1 == jobs0, f"bal {b0} -> {balance()}, jobs {jobs0} -> {jobs1}")
    cert = client.post("/billing/issue-certificate", json={"documentId": doc_id}, headers=headers)
    check("issue-certificate after a free analyze -> 200", cert.status_code == 200, cert.text[:160])
    cj = cert.json()
    check("...but it claims ZERO automatic fixes (no bytes were ever written)",
          int(cj.get("fixedCount") or 0) == 0, str(cj.get("fixedCount")))
    check("...and the claim text does not assert remediation happened",
          "0 issue(s) were automatically remediated" in (cj.get("conformanceClaim") or ""),
          (cj.get("conformanceClaim") or "")[:160])

    # A PAID remediation of the same document does put a real number on it.
    set_balance(100)
    hids = [v["id"] for v in analyze("statement.html", page, HTML)["violations"]]
    rr = remediate("statement.html", page, HTML, hids)
    rb = rr.json()
    check("paid remediation of the same doc -> charged with a persisted fix",
          rr.status_code == 200 and rb.get("charged") is True and rb.get("persistedFixes") >= 1,
          f"{rr.status_code} {str(rb)[:160]}")
    cert2 = client.post("/billing/issue-certificate", json={"documentId": rb["documentId"]}, headers=headers)
    check("...and NOW the certificate may claim it",
          cert2.status_code == 200 and int(cert2.json().get("fixedCount") or 0) == int(rb["persistedFixes"]),
          f"{cert2.status_code} {cert2.text[:160]}")

    # ---- 5. the free URL scan answers at all ------------------------------
    import app.api.pipeline as _pipeline

    _real_fetch = _pipeline.fetch_url_html
    _pipeline.fetch_url_html = lambda url, *a, **k: (  # no network is touched
        b"<!DOCTYPE html><html><head><title>Example</title></head><body>"
        b"<h1>Example</h1><img src='a.png'>"
        b"<p>A paragraph with enough words to look like a real page.</p></body></html>",
        url,
    )
    try:
        rr = client.post("/pipeline/analyze-url", json={"url": "https://example.com/"}, headers=headers)
    finally:
        _pipeline.fetch_url_html = _real_fetch
    check("POST /pipeline/analyze-url -> 200 (it used to NameError on every request)",
          rr.status_code == 200, f"{rr.status_code} {rr.text[:200]}")
    if rr.status_code == 200:
        sb = rr.json()
        check("...with findings for the scanned page", len(sb.get("violations") or []) >= 1, str(sb)[:200])

    # ---- 6. a short batch-zip says it is short ---------------------------
    from starlette.requests import Request as _Req

    _real_disc = _Req.is_disconnected

    async def _gone(self):  # noqa: ANN001
        return True

    def pending(name: str) -> dict:
        _Req.is_disconnected = _gone
        try:
            return remediate(name, pdf_bytes, PDF, ids).json()
        finally:
            _Req.is_disconnected = _real_disc

    set_balance(100)
    j1, j2 = pending("Z1.pdf"), pending("Z2.pdf")
    check("two deferred-charge jobs exist", j1.get("chargePending") is True and j2.get("chargePending") is True)
    set_balance(5)  # enough for exactly one 5-credit PDF
    z = client.post(
        "/pipeline/batch-zip",
        json={"jobs": [
            {"jobId": j1["jobId"], "filename": j1["filename"]},
            {"jobId": j2["jobId"], "filename": j2["filename"]},
        ]},
        headers=headers,
    )
    check("batch-zip with credits for one of two -> 200 with a partial archive", z.status_code == 200, str(z.status_code))
    if z.status_code == 200:
        names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
        check("...the headers count the batch honestly",
              z.headers.get("X-Batch-Requested") == "2"
              and z.headers.get("X-Batch-Delivered") == "1"
              and z.headers.get("X-Batch-Withheld") == "1",
              str(dict(z.headers)))
        check("...and name the withheld job", z.headers.get("X-Batch-Withheld-Jobs") == j2["jobId"],
              str(z.headers.get("X-Batch-Withheld-Jobs")))
        check("...and the archive itself carries a manifest", "NOT-INCLUDED.txt" in names, str(names))
        check("...one file delivered, one withheld", len([n for n in names if n.endswith(".pdf")]) == 1, str(names))
    check("...the withheld job stays payable, not lost", balance() == 0)
    set_balance(20)
    dl = client.get(j2["downloadUrl"], headers=headers)
    check("...and after a top-up it downloads and pays exactly once",
          dl.status_code == 200 and balance() == 15, f"{dl.status_code} bal={balance()}")

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except AssertionError as exc:
        print(f"FAIL delivery-honesty smoke: {exc}")
        sys.exit(1)
    except Exception as exc:  # pragma: no cover
        print(f"FAIL delivery-honesty smoke: {exc.__class__.__name__}: {exc}")
        sys.exit(1)
