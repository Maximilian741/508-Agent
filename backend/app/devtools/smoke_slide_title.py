"""Smoke: untitled slides get a REAL title (SLIDE_TITLE_MISSING auto-fix).

A slide with no title is the #1 PowerPoint accessibility failure — screen-reader
users navigate a deck by its slide-title list, and an untitled slide is invisible
there. Previously we only flagged it (detect-only). This pins the new auto-fix:

  parser   -> each untitled slide's SectionNode carries missing_title
  analyzer -> SLIDE_TITLE_MISSING fires once per untitled slide
  executor -> SET_SLIDE_TITLE derives a title from the slide's own text
  writer   -> a real title PLACEHOLDER is inserted (cloned from the layout, or
              built from scratch on a Blank layout) carrying that text
  honesty  -> SET_SLIDE_TITLE persists for pptx; re-parsing the OUTPUT shows a
              titled slide and the flag clears (idempotent), and the title text
              equals the slide's most prominent line.

Usage:
    python -m app.devtools.smoke_slide_title
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_st_')}/s.db")

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.api.pipeline import _action_persists  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.pptx_writer import write_remediated_pptx  # noqa: E402

FLAG = "SLIDE_TITLE_MISSING"


def _flag_count(tree) -> int:
    n = 0

    def walk(node):
        nonlocal n
        n += sum(1 for f in node.accessibility_flags if f.code.value == FLAG)
        for ch in node.children:
            walk(ch)

    walk(tree.root)
    return n


def _slide_titles(path) -> list[str]:
    prs = Presentation(str(path))
    out = []
    for slide in prs.slides:
        try:
            t = slide.shapes.title
            out.append((t.text or "").strip() if t is not None else "")
        except Exception:
            out.append("")
    return out


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="slide_title_"))
    apply_policy = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)

    prs = Presentation()
    # Slide 1 — "Title Only" layout (has an empty title placeholder) + body text.
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    tb1 = s1.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    tb1.text_frame.text = "Quarterly Revenue Overview"
    # Slide 2 — "Blank" layout (NO title placeholder at all) + body text.
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    tb2 = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    tb2.text_frame.text = "Market Trends 2026"
    # Slide 3 — properly titled (control: must NOT be flagged or changed).
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Already Has A Title"
    tb3 = s3.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    tb3.text_frame.text = "Some body content."
    src = tmp / "deck.pptx"
    prs.save(str(src))

    # Sanity on the fixture: slides 1 & 2 have no title text, slide 3 does.
    titles_before = _slide_titles(src)
    check("fixture: slides 1 & 2 untitled, slide 3 titled",
          titles_before[0] == "" and titles_before[1] == "" and titles_before[2] == "Already Has A Title",
          str(titles_before))

    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    check("two slides flagged SLIDE_TITLE_MISSING", _flag_count(res.tree) == 2, f"got {_flag_count(res.tree)}")
    check("SET_SLIDE_TITLE persists for pptx (honesty matrix)", _action_persists("SET_SLIDE_TITLE", "pptx"))

    plans = [p for p in plan_remediations(res.tree, apply_policy) if p.flag.code.value == FLAG]
    check("two SET_SLIDE_TITLE plans", len(plans) == 2, f"got {len(plans)}")
    execs = execute_plans(res.tree, plans)
    ok = [e for e in execs if e.status.value == "success"]
    check("both title-sets execute successfully", len(ok) == 2, str([(e.status.value, e.notes) for e in execs]))

    out = tmp / "deck_fixed.pptx"
    result = write_remediated_pptx(src, res.tree, out)
    titled = [a for a in result["applied"] if a.get("kind") == "slide_title"]
    check("writer applied 2 slide titles", len(titled) == 2, str(result))
    check("no slide-title skips (insertion always succeeded)",
          not any(s.get("reason", "").startswith(("could_not", "slide_not_found")) for s in result["skipped"]),
          str(result["skipped"]))

    titles_after = _slide_titles(out)
    check("slide 1 title derived from its text (Title-Only layout, tier 1)",
          titles_after[0] == "Quarterly Revenue Overview", str(titles_after))
    check("slide 2 title derived from its text (Blank layout, tier 3 from scratch)",
          titles_after[1] == "Market Trends 2026", str(titles_after))
    check("slide 3 title untouched", titles_after[2] == "Already Has A Title", str(titles_after))

    # Self-consistency: re-parse the OUTPUT -> no slide is missing a title now.
    res2 = parse_to_tree(str(out))
    run_analyzers(res2.tree)
    check("re-analysis of output raises NO slide-title flag (idempotent)",
          _flag_count(res2.tree) == 0, f"got {_flag_count(res2.tree)}")

    # Re-remediating the already-fixed deck does nothing (no plans, no changes).
    plans2 = [p for p in plan_remediations(res2.tree, apply_policy) if p.flag.code.value == FLAG]
    check("fixed deck: no further SET_SLIDE_TITLE plans", len(plans2) == 0, f"got {len(plans2)}")

    # Clean deck (all titled): nothing flagged, nothing changed.
    prs2 = Presentation()
    c = prs2.slides.add_slide(prs2.slide_layouts[5])
    c.shapes.title.text = "Clean Slide"
    csrc = tmp / "clean.pptx"
    prs2.save(str(csrc))
    rc = parse_to_tree(str(csrc))
    run_analyzers(rc.tree)
    check("clean deck: no slide-title flag", _flag_count(rc.tree) == 0)

    # =========================================================================
    # Title derivation must use POSITION (topmost), not XML order — a footer
    # added first in the XML must NOT hijack the title (the MEDIUM the
    # adversarial review found). Also: a leading date box is skipped.
    # =========================================================================
    prs3 = Presentation()
    sA = prs3.slides.add_slide(prs3.slide_layouts[6])  # blank
    # Footer added FIRST (so it's first in XML) but positioned at the BOTTOM.
    fa = sA.shapes.add_textbox(Inches(1), Inches(6.8), Inches(8), Inches(0.4))
    fa.text_frame.text = "Confidential - Acme Corp 2026"
    # Real title added SECOND but positioned at the TOP.
    ta = sA.shapes.add_textbox(Inches(1), Inches(0.4), Inches(8), Inches(1))
    ta.text_frame.text = "Strategic Priorities"
    # Second slide: a DATE box at the very top, real heading just below.
    sB = prs3.slides.add_slide(prs3.slide_layouts[6])
    db = sB.shapes.add_textbox(Inches(1), Inches(0.3), Inches(4), Inches(0.4))
    db.text_frame.text = "June 16, 2026"
    hb = sB.shapes.add_textbox(Inches(1), Inches(1.2), Inches(8), Inches(1))
    hb.text_frame.text = "Quarterly Results"
    psrc = tmp / "footer_deck.pptx"
    prs3.save(str(psrc))

    rp = parse_to_tree(str(psrc))
    run_analyzers(rp.tree)
    pplans = [p for p in plan_remediations(rp.tree, apply_policy) if p.flag.code.value == FLAG]
    execute_plans(rp.tree, pplans)
    pout = tmp / "footer_deck_fixed.pptx"
    write_remediated_pptx(psrc, rp.tree, pout)
    ptitles = _slide_titles(pout)
    check("title is the TOPMOST text, not the footer added first in XML",
          ptitles[0] == "Strategic Priorities", str(ptitles))
    check("a leading date box is skipped; the real heading becomes the title",
          ptitles[1] == "Quarterly Results", str(ptitles))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
