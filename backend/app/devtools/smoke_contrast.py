"""Smoke: WCAG colour-contrast analysis (SC 1.4.3) for DOCX and PPTX.

Verifies the contrast math, and that real documents are flagged honestly:
explicitly-coloured low-contrast text is flagged; sufficient-contrast text and
text whose colour/background can't be determined are NOT flagged (no guessing,
no false positives).

Usage:
    python -m app.devtools.smoke_contrast
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_contrast_')}/s.db"

from docx import Document  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor as PRGB  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt as PPt  # noqa: E402

from app.analyzers.contrast import contrast_ratio, is_large_text, required_ratio  # noqa: E402
from app.analyzers.registry import run_analyzers  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_engine import RemediationEngine  # noqa: E402


def _contrast_violations(path: Path) -> int:
    res = parse_to_tree(str(path))
    run_analyzers(res.tree)
    return len([v for v in RemediationEngine().detect_violations(res.tree) if v.rule_id == "LOW_CONTRAST_TEXT"])


def main() -> int:
    failures = 0

    def check(name: str, cond: bool) -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name)
        if not cond:
            failures += 1

    # --- math ---
    check("black on white = 21:1", abs(contrast_ratio("000000", "FFFFFF") - 21.0) < 0.1)
    check("white on white = 1:1", abs(contrast_ratio("FFFFFF", "FFFFFF") - 1.0) < 0.01)
    check("#777 on white fails AA (<4.5)", contrast_ratio("777777", "FFFFFF") < 4.5)
    check("large text threshold is 3.0", required_ratio(18.0) == 3.0 and is_large_text(18.0, False))
    check("normal text threshold is 4.5", required_ratio(12.0) == 4.5)
    check("unparseable colour -> None", contrast_ratio("nope", "FFFFFF") is None)

    tmp = Path(tempfile.mkdtemp())

    # --- DOCX ---
    docx_path = tmp / "c.docx"
    d = Document()
    p1 = d.add_paragraph(); r1 = p1.add_run("Light gray text that fails contrast")
    r1.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA); r1.font.size = Pt(11)
    p2 = d.add_paragraph(); r2 = p2.add_run("Black text that passes")
    r2.font.color.rgb = RGBColor(0, 0, 0); r2.font.size = Pt(11)
    d.add_paragraph("Inherited-colour text (skipped, not guessed)")
    d.save(str(docx_path))
    check("docx: exactly 1 contrast violation (the gray run)", _contrast_violations(docx_path) == 1)

    # --- PPTX ---
    pptx_path = tmp / "c.pptx"
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    b1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(1))
    b1.fill.solid(); b1.fill.fore_color.rgb = PRGB(0x33, 0x33, 0x33)
    b1.text_frame.text = "Hard to read"
    rr1 = b1.text_frame.paragraphs[0].runs[0]; rr1.font.color.rgb = PRGB(0x55, 0x55, 0x55); rr1.font.size = PPt(12)
    b2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(3), Inches(4), Inches(1))
    b2.fill.solid(); b2.fill.fore_color.rgb = PRGB(0x00, 0x33, 0x66)
    b2.text_frame.text = "Easy to read"
    rr2 = b2.text_frame.paragraphs[0].runs[0]; rr2.font.color.rgb = PRGB(0xFF, 0xFF, 0xFF); rr2.font.size = PPt(12)
    tb = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    tb.text_frame.text = "On the slide background"
    tb.text_frame.paragraphs[0].runs[0].font.color.rgb = PRGB(0xCC, 0xCC, 0xCC)
    prs.save(str(pptx_path))
    check("pptx: exactly 1 contrast violation (dark-on-dark box; good + unknown-bg skipped)", _contrast_violations(pptx_path) == 1)

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
