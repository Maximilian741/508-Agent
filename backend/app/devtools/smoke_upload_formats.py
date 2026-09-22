"""Smoke: every accepted type is checked by its bytes, and every refusal is a sentence.

The upload door now takes PDF, Word (.docx .doc .rtf .odt), PowerPoint
(.pptx .ppt .odp), Excel (.xlsx .xls .ods), HTML and images. Pinned here:

  * magic bytes are checked for EVERY accepted type, and a mismatch names
    what the file really is ("named .png but it is really a JPEG image");
  * an OOXML file is checked for the RIGHT package (a workbook renamed
    .docx is caught), a password-protected Office file (an OLE2 container
    named .docx) is recognised as such, a plain zip and a zip bomb are
    refused — all as sentences with no machine codes (``empty_upload``,
    ``file_content_mismatch``, ``invalid_ooxml`` used to reach the UI);
  * unknown and missing extensions are refused listing what we accept, and
    the accepted list the sentence gives covers every accepted suffix;
  * .htm is html end to end. The route used to gate persisted fixes on
    "htm", which is in no persisted-action set, so every real fix made to a
    .htm upload was thrown away and the untouched file handed back;
  * the effective format (what a type is fixed, delivered and priced as) is
    what the money path uses.

Usage:
    python -m app.devtools.smoke_upload_formats
"""

from __future__ import annotations

import io
import json
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="508_smoke_formats_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/s.db"
os.environ["MATERIALIZED_ROOT"] = str(Path(_TMP) / "materialized")
os.environ.setdefault("APP_SECRET", "x" * 64)
os.environ["SEMANTIC_PROVIDER"] = "heuristic"
for _key in ("OPENAI_API_KEY", "ANTHROPIC_API_KEY"):
    os.environ.pop(_key, None)

from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import select  # noqa: E402

_MACHINE_CODE = re.compile(r"^[a-z_]+(:|$)")


def _docx_bytes() -> bytes:
    from docx import Document

    d = Document()
    d.add_paragraph("hello")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    import xlsxwriter

    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    wb.add_worksheet("Data").write("A1", "x")
    wb.close()
    return buf.getvalue()


def main() -> int:  # noqa: PLR0915
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.api.credits import DOC_FORMAT_COSTS
    from app.db.models import UserRow
    from app.db.session_sqlalchemy import session_scope
    from app.intake import ACCEPTED_SENTENCE, ACCEPTED_SUFFIXES, effective_format
    from app.main import app

    client = TestClient(app, raise_server_exceptions=False)
    email = "formats-smoke@example.com"
    r = client.post("/auth/sign-in", json={"email": email, "displayName": "F", "password": "formatsmokepass1"})
    assert r.status_code == 200, r.text
    headers = {"Authorization": f"Bearer {r.json()['token']}"}

    def set_balance(n: int) -> None:
        with session_scope() as s:
            row = s.execute(select(UserRow).where(UserRow.email == email)).scalars().first()
            row.credits_balance = n

    def balance() -> int:
        return int(client.get("/credits/balance", headers=headers).json()["balance"])

    def analyze(name: str, data: bytes):
        return client.post("/pipeline/analyze", files={"file": (name, data, "application/octet-stream")}, headers=headers)

    def refused(label: str, rr, status: int, must: str) -> None:
        detail = ""
        try:
            detail = rr.json().get("detail") or ""
        except Exception:
            pass
        check(
            f"{label} -> {status}, a sentence containing {must!r}",
            rr.status_code == status and must.lower() in detail.lower() and not _MACHINE_CODE.match(detail) and detail.endswith("."),
            f"{rr.status_code} {detail[:200]!r}",
        )

    # ---- 1. the accepted list is complete and priced --------------------
    for suffix in sorted(ACCEPTED_SUFFIXES):
        shown = {".jpeg": ".jpg", ".tif": ".tiff", ".htm": ".html", ".pdf": "PDF"}.get(suffix, suffix)
        check(f"accepted list mentions {suffix}", shown in ACCEPTED_SENTENCE)
        fmt = effective_format(suffix)
        check(f"{suffix} is fixed and priced as {fmt}", fmt in DOC_FORMAT_COSTS, fmt)
    check("effective formats", [effective_format(s) for s in (".htm", ".doc", ".xls", ".ppt", ".png", ".odp")]
          == ["html", "docx", "xlsx", "pptx", "pdf", "pptx"])

    # ---- 2. extension refusals --------------------------------------------
    refused("unsupported .txt", analyze("notes.txt", b"hello"), 400, "We can't check .txt files")
    refused("unsupported .exe", analyze("setup.exe", b"MZ\x90\x00"), 400, "Upload a PDF")
    refused("no extension", analyze("README", b"hello"), 400, "no extension")

    # ---- 3. signature refusals, each naming what the file really is ------
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
    jpg = b"\xff\xd8\xff\xe0" + b"\x00" * 64
    pdf = b"%PDF-1.7\n" + b"\x00" * 64
    refused("empty file", analyze("empty.pdf", b""), 400, "empty")
    refused("PDF named .docx", analyze("really.docx", pdf), 400, "really a PDF")
    refused("JPEG named .png", analyze("photo.png", jpg), 400, "really a JPEG image")
    refused("PNG named .gif", analyze("anim.gif", png), 400, "really a PNG image")
    refused("PDF named .bmp", analyze("x.bmp", pdf), 400, "really a PDF")
    refused("text named .tiff", analyze("x.tiff", b"hello world" * 3), 400, "not a TIFF image")
    refused("PNG named .webp", analyze("x.webp", png), 400, "really a PNG image")
    refused("PDF named .rtf", analyze("x.rtf", pdf), 400, "really a PDF")
    refused("PDF named .doc", analyze("x.doc", pdf), 400, "really a PDF")
    refused("zip named .xls", analyze("x.xls", _docx_bytes()), 400, "zip archive or an Office file")
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504
    refused("password-protected .docx (OLE2 inside)", analyze("locked.docx", ole), 400, "password-protected")
    refused("password-protected .xlsx", analyze("locked.xlsx", ole), 400, "password-protected")
    refused("workbook renamed .docx", analyze("budget.docx", _xlsx_bytes()), 400, "really an Excel workbook")
    refused("Word file renamed .xlsx", analyze("memo.xlsx", _docx_bytes()), 400, "really a Word document")
    plain = io.BytesIO()
    with zipfile.ZipFile(plain, "w") as z:
        z.writestr("hello.txt", "not an office file")
    refused("plain zip named .pptx", analyze("deck.pptx", plain.getvalue()), 400, "This is a zip file")
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        with z.open("xl/zeros.bin", "w", force_zip64=True) as fh:  # streamed: never 520 MB in memory
            chunk = b"\x00" * (1024 * 1024)
            for _ in range(520):
                fh.write(chunk)
    refused("zip bomb named .xlsx", analyze("bomb.xlsx", bomb.getvalue()), 413, "more than 500 MB")
    refused("truncated .xlsx", analyze("cut.xlsx", _xlsx_bytes()[:200]), 400, "damaged")

    # ---- 4. .htm is html, end to end ---------------------------------------
    set_balance(20)
    page = (
        b"<html><head></head><body><h1>Benefits guide</h1>"
        b"<p>Our benefits explained in plain words for every employee.</p></body></html>"
    )
    rep = client.post("/pipeline/analyze", files={"file": ("benefits.htm", page, "text/html")}, headers=headers).json()
    ids = [v["id"] for v in rep["violations"] if v["ruleId"] == "DOCUMENT_TITLE_MISSING"]
    check(".htm analyses as html with a title finding", rep["summary"]["sourceFormat"] == "html" and len(ids) == 1, str(rep["summary"]))
    b0 = balance()
    rr = client.post(
        "/pipeline/remediate",
        files={"file": ("benefits.htm", page, "text/html")},
        data={"approved_violations": json.dumps(ids), "rejected_violations": "[]"},
        headers=headers,
    )
    body = rr.json()
    check(".htm: the approved title fix persists (it used to be discarded)", rr.status_code == 200 and body.get("persistedFixes") == 1,
          f"{rr.status_code} {str(body)[:200]}")
    check(".htm: charged at the html price", body.get("charged") is True and b0 - balance() == DOC_FORMAT_COSTS["html"], f"{b0}->{balance()}")
    out = client.get(body["downloadUrl"], headers=headers).content
    check(".htm: the delivered page has the title", b"<title>Benefits guide</title>" in out, out[:200])
    check(".htm: delivered under its own extension", body.get("filename") == "benefits-remediated.htm", str(body.get("filename")))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
