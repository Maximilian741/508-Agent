"""Smoke: every image on a slide is checked for alt text — and decorative ones are not.

The PPTX parser treated a shape as an image only when python-pptx said
``shape_type == PICTURE``. A picture dropped into a layout's picture/content
PLACEHOLDER (the "insert picture" icon — how most pictures get onto slides) is
a ``p:pic`` whose shape_type is PLACEHOLDER, and charts/SmartArt are graphic
frames: none of them was ever checked, so a deck of unlabeled photos and
charts reported "no missing alt text". And an image the author marked
decorative (PowerPoint's "Mark as decorative", the Office adec extension —
also what our own writer emits) was still flagged for its leftover "image.png"
descr, including on our own output after a mark-decorative fix.

Pinned here, parser and writer in lock-step (the writer re-mints the same ids):
  * a filled picture placeholder is an image: its "image.png" descr raises
    ALT_TEXT_NOT_DESCRIPTIVE and an approved alt lands on THAT shape;
  * a chart is an image: MISSING_ALT_TEXT, with the chart's own title offered
    as the caption; the fix writes descr on the graphic frame;
  * a SmartArt graphic is an image too;
  * a decorative-marked picture raises nothing and the writer leaves it
    byte-for-byte alone;
  * ids stay aligned when these shapes sit BEFORE other ids on the slide
    (a link after them still gets its own text rewritten, not a neighbour's);
  * re-analysis of the output is clean for everything that was fixed.

Usage:
    python -m app.devtools.smoke_pptx_image_kinds
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_pik_')}/s.db")

from lxml import etree  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.models.accessibility import ImageNode, LinkNode, iter_reading_order  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.writers.pptx_writer import _mark_shape_decorative, write_remediated_pptx  # noqa: E402

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_SMARTART = (
    '<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<p:nvGraphicFramePr><p:cNvPr id="90" name="Diagram 1"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
    '<p:xfrm><a:off x="914400" y="4572000"/><a:ext cx="2743200" cy="914400"/></p:xfrm>'
    '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
    '<dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" r:dm="rId90" r:lo="rId91" r:qs="rId92" r:cs="rId93"/>'
    '</a:graphicData></a:graphic></p:graphicFrame>'
)


def _png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 10, 10)).save(buf, "PNG")
    return buf.getvalue()


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="508_pik_"))
    prs = Presentation()
    # Slide 1: "Picture with Caption" layout, picture inserted into its placeholder.
    s1 = prs.slides.add_slide(prs.slide_layouts[8])
    s1.shapes.title.text = "Team photo"
    s1.placeholders[1].insert_picture(io.BytesIO(_png()))
    # Slide 2: chart (with a title) + decorative picture + SmartArt frame + a link after them.
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Permits"
    cd = CategoryChartData()
    cd.categories = ["North", "South"]
    cd.add_series("2026", (120, 98))
    gf = s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(4), Inches(3), cd)
    gf.chart.has_title = True
    gf.chart.chart_title.text_frame.text = "Permits by region"
    deco = s2.shapes.add_picture(io.BytesIO(_png()), Inches(6), Inches(1.5))
    _mark_shape_decorative(deco)
    s2.shapes._spTree.append(etree.fromstring(_SMARTART))  # noqa: SLF001
    tb = s2.shapes.add_textbox(Inches(1), Inches(6), Inches(6), Inches(0.5))
    tb.text_frame.text = "click here"
    tb.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.gov/permits"
    src = tmp / "deck.pptx"
    prs.save(str(src))
    slide2_before = zipfile.ZipFile(src).read("ppt/slides/slide2.xml")

    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    imgs = {n.id: n for n in iter_reading_order(res.tree.root) if isinstance(n, ImageNode)}

    def flags(node_id):
        return [f.code.value for f in imgs[node_id].accessibility_flags] if node_id in imgs else None

    check("picture in a placeholder is an image node", "slide-1-img-1" in imgs, str(sorted(imgs)))
    check("placeholder picture's 'image.png' descr -> ALT_TEXT_NOT_DESCRIPTIVE", flags("slide-1-img-1") == ["ALT_TEXT_NOT_DESCRIPTIVE"], str(flags("slide-1-img-1")))
    chart = imgs.get("slide-2-img-1")
    check("chart is an image node with no alt -> MISSING_ALT_TEXT", chart is not None and flags("slide-2-img-1") == ["MISSING_ALT_TEXT"], str(flags("slide-2-img-1")))
    check("chart's own title is offered as its caption", chart is not None and (chart.metadata.properties or {}).get("caption") == "Permits by region")
    deco_node = imgs.get("slide-2-img-2")
    check("decorative-marked picture is recognised and raises nothing",
          deco_node is not None and deco_node.is_decorative and not deco_node.alt_text and flags("slide-2-img-2") == [], str(flags("slide-2-img-2")))
    smart = imgs.get("slide-2-img-3")
    check("SmartArt graphic is an image node (MISSING_ALT_TEXT)",
          smart is not None and (smart.metadata.properties or {}).get("image_kind") == "smartart" and flags("slide-2-img-3") == ["MISSING_ALT_TEXT"],
          str(flags("slide-2-img-3")))

    # Approve: alt for the placeholder picture + chart + SmartArt; rewrite the link text.
    imgs["slide-1-img-1"].alt_text = "Five team members outside City Hall"
    imgs["slide-2-img-1"].alt_text = "Column chart: North 120 permits, South 98"
    imgs["slide-2-img-3"].alt_text = "Process: apply, review, approve"
    link = next(n for n in iter_reading_order(res.tree.root) if isinstance(n, LinkNode))
    link.content.text = "Permit application portal"
    out = tmp / "deck_fixed.pptx"
    rep = write_remediated_pptx(src, res.tree, out)
    alt_applied = {a["target_id"] for a in rep["applied"] if a.get("kind") == "image_alt_text"}
    check("writer wrote the three approved alts", {"slide-1-img-1", "slide-2-img-1", "slide-2-img-3"} <= alt_applied, str(rep))
    check("writer touched nothing decorative", not any(a.get("kind") == "image_decorative" for a in rep["applied"]), str(rep["applied"]))

    root = etree.fromstring(zipfile.ZipFile(out).read("ppt/slides/slide2.xml"))
    names = {c.get("name"): c.get("descr") for c in root.iter(P + "cNvPr")}
    check("chart graphic frame carries the alt", names.get(gf.name) == "Column chart: North 120 permits, South 98", str(names))
    check("SmartArt frame carries the alt", names.get("Diagram 1") == "Process: apply, review, approve", str(names))
    deco_before = [c for c in etree.fromstring(slide2_before).iter(P + "cNvPr") if c.get("name") == deco.name][0]
    deco_after = [c for c in root.iter(P + "cNvPr") if c.get("name") == deco.name][0]
    check("decorative picture byte-for-byte unchanged", etree.tostring(deco_after) == etree.tostring(deco_before))
    p1 = Presentation(str(out))
    pic = [sh for sh in p1.slides[0].shapes if sh._element.tag == P + "pic"][0]  # noqa: SLF001
    check("placeholder picture carries the alt", pic._element.find(f"{P}nvPicPr/{P}cNvPr").get("descr") == "Five team members outside City Hall")  # noqa: SLF001
    link_texts = ["".join(r.text for r in para.runs) for sh in p1.slides[1].shapes if getattr(sh, "has_text_frame", False) for para in sh.text_frame.paragraphs]
    check("ids stayed aligned: the link after the new image kinds got ITS new text", "Permit application portal" in link_texts, str(link_texts))

    res2 = parse_to_tree(str(out))
    run_analyzers(res2.tree)
    left = [(n.id, [f.code.value for f in n.accessibility_flags]) for n in iter_reading_order(res2.tree.root)
            if isinstance(n, ImageNode) and n.accessibility_flags]
    check("re-analysis: no image finding left", left == [], str(left))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
