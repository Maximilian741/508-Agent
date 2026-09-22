"""Smoke: a promoted fake heading gets the level its LOOK says, never H1-by-default.

PROMOTE_HEADING used to make every fake heading a sibling of the heading
before it. In a manual whose real headings are "Part N" (Heading 1) and whose
500 procedures are typed as 14pt bold "N. Procedure N", that meant 526 Heading
1s and no Heading 2 — a flat outline a screen-reader user cannot navigate by
level. It also assumed the "Heading N" style existed: a real Word file that
never used Heading 2 has no such style in styles.xml, python-docx raised, the
writer skipped it, and the pipeline still counted and charged the promotion.

Pinned here, through parser -> analyzers -> planner -> executor -> writer ->
re-parse of the OUTPUT:

  * level follows the document's own ladder: 20pt title -> H1, 14pt numbered
    sections under it -> H2; "Part N" (Heading 1) with same-size plain-numbered
    procedures under it -> the procedures are H2, the next Part is H1 again
  * never skips a level: no HEADING_LEVEL_JUMP after promotion, ever — and a
    gap the source ALREADY has (H2 -> H4) does not drag a line that looks
    like the H2 underneath it; when the look and the no-skip rule disagree
    (a 20pt line between an H3 and an H4) the promotion is declined
  * a bold, body-size numbered subsection ("2.1 Data Sources") is detected
    and lands one level under its section; a bold list step, a number
    ("2.1 million residents"), a non-bold line and a TOC line are not
  * a heading style missing from styles.xml is CREATED (with its outline
    level) so the promotion lands, and the writer confirms it
  * appearance is pinned: the promoted line keeps its size, weight, colour,
    font and alignment (a Title keeps its 26pt; nothing turns blue)
  * a heading style that auto-numbers does NOT add a number to the text
  * ambiguous signals (bigger but not bold vs smaller but bold) are DECLINED
    with a plain-English reason, and are neither counted nor charged
  * a promotion the writer could not place is not counted (writer-confirmed)

Usage:
    python -m app.devtools.smoke_docx_heading_levels
"""

from __future__ import annotations

import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_dhl_')}/s.db")
os.environ.pop("OPENAI_API_KEY", None)
os.environ.pop("ANTHROPIC_API_KEY", None)
os.environ["SEMANTIC_PROVIDER"] = "heuristic"

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from lxml import etree  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.api.pipeline import _count_persisted_fixes  # noqa: E402
from app.models.accessibility import HeadingNode, ParagraphNode, iter_reading_order  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.docx_writer import write_remediated_docx  # noqa: E402

FAKE_H = "TEXT_STYLED_AS_HEADING"
JUMP = "HEADING_LEVEL_JUMP"
POL = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _fake(doc, text, size, bold=True, color=None, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if align is not None:
        p.alignment = align
    return p


def _body(doc, n=1):
    for _ in range(n):
        doc.add_paragraph(
            "This is ordinary body text that explains the procedure in enough words to be prose."
        )


def _flags(tree, code):
    return sum(
        1 for n in iter_reading_order(tree.root) for f in n.accessibility_flags if f.code.value == code
    )


def _outline(tree):
    return [
        ((n.content.text or "").strip(), n.level)
        for n in iter_reading_order(tree.root)
        if isinstance(n, HeadingNode)
    ]


def _fix_fake_headings(src: Path, out: Path):
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    plans = [p for p in plan_remediations(res.tree, POL) if p.flag.code.value == FAKE_H]
    execs = execute_plans(res.tree, plans)
    wr = write_remediated_docx(src, res.tree, out)
    after = parse_to_tree(str(out))
    run_analyzers(after.tree)
    return res, execs, wr, after


def _drop_styles(path: Path, style_ids) -> None:
    """Remove styles from styles.xml, the way a real Word file that never used
    them looks (Word only writes the styles a document uses)."""
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    root = etree.fromstring(items["word/styles.xml"])
    for st in list(root.iterfind(f"{W}style")):
        if st.get(f"{W}styleId") in style_ids:
            root.remove(st)
    items["word/styles.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in items.items():
            z.writestr(n, b)


def _para_xml(path: Path, text: str):
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    for p in root.iter(f"{W}p"):
        if "".join(t.text or "" for t in p.iter(f"{W}t")).strip() == text:
            return p
    return None


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="508_dhl_"))

    # ===== A. The 500-page-manual shape (scaled down) ========================
    d = Document()
    d.core_properties.title = "Consolidated Procedures Manual"
    d.add_heading("Consolidated Procedures Manual", 1)
    n = 0
    for part in range(1, 4):
        d.add_heading(f"Part {part}: Operations Area {part}", 1)
        for _ in range(4):
            n += 1
            _fake(d, f"{n}. Procedure {n}", 14)
            _body(d, 2)
    src = tmp / "manual.docx"
    d.save(str(src))
    res, execs, wr, after = _fix_fake_headings(src, tmp / "manual_fixed.docx")
    ok = [e for e in execs if e.status.value == "success"]
    check("manual: all 12 procedures promoted", len(ok) == 12, str([(e.status.value, e.notes) for e in execs]))
    outline = _outline(after.tree)
    procs = [lvl for t, lvl in outline if ". Procedure " in t]
    parts = [lvl for t, lvl in outline if t.startswith("Part ")]
    check("manual: procedures under a 'Part' Heading 1 become Heading 2 (not 12 more H1s)",
          procs == [2] * 12, str(outline))
    check("manual: the Part headings are still Heading 1", parts == [1, 1, 1], str(outline))
    check("manual: no heading-level jump after promotion", _flags(after.tree, JUMP) == 0)
    check("manual: no fake-heading flag left", _flags(after.tree, FAKE_H) == 0)
    check("manual: every promotion names its level and the heading it was placed against",
          all("Heading 2" in (e.notes or "") and "'" in (e.notes or "") for e in ok)
          and "Part 1" in (ok[0].notes or ""),
          str([e.notes for e in ok][:2]))

    # ===== B. The everyday report: 20pt title, 14pt numbered sections ========
    d = Document()
    d.core_properties.title = "Community Health Needs Assessment"
    _fake(d, "Community Health Needs Assessment 2026", 20)
    _body(d)
    for t in ("1. Executive Summary", "2. Priority Areas", "3. Recommendations"):
        _fake(d, t, 14)
        _body(d, 2)
    src = tmp / "report.docx"
    d.save(str(src))
    res, execs, wr, after = _fix_fake_headings(src, tmp / "report_fixed.docx")
    outline = dict(_outline(after.tree))
    check("report: 20pt title -> Heading 1", outline.get("Community Health Needs Assessment 2026") == 1, str(outline))
    check("report: 14pt sections under it -> Heading 2 (siblings of each other)",
          [outline.get(t) for t in ("1. Executive Summary", "2. Priority Areas", "3. Recommendations")] == [2, 2, 2],
          str(outline))
    check("report: no heading-level jump", _flags(after.tree, JUMP) == 0)

    # ===== B2. Bold body-size numbered subsections ("2.1 Data Sources") =======
    # The 14pt rule never saw a subsection typed in bold at text size, so the
    # report's sections were promoted and its subsections stayed plain text.
    # Precision: a bold one-level "1. Submit the form" (a list step), a
    # lower-case "2.1 million residents", a non-bold "3.2 Scope" and a TOC
    # line are NOT headings.
    d = Document()
    d.core_properties.title = "Subsections"
    _fake(d, "Annual Data Report", 20)
    _body(d)
    _fake(d, "2. Methods", 14)
    _body(d)
    _fake(d, "2.1 Data Sources", 12)
    _body(d)
    _fake(d, "2.2 Survey Design", 12)
    _body(d)
    _fake(d, "1. Submit the form", 12)          # a bold list step, not a heading
    _fake(d, "2.1 million residents", 12)       # a number, not a section
    _fake(d, "3.2 Scope", 12, bold=False)       # not bold
    toc = d.add_paragraph()
    toc.add_run("4.1 Retention").bold = True
    toc.style = d.styles.add_style("TOC 2", 1)  # a table-of-contents line
    _body(d)
    src = tmp / "subsections.docx"
    d.save(str(src))
    res, execs, wr, after = _fix_fake_headings(src, tmp / "subsections_fixed.docx")
    flagged = sorted(
        (n.content.text or "") for n in iter_reading_order(res.tree.root)
        for f in n.accessibility_flags if f.code.value == FAKE_H
    )
    check("subsections: the two bold numbered subsections are flagged; list step / number / "
          "plain / TOC lines are not",
          flagged == sorted(["Annual Data Report", "2. Methods", "2.1 Data Sources", "2.2 Survey Design"]),
          str(flagged))
    outline = dict(_outline(after.tree))
    check("subsections: title H1, section H2, subsections H3 under it",
          [outline.get(t) for t in ("Annual Data Report", "2. Methods", "2.1 Data Sources", "2.2 Survey Design")]
          == [1, 2, 3, 3], str(outline))
    check("subsections: no heading-level jump", _flags(after.tree, JUMP) == 0)

    # ===== C. Heading 2 is NOT defined in styles.xml (a real Word file) =======
    d = Document()
    d.core_properties.title = "Missing styles"
    d.styles["Heading 1"].font.size = Pt(18)   # the H1 is visibly bigger than the fake
    d.add_heading("Overview", 1)
    _body(d)
    _fake(d, "Scope of this policy", 14, color="1F1F1F")
    _body(d)
    src = tmp / "nostyle.docx"
    d.save(str(src))
    _drop_styles(src, {f"Heading{i}" for i in range(2, 10)})
    with zipfile.ZipFile(src) as z:
        check("fixture: Heading 2 really is absent from styles.xml",
              b'w:styleId="Heading2"' not in z.read("word/styles.xml"))
    res, execs, wr, after = _fix_fake_headings(src, tmp / "nostyle_fixed.docx")
    out = tmp / "nostyle_fixed.docx"
    check("missing style: promotion executes", [e.status.value for e in execs] == ["success"],
          str([(e.status.value, e.notes) for e in execs]))
    outline = dict(_outline(after.tree))
    check("missing style: re-parse of the OUTPUT sees a real Heading 2",
          outline.get("Scope of this policy") == 2, str(outline))
    with zipfile.ZipFile(out) as z:
        styles = etree.fromstring(z.read("word/styles.xml"))
    h2 = [s for s in styles.iterfind(f"{W}style")
          if (s.find(f"{W}name") is not None and s.find(f"{W}name").get(f"{W}val") == "heading 2")]
    check("missing style: a 'heading 2' style now exists in styles.xml", len(h2) == 1)
    lvl = h2[0].find(f"{W}pPr/{W}outlineLvl") if h2 else None
    check("missing style: the created style carries outline level 1 (Word's Heading 2)",
          lvl is not None and lvl.get(f"{W}val") == "1")
    promos = [a for a in wr["applied"] if a.get("action") == "PROMOTE_HEADING"]
    check("missing style: writer confirms the promotion (action entry)", len(promos) == 1, str(wr))
    check("missing style: counted exactly once", _count_persisted_fixes(execs, wr["applied"], "docx") == 1)
    p_el = _para_xml(out, "Scope of this policy")
    rpr = p_el.find(f"{W}r/{W}rPr") if p_el is not None else None
    color = rpr.find(f"{W}color") if rpr is not None else None
    check("missing style: the run keeps its own colour", color is not None and color.get(f"{W}val") == "1F1F1F")

    # ===== D. Appearance is pinned (Title style -> Heading 1) =================
    d = Document()
    d.core_properties.title = "Pinned look"
    t = d.add_paragraph("Strategic Plan 2027")
    t.style = d.styles["Title"]
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _body(d)
    d.add_heading("Goals", 1)
    _body(d)
    src = tmp / "title.docx"
    d.save(str(src))
    res, execs, wr, after = _fix_fake_headings(src, tmp / "title_fixed.docx")
    out = tmp / "title_fixed.docx"
    check("title: promoted to Heading 1", dict(_outline(after.tree)).get("Strategic Plan 2027") == 1,
          str(_outline(after.tree)))
    od = Document(str(out))
    tp = next(p for p in od.paragraphs if p.text == "Strategic Plan 2027")
    run = tp.runs[0]
    check("title: style is now Heading 1", tp.style.name == "Heading 1", tp.style.name)
    check("title: keeps its 26pt size (Title's), not Heading 1's 14pt",
          run.font.size is not None and abs(run.font.size.pt - 26) < 0.01, str(run.font.size))
    check("title: keeps its alignment", tp.alignment == WD_ALIGN_PARAGRAPH.CENTER, str(tp.alignment))
    check("title: not made bold by the heading style", run.font.bold is False, str(run.font.bold))
    rpr = run._r.find(f"{W}rPr")
    col = rpr.find(f"{W}color") if rpr is not None else None
    check("title: keeps Title's colour (text2), does not turn Heading-1 blue",
          col is not None and col.get(f"{W}themeColor") == "text2", etree.tostring(rpr).decode() if rpr is not None else "")

    # ===== E. A numbered heading style must not add a number to the text ======
    d = Document()
    d.core_properties.title = "Numbered styles"
    d.add_heading("Introduction", 1)
    _body(d)
    _fake(d, "Background", 12)
    _body(d)
    src = tmp / "numbered.docx"
    d.save(str(src))
    d = Document(str(src))
    d.styles["Heading 1"].font.size = Pt(16)
    h2 = d.styles["Heading 2"]._element
    ppr = h2.find(qn("w:pPr"))
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "1"); numpr.append(ilvl)
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), "1"); numpr.append(nid)
    ppr.insert(0, numpr)
    for p in d.paragraphs:
        if p.text == "Background":
            p.runs[0].font.size = Pt(14)
    d.save(str(src))
    res, execs, wr, after = _fix_fake_headings(src, tmp / "numbered_fixed.docx")
    p_el = _para_xml(tmp / "numbered_fixed.docx", "Background")
    ppr = p_el.find(f"{W}pPr") if p_el is not None else None
    num = ppr.find(f"{W}numPr/{W}numId") if ppr is not None else None
    check("numbered style: promoted to Heading 2", dict(_outline(after.tree)).get("Background") == 2,
          str(_outline(after.tree)))
    check("numbered style: numbering switched off on the paragraph (numId 0), text unchanged",
          num is not None and num.get(f"{W}val") == "0", etree.tostring(ppr).decode() if ppr is not None else "")

    # ===== F. Ambiguous look is declined, not guessed =========================
    d = Document()
    d.core_properties.title = "Ambiguous"
    h = d.add_heading("Overview", 1)          # 14pt bold
    h.runs[0].font.size = Pt(14)
    _body(d)
    # One point BIGGER than the H1 but NOT bold: the size says "above it",
    # the weight says "below it". Subtitle style is what gets it flagged.
    _fake(d, "Notes", 15, bold=False)
    d.paragraphs[-1].style = d.styles["Subtitle"]
    _body(d)
    src = tmp / "ambiguous.docx"
    d.save(str(src))
    res, execs, wr, after = _fix_fake_headings(src, tmp / "ambiguous_fixed.docx")
    check("ambiguous: exactly one promotion attempted", len(execs) == 1, str(execs))
    e = execs[0] if execs else None
    check("ambiguous: declined (skipped), not guessed", e is not None and e.status.value == "skipped",
          str((e.status.value, e.notes)) if e else "")
    check("ambiguous: the reason is in plain words", e is not None and "left for you" in (e.notes or "").lower(),
          (e.notes if e else ""))
    check("ambiguous: nothing counted", _count_persisted_fixes(execs, wr["applied"], "docx") == 0)
    check("ambiguous: the line is untouched in the output",
          dict(_outline(after.tree)).get("Notes") is None, str(_outline(after.tree)))

    # ===== G. A pre-existing jump is not made worse ===========================
    d = Document()
    d.core_properties.title = "Pre-existing jump"
    d.add_heading("Top", 1)
    _body(d)
    _fake(d, "Middle part", 14)
    _body(d)
    d.add_heading("Deep", 4)
    _body(d)
    src = tmp / "jump.docx"
    d.save(str(src))
    before = parse_to_tree(str(src))
    run_analyzers(before.tree)
    res, execs, wr, after = _fix_fake_headings(src, tmp / "jump_fixed.docx")
    outline = _outline(after.tree)
    lv = dict(outline).get("Middle part")
    check("pre-existing jump: promotion never goes deeper than previous+1", lv in (1, 2), str(outline))
    check("pre-existing jump: no MORE jumps than the source had",
          _flags(after.tree, JUMP) <= _flags(before.tree, JUMP), f"{_flags(before.tree, JUMP)} -> {_flags(after.tree, JUMP)}")

    # ===== G2. A pre-existing gap does not drag a heading below its look ======
    # "Eligibility" looks exactly like the Heading 2 before it. The heading
    # after it is a Heading 4 (the author skipped H3). Forcing "no gap before
    # the next heading" made it an H3 CHILD of the identical-looking H2; the
    # gap was already there, so the look decides: H2, and no more jumps.
    d = Document()
    d.core_properties.title = "Look decides"
    d.styles["Heading 1"].font.size = Pt(18)
    d.styles["Heading 2"].font.size = Pt(14)
    d.add_heading("Handbook", 1)
    _body(d)
    d.add_heading("Overview", 2)
    _body(d)
    _fake(d, "Eligibility", 14)   # exactly the Heading 2's look: 14pt bold
    _body(d)
    d.add_heading("Detail", 4)
    _body(d)
    src = tmp / "gap_look.docx"
    d.save(str(src))
    before = parse_to_tree(str(src))
    run_analyzers(before.tree)
    res, execs, wr, after = _fix_fake_headings(src, tmp / "gap_look_fixed.docx")
    outline = _outline(after.tree)
    check("pre-existing gap: a line that looks like the H2 above it is an H2 (sibling), not pushed to H3",
          dict(outline).get("Eligibility") == 2, str((outline, [e.notes for e in execs])))
    check("pre-existing gap: no more jumps than the source had",
          _flags(after.tree, JUMP) <= _flags(before.tree, JUMP), f"{_flags(before.tree, JUMP)} -> {_flags(after.tree, JUMP)}")

    # ===== G3. Look and no-skip rule disagree -> declined =====================
    # A 20pt bold "Appendix" (larger than every heading: a top-level heading)
    # sits between an H3 and an H4. As an H1 it would open a gap before the
    # H4 that the source does not have; as an H3 it would misstate how it
    # looks. Neither is clearly better than leaving it for a person.
    d = Document()
    d.core_properties.title = "Conflict"
    d.add_heading("Guide", 1)
    _body(d)
    d.add_heading("Part A", 2)
    _body(d)
    d.add_heading("Step one", 3)
    _body(d)
    _fake(d, "Appendix", 20)
    _body(d)
    d.add_heading("Detail", 4)
    _body(d)
    src = tmp / "conflict.docx"
    d.save(str(src))
    before = parse_to_tree(str(src))
    run_analyzers(before.tree)
    res, execs, wr, after = _fix_fake_headings(src, tmp / "conflict_fixed.docx")
    e = execs[0] if execs else None
    check("conflict: declined (skipped), with a plain-English reason naming the next heading",
          e is not None and e.status.value == "skipped" and "left for you" in (e.notes or "").lower()
          and "Heading 4" in (e.notes or ""), str((e.status.value, e.notes)) if e else "")
    check("conflict: nothing counted or charged", _count_persisted_fixes(execs, wr["applied"], "docx") == 0)
    check("conflict: the line is untouched and no jump was added",
          dict(_outline(after.tree)).get("Appendix") is None
          and _flags(after.tree, JUMP) == _flags(before.tree, JUMP), str(_outline(after.tree)))

    # ===== H. A promotion the writer cannot place is not counted ==============
    d = Document()
    d.core_properties.title = "Unplaced"
    d.add_heading("Top", 1)
    _fake(d, "Lonely section", 12)
    d.paragraphs[-1].runs[0].font.size = Pt(14)
    src = tmp / "unplaced.docx"
    d.save(str(src))
    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    plans = [p for p in plan_remediations(res.tree, POL) if p.flag.code.value == FAKE_H]
    execs = execute_plans(res.tree, plans)
    for node in iter_reading_order(res.tree.root):
        if isinstance(node, ParagraphNode) and (node.metadata.properties or {}).get("promote_to_heading_level"):
            node.id = "docx-p-999"  # no such paragraph in the source
            for e in execs:
                e.target_node_id = "docx-p-999"
    wr = write_remediated_docx(src, res.tree, tmp / "unplaced_fixed.docx")
    check("unplaced: executor said success", [e.status.value for e in execs] == ["success"])
    check("unplaced: writer did NOT confirm it", not [a for a in wr["applied"] if a.get("action") == "PROMOTE_HEADING"])
    check("unplaced: so it is NOT counted or charged", _count_persisted_fixes(execs, wr["applied"], "docx") == 0)

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
