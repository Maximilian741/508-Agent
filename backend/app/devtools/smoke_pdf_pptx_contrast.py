"""Smoke: colour-contrast coverage for PDF text + PPTX theme colours.

- PDF: text drawn with a low-contrast fill colour (rg/g/k in the content stream)
  is flagged vs the white page background; black/default text is not.
- PPTX: a theme scheme colour (resolved from the deck theme, with
  lumMod/lumOff/tint/shade) is now contrast-checked, not skipped.

Usage:
    python -m app.devtools.smoke_pdf_pptx_contrast
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_c3_')}/s.db"

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from pypdf import PdfWriter  # noqa: E402
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_engine import RemediationEngine  # noqa: E402

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _flagged(path: Path) -> bool:
    res = parse_to_tree(str(path))
    run_analyzers(res.tree)
    return any(v.rule_id == "LOW_CONTRAST_TEXT" for v in RemediationEngine().detect_violations(res.tree))


def main() -> int:
    failures = 0

    def check(name: str, cond: bool) -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name)
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp())

    # --- PDF text contrast ---
    def build_pdf(path: Path, content: bytes) -> None:
        w = PdfWriter(); pg = w.add_blank_page(width=400, height=300)
        cs = DecodedStreamObject(); cs.set_data(content)
        pg[NameObject("/Contents")] = w._add_object(cs)  # noqa: SLF001
        f = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        pg[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): w._add_object(f)})})  # noqa: SLF001
        with open(path, "wb") as fh:
            w.write(fh)

    build_pdf(tmp / "gray.pdf", b"0.6 0.6 0.6 rg BT /F1 12 Tf 20 200 Td (Light gray text) Tj ET")
    build_pdf(tmp / "black.pdf", b"0 0 0 rg BT /F1 12 Tf 20 200 Td (Black text) Tj ET")
    build_pdf(tmp / "default.pdf", b"BT /F1 12 Tf 20 200 Td (Default colour) Tj ET")
    check("PDF gray text flagged", _flagged(tmp / "gray.pdf"))
    check("PDF black text not flagged", not _flagged(tmp / "black.pdf"))
    check("PDF default(black) text not flagged", not _flagged(tmp / "default.pdf"))

    # --- Coloured-background false positive (fixed) ---------------------------
    # Dark navy band painted first, then WHITE text on it: perfectly readable,
    # but the white-bg assumption would call it 1:1 contrast. The page paints a
    # non-white fill, so contrast scanning skips it -> no flag.
    build_pdf(
        tmp / "darkband.pdf",
        b"0.05 0.10 0.30 rg 0 0 400 300 re f "
        b"1 1 1 rg BT /F1 14 Tf 20 200 Td (White on navy reads fine) Tj ET",
    )
    check("PDF white-on-dark band NOT flagged (FP fixed)", not _flagged(tmp / "darkband.pdf"))

    # A WHITE painted background keeps the assumption valid — light gray text
    # on it must still be caught.
    build_pdf(
        tmp / "whitebg.pdf",
        b"1 1 1 rg 0 0 400 300 re f "
        b"0.6 0.6 0.6 rg BT /F1 12 Tf 20 200 Td (Gray on painted white) Tj ET",
    )
    check("PDF gray text on painted-white bg still flagged", _flagged(tmp / "whitebg.pdf"))

    # Gradient (sh) pages are unknowable -> skipped, no flag either way.
    build_pdf(
        tmp / "shading.pdf",
        b"sh 0.6 0.6 0.6 rg BT /F1 12 Tf 20 200 Td (Gray over gradient) Tj ET",
    )
    check("PDF gradient page skipped (no flag)", not _flagged(tmp / "shading.pdf"))

    # --- PPTX theme-colour contrast ---
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(1))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    box.text_frame.text = "Light scheme text"
    run = box.text_frame.paragraphs[0].runs[0]; run.font.size = Pt(12)
    rpr = run._r.get_or_add_rPr()
    sf = etree.SubElement(rpr, f"{_A}solidFill")
    sc = etree.SubElement(sf, f"{_A}schemeClr"); sc.set("val", "accent1")
    etree.SubElement(sc, f"{_A}lumMod").set("val", "40000")
    etree.SubElement(sc, f"{_A}lumOff").set("val", "60000")
    pth = tmp / "theme.pptx"; prs.save(str(pth))
    check("PPTX light theme-colour text flagged", _flagged(pth))

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
