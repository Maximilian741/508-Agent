"""Smoke: Excel workbooks are checked and fixed, and only fixes in the bytes count.

XLSX is a first-class pipeline format. A workbook built the way Excel builds
one (shared strings, styles, drawings, a chart, a text box, an existing table)
goes through /pipeline/analyze and /pipeline/remediate over HTTP:

  * analyze finds: no title / no language; a default "Sheet1" tab; a clear
    header + data range (TABLE_MISSING_HEADERS, auto-fixable) vs blocks whose
    header row cannot be identified (DATA_RANGE_HEADERS_UNCLEAR, manual —
    names over names, merged cells, a label/value form), each reported ONCE;
    a chart and a picture with no alt; a decorative picture that has alt;
  * remediate with everything approved: the title (from the sheet's own title
    line), the chart's alt text (grounded in the chart's own title), the
    decorative picture's alt removal and the Excel table persist, are counted
    and charged ONCE at the xlsx price (3); the caption-less picture is
    REFUSED (no placeholder written); the file reopens in openpyxl with the
    table where we put it; alt text the author wrote, the text box, and every
    part we did not touch are exactly as they were; re-analysing the output
    clears the fixed findings and raises nothing new;
  * approving only manual findings charges nothing and returns the upload;
  * the writer refuses a made-up header row, keeps the metadata fixes when a
    table conversion fails verification, and a failed save is a 422 with no
    charge;
  * a sheet longer than the scan window keeps its data block whole, and data
    we did not read is disclosed (ANALYSIS_TRUNCATED), never silently dropped;
  * only a line that stands on its own (a blank row under it, or the data it
    labels right below) is offered as the workbook title: "Name" over a
    column of names is a heading, and writing it as the title is worse than
    the filename.

Usage:
    python -m app.devtools.smoke_xlsx_pipeline
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import warnings
import zipfile
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="508_smoke_xlsx_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/s.db"
os.environ["MATERIALIZED_ROOT"] = str(Path(_TMP) / "materialized")
os.environ.setdefault("APP_SECRET", "x" * 64)
os.environ["SEMANTIC_PROVIDER"] = "heuristic"
for _key in ("OPENAI_API_KEY", "ANTHROPIC_API_KEY"):
    os.environ.pop(_key, None)

import openpyxl  # noqa: E402
import xlsxwriter  # noqa: E402
from PIL import Image  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import select  # noqa: E402

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GOOD_ALT = "Company logo: a blue square with the letters ACME"


def _patch_zip(data: bytes, part: str, old: str, new: str) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info)
            if info.filename == part:
                text = body.decode("utf-8")
                assert old in text, f"{old!r} not in {part}"
                body = text.replace(old, new, 1).encode("utf-8")
            dst.writestr(info, body)
    return out.getvalue()


def build_workbook(tmp: Path) -> bytes:
    png = tmp / "logo.png"
    Image.new("RGB", (60, 40), (20, 90, 200)).save(png)
    path = tmp / "regional.xlsx"
    wb = xlsxwriter.Workbook(str(path))
    bold = wb.add_format({"bold": True})

    ws = wb.add_worksheet()  # "Sheet1": a default tab name
    ws.write("A1", "Q3 2026 Regional Sales")
    ws.write_row("A3", ["Region", "Q1", "Q2", "Q3"], bold)
    for i, row in enumerate([("North", 10, 20, 30), ("South", 11, 21, 31), ("East", 12, 22, 32), ("West", 13, 23, 33)]):
        ws.write_row(3 + i, 0, list(row))
    ws.insert_image("G1", str(png), {"description": GOOD_ALT})          # alt the author wrote
    ws.insert_image("G5", str(png), {"description": ""})                # no alt at all
    ws.insert_image("G8", str(png), {"decorative": True})               # decorative (patched below)
    chart = wb.add_chart({"type": "column"})
    chart.add_series({"name": "Q1", "categories": "=Sheet1!$A$4:$A$7", "values": "=Sheet1!$B$4:$B$7"})
    chart.add_series({"name": "Q2", "categories": "=Sheet1!$A$4:$A$7", "values": "=Sheet1!$C$4:$C$7"})
    chart.set_title({"name": "Quarterly sales by region"})
    ws.insert_chart("G12", chart)
    ws.insert_textbox("M1", "Prepared by the finance team")

    staff = wb.add_worksheet("Staff")       # names over names: header not identifiable
    staff.write_row("A1", ["Name", "Email", "Team"])
    for i, row in enumerate([("Ann", "ann@example.com", "Ops"), ("Bob", "bob@example.com", "Sales"), ("Cy", "cy@example.com", "Ops")]):
        staff.write_row(1 + i, 0, list(row))

    merged = wb.add_worksheet("Budget")     # merged cells inside the data
    merged.write_row("A1", ["Item", "Cost", "Owner"], bold)
    merged.write_row("A2", ["Paper", 12, "Ann"])
    merged.merge_range("A3:B3", "Shared services")
    merged.write("C3", "Bob")
    merged.write_row("A4", ["Ink", 30, "Cy"])

    form = wb.add_worksheet("Form")         # a label/value form, not a table
    for i, (label, value) in enumerate([("Name", "Dana Smith"), ("Start date", 45000), ("Salary", 52000), ("Office", "Leeds")]):
        form.write(i, 0, label, bold)
        form.write(i, 1, value)

    declared = wb.add_worksheet("Inventory")  # an Excel table with its header row on
    declared.add_table(
        "B2:D5",
        {"columns": [{"header": "Item"}, {"header": "Qty"}, {"header": "Cost"}], "data": [["a", 1, 2], ["b", 3, 4], ["c", 5, 6]]},
    )
    wb.close()
    data = path.read_bytes()
    # Decorative AND described: the contradiction a person has to resolve.
    return _patch_zip(data, "xl/drawings/drawing1.xml", 'name="Picture 3">', 'name="Picture 3" descr="Blue swirl">')


def main() -> int:  # noqa: PLR0915
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    from app.api.credits import DOC_FORMAT_COSTS
    from app.api.pipeline import _PERSISTED_ACTIONS, _count_persisted_fixes
    from app.db.models import UserRow
    from app.db.session_sqlalchemy import session_scope
    from app.main import app
    from app.parsers import parse_to_tree
    from app.parsers import xlsx_parser
    from app.services.remediation_engine import RemediationEngine

    tmp = Path(_TMP)
    src = build_workbook(tmp)
    check("xlsx is priced (3 credits, docx tier)", DOC_FORMAT_COSTS.get("xlsx") == 3, str(DOC_FORMAT_COSTS))
    check("xlsx has a persisted-action set", "ADD_TABLE_HEADERS" in _PERSISTED_ACTIONS.get("xlsx", set()))

    client = TestClient(app, raise_server_exceptions=False)
    email = "xlsx-smoke@example.com"
    r = client.post("/auth/sign-in", json={"email": email, "displayName": "X", "password": "xlsxsmokepass1"})
    assert r.status_code == 200, r.text
    headers = {"Authorization": f"Bearer {r.json()['token']}"}

    def set_balance(n: int) -> None:
        with session_scope() as s:
            row = s.execute(select(UserRow).where(UserRow.email == email)).scalars().first()
            row.credits_balance = n

    def balance() -> int:
        return int(client.get("/credits/balance", headers=headers).json()["balance"])

    def analyze(name: str, data: bytes) -> dict:
        rr = client.post("/pipeline/analyze", files={"file": (name, data, XLSX)}, headers=headers)
        assert rr.status_code == 200, rr.text
        return rr.json()

    def remediate(name: str, data: bytes, ids: list):
        return client.post(
            "/pipeline/remediate",
            files={"file": (name, data, XLSX)},
            data={"approved_violations": json.dumps(ids), "rejected_violations": "[]"},
            headers=headers,
        )

    set_balance(50)

    # ---- 1. analyze ---------------------------------------------------------
    rep = analyze("regional.xlsx", src)
    check("analyze: sourceFormat xlsx", rep["summary"]["sourceFormat"] == "xlsx", rep["summary"]["sourceFormat"])
    check("analyze: a native upload carries no intake record", rep.get("intake") is None)
    found = {(v["ruleId"], v["nodeId"]) for v in rep["violations"]}
    rules = [v["ruleId"] for v in rep["violations"]]
    for rule in ("DOCUMENT_TITLE_MISSING", "DOCUMENT_LANGUAGE_MISSING"):
        check(f"analyze: {rule}", rule in rules)
    check("analyze: default 'Sheet1' tab flagged", ("SHEET_NAME_DEFAULT", "xlsx-s1") in found, str(found))
    check("analyze: named tabs not flagged", sum(1 for r_ in rules if r_ == "SHEET_NAME_DEFAULT") == 1)
    check("analyze: clear header range -> TABLE_MISSING_HEADERS (auto)", ("TABLE_MISSING_HEADERS", "xlsx-s1-t1") in found, str(found))
    for sheet_no, what in ((2, "names over names"), (3, "merged cells"), (4, "label/value form")):
        check(
            f"analyze: {what} -> DATA_RANGE_HEADERS_UNCLEAR, not TABLE_MISSING_HEADERS",
            ("DATA_RANGE_HEADERS_UNCLEAR", f"xlsx-s{sheet_no}-t1") in found
            and ("TABLE_MISSING_HEADERS", f"xlsx-s{sheet_no}-t1") not in found,
            str(sorted(f for f in found if f[1].startswith(f"xlsx-s{sheet_no}"))),
        )
    check("analyze: declared Excel table raises no header finding",
          not any(f[1] == "xlsx-s5-t1" and "HEADER" in f[0] for f in found), str(found))
    unclear = [v for v in rep["violations"] if v["ruleId"] == "DATA_RANGE_HEADERS_UNCLEAR"]
    check("analyze: unclear finding is manual-only", all(v["recommendedActions"] == ["FLAG_FOR_MANUAL_REVIEW"] for v in unclear))
    check("analyze: chart without alt -> MISSING_ALT_TEXT", ("MISSING_ALT_TEXT", "xlsx-s1-chart1") in found)
    check("analyze: picture without alt -> MISSING_ALT_TEXT", ("MISSING_ALT_TEXT", "xlsx-s1-img2") in found)
    check("analyze: picture with real alt is not flagged", not any(f[1] == "xlsx-s1-img1" for f in found), str(found))
    check("analyze: decorative picture with alt -> DECORATIVE_IMAGE_WITH_ALT", ("DECORATIVE_IMAGE_WITH_ALT", "xlsx-s1-img3") in found)
    # The large-grid "needs a summary" rule is not fed every spreadsheet row.
    complex_nodes = {f[1] for f in found if f[0] == "TABLE_COMPLEX_NEEDS_SUMMARY"}
    check("analyze: 'needs a summary' only where cells are merged, not on plain data",
          complex_nodes <= {"xlsx-s3-t1"}, str(complex_nodes))

    # ---- 2. remediate everything --------------------------------------------
    all_ids = [v["id"] for v in rep["violations"]]
    b0 = balance()
    rr = remediate("regional.xlsx", src, all_ids)
    check("remediate: 200", rr.status_code == 200, rr.text[:300])
    body = rr.json()
    applied = (body.get("writer") or {}).get("applied") or []
    by_action = {}
    for a in applied:
        by_action.setdefault(a.get("action"), []).append(a.get("target_id"))
    check("remediate: title written", by_action.get("SET_DOCUMENT_TITLE") == ["doc-1"], str(applied))
    # The chart's alt comes only from its own title. Whether the offline
    # provider hands that title back cleanly or glued to our node id
    # ("Image xlsx-s1-chart1 — ...") is the provider's business; the writer's
    # job is that the id never reaches the file, and that what lands is the
    # chart's own words.
    chart_written = "xlsx-s1-chart1" in (by_action.get("GENERATE_ALT_TEXT") or [])
    chart_refused = any(s.get("target_id") == "xlsx-s1-chart1" and "alt_text_refused" in s.get("reason", "")
                        for s in (body.get("writer") or {}).get("skipped") or [])
    check("remediate: chart alt is written from its own title, or refused with a reason",
          chart_written != chart_refused, str(applied))
    check("remediate: caption-less picture NOT given a placeholder", "xlsx-s1-img2" not in (by_action.get("GENERATE_ALT_TEXT") or []))
    check("remediate: decorative alt removed", by_action.get("REMOVE_DECORATIVE_ALT_TEXT") == ["xlsx-s1-img3"], str(applied))
    check("remediate: clear range became an Excel table", by_action.get("ADD_TABLE_HEADERS") == ["xlsx-s1-t1"], str(applied))
    check("remediate: persistedFixes == what the writer confirmed", body["persistedFixes"] == len(applied), f"{body['persistedFixes']} vs {len(applied)}")
    check("remediate: charged once at the xlsx price", body["charged"] is True and b0 - balance() == 3, f"{b0} -> {balance()}")
    check("remediate: output named .xlsx", body["filename"] == "regional-remediated.xlsx", body["filename"])
    ex = {(e["actionCode"], e["targetNodeId"]): e for e in body["executions"]}
    img_ex = ex.get(("GENERATE_ALT_TEXT", "xlsx-s1-img2"))
    check("remediate: refused picture reported as skipped with a reason",
          img_ex is not None and img_ex["status"] == "skipped" and "placeholder" in img_ex["notes"], str(img_ex))
    check("remediate: no manual item reported as a success",
          all(e["status"] != "success" for e in body["executions"] if e["actionCode"] == "FLAG_FOR_MANUAL_REVIEW"))

    dl = client.get(body["downloadUrl"], headers=headers)
    check("download: 200 with the xlsx media type", dl.status_code == 200 and "spreadsheetml" in dl.headers.get("content-type", ""), dl.headers.get("content-type"))
    out = dl.content
    out_path = tmp / "out.xlsx"
    out_path.write_bytes(out)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(out_path)
    tables = dict(wb["Sheet1"].tables.items())
    check("openpyxl: output reopens with the new table at A3:D7", list(tables.values()) == ["A3:D7"], str(tables))
    tbl = wb["Sheet1"].tables[list(tables)[0]] if tables else None
    check("openpyxl: table columns are the header cells", tbl is not None and [c.name for c in tbl.tableColumns] == ["Region", "Q1", "Q2", "Q3"])
    check("openpyxl: the existing table is untouched", dict(wb["Inventory"].tables.items()) == {"Table1": "B2:D5"})
    check("openpyxl: title reads back", wb.properties.title == "Q3 2026 Regional Sales", str(wb.properties.title))
    wb.close()

    zin, zout = zipfile.ZipFile(io.BytesIO(src)), zipfile.ZipFile(io.BytesIO(out))
    drawing = zout.read("xl/drawings/drawing1.xml").decode("utf-8")
    check("alt the author wrote is preserved", f'descr="{GOOD_ALT}"' in drawing)
    check("the text box survives", "Prepared by the finance team" in drawing and drawing.count("<xdr:sp ") == zin.read("xl/drawings/drawing1.xml").decode().count("<xdr:sp "))
    check("the decorative picture no longer has alt", 'descr="Blue swirl"' not in drawing)
    chart_drawing = zout.read("xl/drawings/drawing1.xml").decode("utf-8")
    check("no internal node id ever reaches the workbook", "xlsx-s1-" not in chart_drawing)
    if chart_written:
        check("the chart's alt is its own title", 'descr="Column chart: Quarterly sales by region (Q1, Q2)"' in chart_drawing,
              chart_drawing[chart_drawing.find("Chart"):][:300])
    touched = {"xl/drawings/drawing1.xml", "docProps/core.xml", "[Content_Types].xml", "xl/worksheets/sheet1.xml",
               "xl/worksheets/_rels/sheet1.xml.rels"}
    same = [n for n in zin.namelist() if n not in touched and zin.read(n) == zout.read(n)]
    diff = [n for n in zin.namelist() if n not in touched and zin.read(n) != zout.read(n)]
    check("every part we did not edit is byte-identical", not diff and len(same) >= 10, str(diff))
    check("the new table part is the only new part", sorted(set(zout.namelist()) - set(zin.namelist())) == ["xl/tables/table2.xml"],
          str(sorted(set(zout.namelist()) - set(zin.namelist()))))

    after = analyze("regional-remediated.xlsx", out)
    after_found = {(v["ruleId"], v["nodeId"]) for v in after["violations"]}
    cleared = [("TABLE_MISSING_HEADERS", "xlsx-s1-t1"), ("DECORATIVE_IMAGE_WITH_ALT", "xlsx-s1-img3")]
    if chart_written:
        cleared.append(("MISSING_ALT_TEXT", "xlsx-s1-chart1"))
    else:
        check("re-analysis: a refused chart alt keeps its finding", ("MISSING_ALT_TEXT", "xlsx-s1-chart1") in after_found)
    for gone in cleared:
        check(f"re-analysis: {gone[0]} cleared", gone not in after_found)
    check("re-analysis: title finding cleared", not any(f[0] == "DOCUMENT_TITLE_MISSING" for f in after_found))
    check("re-analysis: nothing NEW after fix-everything", after_found <= found, str(sorted(after_found - found)))

    # ---- 3. only manual items approved: free, and the upload comes back ----
    manual_ids = [v["id"] for v in rep["violations"] if v["ruleId"] in ("SHEET_NAME_DEFAULT", "DATA_RANGE_HEADERS_UNCLEAR")]
    b0 = balance()
    rr = remediate("regional.xlsx", src, manual_ids)
    body = rr.json()
    check("manual-only: not charged", rr.status_code == 200 and body["charged"] is False and balance() == b0, str(body.get("charged")))
    dl = client.get(body["downloadUrl"], headers=headers)
    check("manual-only: the original bytes come back", dl.content == src)

    # ---- 4. writer refusals and failure handling ---------------------------
    from app.models.accessibility import (
        ContentKind,
        NodeContent,
        NodeMetadata,
        TableCellNode,
        TableCellType,
        TableHeaderScope,
        TableNode,
        TableRowNode,
        iter_reading_order,
    )
    from app.writers import xlsx_writer
    from app.writers.xlsx_writer import write_remediated_xlsx

    src_path = tmp / "regional.xlsx"
    src_path.write_bytes(src)

    res = parse_to_tree(str(src_path))
    table = next(n for n in iter_reading_order(res.tree.root) if isinstance(n, TableNode) and n.id == "xlsx-s1-t1")
    fake = TableRowNode(
        id=f"{table.id}-thead",
        content=NodeContent(kind=ContentKind.NONE),
        metadata=NodeMetadata(properties={"synthesized": True}),
        children=[
            TableCellNode(id=f"{table.id}-th-{i}", cell_type=TableCellType.HEADER, header_scope=TableHeaderScope.COLUMN,
                          content=NodeContent(kind=ContentKind.TEXT, text=f"Column {i}"),
                          metadata=NodeMetadata(properties={"synthesized": True}), children=[], accessibility_flags=[])
            for i in range(1, 5)
        ],
        accessibility_flags=[],
    )
    table.children.insert(0, fake)
    wr = write_remediated_xlsx(src_path, res.tree, tmp / "fake.xlsx")
    check("writer: a made-up header row never becomes an Excel table",
          not wr["applied"] and any(s.get("reason", "").startswith("no_clear_header_row") for s in wr["skipped"]), str(wr))

    def chart_node(tree):
        from app.models.accessibility import ImageNode

        return next(n for n in iter_reading_order(tree.root) if isinstance(n, ImageNode) and n.id == "xlsx-s1-chart1")

    res = parse_to_tree(str(src_path))
    cprops = chart_node(res.tree).metadata.properties
    check("parser: the chart's caption is its own title, marked as authored",
          cprops.get("caption") == "Column chart: Quarterly sales by region (Q1, Q2)" and cprops.get("caption_source") == "title",
          f"{cprops.get('caption')!r} {cprops.get('caption_source')!r}")
    chart_node(res.tree).alt_text = "Image xlsx-s1-chart1 — Column chart: Quarterly sales by region (Q1, Q2)"
    wr = write_remediated_xlsx(src_path, res.tree, tmp / "idalt.xlsx")
    check("writer: alt text carrying our node id is refused, not written",
          not wr["applied"] and any(s.get("target_id") == "xlsx-s1-chart1" and "internal id" in s.get("reason", "")
                                    for s in wr["skipped"]), str(wr))
    with zipfile.ZipFile(tmp / "idalt.xlsx") as z:
        check("writer: ...and the file carries no trace of it", b"xlsx-s1-chart1" not in z.read("xl/drawings/drawing1.xml"))

    res = parse_to_tree(str(src_path))
    chart_node(res.tree).alt_text = "Column chart: Quarterly sales by region (Q1, Q2)"
    wr = write_remediated_xlsx(src_path, res.tree, tmp / "cleanalt.xlsx")
    with zipfile.ZipFile(tmp / "cleanalt.xlsx") as z:
        check("writer: the chart's own title lands as its alt text",
              [a["action"] for a in wr["applied"]] == ["GENERATE_ALT_TEXT"]
              and b'descr="Column chart: Quarterly sales by region (Q1, Q2)"' in z.read("xl/drawings/drawing1.xml"), str(wr))

    res = parse_to_tree(str(src_path))
    res.tree.root.metadata.properties["title"] = "Regional sales"
    chart_node(res.tree).alt_text = "Sales\x01chart"
    wr = write_remediated_xlsx(src_path, res.tree, tmp / "ctrl.xlsx")
    check("writer: an alt text a workbook cannot store is refused on its own; the title still lands",
          [a["action"] for a in wr["applied"]] == ["SET_DOCUMENT_TITLE"]
          and any("cannot store" in s.get("reason", "") for s in wr["skipped"]), str(wr))

    res = parse_to_tree(str(src_path))
    res.tree.root.metadata.properties["title"] = "Regional sales"
    table = next(n for n in iter_reading_order(res.tree.root) if isinstance(n, TableNode) and n.id == "xlsx-s1-t1")
    for cell in table.children[0].children:
        cell.cell_type = TableCellType.HEADER
    real_verify = xlsx_writer._verify

    def _tables_fail(path, edits):
        return ["simulated reopen failure"] if edits.tables else real_verify(path, edits)

    xlsx_writer._verify = _tables_fail
    try:
        wr = write_remediated_xlsx(src_path, res.tree, tmp / "fallback.xlsx")
    finally:
        xlsx_writer._verify = real_verify
    acts = [a["action"] for a in wr["applied"]]
    check("writer: a table that fails verification is dropped, the title still lands",
          acts == ["SET_DOCUMENT_TITLE"] and any("output_failed_verification" in s.get("reason", "") for s in wr["skipped"]), str(wr))
    with zipfile.ZipFile(tmp / "fallback.xlsx") as z:
        check("writer: ...and the fallback file has no new table part", not any(n.startswith("xl/tables/table2") for n in z.namelist()))

    class _Exec:  # minimal shape _count_persisted_fixes reads
        def __init__(self, code, target):
            from app.models.accessibility import ActionCode
            from app.services.remediators.base import ExecutionStatus

            self.action_code = ActionCode(code)
            self.target_node_id = target
            self.status = ExecutionStatus.SUCCESS

    count = _count_persisted_fixes(
        [_Exec("SET_DOCUMENT_TITLE", "doc-1"), _Exec("ADD_TABLE_HEADERS", "xlsx-s1-t1")], wr["applied"], "xlsx", wr["skipped"]
    )
    check("honesty gate: only the writer-confirmed title is counted", count == 1, str(count))

    real_rewrite = xlsx_writer._rewrite_zip

    def _enospc(*a, **k):
        raise OSError(28, "No space left on device")

    b0 = balance()
    xlsx_writer._rewrite_zip = _enospc
    try:
        rr = remediate("regional.xlsx", src, all_ids)
    finally:
        xlsx_writer._rewrite_zip = real_rewrite
    check("failed save -> 422 and nothing charged", rr.status_code == 422 and balance() == b0, f"{rr.status_code} {b0}->{balance()}")
    check("failed save -> the reason is a sentence", "not charged" in (rr.json().get("detail") or ""), rr.text[:200])

    # ---- 5. header judgement and the scan window ---------------------------
    grid = {}

    def put(r, c, text, kind="text", styled=False):
        grid[(r, c)] = xlsx_parser.CellInfo(text=text, kind=kind, styled=styled)

    for c, t in enumerate(["Name", "Email"], 1):
        put(1, c, t, styled=True)
    for r in (2, 3, 4):
        put(r, 1, f"Person {r}")
        put(r, 2, f"p{r}@example.com")
    blk = xlsx_parser.DataBlock(r1=1, c1=1, r2=4, c2=2)
    check("judge: bold text headings over plain text are clear", xlsx_parser.judge_header(blk, grid)[0] is True)
    grid.clear()
    for c, t in enumerate(["Name:", "Email:"], 1):
        put(1, c, t, styled=True)
    for r in (2, 3):
        put(r, 1, "x")
        put(r, 2, "y")
    check("judge: labels ending in ':' are not headings", xlsx_parser.judge_header(xlsx_parser.DataBlock(1, 1, 3, 2), grid)[0] is False)
    grid.clear()
    for c, t in enumerate(["Cost", "Cost"], 1):
        put(1, c, t, styled=True)
    for r in (2, 3):
        put(r, 1, "1", "number")
        put(r, 2, "2", "number")
    ok, why = xlsx_parser.judge_header(xlsx_parser.DataBlock(1, 1, 3, 2), grid)
    check("judge: duplicate headings are refused, and the reason says so", ok is False and "same text" in (why or ""), str(why))
    grid[(1, 1)] = xlsx_parser.CellInfo(text="Cost_x0041_", kind="text", styled=True)
    check("judge: a heading that would not round-trip as a column name is refused",
          xlsx_parser.judge_header(xlsx_parser.DataBlock(1, 1, 3, 2), grid)[0] is False)

    big = tmp / "long.xlsx"
    wbk = xlsxwriter.Workbook(str(big))
    sh = wbk.add_worksheet("Ledger")
    sh.write_row(0, 0, ["Date", "Amount", "Memo"], wbk.add_format({"bold": True}))
    for r in range(1, 120):
        sh.write_row(r, 0, [45000 + r, r * 3, f"entry {r}"])
    sh.write_row(199, 0, ["Late", "note"])  # a second block entirely below the window
    sh.write_row(200, 0, ["x", "y"])
    wbk.close()
    saved = xlsx_parser.SCAN_ROWS
    xlsx_parser.SCAN_ROWS = 50
    try:
        res = parse_to_tree(str(big))
        viols = RemediationEngine().detect_violations(res.tree)
    finally:
        xlsx_parser.SCAN_ROWS = saved
    t1 = next(n for n in iter_reading_order(res.tree.root) if isinstance(n, TableNode))
    props = t1.metadata.properties
    check("window: a block longer than the window keeps its full extent", props.get("cell_range") == "A1:C120", str(props.get("cell_range")))
    check("window: ...and is still offered the fix", props.get("header_detection") == "candidate", str(props.get("header_detection")))
    check("window: data below the window is disclosed (ANALYSIS_TRUNCATED)", any(v.rule_id == "ANALYSIS_TRUNCATED" for v in viols))

    # ---- 6. which line may become the workbook title -----------------------
    # The title fix prefers the sheet's own title line over the filename, so
    # a column heading mistaken for a title would be written into the file.
    def title_of(name: str, rows) -> object:
        p = tmp / name
        book = xlsxwriter.Workbook(str(p))
        sheet = book.add_worksheet("Data")
        for (r, c, v) in rows:
            sheet.write(r, c, v)
        book.close()
        return (parse_to_tree(str(p)).tree.root.metadata.properties or {}).get("title_candidate")

    one_col = [(0, 0, "Name")] + [(i, 0, n) for i, n in enumerate(["Alice Jones", "Bob Smith", "Carol White"], 1)]
    check("title: 'Name' over a column of names is NOT the workbook title", title_of("names.xlsx", one_col) is None,
          repr(title_of("names.xlsx", one_col)))
    listed = [(0, 0, "Contacts"), (1, 0, "Alice"), (2, 0, "Bob"), (5, 0, "Item"), (5, 1, "Cost")] + [
        (6 + i, c, v) for i in range(4) for c, v in ((0, f"thing {i}"), (1, i))
    ]
    check("title: the first entry of a list above the data is not a title", title_of("listed.xlsx", listed) is None,
          repr(title_of("listed.xlsx", listed)))
    titled = [(0, 0, "Department budget 2025"), (2, 0, "Department"), (2, 1, "Q1")] + [
        (3 + i, c, v) for i in range(4) for c, v in ((0, f"Dept {i}"), (1, i))
    ]
    check("title: a line with a blank row under it still is", title_of("titled.xlsx", titled) == "Department budget 2025",
          repr(title_of("titled.xlsx", titled)))
    labelled = [(0, 0, "Department budget 2025"), (1, 0, "Department"), (1, 1, "Q1")] + [
        (2 + i, c, v) for i in range(4) for c, v in ((0, f"Dept {i}"), (1, i))
    ]
    check("title: ...and so does the label right above the data", title_of("labelled.xlsx", labelled) == "Department budget 2025",
          repr(title_of("labelled.xlsx", labelled)))

    # ---- 7. default tab names, including openpyxl's number-less ones -------
    defaults = {"Sheet1": True, "Sheet": True, "Chart": True, "Chart2": True, "Feuil3": True, "Tabelle1": True,
                "Summary": False, "Sheets": False, "List": False, "Q3 Sheet": False, "Sheet 2 notes": False}
    wrong = {n: want for n, want in defaults.items() if xlsx_parser.is_default_sheet_name(n) != want}
    check("tabs: default names recognised (openpyxl's 'Sheet'/'Chart' too), real names left alone", not wrong, str(wrong))

    # ---- 8. inline strings (streaming writers store text in the cell) -----
    inline = tmp / "inline.xlsx"
    book = xlsxwriter.Workbook(str(inline), {"constant_memory": True})
    sheet = book.add_worksheet("Stock")
    sheet.write_row(0, 0, ["Part", "Count"], book.add_format({"bold": True}))
    for i in range(1, 5):
        sheet.write_row(i, 0, [f"part {i}", i * 7])
    book.close()
    with zipfile.ZipFile(inline) as z:
        is_inline = b't="inlineStr"' in z.read("xl/worksheets/sheet1.xml")
    t_inline = next(n for n in iter_reading_order(parse_to_tree(str(inline)).tree.root) if isinstance(n, TableNode))
    check("inline strings: read as text, and the heading row is recognised",
          is_inline and t_inline.metadata.properties.get("header_detection") == "candidate"
          and [c.content.text for c in t_inline.children[0].children] == ["Part", "Count"],
          f"{is_inline} {t_inline.metadata.properties.get('header_detection')}")

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
