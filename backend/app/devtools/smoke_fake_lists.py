"""Smoke: typed fake-list detection + conversion to REAL Word lists.

Plain paragraphs typed as "- item" / "1. item" carry no list semantics for
screen readers. This pins the full path:

  parser  -> consecutive same-marker runs grouped, one LIST_STRUCTURE_INVALID
             per typed list (precision guards: prose, single lines, real
             lists, year-prefix sentences all stay quiet)
  executor-> FIX_LIST_STRUCTURE succeeds, strips markers in the tree
  writer  -> output bytes gain w:numPr per paragraph + numbering.xml
             definitions; literal markers stripped from runs
  honesty -> FIX_LIST_STRUCTURE persists for docx; re-parsing the OUTPUT
             yields a real ListNode and re-analysis raises NO fake-list flag
             (self-consistent, idempotent).

Usage:
    python -m app.devtools.smoke_fake_lists
"""

from __future__ import annotations

import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp(prefix='508_smoke_fl_')}/s.db")

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from app.analyzers.registry import run_analyzers  # noqa: E402
from app.api.pipeline import _action_persists  # noqa: E402
from app.models.accessibility import ListNode  # noqa: E402
from app.parsers import parse_to_tree  # noqa: E402
from app.services.remediation_planner import RemediationPolicy, plan_remediations  # noqa: E402
from app.services.remediators.registry import execute_plans  # noqa: E402
from app.writers.docx_writer import write_remediated_docx  # noqa: E402

FLAG = "LIST_STRUCTURE_INVALID"


def _flag_count(tree) -> int:
    n = 0

    def walk(node):
        nonlocal n
        n += sum(1 for f in node.accessibility_flags if f.code.value == FLAG)
        for ch in node.children:
            walk(ch)

    walk(tree.root)
    return n


def main() -> int:
    failures = 0

    def check(name: str, cond: bool, extra: str = "") -> None:
        nonlocal failures
        print(("PASS" if cond else "FAIL"), "-", name, extra if not cond else "")
        if not cond:
            failures += 1

    tmp = Path(tempfile.mkdtemp(prefix="fake_lists_"))
    apply_policy = RemediationPolicy(allow_ai_actions=True, require_human_review_for_all=False)

    # --- fixture: 2 fake lists + every false-positive trap -------------------
    d = Document()
    d.core_properties.title = "Fake Lists"
    d.add_paragraph("An introduction paragraph that is ordinary prose.")
    d.add_paragraph("- apples are red")          # fake bullet run (3)
    d.add_paragraph("- bananas are yellow")
    d.add_paragraph("- cherries are dark")
    d.add_paragraph("Now some steps follow below.")
    d.add_paragraph("1. open the door")          # fake numbered run (3)
    d.add_paragraph("2. step inside")
    d.add_paragraph("3. close the door")
    d.add_paragraph("- a single stray dash line")  # run of 1 -> never flagged
    d.add_paragraph("1986. It was a very good year, says the prose.")  # year sentence
    d.add_paragraph("5. a stray mid-sequence number line")  # doesn't start at 1
    real = d.add_paragraph("first real item", style="List Bullet")  # real list
    d.add_paragraph("second real item", style="List Bullet")
    assert real._p.find(qn("w:pPr")) is not None
    src = tmp / "fake.docx"
    d.save(str(src))

    res = parse_to_tree(str(src))
    run_analyzers(res.tree)
    n_flags = _flag_count(res.tree)
    check("exactly TWO fake lists flagged (bullet run + numbered run)", n_flags == 2, f"got {n_flags}")

    # --- executor: both runs convert -----------------------------------------
    plans = [p for p in plan_remediations(res.tree, apply_policy) if p.flag.code.value == FLAG]
    check("two FIX_LIST_STRUCTURE plans", len(plans) == 2, f"got {len(plans)}")
    execs = execute_plans(res.tree, plans)
    ok = [e for e in execs if e.status.value == "success"]
    check("both conversions execute successfully", len(ok) == 2, str([(e.status.value, e.notes) for e in execs]))
    check("FIX_LIST_STRUCTURE persists for docx (honesty matrix)", _action_persists("FIX_LIST_STRUCTURE", "docx"))

    # --- writer: numbering + markers in the BYTES -----------------------------
    out = tmp / "fixed.docx"
    result = write_remediated_docx(src, res.tree, out)
    conversions = [a for a in result["applied"] if a.get("kind") == "list_conversion"]
    check("writer applied 6 list conversions", len(conversions) == 6, str(result))

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        check("output contains numbering.xml", "word/numbering.xml" in names)
        numbering_xml = z.read("word/numbering.xml").decode("utf-8", "replace")
        doc_xml = z.read("word/document.xml").decode("utf-8", "replace")
    check("bullet + decimal numFmt defined", 'w:val="bullet"' in numbering_xml and 'w:val="decimal"' in numbering_xml)
    check("markers stripped from document.xml", "- apples" not in doc_xml and "1. open" not in doc_xml)
    check("item text preserved", "apples are red" in doc_xml and "open the door" in doc_xml)
    check("prose untouched", "1986. It was a very good year" in doc_xml)
    check("single stray dash untouched", "- a single stray dash line" in doc_xml)

    fixed = Document(str(out))
    numbered = [
        p for p in fixed.paragraphs
        if p._p.find(qn("w:pPr")) is not None
        and p._p.find(qn("w:pPr")).find(qn("w:numPr")) is not None
    ]
    # 6 converted + 2 real List Bullet style items (style-based, no direct numPr)
    check("6 paragraphs gained direct w:numPr", len(numbered) == 6, f"got {len(numbered)}")

    # --- self-consistency: re-parse the OUTPUT --------------------------------
    res2 = parse_to_tree(str(out))
    run_analyzers(res2.tree)
    check("re-analysis of output raises NO fake-list flag (idempotent)", _flag_count(res2.tree) == 0)

    def _lists(node, acc):
        if isinstance(node, ListNode):
            acc.append(node)
        for ch in node.children:
            _lists(ch, acc)
        return acc

    out_lists = _lists(res2.tree.root, [])
    check("output re-parses with real ListNodes", len(out_lists) >= 2, f"got {len(out_lists)}")
    sizes = sorted(len(l.children) for l in out_lists)
    check("converted lists carry their 3 items each", sizes.count(3) >= 2, f"sizes {sizes}")

    # --- clean document: nothing flagged, nothing converted -------------------
    c = Document()
    c.core_properties.title = "Clean"
    c.add_paragraph("Just a paragraph - with an inline dash, not a list.")
    c.add_paragraph("Another ordinary line follows it.")
    clean_src = tmp / "clean.docx"
    c.save(str(clean_src))
    res3 = parse_to_tree(str(clean_src))
    run_analyzers(res3.tree)
    check("clean doc: no fake-list flag", _flag_count(res3.tree) == 0)

    # =========================================================================
    # PPTX: typed markers in a plain text box -> real a:buChar bullets
    # =========================================================================
    from pptx import Presentation
    from pptx.oxml.ns import qn as pqn
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
    tf = tb.text_frame
    tf.text = "- alpha item"
    tf.add_paragraph().text = "- beta item"
    tf.add_paragraph().text = "- gamma item"
    tf.add_paragraph().text = "Closing prose sentence here."
    # second text box: ordinary prose, must NOT flag
    tb2 = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
    tb2.text_frame.text = "Speaker notes style - prose with a dash inside."
    psrc = tmp / "deck.pptx"
    prs.save(str(psrc))

    pres = parse_to_tree(str(psrc))
    run_analyzers(pres.tree)
    check("pptx: typed text box flagged exactly once", _flag_count(pres.tree) == 1)
    check("pptx: FIX_LIST_STRUCTURE persists (honesty matrix)", _action_persists("FIX_LIST_STRUCTURE", "pptx"))

    pplans = [p for p in plan_remediations(pres.tree, apply_policy) if p.flag.code.value == FLAG]
    pexecs = execute_plans(pres.tree, pplans)
    check(
        "pptx: conversion executes successfully",
        len([e for e in pexecs if e.status.value == "success"]) == 1,
        str([(e.status.value, e.notes) for e in pexecs]),
    )

    pout = tmp / "deck_fixed.pptx"
    from app.writers import write_remediated

    write_remediated(psrc, pres.tree, pout, source_format="pptx")
    fixed_prs = Presentation(str(pout))
    sl = fixed_prs.slides[0]
    boxes = [sh for sh in sl.shapes if sh.has_text_frame]
    target_tf = boxes[0].text_frame
    bullet_paras = [
        p for p in target_tf.paragraphs
        if p._p.find(pqn("a:pPr")) is not None
        and p._p.find(pqn("a:pPr")).find(pqn("a:buChar")) is not None
    ]
    check("pptx: 3 lines gained real a:buChar bullets", len(bullet_paras) == 3, f"got {len(bullet_paras)}")
    texts = [p.text for p in target_tf.paragraphs]
    check("pptx: markers stripped from output lines", texts[:3] == ["alpha item", "beta item", "gamma item"], str(texts))
    check("pptx: prose line untouched (no bullet)", texts[3] == "Closing prose sentence here.")
    prose_tf = boxes[1].text_frame
    check(
        "pptx: prose text box untouched",
        prose_tf.paragraphs[0].text.startswith("Speaker notes style"),
    )

    # Idempotency: re-parse output -> bullets are real now -> no flag.
    pres2 = parse_to_tree(str(pout))
    run_analyzers(pres2.tree)
    check("pptx: re-analysis of output raises NO fake-list flag", _flag_count(pres2.tree) == 0)

    print(f"\nRESULT: {'all passed' if failures == 0 else str(failures) + ' FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
