"""WHERE each finding is in the document — the ``location`` block of the API.

The shared contract (``/pipeline/analyze`` and ``/pipeline/remediate``)::

    location = {
        "kind":      "pdf-region" | "image" | "text" | "table" | "document",
        "page":      int | None,          # 1-based
        "bbox":      [x0, y0, x1, y1] | None,  # PDF user space, origin bottom-left
        "pageSize":  [w, h] | None,
        "snippet":   str | None,          # <= 200 chars of the surrounding text
        "highlight": str | None,          # the exact offending substring OF snippet
        "thumbnail": "data:image/png;base64,..." | None,  # <= 240 px, images only
    }

Rules this module keeps, because a location that points at the wrong thing is
worse than none:

* Every string it returns is text that really is in the document. It never
  stitches pieces together into a sentence the document doesn't contain, and
  ``highlight`` is only set when it is a literal substring of ``snippet``.
* It reads the tree AS FOUND. Callers must build locations BEFORE any executor
  mutates the tree (a rewritten link text would otherwise be shown as the
  problem).
* Geometry comes from what the parser recorded (``properties["bbox"]``,
  ``properties["page_size"]``) or, for a few PDF objects whose geometry lives
  in the file itself (link annotations, form-field widgets, page boxes), from
  the source bytes — never from a guess.
* Thumbnails decode the image bytes the parser kept, or read the same image
  back from the source file by the reference the parser recorded (DOCX
  relationship id, PPTX shape id, PDF XObject name, HTML ``data:`` src). It
  NEVER fetches a URL. Pillow does the work; the count, pixel area and wall
  time are all capped, and any failure is simply ``thumbnail: None``.
* Nothing here may fail a request: every finding degrades to
  ``{"kind": "document", ...None}`` on any error.
"""

from __future__ import annotations

import base64
import io
import logging
import re
import time
import zipfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from app.models.accessibility import (
    AccessibilityTree,
    ContentKind,
    DocumentNode,
    HeadingNode,
    ImageNode,
    LinkNode,
    ListItemNode,
    ListNode,
    ParagraphNode,
    SectionNode,
    TableCellNode,
    TableNode,
    TableRowNode,
)

logger = logging.getLogger(__name__)

SNIPPET_MAX = 200
THUMB_MAX_PX = 240
MAX_THUMBNAILS = 24
# Budget for ALL thumbnails of one response. A scan-like PDF can hold hundreds
# of full-page images; the location block must not become the slow part.
THUMB_BUDGET_SECONDS = 4.0
# Refuse to decode an image bigger than this many pixels (decompression-bomb
# guard, well under Pillow's own warning threshold).
MAX_THUMB_SOURCE_PIXELS = 40_000_000
MAX_THUMB_SOURCE_BYTES = 12 * 1024 * 1024
MAX_THUMB_PNG_BYTES = 160 * 1024

_TABLE_FLAGS = {
    "TABLE_MISSING_HEADERS",
    "TABLE_HEADER_SCOPE_INVALID",
    "TABLE_CAPTION_MISSING",
    "TABLE_COMPLEX_NEEDS_SUMMARY",
    "TABLE_NESTED",
}
_ALT_QUALITY_FLAGS = {"ALT_TEXT_NOT_DESCRIPTIVE", "DECORATIVE_IMAGE_WITH_ALT"}
# Root-level COUNT findings that are really about individual elements; the
# location points at the first one.
_DOCUMENT_LEVEL_ELEMENT_RULES = {
    "FORM_FIELD_UNLABELED",
    "INPUT_AUTOCOMPLETE_MISSING",
    "POSITIVE_TABINDEX",
    "IFRAME_TITLE_MISSING",
    "LABEL_IN_NAME_MISMATCH",
}
# A typed list marker at the start of a line ("- ", "• ", "1. ", "a) ").
_LIST_MARKER_RE = re.compile(r"^\s*((?:[-*•‣◦▪–—·o])|(?:\(?\d{1,3}[.)])|(?:\(?[A-Za-z][.)]))(?=\s)")


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def empty_location(kind: str = "document", page: Optional[int] = None) -> Dict[str, Any]:
    return {
        "kind": kind,
        "page": page,
        "bbox": None,
        "pageSize": None,
        "snippet": None,
        "highlight": None,
        "thumbnail": None,
    }


def build_locations(
    tree: AccessibilityTree,
    violations: Iterable[Any],
    source_format: str,
    source_path: Optional[Path] = None,
    *,
    thumbnails: bool = True,
) -> Dict[str, Dict[str, Any]]:
    """``{violation_id: location}`` for every violation. Never raises."""
    fmt = (source_format or "").lower()
    out: Dict[str, Dict[str, Any]] = {}
    try:
        ctx = _Context(tree, fmt, Path(source_path) if source_path else None, thumbnails)
    except Exception:  # pragma: no cover - defensive
        logger.debug("finding_location: context build failed", exc_info=True)
        ctx = None
    for v in violations:
        vid = getattr(v, "violation_id", None)
        if not vid:
            continue
        if ctx is None:
            out[vid] = empty_location()
            continue
        try:
            out[vid] = ctx.locate(v)
        except Exception:
            logger.debug("finding_location: %s failed", vid, exc_info=True)
            out[vid] = empty_location()
    if ctx is not None:
        ctx.close()
    return out


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------


def _norm(text: Optional[str]) -> str:
    return " ".join(str(text or "").split())


def _own_text(node: Any) -> str:
    content = getattr(node, "content", None)
    if content is not None and getattr(content, "kind", None) == ContentKind.TEXT:
        return _norm(getattr(content, "text", None))
    return ""


def _deep_text(node: Any, limit: int = SNIPPET_MAX * 2) -> str:
    """The node's own text, else its descendants' text in reading order."""
    own = _own_text(node)
    if own:
        return own[:limit]
    parts: List[str] = []
    total = 0
    stack = list(reversed(getattr(node, "children", None) or []))
    while stack and total < limit:
        child = stack.pop()
        t = _own_text(child)
        if t:
            parts.append(t)
            total += len(t) + 1
            continue
        stack.extend(reversed(getattr(child, "children", None) or []))
    return " ".join(parts)[:limit]


def _window(text: str, needle: str) -> Tuple[Optional[str], Optional[str]]:
    """``(snippet, highlight)``: <= SNIPPET_MAX chars of ``text`` around ``needle``.

    ``highlight`` is ``needle`` only when it really occurs in the snippet.
    """
    text = _norm(text)
    needle = _norm(needle)
    if not text:
        return None, None
    if not needle:
        return text[:SNIPPET_MAX], None
    if len(needle) >= SNIPPET_MAX:
        cut = needle[:SNIPPET_MAX]
        return cut, cut
    idx = text.find(needle)
    if idx < 0:
        return text[:SNIPPET_MAX], None
    if len(text) <= SNIPPET_MAX:
        return text, needle
    room = SNIPPET_MAX - len(needle)
    start = max(0, idx - room // 2)
    end = start + SNIPPET_MAX
    if end > len(text):
        end = len(text)
        start = max(0, end - SNIPPET_MAX)
    snippet = text[start:end]
    return snippet, (needle if needle in snippet else None)


def _first_row_snippet(table: Any) -> Optional[str]:
    rows = [c for c in (getattr(table, "children", None) or []) if isinstance(c, TableRowNode)]
    if not rows:
        text = _deep_text(table)
        return text[:SNIPPET_MAX] or None
    cells = [c for c in (rows[0].children or []) if isinstance(c, TableCellNode)]
    texts = [_deep_text(c, 80) for c in cells] or [_deep_text(rows[0])]
    joined = " | ".join(t for t in texts if t is not None)
    joined = _norm(joined)
    return joined[:SNIPPET_MAX] or None


# ---------------------------------------------------------------------------
# The per-request context
# ---------------------------------------------------------------------------


class _Context:
    def __init__(self, tree: AccessibilityTree, fmt: str, source: Optional[Path], thumbnails: bool) -> None:
        self.tree = tree
        self.fmt = fmt
        self.source = source if (source is not None and source.exists()) else None
        self.nodes: Dict[str, Any] = {}
        self.parent: Dict[str, Any] = {}
        self.pos: Dict[str, int] = {}  # index of a node among its parent's children
        self._index(tree.root)
        self.thumbs = _Thumbnailer(self.source, fmt) if thumbnails else None
        self._pdf_reader: Any = None
        self._pdf_failed = False
        self._pdf_link_rects: Dict[int, List[Tuple[Optional[str], List[float]]]] = {}
        self._html_doc: Any = None

    # -- tree index ----------------------------------------------------------

    def _index(self, root: Any) -> None:
        stack: List[Tuple[Any, Any]] = [(root, None)]
        while stack:
            node, parent = stack.pop()
            nid = getattr(node, "id", None)
            if nid is not None and nid not in self.nodes:
                self.nodes[nid] = node
                if parent is not None:
                    self.parent[nid] = parent
            for i, child in enumerate(getattr(node, "children", None) or []):
                cid = getattr(child, "id", None)
                if cid is not None and cid not in self.pos:
                    self.pos[cid] = i
                stack.append((child, node))

    def _ancestors(self, node: Any) -> Iterable[Any]:
        cur = self.parent.get(getattr(node, "id", None))
        seen = 0
        while cur is not None and seen < 64:
            yield cur
            cur = self.parent.get(getattr(cur, "id", None))
            seen += 1

    def _position(self, node: Any) -> Tuple[Optional[List[Any]], int]:
        parent = self.parent.get(getattr(node, "id", None))
        if parent is None:
            return None, -1
        kids = getattr(parent, "children", None) or []
        i = self.pos.get(node.id, -1)
        if not (0 <= i < len(kids)) or kids[i] is not node:
            return None, -1
        return kids, i

    def _siblings_before(self, node: Any, limit: int = 4) -> List[Any]:
        kids, i = self._position(node)
        if kids is None:
            return []
        return list(reversed(kids[max(0, i - limit):i]))

    def _next_text_after(self, node: Any, limit: int = 6) -> str:
        kids, i = self._position(node)
        if kids is None:
            return ""
        for kid in kids[i + 1:i + 1 + limit]:
            t = _deep_text(kid)
            if t:
                return t
        return ""

    # -- geometry ------------------------------------------------------------

    def _page(self, node: Any) -> Optional[int]:
        for n in [node, *self._ancestors(node)]:
            page = getattr(getattr(n, "metadata", None), "page", None)
            if isinstance(page, int) and page >= 1:
                return page
        return None

    @staticmethod
    def _quad(value: Any) -> Optional[List[float]]:
        if not isinstance(value, (list, tuple)) or len(value) != 4:
            return None
        try:
            nums = [round(float(x), 2) for x in value]
        except (TypeError, ValueError):
            return None
        if nums[2] < nums[0]:
            nums[0], nums[2] = nums[2], nums[0]
        if nums[3] < nums[1]:
            nums[1], nums[3] = nums[3], nums[1]
        return nums

    @staticmethod
    def _pair(value: Any) -> Optional[List[float]]:
        if not isinstance(value, (list, tuple)) or len(value) != 2:
            return None
        try:
            w, h = round(float(value[0]), 2), round(float(value[1]), 2)
        except (TypeError, ValueError):
            return None
        return [w, h] if w > 0 and h > 0 else None

    def _page_size(self, node: Any, page: Optional[int]) -> Optional[List[float]]:
        for n in [node, *self._ancestors(node)]:
            props = getattr(getattr(n, "metadata", None), "properties", None) or {}
            size = self._pair(props.get("page_size"))
            if size:
                return size
        root_props = getattr(self.tree.root.metadata, "properties", None) or {}
        sizes = root_props.get("page_sizes")
        if page and isinstance(sizes, list) and 0 < page <= len(sizes):
            size = self._pair(sizes[page - 1])
            if size:
                return size
        if self.fmt == "pdf" and page:
            reader = self._pdf()
            if reader is not None:
                try:
                    box = reader.pages[page - 1].mediabox
                    return self._pair([float(box.width), float(box.height)])
                except Exception:
                    return None
        return None

    def _pdf(self) -> Any:
        if self._pdf_reader is not None or self._pdf_failed or self.source is None or self.fmt != "pdf":
            return self._pdf_reader
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(self.source))
            if getattr(reader, "is_encrypted", False):
                try:
                    reader.decrypt("")
                except Exception:
                    pass
            self._pdf_reader = reader
        except Exception:
            self._pdf_failed = True
            self._pdf_reader = None
        return self._pdf_reader

    def _pdf_link_bbox(self, node: Any, page: Optional[int]) -> Optional[List[float]]:
        """The link annotation's /Rect, matched by order on the page AND target.

        The parser emits one LinkNode per /Link annotation, per page, in
        annotation order. The k-th LinkNode on page P is therefore the k-th
        /Link annotation on page P — and we only accept the match when the
        annotation's URI equals the node's target, so a future change in how
        links are extracted can't silently pin a box on the wrong link.
        """
        if self.fmt != "pdf" or not page:
            return None
        kids, idx = self._position(node)
        if kids is None:
            return None
        ordinal = sum(1 for kid in kids[:idx] if isinstance(kid, LinkNode))
        rects = self._pdf_link_rects.get(page)
        if rects is None:
            rects = []
            reader = self._pdf()
            try:
                if reader is not None:
                    pg = reader.pages[page - 1]
                    for ref in (pg.get("/Annots") or []):
                        annot = ref.get_object() if hasattr(ref, "get_object") else ref
                        if not hasattr(annot, "get") or annot.get("/Subtype") != "/Link":
                            continue
                        uri = None
                        action = annot.get("/A")
                        action = action.get_object() if hasattr(action, "get_object") else action
                        if hasattr(action, "get") and action.get("/URI") is not None:
                            uri = str(action.get("/URI"))
                        rects.append((uri, [float(x) for x in annot.get("/Rect") or []]))
            except Exception:
                rects = []
            self._pdf_link_rects[page] = rects
        if ordinal >= len(rects):
            return None
        uri, rect = rects[ordinal]
        target = getattr(node, "target", None)
        if (target or None) != (uri or None):
            return None
        return self._quad(rect)

    # -- the location --------------------------------------------------------

    def locate(self, v: Any) -> Dict[str, Any]:
        rule = str(getattr(v, "rule_id", "") or "")
        node_id = getattr(getattr(v, "location", None), "node_id", None)
        node = self.nodes.get(node_id)
        if node is None:
            return empty_location()
        page = self._page(node)
        loc = empty_location(page=page)
        props = getattr(node.metadata, "properties", None) or {}
        bbox = self._quad(props.get("bbox"))

        if isinstance(node, DocumentNode):
            loc["kind"] = "document"
            if rule in _DOCUMENT_LEVEL_ELEMENT_RULES:
                self._document_level(loc, rule)
            if loc["page"] is not None and self.fmt == "pdf":
                loc["pageSize"] = self._page_size(node, loc["page"])
        elif isinstance(node, ImageNode):
            loc["kind"] = "image"
            if rule in _ALT_QUALITY_FLAGS and getattr(node, "alt_text", None):
                alt = _norm(node.alt_text)[:SNIPPET_MAX]
                loc["snippet"], loc["highlight"] = alt or None, alt or None
            else:
                caption = _norm(props.get("caption"))[:SNIPPET_MAX]
                loc["snippet"] = caption or None
            if self.thumbs is not None:
                loc["thumbnail"] = self.thumbs.for_image(node, page)
        elif isinstance(node, (TableNode, TableRowNode, TableCellNode)) or rule in _TABLE_FLAGS:
            loc["kind"] = "table"
            table = node
            if not isinstance(table, TableNode):
                for anc in self._ancestors(node):
                    if isinstance(anc, TableNode):
                        table = anc
                        break
            loc["snippet"] = _first_row_snippet(table)
        elif isinstance(node, LinkNode):
            self._link(node, loc, rule)
            if bbox is None:
                bbox = self._pdf_link_bbox(node, page)
        elif isinstance(node, HeadingNode):
            text = _own_text(node)
            if text:
                loc["snippet"], loc["highlight"] = _window(text, text)
            else:
                # An empty heading has no text of its own; show what follows it.
                after = self._next_text_after(node)
                loc["snippet"] = after[:SNIPPET_MAX] or None
            loc["kind"] = "text"
        elif isinstance(node, ParagraphNode) and rule == "LIST_STRUCTURE_INVALID":
            self._fake_list(node, loc)
        elif isinstance(node, ListNode):
            items = [_deep_text(c, 80) for c in (node.children or [])]
            joined = "\n".join(t for t in items if t)[:SNIPPET_MAX]
            loc["snippet"] = joined or None
            loc["kind"] = "text" if joined else "document"
        elif isinstance(node, SectionNode):
            # A page / slide level finding: show how that page starts.
            first = ""
            for kid in node.children or []:
                first = _deep_text(kid)
                if first:
                    break
            loc["snippet"] = first[:SNIPPET_MAX] or None
            loc["kind"] = "text" if first else "document"
        else:
            text = _deep_text(node)
            if text:
                loc["snippet"], loc["highlight"] = _window(text, text)
                loc["kind"] = "text"

        if bbox is not None:
            loc["bbox"] = bbox
        if loc["page"] is not None and loc["pageSize"] is None and (bbox is not None or self.fmt == "pdf"):
            loc["pageSize"] = self._page_size(node, loc["page"])
        if loc["bbox"] is not None and loc["pageSize"] is not None:
            loc["kind"] = "pdf-region"
        # Contract invariant: highlight is a literal substring of snippet.
        if loc["highlight"] and (not loc["snippet"] or loc["highlight"] not in loc["snippet"]):
            loc["highlight"] = None
        return loc

    def _link(self, node: LinkNode, loc: Dict[str, Any], rule: str) -> None:
        text = _own_text(node) or _deep_text(node)
        if self.fmt == "pdf":
            # PDF link "text" is the URI or a placeholder, not words on the page.
            text = _own_text(node)
            if text in ("(link)", "link"):
                text = ""
        loc["kind"] = "text"
        # The sentence the link sits in, when the tree has it: an enclosing
        # block (HTML), else — PowerPoint only, where a shape's paragraphs
        # precede its links on the same slide — a preceding paragraph that
        # literally contains the link text.
        context = ""
        for anc in self._ancestors(node):
            if isinstance(anc, (ParagraphNode, ListItemNode, HeadingNode, TableCellNode)):
                t = _deep_text(anc, 2000)
                if t and (not text or text in t):
                    context = t
                break
            if isinstance(anc, (SectionNode, DocumentNode)):
                break
        if not context and text and self.fmt == "pptx":
            for sib in self._siblings_before(node):
                if isinstance(sib, LinkNode):
                    continue
                t = _deep_text(sib, 2000)
                if t and text in t:
                    context = t
                    break
        if context:
            loc["snippet"], loc["highlight"] = _window(context, text)
        elif text:
            loc["snippet"], loc["highlight"] = _window(text, text)
        else:
            loc["kind"] = "document" if loc["page"] is None else "text"

    def _fake_list(self, node: ParagraphNode, loc: Dict[str, Any]) -> None:
        props = node.metadata.properties or {}
        run_ids = props.get("fake_list_run_ids")
        members: List[Any] = []
        if isinstance(run_ids, list):
            members = [self.nodes[i] for i in run_ids if i in self.nodes]
        if not members:
            members = [node]
        lines: List[str] = []
        total = 0
        for m in members:
            t = _own_text(m)
            if not t:
                continue
            if total + len(t) > SNIPPET_MAX and lines:
                break
            lines.append(t)
            total += len(t) + 1
        snippet = "\n".join(lines)[:SNIPPET_MAX]
        loc["kind"] = "text"
        loc["snippet"] = snippet or None
        if lines:
            m = _LIST_MARKER_RE.match(lines[0])
            loc["highlight"] = m.group(1) if m else lines[0][:SNIPPET_MAX]

    # -- document-level findings that are really about ONE element -------------

    def _document_level(self, loc: Dict[str, Any], rule: str) -> None:
        """Point a root-level COUNT finding ("3 form fields have no label") at
        the FIRST offending element, found with the parser's own rules."""
        if self.source is None:
            return
        try:
            if self.fmt == "html":
                el = self._html_first_offender(rule)
                if el is not None:
                    self._html_markup(loc, el)
            elif rule == "FORM_FIELD_UNLABELED" and self.fmt == "pdf":
                self._pdf_form_field(loc)
            elif rule == "FORM_FIELD_UNLABELED" and self.fmt == "docx":
                self._docx_form_field(loc)
        except Exception:
            logger.debug("finding_location: %s lookup failed", rule, exc_info=True)

    def _html_first_offender(self, rule: str) -> Any:
        from app.parsers import html_parser as hp

        if self._html_doc is None:  # parsed once per request, not once per finding
            self._html_doc = hp._parse_document(self.source.read_bytes())
        doc = self._html_doc
        if rule == "FORM_FIELD_UNLABELED":
            labels_for = hp._build_labels_for(doc)
            for ctrl in hp._iter_labelable_controls(doc):
                if not hp._control_has_accessible_name(ctrl, labels_for):
                    return ctrl
        elif rule == "INPUT_AUTOCOMPLETE_MISSING":
            for ctrl, _token in hp.iter_autocomplete_candidates(doc):
                return ctrl
        elif rule == "POSITIVE_TABINDEX":
            for el in hp.iter_positive_tabindex(doc):
                return el
        elif rule == "IFRAME_TITLE_MISSING":
            for tag in ("iframe", "frame"):
                for el in doc.iter(tag):
                    if any((el.get(a) or "").strip() for a in ("title", "aria-label", "aria-labelledby")):
                        continue
                    if (el.get("role") or "").strip().lower() in {"presentation", "none"} or el.get("hidden") is not None:
                        continue
                    return el
        elif rule == "LABEL_IN_NAME_MISMATCH":
            for tag in hp._LABEL_IN_NAME_TAGS:
                for el in doc.iter(tag):
                    label = (el.get("aria-label") or "").strip()
                    visible = hp._speech_normalize(hp._visible_subtree_text(el)) if label else ""
                    if label and len(visible) >= 2 and hp._speech_normalize(label).find(visible) == -1:
                        return el
        return None

    @staticmethod
    def _html_markup(loc: Dict[str, Any], el: Any) -> None:
        """The element's own markup, highlighted inside its parent's markup."""
        from lxml import etree

        tag = _norm(etree.tostring(el, method="html", with_tail=False, encoding="unicode"))
        if len(tag) > SNIPPET_MAX:
            tag = tag[: tag.find(">") + 1] if ">" in tag else tag[:SNIPPET_MAX]
        parent = el.getparent()
        outer = ""
        if parent is not None:
            outer = _norm(etree.tostring(parent, method="html", with_tail=False, encoding="unicode"))
        if tag and outer and tag in outer:
            loc["snippet"], loc["highlight"] = _window(outer, tag)
        elif tag:
            loc["snippet"], loc["highlight"] = _window(tag, tag)
        loc["kind"] = "text"

    def _pdf_form_field(self, loc: Dict[str, Any]) -> None:
        reader = self._pdf()
        if reader is None:
            return
        from app.parsers.pdf_parser import iter_acroform_fields

        root = reader.trailer.get("/Root", {})
        root = root.get_object() if hasattr(root, "get_object") else root
        acro = root.get("/AcroForm") if root else None
        acro = acro.get_object() if hasattr(acro, "get_object") else acro
        if not acro:
            return
        page_of: Dict[int, int] = {}
        for i, pg in enumerate(reader.pages, start=1):
            ref = getattr(pg, "indirect_reference", None)
            if ref is not None:
                page_of[ref.idnum] = i
        for fo in iter_acroform_fields(acro):
            tu = fo.get("/TU")
            if tu and str(tu).strip():
                continue
            name = _norm(str(fo.get("/T") or ""))
            widget = fo
            kids = fo.get("/Kids")
            if "/Rect" not in fo and kids:
                try:
                    first = kids[0]
                    widget = first.get_object() if hasattr(first, "get_object") else first
                except Exception:
                    widget = fo
            if name:
                loc["snippet"], loc["highlight"] = name[:SNIPPET_MAX], name[:SNIPPET_MAX]
                loc["kind"] = "text"
            page_ref = widget.get("/P") if hasattr(widget, "get") else None
            page = page_of.get(getattr(page_ref, "idnum", None)) if page_ref is not None else None
            if page is None:
                wref = getattr(widget, "indirect_reference", None)
                for i, pg in enumerate(reader.pages, start=1):
                    for a in pg.get("/Annots") or []:
                        if wref is not None and getattr(a, "idnum", None) == wref.idnum:
                            page = i
                            break
                    if page is not None:
                        break
            rect = self._quad(list(widget.get("/Rect") or [])) if hasattr(widget, "get") else None
            if page is not None:
                loc["page"] = page
                if rect is not None:
                    loc["bbox"] = rect
            return

    def _docx_form_field(self, loc: Dict[str, Any]) -> None:
        from lxml import etree

        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(self.source) as zf:
            info = zf.getinfo("word/document.xml")
            if info.file_size > 64 * 1024 * 1024:
                return
            root = etree.fromstring(zf.read(info), parser=etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True))
        body = root.find(f"{ns}body")
        if body is None:
            return
        for sdt in body.iter(f"{ns}sdt"):
            alias = sdt.find(f"{ns}sdtPr/{ns}alias")
            val = alias.get(f"{ns}val") if alias is not None else None
            if val and str(val).strip():
                continue
            content = sdt.find(f"{ns}sdtContent")
            ctrl_text = _norm("".join(t.text or "" for t in (content if content is not None else sdt).iter(f"{ns}t")))
            para = sdt.getparent()
            while para is not None and para.tag != f"{ns}p":
                para = para.getparent()
                if para is not None and para.tag == f"{ns}body":
                    para = None
                    break
            context = _norm("".join(t.text or "" for t in para.iter(f"{ns}t"))) if para is not None else ""
            if context and ctrl_text and ctrl_text in context:
                loc["snippet"], loc["highlight"] = _window(context, ctrl_text)
            elif ctrl_text:
                loc["snippet"], loc["highlight"] = _window(ctrl_text, ctrl_text)
            elif context:
                loc["snippet"] = context[:SNIPPET_MAX]
            loc["kind"] = "text" if loc["snippet"] else "document"
            return

    def close(self) -> None:
        if self.thumbs is not None:
            self.thumbs.close()


# ---------------------------------------------------------------------------
# Thumbnails
# ---------------------------------------------------------------------------


def png_thumbnail_data_uri(data: Any) -> Optional[str]:
    """``data:image/png;base64,...`` no wider/taller than THUMB_MAX_PX, or None.

    ``data`` is encoded image bytes or a PIL image. Never raises.
    """
    try:
        from PIL import Image

        if isinstance(data, (bytes, bytearray)):
            if not data or len(data) > MAX_THUMB_SOURCE_BYTES:
                return None
            im = Image.open(io.BytesIO(bytes(data)))
        elif hasattr(data, "size") and hasattr(data, "convert"):
            im = data
        else:
            return None
        w, h = im.size
        if w <= 0 or h <= 0 or w * h > MAX_THUMB_SOURCE_PIXELS:
            return None
        try:
            im.draft("RGB", (THUMB_MAX_PX * 2, THUMB_MAX_PX * 2))  # fast path for JPEG
        except Exception:
            pass
        if getattr(im, "n_frames", 1) > 1:
            try:
                im.seek(0)
            except Exception:
                pass
        if im.mode not in ("RGB", "RGBA", "L", "LA"):
            im = im.convert("RGBA" if ("A" in im.mode or im.mode == "P") else "RGB")
        im = im.copy()
        im.thumbnail((THUMB_MAX_PX, THUMB_MAX_PX))
        buf = io.BytesIO()
        im.save(buf, format="PNG", optimize=True)
        png = buf.getvalue()
        if len(png) > MAX_THUMB_PNG_BYTES:
            buf = io.BytesIO()
            im.convert("RGB").quantize(colors=128).save(buf, format="PNG", optimize=True)
            png = buf.getvalue()
            if len(png) > MAX_THUMB_PNG_BYTES:
                return None
        return "data:image/png;base64," + base64.b64encode(png).decode("ascii")
    except Exception:
        return None


class _Thumbnailer:
    def __init__(self, source: Optional[Path], fmt: str) -> None:
        self.source = source
        self.fmt = fmt
        self.count = 0
        self.deadline = time.monotonic() + THUMB_BUDGET_SECONDS
        self._docx_rels: Optional[Dict[str, str]] = None
        self._zip: Optional[zipfile.ZipFile] = None
        self._pptx: Any = None
        self._pdf: Any = None
        self._html: Any = None
        self._failed: set = set()

    def close(self) -> None:
        try:
            if self._zip is not None:
                self._zip.close()
        except Exception:
            pass

    def for_image(self, node: ImageNode, page: Optional[int]) -> Optional[str]:
        if self.count >= MAX_THUMBNAILS or time.monotonic() > self.deadline:
            return None
        props = node.metadata.properties or {}
        uri: Optional[str] = None
        # PDF image_b64 is the XObject's DECODED sample stream for Flate images
        # (not a PNG file), so for PDFs go to the source first.
        if self.fmt != "pdf":
            uri = self._from_b64(props.get("image_b64"))
        if uri is None and self.source is not None:
            try:
                data = self._from_source(props, page)
            except Exception:
                data = None
            if data is not None:
                uri = png_thumbnail_data_uri(data)
        if uri is None and self.fmt == "pdf":
            uri = self._from_b64(props.get("image_b64"))
        if uri is not None:
            self.count += 1
        return uri

    @staticmethod
    def _from_b64(b64: Any) -> Optional[str]:
        if not isinstance(b64, str) or not b64:
            return None
        if len(b64) > MAX_THUMB_SOURCE_BYTES * 4 // 3 + 4:
            return None
        try:
            raw = base64.b64decode(b64, validate=False)
        except Exception:
            return None
        return png_thumbnail_data_uri(raw)

    def _from_source(self, props: Dict[str, Any], page: Optional[int]) -> Any:
        if self.fmt in self._failed:
            return None
        try:
            if self.fmt == "docx":
                return self._docx(props.get("image_rid"))
            if self.fmt == "pptx":
                return self._pptx_image(props.get("shape_id"), page)
            if self.fmt == "pdf":
                return self._pdf_image(props.get("xobject"), page)
            if self.fmt == "html":
                return self._html_image(props.get("__xpath"))
        except (OSError, zipfile.BadZipFile):
            self._failed.add(self.fmt)
        return None

    # DOCX: the relationship id -> word/media/* bytes, straight from the zip.
    def _docx(self, rid: Any) -> Optional[bytes]:
        if not rid:
            return None
        if self._docx_rels is None:
            from lxml import etree

            self._zip = zipfile.ZipFile(self.source)
            rels: Dict[str, str] = {}
            try:
                xml = self._zip.read("word/_rels/document.xml.rels")
                root = etree.fromstring(xml, parser=etree.XMLParser(resolve_entities=False, no_network=True))
                for rel in root:
                    if (rel.get("TargetMode") or "").lower() == "external":
                        continue  # a linked picture: never fetched
                    rid_ = rel.get("Id")
                    target = rel.get("Target") or ""
                    if rid_ and target:
                        rels[rid_] = target
            except KeyError:
                pass
            self._docx_rels = rels
        target = self._docx_rels.get(str(rid))
        if not target:
            return None
        name = target.lstrip("/") if target.startswith("/") else f"word/{target}"
        # Normalise "word/../media/x.png" style paths.
        parts: List[str] = []
        for part in name.split("/"):
            if part == "..":
                if parts:
                    parts.pop()
            elif part and part != ".":
                parts.append(part)
        name = "/".join(parts)
        try:
            info = self._zip.getinfo(name)
        except KeyError:
            return None
        if info.file_size > MAX_THUMB_SOURCE_BYTES:
            return None
        return self._zip.read(info)

    # PPTX: the picture shape by id on its slide (group shapes included).
    def _pptx_image(self, shape_id: Any, page: Optional[int]) -> Optional[bytes]:
        if shape_id is None or not page:
            return None
        if self._pptx is None:
            from pptx import Presentation

            self._pptx = Presentation(str(self.source))
        slides = self._pptx.slides
        if page < 1 or page > len(slides):
            return None
        stack = list(slides[page - 1].shapes)
        while stack:
            shape = stack.pop()
            if getattr(shape, "shape_id", None) == shape_id:
                try:
                    blob = shape.image.blob
                except Exception:
                    return None
                return blob if blob and len(blob) <= MAX_THUMB_SOURCE_BYTES else None
            inner = getattr(shape, "shapes", None)
            if inner is not None:
                try:
                    stack.extend(list(inner))
                except Exception:
                    pass
        return None

    # PDF: the page's image XObject by name, decoded by pypdf (needs Pillow).
    def _pdf_image(self, name: Any, page: Optional[int]) -> Any:
        if not name or not page:
            return None
        if self._pdf is None:
            from pypdf import PdfReader

            self._pdf = PdfReader(str(self.source))
            if getattr(self._pdf, "is_encrypted", False):
                try:
                    self._pdf.decrypt("")
                except Exception:
                    return None
        pages = self._pdf.pages
        if page < 1 or page > len(pages):
            return None
        pg = pages[page - 1]
        key = "/" + str(name).lstrip("/")
        try:
            xobjects = pg["/Resources"]["/XObject"]
            xo = xobjects[key].get_object()
            w = int(xo.get("/Width") or 0)
            h = int(xo.get("/Height") or 0)
            if w <= 0 or h <= 0 or w * h > MAX_THUMB_SOURCE_PIXELS:
                return None
        except Exception:
            return None
        try:
            img = pg.images[key]
            return getattr(img, "image", None)
        except Exception:
            return None

    # HTML: only a data: URI src is ever decoded. Remote images are NEVER fetched.
    def _html_image(self, xpath: Any) -> Optional[bytes]:
        if not xpath:
            return None
        if self._html is None:
            from app.parsers import html_parser as hp

            self._html = hp._parse_document(self.source.read_bytes())
        found = self._html.getroottree().xpath(str(xpath))
        if not found:
            return None
        src = (found[0].get("src") or "").strip()
        if not src.lower().startswith("data:image/"):
            return None
        header, _, payload = src.partition(",")
        if ";base64" not in header.lower() or not payload:
            return None
        if len(payload) > MAX_THUMB_SOURCE_BYTES * 4 // 3 + 4:
            return None
        try:
            return base64.b64decode(payload, validate=False)
        except Exception:
            return None


__all__ = ["build_locations", "empty_location", "png_thumbnail_data_uri"]
